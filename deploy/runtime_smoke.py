"""Offline synthetic capability checks executed inside the final non-root image."""

import bz2
import io
import json
import lzma
import os
import sqlite3
import subprocess
import tempfile
import uuid
from pathlib import Path


def main():
    assert os.getuid() == 10001, "Runtime must retain the non-root application identity"
    os.environ["ENVIRONMENT"] = "test"
    os.environ["JWT_SECRET"] = "synthetic-runtime-smoke-not-a-production-secret"
    os.environ["SYNTHETIC_DATA_ENABLED"] = "false"
    os.environ["DATABASE_URL"] = (
        "postgresql+asyncpg://synthetic:synthetic@127.0.0.1:1/synthetic_runtime_test"
    )
    os.environ["REDIS_URL"] = "redis://127.0.0.1:1/0"
    assert uuid.uuid4().version == 4 and uuid.uuid1().version == 1
    assert bz2.decompress(bz2.compress(b"runtime")) == b"runtime"
    assert lzma.decompress(lzma.compress(b"runtime")) == b"runtime"
    assert sqlite3.connect(":memory:").execute("SELECT 42").fetchone() == (42,)
    from cryptography.fernet import Fernet

    cipher = Fernet(Fernet.generate_key())
    assert cipher.decrypt(cipher.encrypt(b"synthetic")) == b"synthetic"
    import duckdb

    assert duckdb.connect(":memory:").execute("SELECT 6 * 7").fetchone() == (42,)
    import pypdfium2 as pdfium
    from app.core.pdf_text import render_text_schedule

    # Import the real API and worker command paths without opening network services.
    from app.main import app
    from app.modules.intelligence.documents import extract_document
    from app.workers.celery_app import celery_app
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from PIL import Image, ImageDraw, ImageFont
    from pypdf import PdfReader

    assert app.openapi()["paths"] and celery_app is not None
    document = Document()
    document.add_paragraph("SYNTHETIC UAE India runtime")
    buffer = io.BytesIO()
    document.save(buffer)
    assert (
        Document(io.BytesIO(buffer.getvalue()))
        .paragraphs[0]
        .text.startswith("SYNTHETIC")
    )
    workbook = Workbook()
    workbook.active.append(["SYNTHETIC", "123.00000001"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    assert (
        load_workbook(io.BytesIO(buffer.getvalue())).active["B1"].value
        == "123.00000001"
    )
    content = render_text_schedule(
        [("SYNTHETIC UAE India runtime", True), ("مرحبا भारत 123.45", False)],
        draft=True,
    )
    assert "SYNTHETIC" in PdfReader(io.BytesIO(content)).pages[0].extract_text()
    pdf = pdfium.PdfDocument(content)
    bitmap = pdf[0].render(scale=0.5)
    assert bitmap.to_pil().width > 0
    bitmap.close()
    pdf.close()
    with tempfile.TemporaryDirectory(prefix="runtime-smoke-") as folder:
        image = Image.new("RGB", (1000, 200), "white")
        ImageDraw.Draw(image).text(
            (30, 50),
            "INVOICE 12345 TOTAL 100.00",
            font=ImageFont.load_default(size=48),
            fill="black",
        )
        image_path = Path(folder) / "SYNTHETIC.png"
        image.save(image_path)
        result = extract_document(str(image_path), "image/png", "AE")
        text = " ".join(
            block["text"] for page in result["pages"] for block in page["blocks"]
        )
        assert "12345" in text, (
            "Real offline OCR must recognize the synthetic invoice reference"
        )
        assert (
            result["network_document_transfer"] is False and result["model_artifacts"]
        )
    code = "import * as tf from '@tensorflow/tfjs'; const x=tf.tensor2d([[6]]); const y=tf.tensor2d([[7]]); if(x.matMul(y).dataSync()[0]!==42) throw Error('ML runtime failed'); x.dispose(); y.dispose();"
    subprocess.run(
        ["node", "--input-type=module", "-e", code],
        cwd="/srv/ml-engine",
        check=True,
        timeout=60,
    )
    subprocess.run(
        ["alembic", "--help"], check=True, stdout=subprocess.DEVNULL, timeout=30
    )
    print(
        json.dumps(
            {
                "status": "PASSED",
                "synthetic_only": True,
                "nonroot_uid": os.getuid(),
                "capabilities": [
                    "API imports",
                    "Celery imports",
                    "migration CLI",
                    "UUID fallback",
                    "compression",
                    "SQLite",
                    "cryptography",
                    "DuckDB",
                    "DOCX",
                    "exact XLSX",
                    "Arabic Hindi PDF",
                    "PDF raster",
                    "isolated offline OCR",
                    "Node TensorFlow CPU",
                ],
            }
        )
    )


if __name__ == "__main__":
    main()
