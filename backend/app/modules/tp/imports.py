"""Authorized user-supplied observations only; no commercial scraping or invented APIs."""

import csv
import hashlib
import io
import json
from collections.abc import Iterable
from datetime import date
from decimal import Decimal
from typing import Any

from pydantic import Field, model_validator

from app.modules.audit_execution.schemas import Strict
from app.modules.tp.schemas import Money


class ObservationInput(Strict):
    company_key: str = Field(min_length=1, max_length=100)
    name: str = Field(min_length=1, max_length=300)
    country: str = Field(pattern=r"^[A-Z]{2}$")
    industry: str = Field(min_length=1, max_length=100)
    fiscal_year: int = Field(ge=1900, le=2200)
    currency: str = Field(pattern=r"^[A-Z]{3}$")
    revenue: Money | None = None
    cost_sales: Money | None = None
    opex: Money | None = None
    assets: Money | None = None
    value: Decimal | None = Field(default=None, max_digits=28, decimal_places=12)
    activity: str = Field(min_length=1, max_length=2000)
    independence: str = Field(min_length=1, max_length=2000)
    related_party_exposure: str = Field(min_length=1, max_length=1000)
    extraordinary_events: str = Field(min_length=1, max_length=1000)
    observation_date: date
    tenor: str = Field(default="", max_length=200)
    credit_risk: str = Field(default="", max_length=1000)
    security: str = Field(default="", max_length=1000)
    seniority: str = Field(default="", max_length=200)
    rate_type: str = Field(default="", max_length=100)

    @model_validator(mode="after")
    def finite(self):
        for key in ["revenue", "cost_sales", "opex", "assets", "value"]:
            value = getattr(self, key)
            if value is not None and not value.is_finite():
                raise ValueError("Finite numeric observation required")
        return self


def parse(data, filename):
    rows: Iterable[dict[str, Any]]
    if len(data) > 10 * 1024 * 1024:
        raise ValueError("Comparable import limit is 10MiB")
    if filename.lower().endswith(".csv"):
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
        if not reader.fieldnames or len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise ValueError("Unique nonempty CSV column headers required")
        rows = reader
    elif filename.lower().endswith(".xlsx"):
        import zipfile

        from openpyxl import load_workbook

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds 30MiB")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        cells = wb.active.iter_rows(values_only=True)
        header = next(cells, None)
        if not header or len(header) != len(set(header)):
            wb.close()
            raise ValueError("Unique nonempty XLSX column headers required")
        rows = (dict(zip(header, row, strict=True)) for row in cells)
    else:
        raise ValueError("Only CSV or XLSX user imports are supported")
    output = []
    seen = set()
    for number, row in enumerate(rows, 2):
        if number > 50001:
            raise ValueError("Import at most 50,000 observations per dataset")
        if any(isinstance(value, str) and value.startswith("=") for value in row.values()):
            raise ValueError("Formula cells are not accepted as financial observations")
        values = {
            str(key): value for key, value in row.items() if value is not None and value != ""
        }
        item = ObservationInput.model_validate(values)
        key = (item.company_key, item.fiscal_year)
        if key in seen:
            raise ValueError("Duplicate company/year observation; no silent deduplication")
        seen.add(key)
        body = item.model_dump(mode="json")
        output.append((item, hashlib.sha256(json.dumps(body, sort_keys=True).encode()).hexdigest()))
    if not output:
        raise ValueError("Empty comparable dataset")
    return output
