"""Frozen source-linked narratives and deterministic numbers; never live AI or recalculation."""

import io
import json
from types import SimpleNamespace

from app.modules.demo.guard import DEMO_TENANT
from app.modules.ifrs.exports import render as render_document


def software_data(run):
    return bool(
        run.results.get("synthetic")
        or run.results.get("software_test_data")
        or getattr(run, "tenant_id", None) == DEMO_TENANT
        or any(x.get("synthetic") for x in getattr(run, "dependencies", []))
    )


def rows_for(run, released=False, client=False):
    rows = [
        ["AUDITOS UAE", "Transfer pricing professional workpaper"],
        ["Status", "APPROVED" if released else "DRAFT PROFESSIONAL REVIEW REQUIRED"],
        ["Document", run.kind],
        ["Run and methodology", str(run.id) + " / " + run.methodology_version],
        ["Period", str(run.period_start) + " to " + str(run.period_end)],
        ["SHA256", run.sha256],
    ]
    data = run.results
    if software_data(run):
        rows.append(
            [
                "SYNTHETIC SOFTWARE TEST DATA",
                "Not genuine market comparable evidence or a real economic conclusion",
            ]
        )
    if run.kind in {"localfile", "masterfile"}:
        rows.append(
            [
                "Purpose",
                "Source-linked annual transfer pricing documentation for professional review. Economic conclusions rely on the separately reviewed facts and authorized benchmark sources.",
            ]
        )
        rows.append(["Financial source profit", str(data["source_profit"])])
        for section in data["sections"]:
            rows.append([section["key"].replace("_", " ").title()])
            rows.append(["Analysis", section["narrative"]])
            rows.append(["Sources", ", ".join(section["evidence_ids"])])
            for rid in section["run_ids"]:
                rows.append(["Calculation reference", rid])
        for summary in data.get("calculation_summaries", []):
            rows.append(["Deterministic " + summary["kind"], summary["run_id"]])
            for key, value in summary["results"].items():
                rows.append(
                    [
                        key.replace("_", " "),
                        json.dumps(value, ensure_ascii=True)
                        if isinstance(value, (dict, list))
                        else str(value),
                    ]
                )
    elif not client:
        for key, value in data.items():
            if key in {"rules", "population", "company_results", "missing_observation_ids"}:
                continue
            rows.append([key.replace("_", " ").title()])
            if isinstance(value, list):
                for row in value:
                    rows.append([json.dumps(row, ensure_ascii=True)])
            elif isinstance(value, dict):
                for name, item in value.items():
                    rows.append([name, str(item)])
            else:
                rows.append([str(value)])
    return rows


def render(run, format, released=False, client=False):
    if client and format != "pdf":
        raise ValueError("Client releases contain the reviewed PDF only")
    if format == "docx" and run.kind in {"localfile", "masterfile"}:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        doc = Document()
        for border in doc.styles.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        for name in ["Title", "Heading 1", "Heading 2"]:
            doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
            doc.styles[name].font.name = "Arial"
            doc.styles[name].paragraph_format.keep_with_next = True
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(10)
        doc.styles["Normal"].paragraph_format.space_after = Pt(6)
        doc.styles["Normal"].paragraph_format.line_spacing = 1.15
        doc.styles["Normal"].paragraph_format.widow_control = True
        section = doc.sections[0]
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.left_margin = section.right_margin = Inches(0.75)
        section.top_margin = section.bottom_margin = Inches(0.75)
        section.header.paragraphs[0].text = "AUDITOS UAE | TRANSFER PRICING"
        doc.add_paragraph(
            "Transfer Pricing " + ("Local File" if run.kind == "localfile" else "Master File"),
            "Title",
        )
        doc.add_paragraph("APPROVED" if released else "DRAFT — PROFESSIONAL REVIEW REQUIRED")
        if software_data(run):
            doc.add_paragraph("SYNTHETIC SOFTWARE TEST DATA — NOT MARKET BENCHMARK EVIDENCE")
        doc.add_paragraph(
            f"Period: {run.period_start} to {run.period_end}. Annual version: {run.results['annual_version']}."
        )
        doc.add_paragraph(
            f"Immutable calculation: {run.id}. Methodology: {run.methodology_version}."
        )
        doc.add_paragraph(
            f"Approved financial source: {run.results['financial_source_id']}. Source profit: {run.results['source_profit']}."
        )
        doc.add_paragraph(
            f"Financial reconciliation: {run.results['financial_reconciliation_run_id']}."
        )
        for item in run.results["sections"]:
            doc.add_heading(item["key"].replace("_", " ").title(), 1)
            paragraphs = [p for p in item["narrative"].split("\n") if p.strip()]
            for index, paragraph in enumerate(paragraphs):
                text = doc.add_paragraph(paragraph)
                text.paragraph_format.keep_with_next = index == len(paragraphs) - 1
            sources = doc.add_paragraph(
                "Reviewed source references: " + ", ".join(item["evidence_ids"])
            )
            sources.paragraph_format.keep_with_next = bool(item["run_ids"])
            if item["run_ids"]:
                doc.add_paragraph("Calculation references: " + ", ".join(item["run_ids"]))
        doc.add_heading("Deterministic calculation summaries", 1)
        for item in run.results.get("calculation_summaries", []):
            doc.add_heading(item["kind"].replace("_", " ").title(), 2)
            for key, value in item["results"].items():
                doc.add_paragraph(key.replace("_", " ") + ": " + str(value))
        footer = section.footer.paragraphs[0]
        footer.text = f"AuditOS UAE | {str(run.id)[:8]} | {run.sha256[:16]} | Page "
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        output = io.BytesIO()
        doc.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    safe = SimpleNamespace(
        id=run.id,
        sha256=run.sha256,
        methodology_version=run.methodology_version,
        created_by=run.created_by,
        created_at=run.created_at,
        inputs={key: value for key, value in run.inputs.items() if key != "observation_ids"},
        dependencies=run.dependencies,
        results={
            **{
                key: value
                for key, value in run.results.items()
                if key not in {"population", "company_results", "missing_observation_ids"}
            },
            "software_test_data": software_data(run),
        },
    )
    return render_document(
        safe,
        format,
        released,
        rows_override=rows_for(run, released, client),
        document_title="Transfer Pricing Local File"
        if run.kind == "localfile"
        else "Transfer Pricing Master File"
        if run.kind == "masterfile"
        else "Transfer Pricing Analysis",
    )
