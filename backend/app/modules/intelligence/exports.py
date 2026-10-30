"""Exact source-value exports; human conclusions remain distinct from generated tables."""

import io
import json


def render(event, format):
    if event["kind"] == "statutory-report":
        from app.modules.regulatory.report_exports import render as report_render

        return report_render(event, format)
    body = event["body"]
    rows = (
        body.get("rows")
        or body.get("items")
        or body.get("suggestions")
        or body.get("matches")
        or []
    )
    metadata = {
        "Firm": event.get("firm_name", ""),
        "Reference": event["id"],
        "Result kind": event["kind"],
        "Version": event["version"],
        "SHA-256": event["sha256"],
        "Country": body.get("country"),
        "Framework": body.get("reporting_framework"),
        "Review status": body.get("review_status", "UNVERIFIED"),
        "Export timestamp": event.get("exported_at", "NOT_RECORDED"),
        "Source freshness at export": event.get("freshness", "NOT_EVALUATED"),
        "Export status": "STALE_HISTORICAL_RESULT_REQUIRES_RECALCULATION"
        if event.get("freshness") == "STALE"
        else "CURRENT_RETAINED_RESULT"
        if event.get("freshness") == "CURRENT"
        else "FRESHNESS_NOT_EVALUATED",
        "Prepared by": ", ".join(body.get("contributors", [])),
        "Reviewed by": body.get("reviewer"),
        "Generated result timestamp": event["created_at"],
        "Financial convention": "Exact Decimal strings; no audit opinion or evidence sufficiency inferred",
    }
    if format == "xlsx":
        metadata["XLSX text contract"] = "exact-text-with-continuations/1"
    if event["kind"] in {"workflow-source", "country-workflow"}:
        metadata.update(
            {
                "Workflow": body.get("workflow", body.get("table_type")),
                "Population completeness": body.get("population_completeness"),
                "Statutory calculation qualified": body.get(
                    "statutory_calculation_qualified", False
                ),
                "Professional conclusion": body.get("professional_conclusion", "NOT_DETERMINED"),
                "Workflow limitations": body.get("limitations", []),
            }
        )
    if body.get("box_comparison_scope", {}).get("mode") == "SELECTED_REQUIRED_BOXES":
        metadata["Comparison scope"] = (
            "Selected VAT201 boxes only: "
            + ", ".join(body["box_comparison_scope"]["required_boxes"])
            + ". Whole-return completeness is not determined."
        )
    if body.get("excluded_box_rows"):
        metadata["Excluded VAT box rows"] = (
            f"{len(body['excluded_box_rows'])} outside the selected box scope; full details are retained in the XLSX Excluded VAT boxes sheet and original source result"
        )
    if body.get("obligation_population_scope") == "ARREARS_ONLY_NOT_COMPLETE_MONTHLY_RETURN":
        metadata["Comparison scope"] = (
            "Arrears only; regular monthly obligations, returns and payments are outside this comparison. Salary disbursement does not establish PF contribution payment."
        )
    summary = {
        k: v
        for k, v in body.items()
        if k
        in {
            "media_review_contract",
            "media_document_binding",
            "media_scope",
            "identity_authenticity_established",
            "media_clearance_issued",
            "automatic_rejection",
            "grant_utilisation_contract",
            "grant_document_binding",
            "cash_availability_established",
            "legal_eligibility_established",
            "certificate_issued",
            "source_amount_basis",
            "population_basis",
            "currency_scope",
            "review_window",
            "totals",
            "totals_basis",
            "totals_by_currency",
            "unique_findings",
            "workpaper_reference_count",
            "reference_count_basis",
            "duplicate_references_eliminated",
            "issues",
            "blocked",
            "contract",
            "control_comparison",
            "booking_date_review_count",
            "import_nil_scope",
            "document_series_contract",
            "document_series_scope",
            "document_series_totals",
            "document_series_controls",
            "count_unit",
            "bank_account_contract",
            "bank_observation_scope",
            "bank_account_evidence",
            "bank_account_scope",
            "bank_profile_references",
            "bank_authentication_status",
            "allocation_controls",
            "bank_consumption_manifest",
            "bank_consumption_review",
            "payment_scope",
            "matching_contract",
            "candidate_checks",
            "candidate_check_limit",
            "ambiguity_link_count",
            "source_totals",
            "destination_totals",
            "section",
            "sampling_method",
            "seed",
            "source_batch_id",
            "source_snapshot_id",
            "source_filename",
            "coverage",
            "adjusted_book_balance",
            "country_guidance",
            "source_roles",
            "exception_count",
            "blocked",
            "parameters",
            "bank_account_binding",
            "bank_account_scope",
            "nil_population",
            "nil_declarations",
            "execution_mode",
            "ml_execution_mode",
            "language_model_dependency",
            "calculation_engine",
            "total_ecl_by_currency",
            "count_bridges",
            "payroll_ledger_scope",
            "native_format",
            "native_metadata",
            "workbook_metadata",
            "csv_mapping",
            "supporting_table_id",
            "supporting_table_metadata",
            "supporting_scope_assessment",
            "supporting_table_basis",
            "supporting_side",
            "supporting_balance_date",
            "supporting_amount_sign",
            "account_identity_verified",
            "native_controls_verified",
            "expected_total",
            "expected_total_complete",
            "msme_interest_contract",
            "msme_interest_document_binding",
            "msme_evidence_contract",
            "msme_document_binding",
            "principal_timing_exception_count",
            "msme_tax_contract",
            "msme_tax_document_binding",
            "tax_reference",
            "tax_year_start",
            "tax_observed_through",
            "net_expected_tax_adjustment",
            "trace_record_count",
            "expected_total_basis",
            "rule_document_binding",
            "rule_document_version_id",
            "reported_total",
            "currency",
            "withholding_role",
            "comparison_basis",
            "selected_current_totals",
            "assessment_contract_version",
            "vat_reporting_currency_contract",
            "vat_reporting_currency",
            "vat_fx_basis",
            "annual_gst_assessment",
            "gst_reporting_currency_contract",
            "gst_reporting_currency",
            "gst_fx_basis",
            "merchant_reserve_contract",
            "merchant_reserve_controls",
            "benefit_service_controls",
            "benefit_service_calendar_contract",
            "service_calendar_basis",
            "reserve_control_receipts",
            "inventory_fifo_contract",
            "inventory_fifo_scope",
            "inventory_fifo_document_binding",
            "annual_statement_scope_contract",
            "annual_statement_scope_basis",
            "annual_statement_scopes",
            "box_comparison_scope",
            "turnover_contract",
            "turnover_basis",
            "turnover_quantification",
            "turnover_difference_totals",
            "vat201_layout",
            "vat201_arithmetic_contract",
            "vat201_arithmetic_scope",
            "vat201_arithmetic_source",
            "vat201_arithmetic_error_count",
            "invoice_evidence_scope",
            "outward_gst_scope",
            "excluded_box_controls",
            "included_box_controls",
        }
    }

    if body.get("turnover_adjustment_rows"):
        metadata["Turnover adjustment detail"] = (
            f"{len(body['turnover_adjustment_rows'])} signed adjustments retained with full reasons and assessments in the Turnover adjustments XLSX sheet and original result"
        )
    if body.get("vat201_arithmetic_checks"):
        metadata["VAT201 arithmetic detail"] = (
            f"{len(body['vat201_arithmetic_checks'])} independent checks retained in the VAT201 arithmetic XLSX sheet and original result; scalar Boxes 12-14 use tax only, with zero base and adjustment."
        )

    if isinstance(body.get("count_bridges"), list):
        summary["count_bridges"] = (
            f"{len(body['count_bridges'])} bridge rows; complete original count details and individual movement links are retained in the XLSX Count date bridges and Inventory movement links sheets"
        )
    for field in (
        "fifo_allocations",
        "fifo_event_costs",
        "fifo_closing_layers",
        "inventory_fifo_controls",
    ):
        if isinstance(body.get(field), list):
            summary[field] = (
                f"{len(body[field])} records; complete detail is retained in its dedicated FIFO XLSX sheet"
            )
    if isinstance(body.get("inventory_quantity_controls"), list):
        summary["inventory_quantity_controls"] = (
            f"{len(body['inventory_quantity_controls'])} controlled partitions; complete row and directional quantity controls are retained in the XLSX Inventory quantity controls sheet"
        )

    if isinstance(body.get("transaction_dates"), list):
        summary["transaction_dates"] = (
            f"{len(body['transaction_dates'])} retained transaction date records; complete detail is retained in the XLSX Bank transaction dates sheet"
        )
    if isinstance(body.get("payroll_posting_links"), list):
        summary["payroll_posting_links"] = (
            f"{len(body['payroll_posting_links'])} individual source posting links; complete detail is retained in the XLSX Payroll posting links sheet"
        )

    def text(value):
        return (
            json.dumps(value, ensure_ascii=False, sort_keys=True)
            if isinstance(value, (dict, list))
            else ""
            if value is None
            else str(value)
        )

    native_metadata = body.get("native_metadata")
    message_controls = (
        native_metadata.get("message_controls", [])
        if body.get("native_format") == "swift-mt940" and isinstance(native_metadata, dict)
        else []
    )
    if message_controls:
        summary["native_metadata"] = {
            **body["native_metadata"],
            "message_controls": f"{len(message_controls)} message control records retained in the source result and full XLSX MT940 message controls sheet",
        }
    page_controls = (
        native_metadata.get("page_controls", [])
        if body.get("native_format") == "camt053" and isinstance(native_metadata, dict)
        else []
    )
    if page_controls:
        summary["native_metadata"] = {
            **body["native_metadata"],
            "page_controls": f"{len(page_controls)} page control records retained in the source result and full XLSX CAMT page controls sheet",
        }

    buffer = io.BytesIO()
    if format == "xlsx":
        from openpyxl import Workbook

        from app.modules.intelligence.xlsx_values import ExactTextWriter, XlsxValueLimitError

        workbook = Workbook(write_only=True)
        writer = ExactTextWriter(workbook, text)
        append = writer.append

        try:
            cover = workbook.create_sheet("Source and review")
            for key, value in metadata.items():
                append(cover, [key, value])
            for key, value in summary.items():
                append(cover, [key, value])
            if body.get("matches") and (
                event["kind"] == "reconciliation" or body.get("workflow") == "bank-brs"
            ):
                proposals = workbook.create_sheet("Matched proposals")
                fields = [
                    "left_ids",
                    "right_ids",
                    "rule",
                    "variance",
                    "auditor_status",
                    "limitation",
                    "auditor",
                    "auditor_comments",
                    "auditor_at",
                ]
                append(proposals, fields)
                for match in body["matches"]:
                    append(proposals, [match.get(field) for field in fields])
            if event["kind"] == "reconciliation":
                from app.modules.intelligence.matching_review import review_summary

                decisions = workbook.create_sheet("Reconciliation decisions")
                for key, value in review_summary(body).items():
                    append(decisions, [key, value])
                append(decisions, [])
                fields = ["item_id", "record_id", "kind", "reason", "actor"]
                append(decisions, fields)
                for link in body.get("exceptions", []):
                    append(decisions, [link.get(field) for field in fields])
            if body.get("ambiguous"):
                from app.modules.intelligence.matching_review import saved_candidates_complete

                candidates = workbook.create_sheet("Matching candidates")
                append(
                    candidates,
                    ["left_id", "candidate_id", "candidate_count", "retention", "reason"],
                )
                retention = (
                    "COMPLETE_RETAINED_CANDIDATES"
                    if saved_candidates_complete(body)
                    else "LEGACY_OR_INCOMPLETE_CANDIDATES_REGENERATE"
                )
                for group in body["ambiguous"]:
                    for candidate in group["candidate_ids"]:
                        append(
                            candidates,
                            [
                                group["left_id"],
                                candidate,
                                group.get("candidate_count"),
                                retention,
                                group.get("reason"),
                            ],
                        )
            if body.get("workflow") == "adverse-media-review":
                documents = workbook.create_sheet("Media documents")
                fields = ["type", "id", "record_id", "number", "sha256"]
                append(documents, fields)
                for document in body.get("retained_media_documents", []):
                    append(documents, [document.get(field) for field in fields])
            if body.get("workflow") == "grant-utilisation":
                documents = workbook.create_sheet("Grant documents")
                fields = ["type", "id", "record_id", "number", "sha256"]
                append(documents, fields)
                for document in body.get("retained_grant_documents", []):
                    append(documents, [document.get(field) for field in fields])
            if body.get("turnover_adjustment_rows"):
                adjustments = workbook.create_sheet("Turnover adjustments")
                fields = list(
                    dict.fromkeys(key for row in body["turnover_adjustment_rows"] for key in row)
                )
                append(adjustments, fields)
                for row in body["turnover_adjustment_rows"]:
                    append(adjustments, [row.get(field) for field in fields])
            if body.get("vat201_arithmetic_checks"):
                arithmetic = workbook.create_sheet("VAT201 arithmetic")
                fields = [
                    "id",
                    "source_role",
                    "box",
                    "field",
                    "formula",
                    "expected",
                    "observed",
                    "difference",
                    "status",
                    "missing_boxes",
                    "source_row_ids",
                ]
                append(arithmetic, fields)
                for check in body["vat201_arithmetic_checks"]:
                    append(arithmetic, [check.get(field) for field in fields])
            source = workbook.create_sheet("Rows")
            keys = list(dict.fromkeys(key for row in rows for key in row))
            append(source, keys)
            for row in rows:
                exported_row = row
                if body.get("workflow") == "india-msme-interest":
                    exported_row = dict(row)
                    for field, sheet in (
                        ("segments", "MSE interest intervals"),
                        ("balance_events", "MSE interest events"),
                        ("monthly_rests", "MSE monthly rests"),
                        ("allocations", "MSE gross principal links"),
                    ):
                        exported_row[field] = (
                            f"{len(row.get(field, []))} records; complete detail in {sheet} sheet"
                        )
                    if row.get("interest_policy"):
                        exported_row["interest_policy"] = (
                            "Complete assessed convention in MSE interest policies sheet"
                        )
                if body.get("workflow") == "india-establishment-charges":
                    exported_row = {
                        **row,
                        "member_trace": f"{len(row.get('member_trace', []))} records; complete detail in Establishment charge members sheet",
                        "member_source_ids": f"{len(row.get('member_source_ids', []))} IDs; complete detail in Establishment charge members sheet",
                    }
                if body.get("workflow") == "india-payroll-basis-proration":
                    exported_row = {
                        **row,
                        "segment_trace": f"{len(row.get('segment_trace', []))} records; complete detail in Payroll basis segments sheet",
                    }
                if body.get("workflow") == "uae-benefit-review" and isinstance(
                    row.get("service_calendar_records"), list
                ):
                    exported_row = {
                        **row,
                        "service_calendar_records": f"{len(row['service_calendar_records'])} records; complete detail in Benefit service exclusions sheet",
                        "service_calendar_source_row_ids": f"{len(row['service_calendar_records'])} source IDs retained in Benefit service exclusions sheet",
                    }
                if body.get("workflow") == "merchant-reserve-rollforward":
                    exported_row = dict(row)
                    for field, sheet in (
                        ("movement_trace", "Reserve movements"),
                        ("movement_ids", "Reserve movements"),
                        ("source_balance_ids", "Reserve balance links"),
                    ):
                        if isinstance(row.get(field), list):
                            exported_row[field] = (
                                f"{len(row[field])} records; complete detail in {sheet} sheet"
                            )
                append(source, [exported_row.get(key) for key in keys])
            if body.get("workflow") == "india-payroll-basis-proration":
                segment_sheet = workbook.create_sheet("Payroll basis segments")
                records = [
                    {"result_row_id": result["id"], **segment}
                    for result in body.get("rows", [])
                    for segment in result.get("segment_trace", [])
                ]
                fields = list(dict.fromkeys(field for record in records for field in record))
                append(segment_sheet, fields)
                for record in records:
                    append(segment_sheet, [record.get(field) for field in fields])
            if body.get("workflow") == "india-establishment-charges":
                members_sheet = workbook.create_sheet("Establishment charge members")
                records = [
                    {"result_row_id": result["id"], **member}
                    for result in body.get("rows", [])
                    for member in result.get("member_trace", [])
                ]
                fields = list(dict.fromkeys(field for record in records for field in record))
                append(members_sheet, fields)
                for record in records:
                    append(members_sheet, [record.get(field) for field in fields])
            if body.get("workflow") == "india-msme-annual-tax":
                populations = {
                    "MSE annual tax totals": [
                        {"metric": metric, **values}
                        for metric, values in body.get("tax_totals", {}).items()
                    ],
                    "MSE tax bank residuals": body.get("unallocated_bank_rows", []),
                }
                for title, key in (
                    ("MSE tax bases", "tax_basis"),
                    ("MSE tax PNL interest", "interest_pnl_assessment"),
                ):
                    populations[title] = [
                        {"invoice_id": invoice["id"], **invoice[key]}
                        for invoice in body.get("invoice_tax_details", [])
                        if invoice.get(key)
                    ]
                populations["MSE tax invoice scope"] = [
                    {
                        key: value
                        for key, value in invoice.items()
                        if not isinstance(value, (dict, list))
                    }
                    for invoice in body.get("invoice_tax_details", [])
                ]
                for title, key in (
                    ("MSE tax settlements", "tax_settlements"),
                    ("MSE tax principal links", "allocations"),
                ):
                    records = []
                    for invoice in body.get("invoice_tax_details", []):
                        for record in invoice.get(key, []):
                            records.append(
                                {
                                    "result_invoice_id": invoice["id"],
                                    **{
                                        field: value
                                        for field, value in record.items()
                                        if field != "original_bank_allocation"
                                    },
                                    **{
                                        "original_" + field: value
                                        for field, value in record.get(
                                            "original_bank_allocation", {}
                                        ).items()
                                    },
                                }
                            )
                    populations[title] = records
                for title, key in (
                    ("MSE tax invoice facts", "retained_invoice_facts"),
                    ("MSE tax supplier facts", "retained_supplier_assessment"),
                ):
                    populations[title] = [
                        {"result_invoice_id": invoice["id"], **invoice[key]}
                        for invoice in body.get("invoice_tax_details", [])
                        if invoice.get(key)
                    ]
                for title, records in populations.items():
                    sheet = workbook.create_sheet(title)
                    fields = list(dict.fromkeys(field for record in records for field in record))
                    append(sheet, fields)
                    for record in records:
                        append(sheet, [record.get(field) for field in fields])
            if body.get("workflow") == "india-msme-interest":
                for key, title in (
                    ("segments", "MSE interest intervals"),
                    ("balance_events", "MSE interest events"),
                    ("monthly_rests", "MSE monthly rests"),
                    ("allocations", "MSE gross principal links"),
                ):
                    sheet = workbook.create_sheet(title)
                    records = []
                    for result in body.get("rows", []):
                        for record in result.get(key, []):
                            flat = {
                                "result_row_id": result["id"],
                                **{
                                    field: value
                                    for field, value in record.items()
                                    if field != "source_record"
                                },
                            }
                            flat.update(
                                {
                                    "source_" + field: value
                                    for field, value in record.get("source_record", {}).items()
                                }
                            )
                            records.append(flat)
                    fields = list(dict.fromkeys(field for record in records for field in record))
                    append(sheet, fields)
                    for record in records:
                        append(sheet, [record.get(field) for field in fields])
                for key, title in (
                    ("retained_interest_rate_rows", "MSE interest rate sources"),
                    ("interest_bank_controls", "MSE interest bank controls"),
                    ("interest_policy", "MSE interest policies"),
                ):
                    records = (
                        [
                            {"result_row_id": row["id"], **row[key]}
                            for row in body.get("rows", [])
                            if row.get(key)
                        ]
                        if key == "interest_policy"
                        else body.get(key, [])
                    )
                    sheet = workbook.create_sheet(title)
                    fields = list(dict.fromkeys(field for record in records for field in record))
                    append(sheet, fields)
                    for record in records:
                        append(sheet, [record.get(field) for field in fields])
            if (
                body.get("workflow") == "uae-benefit-review"
                and body.get("service_calendar_basis") == "DATED_CONTROLLED_SOURCE"
            ):
                calendar = workbook.create_sheet("Benefit service exclusions")
                fields = [
                    "id",
                    "employee_id",
                    "service_start",
                    "service_end",
                    "regime",
                    "cohort",
                    "work_model",
                    "scheme",
                    "kind",
                    "exclusion_start",
                    "exclusion_end",
                    "days",
                    "evidence_reference",
                    "assessment",
                ]
                append(calendar, ["result_row_id", *fields])
                for employee in body.get("rows", []):
                    for record in employee.get("service_calendar_records", []):
                        append(calendar, [employee["id"], *[record.get(field) for field in fields]])
            if body.get("workflow") == "merchant-reserve-rollforward":
                movements = workbook.create_sheet("Reserve movements")
                identity = ["merchant", "processor", "currency", "source_balance_id"]
                fields = [
                    "id",
                    "settlement_id",
                    "date",
                    "sequence",
                    "kind",
                    "amount",
                    "reference",
                    "evidence_reference",
                    "balance_after",
                ]
                append(movements, ["result_row_id", *identity, *fields])
                links = workbook.create_sheet("Reserve balance links")
                append(links, ["result_row_id", "account", "currency", "source_balance_id"])
                for position in body.get("rows", []):
                    for movement in position.get("movement_trace", []):
                        append(
                            movements,
                            [
                                position["id"],
                                *[position.get(field) for field in identity],
                                *[movement.get(field) for field in fields],
                            ],
                        )
                    for identifier in position.get("source_balance_ids", []):
                        append(
                            links,
                            [position["id"], position["account"], position["currency"], identifier],
                        )
            if body.get("workflow") == "party-balance-confirmations":
                links = workbook.create_sheet("Confirmation item links")
                documents = workbook.create_sheet("Confirmation documents")
                identity_fields = ["id", "account", "counterparty", "side", "currency"]
                append(
                    links,
                    [
                        "result_row_id",
                        *identity_fields[1:],
                        "reply_scope_membership",
                        "open_item_source_id",
                    ],
                )
                append(
                    documents,
                    [
                        "result_row_id",
                        *identity_fields[1:],
                        "confirmation_status",
                        "request_basis",
                        "scope_balance",
                        "requested_amount",
                        "document_role",
                        "document_date",
                        "evidence_version_id",
                        "source_sha256",
                        "reference",
                        "contact_verified_declaration",
                        "independent_receipt_declaration",
                        "scope_assessment",
                        "followup_kind",
                        "followup_assessment",
                        "evidence_status",
                    ],
                )
                for position in body.get("rows", []):
                    if position.get("kind") != "party-confirmation":
                        continue
                    identity = [position.get(field) for field in identity_fields]
                    for identifier in position.get("open_item_source_ids", []):
                        membership = "FULL_BALANCE_SCOPE"
                        if position.get("confirmation_status") == "partially-received":
                            covered = position.get("covered_item_source_ids")
                            membership = (
                                "UNASSESSED_PARTIAL_SCOPE"
                                if covered is None
                                else "DECLARED_REPLY_COVERED"
                                if identifier in covered
                                else "NOT_REPLY_COVERED"
                            )
                        append(links, [*identity, membership, identifier])
                    evidence = position.get("confirmation_evidence") or {}
                    for prefix, date_field in (
                        ("request", "requested_on"),
                        ("response", "responded_on"),
                        ("followup", "followup_on"),
                    ):
                        if evidence.get(prefix + "_evidence_version_id"):
                            append(
                                documents,
                                [
                                    *identity,
                                    position.get("confirmation_status"),
                                    evidence.get("request_basis"),
                                    evidence.get("scope_balance"),
                                    evidence.get("requested_amount"),
                                    prefix,
                                    evidence.get(date_field),
                                    evidence[prefix + "_evidence_version_id"],
                                    evidence[prefix + "_source_sha256"],
                                    evidence[prefix + "_reference"],
                                    evidence.get("contact_verified"),
                                    evidence.get("independently_received"),
                                    evidence.get("scope_assessment"),
                                    evidence.get("followup_kind") if prefix == "followup" else None,
                                    evidence.get("followup_assessment")
                                    if prefix == "followup"
                                    else None,
                                    "RETAINED_IDENTITY_ONLY",
                                ],
                            )
            if body.get("workflow") == "inventory-custody-confirmations":
                documents = workbook.create_sheet("Custody documents")
                append(
                    documents,
                    [
                        "result_row_id",
                        "position_source_id",
                        "custodian",
                        "site",
                        "sku",
                        "lot",
                        "uom",
                        "document_role",
                        "document_date",
                        "evidence_version_id",
                        "source_sha256",
                        "reference",
                        "evidence_status",
                    ],
                )
                for position in body.get("rows", []):
                    evidence = position.get("confirmation_evidence") or {}
                    for role, date_field in (
                        ("request", "requested_on"),
                        ("response", "responded_on"),
                        ("followup", "followup_on"),
                    ):
                        if evidence.get(role + "_evidence_version_id"):
                            append(
                                documents,
                                [
                                    *[
                                        position.get(field)
                                        for field in (
                                            "id",
                                            "position_source_id",
                                            "custodian",
                                            "site",
                                            "sku",
                                            "lot",
                                            "uom",
                                        )
                                    ],
                                    role,
                                    evidence.get(date_field),
                                    evidence[role + "_evidence_version_id"],
                                    evidence[role + "_source_sha256"],
                                    evidence[role + "_reference"],
                                    "RETAINED_IDENTITY_ONLY",
                                ],
                            )
            for key, title, fields in (
                (
                    "custody_quantity_controls",
                    "Custody quantity controls",
                    [
                        "source_role",
                        "position_id",
                        "custodian",
                        "site",
                        "sku",
                        "lot",
                        "uom",
                        "quantity",
                        "response_quantity",
                    ],
                ),
                (
                    "ownership_paths",
                    "Ownership paths",
                    [
                        "path_id",
                        "subject_company",
                        "person_id",
                        "kind",
                        "economic_fraction",
                        "reach_fraction",
                        "gap_reason",
                        "gap_at_company",
                        "observed_outgoing_fraction",
                        "edge_count",
                        "edge_source_ids",
                        "company_owner_chain",
                        "source_graph_complete",
                    ],
                ),
                (
                    "ownership_path_steps",
                    "Ownership path steps",
                    [
                        "path_id",
                        "subject_company",
                        "step",
                        "edge_source_id",
                        "company",
                        "owner",
                        "owner_kind",
                        "fraction",
                        "effective_from",
                        "effective_to",
                        "evidence_reference",
                    ],
                ),
                (
                    "cutoff_document_controls",
                    "Cutoff document controls",
                    [
                        "document_source_id",
                        "counterparty",
                        "reference",
                        "currency",
                        "document_kind",
                        "document_amount",
                        "allocated_amount",
                        "difference",
                        "allocation_rows",
                        "state",
                    ],
                ),
                (
                    "cutoff_posting_links",
                    "Cutoff posting links",
                    [
                        "posting_source_id",
                        "allocation_id",
                        "document_id",
                        "counterparty",
                        "currency",
                        "account",
                        "amount",
                        "posting_date",
                        "journal_reference",
                        "evidence_reference",
                    ],
                ),
            ):
                if body.get(key) is not None:
                    sheet = workbook.create_sheet(title)
                    append(sheet, fields)
                    for record in body[key]:
                        append(sheet, [record.get(field) for field in fields])
            if body.get("confirmation_coverage_links") is not None:
                coverage = workbook.create_sheet("Confirmation coverage links")
                fields = [
                    "source_row_id",
                    "confirmation_source_row_id",
                    "item_id",
                    "as_of",
                    "account",
                    "counterparty",
                    "side",
                    "currency",
                    "book_amount",
                    "evidence_reference",
                ]
                append(coverage, fields)
                for link in body["confirmation_coverage_links"]:
                    append(coverage, [link.get(field) for field in fields])
            if body.get("workflow") == "sanctions-name-review":
                candidates = workbook.create_sheet("Screening candidates")
                candidate_fields = [
                    "source_id",
                    "list_subject_id",
                    "list_name",
                    "list_version",
                    "published_on",
                    "evidence_reference",
                    "name_similarity",
                ]
                append(
                    candidates,
                    [
                        "subject_id",
                        "subject_name",
                        "subject_evidence_reference",
                        "candidate_retention",
                        *candidate_fields,
                    ],
                )
                for subject in body.get("rows", []):
                    incomplete = subject.get("candidate_preview_truncated") or (
                        subject.get("candidate_count") is not None
                        and subject["candidate_count"] != len(subject.get("candidates", []))
                    )
                    retained = (
                        "LEGACY_TRUNCATED"
                        if incomplete
                        else body.get("candidate_retention_contract", "LEGACY_UNQUALIFIED")
                    )
                    for candidate in subject.get("candidates", []):
                        append(
                            candidates,
                            [
                                subject.get("id"),
                                subject.get("subject_name"),
                                subject.get("subject_evidence_reference"),
                                retained,
                                *[candidate.get(key) for key in candidate_fields],
                            ],
                        )
            if body.get("transaction_dates") is not None:
                dates = workbook.create_sheet("Bank transaction dates")
                date_fields = [
                    "source_role",
                    "source_row_id",
                    "transaction_date",
                    "bank_booking_date",
                    "bank_value_date",
                    "bank_date_basis",
                    "booking_date_review_required",
                ]
                append(dates, date_fields)
                for record in body["transaction_dates"]:
                    append(dates, [record.get(key) for key in date_fields])
            if body.get("payroll_posting_links") is not None:
                posting_links = workbook.create_sheet("Payroll posting links")
                posting_fields = list(
                    dict.fromkeys(key for row in body["payroll_posting_links"] for key in row)
                )
                append(posting_links, posting_fields)
                for link in body["payroll_posting_links"]:
                    append(posting_links, [link.get(key) for key in posting_fields])
            for field, title in (
                ("fifo_allocations", "FIFO allocations"),
                ("fifo_event_costs", "FIFO events"),
                ("fifo_closing_layers", "FIFO closing layers"),
                ("inventory_fifo_controls", "FIFO source controls"),
            ):
                if body.get(field) is not None:
                    sheet = workbook.create_sheet(title)
                    keys = list(dict.fromkeys(key for item in body[field] for key in item))
                    append(sheet, keys)
                    for item in body[field]:
                        append(sheet, [item.get(key) for key in keys])
            if body.get("inventory_quantity_controls") is not None:
                controls = workbook.create_sheet("Inventory quantity controls")
                fields = [
                    "site",
                    "sku",
                    "lot",
                    "uom",
                    "currency",
                    "expected_rows",
                    "positive_quantity",
                    "negative_quantity",
                ]
                append(controls, fields)
                for control in body["inventory_quantity_controls"]:
                    append(controls, [control.get(key) for key in fields])
            if body.get("purchase_quantity_controls") is not None:
                controls = workbook.create_sheet("Purchase quantity controls")
                fields = [
                    "order_id",
                    "line_reference",
                    "vendor",
                    "currency",
                    "uom",
                    "document_type",
                    "expected_rows",
                    "quantity",
                ]
                append(controls, fields)
                for control in body["purchase_quantity_controls"]:
                    append(controls, [control.get(key) for key in fields])
            if body.get("count_bridges") is not None:
                bridges = workbook.create_sheet("Count date bridges")
                links = workbook.create_sheet("Inventory movement links")
                bridge_keys = list(
                    dict.fromkeys(
                        key
                        for bridge in body["count_bridges"]
                        for key in bridge
                        if key != "movement_source_ids"
                    )
                )
                append(bridges, bridge_keys)
                append(
                    links,
                    [
                        "count_source_id",
                        "movement_source_id",
                        "site",
                        "sku",
                        "lot",
                        "uom",
                        "currency",
                    ],
                )
                for bridge in body["count_bridges"]:
                    append(bridges, [bridge.get(key) for key in bridge_keys])
                    for movement_id in bridge.get("movement_source_ids", []):
                        append(
                            links,
                            [
                                bridge["count_source_id"],
                                movement_id,
                                *[
                                    bridge.get(key)
                                    for key in ["site", "sku", "lot", "uom", "currency"]
                                ],
                            ],
                        )
            if body.get("excluded_box_rows"):
                excluded = workbook.create_sheet("Excluded VAT boxes")
                excluded_keys = list(
                    dict.fromkeys(key for row in body["excluded_box_rows"] for key in row)
                )
                append(excluded, excluded_keys)
                for row in body["excluded_box_rows"]:
                    append(excluded, [row.get(key) for key in excluded_keys])
            if message_controls:
                messages = workbook.create_sheet("MT940 message controls")
                message_keys = list(
                    dict.fromkeys(key for message in message_controls for key in message)
                )
                append(messages, message_keys)
                for message in message_controls:
                    append(messages, [message.get(key) for key in message_keys])
            if page_controls:
                pages = workbook.create_sheet("CAMT page controls")
                page_keys = list(dict.fromkeys(key for page in page_controls for key in page))
                append(pages, page_keys)
                for page in page_controls:
                    append(pages, [page.get(key) for key in page_keys])
            cash_manifest = body.get("bank_consumption_manifest")
            if isinstance(cash_manifest, dict) and isinstance(cash_manifest.get("rows"), list):
                cash = workbook.create_sheet("Bank cash claims")
                cash_keys = list(
                    dict.fromkeys(
                        [
                            "bank_source_row_id",
                            "account",
                            "account_basis",
                            "currency",
                            "capacity",
                            "allocated",
                            "source_sha256",
                            "cash_key",
                            "economic_key",
                            "origin_key",
                            *[key for claim in cash_manifest["rows"] for key in claim],
                        ]
                    )
                )
                from openpyxl.utils import get_column_letter

                cash.freeze_panes = "A2"
                for index, key in enumerate(cash_keys, 1):
                    cash.column_dimensions[get_column_letter(index)].width = (
                        68
                        if key.endswith("key") or key.endswith("sha256")
                        else 42
                        if key in {"capacity", "allocated"}
                        else 36
                        if key in {"account", "account_basis"}
                        else max(24, len(key) + 2)
                    )
                append(cash, cash_keys)
                for claim in cash_manifest["rows"]:
                    append(cash, [claim.get(key) for key in cash_keys])
            receipts = workbook.create_sheet("Dependency receipts")
            for dep in event["dependencies"]:
                append(receipts, [dep["type"], dep])
            if writer.long_values:
                append(
                    cover,
                    [
                        "Oversized cells",
                        f"{len(writer.long_values)} exact values retained in ordered parts on Long values; source sheet/cell and UTF-8 SHA-256 identify each original value",
                    ],
                )
            writer.finish()
            workbook.save(buffer)
            return (
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        except XlsxValueLimitError:
            # Public save closes streaming XML writers and removes their temp
            # files. This partial cleanup archive is discarded, never returned.
            from contextlib import suppress

            with suppress(Exception):
                workbook.save(io.BytesIO())
            raise
    if format == "docx":
        from docx import Document

        document = Document()
        document.add_heading("Mr Auditor — reviewable audit schedule", 0)
        for key, value in metadata.items():
            document.add_paragraph(f"{key}: {text(value)}")
        document.add_heading("System-generated calculations", 1)
        for key, value in summary.items():
            document.add_paragraph(f"{key}: {text(value)}")
        for index, row in enumerate(rows[:500], 1):
            document.add_paragraph(f"{index}. {text(row)}")
        if len(rows) > 500:
            document.add_paragraph(
                f"Bounded document preview: first 500 of {len(rows)} rows. The XLSX export includes the full result population."
            )
        document.add_heading("Source receipts", 1)
        for dep in event["dependencies"]:
            document.add_paragraph(text(dep))
        document.save(buffer)
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    from app.core.pdf_text import render_text_schedule

    sections = [("Mr Auditor — reviewable audit schedule", True)]
    for key, value in {**metadata, **summary}.items():
        sections.append((f"{key}: {text(value)}", False))
    sections.append(("System-generated source rows", True))
    for index, row in enumerate(rows[:500], 1):
        sections.append((f"{index}. {text(row)}", False))
    if len(rows) > 500:
        sections.append(
            (
                f"First 500 of {len(rows)} rows. Use the full XLSX export for the complete result.",
                False,
            )
        )
    sections.append(("Dependency receipts", True))
    for dep in event["dependencies"]:
        sections.append((text(dep), False))
    return render_text_schedule(sections), "application/pdf"
