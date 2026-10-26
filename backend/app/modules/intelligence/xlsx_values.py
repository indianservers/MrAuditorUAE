"""Preserve full text when an XLSX cell cannot hold the retained value.

Existing cells remain exact text. Oversized cells carry a pointer; ordered parts
and the original UTF-8 hash are stored in a separate, bounded continuation sheet.
"""

import hashlib
from collections import defaultdict

from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter

CELL_CHARACTERS = 32767
CELL_LINE_FEEDS = 253
PART_CHARACTERS = 16000  # Also fits 32,000 UTF-16 units for non-BMP text.
PART_LINE_FEEDS = 250
MAX_LONG_BYTES = 64 * 1024 * 1024
MAX_ROWS = 1048576
MAX_COLUMNS = 16384


class XlsxValueLimitError(ValueError):
    """A bounded complete export cannot be represented; never truncate it."""


def parts(value):
    start, feeds = 0, 0
    for index, character in enumerate(value):
        if index - start >= PART_CHARACTERS or character == "\n" and feeds >= PART_LINE_FEEDS:
            yield value[start:index]
            start, feeds = index, 0
        feeds += character == "\n"
    if start < len(value):
        yield value[start:]


class ExactTextWriter:
    def __init__(self, workbook, text):
        self.workbook, self.text = workbook, text
        self.rows = defaultdict(int)
        self.long_values = []
        self.long_bytes = 0

    def append(self, sheet, values, *, integer_columns=frozenset()):
        row_number = self.rows[sheet.title] + 1
        if row_number > MAX_ROWS:
            raise XlsxValueLimitError("Excel worksheet row limit exceeded")
        cells = []
        for column, value in enumerate(values, 1):
            if column > MAX_COLUMNS:
                raise XlsxValueLimitError("Excel worksheet column limit exceeded")
            if column in integer_columns:
                if type(value) is not int or not 0 <= value <= MAX_ROWS * PART_CHARACTERS:
                    raise XlsxValueLimitError("Invalid generated continuation count")
                cell = WriteOnlyCell(sheet, value=value)
                cell.number_format = "0"
                cells.append(cell)
                continue
            original = self.text(value)
            if (
                len(original) > CELL_CHARACTERS
                or len(original.encode("utf-16-le")) // 2 > CELL_CHARACTERS
                or original.count("\n") > CELL_LINE_FEEDS
            ):
                encoded = original.encode("utf-8")
                if self.long_bytes + len(encoded) > MAX_LONG_BYTES:
                    raise XlsxValueLimitError(
                        "Full oversized export values exceed the 64 MiB continuation limit"
                    )
                self.long_bytes += len(encoded)
                identifier = f"LV{len(self.long_values) + 1:06d}"
                digest = hashlib.sha256(encoded).hexdigest()
                count = sum(1 for _ in parts(original))
                coordinate = f"{get_column_letter(column)}{row_number}"
                self.long_values.append(
                    (identifier, sheet.title, coordinate, count, digest, original)
                )
                original = f"Full value in Long values: {identifier}; {count} parts; UTF-8 SHA-256 {digest}"
            cell = WriteOnlyCell(sheet, value=original)
            cell.data_type = "s"  # Including parts beginning with '=', '+', '-' or '@'.
            cells.append(cell)
        sheet.append(cells)
        self.rows[sheet.title] = row_number

    def finish(self):
        if not self.long_values:
            return
        sheet = self.workbook.create_sheet("Long values")
        sheet.freeze_panes = "A2"
        for column, width in {
            "A": 16,
            "B": 28,
            "C": 14,
            "D": 14,
            "E": 14,
            "F": 68,
            "G": 20,
            "H": 80,
        }.items():
            sheet.column_dimensions[column].width = width
        self.append(
            sheet,
            [
                "value_id",
                "source_sheet",
                "source_cell",
                "part_number",
                "part_count",
                "value_utf8_sha256",
                "value_characters",
                "part_text",
            ],
        )
        for identifier, source_sheet, coordinate, count, digest, original in self.long_values:
            for index, part in enumerate(parts(original), 1):
                self.append(
                    sheet,
                    [
                        identifier,
                        source_sheet,
                        coordinate,
                        index,
                        count,
                        digest,
                        len(original),
                        part,
                    ],
                    integer_columns=frozenset({4, 5, 7}),
                )
