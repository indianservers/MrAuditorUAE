"""Draft report composition with exact source figures and explicit bounded paper views."""

import io
import json
from xml.sax.saxutils import escape


def value(item):
    return (
        json.dumps(item, ensure_ascii=False, sort_keys=True)
        if isinstance(item, (dict, list))
        else ""
        if item is None
        else str(item)
    )


def tables(data):
    if not isinstance(data, dict):
        return []
    result = []
    for key in ["rows", "lines", "items"]:
        if (
            isinstance(data.get(key), list)
            and data[key]
            and all(isinstance(r, dict) for r in data[key])
        ):
            result.append((key.replace("_", " "), data[key]))
    for table in data.get("financial_tables", []):
        if not isinstance(table, dict):
            continue
        if isinstance(table.get("rows"), list):
            normalized = [
                row
                if isinstance(row, dict)
                else {f"column_{i + 1}": v for i, v in enumerate(row)}
                if isinstance(row, list)
                else {"value": row}
                for row in table["rows"]
            ]
            result.append((str(table.get("title", "Financial table")), normalized))
        elif isinstance(table.get("values"), dict):
            result.append(
                (
                    str(table.get("title", "Financial values")),
                    [{"label": k, "value": v} for k, v in table["values"].items()],
                )
            )
    for key in ["totals", "fields"]:
        if isinstance(data.get(key), dict):
            result.append((key, [{"label": k, "value": v} for k, v in data[key].items()]))
    return result


def render(event, format):
    body = event["body"]
    metadata = {
        "Firm": event.get("firm_name", ""),
        "Report": body["title"],
        "Engagement": event["engagement_id"],
        "Reference": event["id"],
        "Version": event["version"],
        "Country": body["country"],
        "Framework": body["reporting_framework"],
        "SHA-256": event["sha256"],
        "Template version": body["template_version"],
        "Review status": body["review_status"],
        "Status": "DRAFT — NOT SIGNED OR FILED",
        "Created": event["created_at"],
        "Export timestamp": event.get("exported_at", "NOT_RECORDED"),
        "Source freshness at export": event.get("freshness", "NOT_EVALUATED"),
        "Export status": "STALE_HISTORICAL_RESULT_REQUIRES_RECALCULATION"
        if event.get("freshness") == "STALE"
        else "CURRENT_RETAINED_RESULT"
        if event.get("freshness") == "CURRENT"
        else "FRESHNESS_NOT_EVALUATED",
    }
    buffer = io.BytesIO()
    if format == "xlsx":
        from openpyxl import Workbook

        from app.modules.intelligence.xlsx_values import ExactTextWriter, XlsxValueLimitError

        workbook = Workbook(write_only=True)
        writer = ExactTextWriter(workbook, value)
        append = writer.append
        metadata["XLSX text contract"] = "exact-text-with-continuations/1"

        try:
            cover = workbook.create_sheet("Report and review")
            cover.column_dimensions["A"].width = 34
            cover.column_dimensions["B"].width = 120
            for key, item in metadata.items():
                append(cover, [key, item])
            blocks = workbook.create_sheet("Ordered blocks")
            blocks.freeze_panes = "A2"
            for column, width in {"A": 10, "B": 36, "C": 24, "D": 80, "E": 80}.items():
                blocks.column_dimensions[column].width = width
            append(blocks, ["Number", "Title", "Kind", "Narrative", "Source receipt"])
            source = workbook.create_sheet("Complete source chunks")
            source.freeze_panes = "A2"
            for column, width in {"A": 12, "B": 12, "C": 120}.items():
                source.column_dimensions[column].width = width
            append(
                source,
                [
                    "Block",
                    "Chunk",
                    "JSON source (resolve Long values, then concatenate chunks per block)",
                ],
            )
            count = 0
            for block in body["blocks"]:
                append(
                    blocks,
                    [
                        block["number"],
                        block["title"],
                        block["kind"],
                        block["text"],
                        block["source_receipt"],
                    ],
                )
                raw = value(block.get("data"))
                for index in range(0, len(raw), 30000):
                    append(
                        source,
                        [block["number"], index // 30000 + 1, raw[index : index + 30000]],
                        integer_columns=frozenset({2}),
                    )
                for title, rows in tables(block.get("data")):
                    count += 1
                    sheet = workbook.create_sheet(f"Table {count}")
                    keys = list(dict.fromkeys(k for row in rows for k in row))
                    append(sheet, [title])
                    append(sheet, keys)
                    for row in rows:
                        append(sheet, [row.get(k) for k in keys])
                    append(
                        sheet,
                        [
                            "Complete original JSON is in Complete source chunks; oversized cells are retained in Long values"
                        ],
                    )
            if writer.long_values:
                append(
                    cover,
                    [
                        "Oversized cells",
                        f"Full values retained: {len(writer.long_values)}. Long values contains original coordinates, ordered parts and UTF-8 checksums",
                    ],
                )
            writer.finish()
            workbook.save(buffer)
            return (
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        except XlsxValueLimitError:
            from contextlib import suppress

            with suppress(Exception):
                workbook.save(io.BytesIO())
            raise
    if format == "docx":
        from docx import Document
        from docx.shared import Inches, Pt

        document = Document()
        document.styles["Normal"].font.size = Pt(9)
        document.sections[0].header.paragraphs[0].text = "MR AUDITOR • DRAFT — NOT SIGNED OR FILED"
        document.sections[0].left_margin = document.sections[0].right_margin = Inches(0.65)
        document.add_heading(body["title"], 0)
        for key, item in metadata.items():
            document.add_paragraph(f"{key}: {value(item)}")
        for block in body["blocks"]:
            document.add_heading(f"{block['number']}. {block['title']}", 1)
            document.add_paragraph(block["text"])
            document.add_paragraph("Source: " + value(block["source_receipt"]))
            for title, rows in tables(block.get("data")):
                document.add_heading(title, 2)
                keys = list(dict.fromkeys(k for row in rows for k in row))
                for offset in range(0, len(keys), 6):
                    columns = keys[offset : offset + 6]
                    table = document.add_table(rows=1, cols=len(columns) + 1)
                    table.style = "Table Grid"
                    for cell, label in zip(table.rows[0].cells, ["Row", *columns]):
                        cell.text = label
                    for number, row in enumerate(rows[:500], 1):
                        for cell, item in zip(
                            table.add_row().cells, [number, *[row.get(k) for k in columns]]
                        ):
                            cell.text = value(item)[:2000]
                document.add_paragraph(
                    f"Paper view: first {min(500, len(rows))} of {len(rows)} rows; long text fields abbreviated at 2,000 characters. Full source values are retained in XLSX."
                )
        document.save(buffer)
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    styles = getSampleStyleSheet()
    styles["BodyText"].fontSize = 8
    styles["BodyText"].leading = 10

    def paragraph(text):
        return Paragraph(escape(value(text)), styles["BodyText"])

    story = [Paragraph(escape(body["title"]), styles["Title"])]
    for key, item in metadata.items():
        story.extend([paragraph(f"{key}: {value(item)}"), Spacer(1, 5)])
    for block in body["blocks"]:
        story.extend(
            [
                Paragraph(escape(f"{block['number']}. {block['title']}"), styles["Heading1"]),
                paragraph(block["text"]),
                paragraph("Source: " + value(block["source_receipt"])),
            ]
        )
        for title, rows in tables(block.get("data")):
            story.append(Paragraph(escape(title), styles["Heading2"]))
            keys = list(dict.fromkeys(k for row in rows for k in row))
            for offset in range(0, len(keys), 6):
                columns = keys[offset : offset + 6]
                matrix = [[paragraph("Row"), *[paragraph(k) for k in columns]]]
                matrix.extend(
                    [
                        [paragraph(i), *[paragraph(value(row.get(k))[:600]) for k in columns]]
                        for i, row in enumerate(rows[:500], 1)
                    ]
                )
                table = LongTable(
                    matrix,
                    repeatRows=1,
                    colWidths=[28, *[700 / max(1, len(columns)) for _ in columns]],
                )
                table.setStyle(
                    TableStyle(
                        [
                            ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ]
                    )
                )
                story.append(table)
            story.append(
                paragraph(
                    f"Paper view: first {min(500, len(rows))} of {len(rows)} rows; long text abbreviated at 600 characters. XLSX retains complete source JSON and exact monetary values."
                )
            )

    def watermark(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#eeeeee"))
        canvas.setFont("Helvetica-Bold", 60)
        canvas.translate(420, 290)
        canvas.rotate(30)
        canvas.drawCentredString(0, 0, "DRAFT")
        canvas.restoreState()
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(790, 20, f"Page {doc.page} • Not signed or filed")

    SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=40, rightMargin=40).build(
        story, onFirstPage=watermark, onLaterPages=watermark
    )
    return buffer.getvalue(), "application/pdf"
