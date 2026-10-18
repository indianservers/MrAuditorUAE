"""Bounded extraction from the existing private evidence repository."""

import csv
import io
import json
import re
import uuid
from datetime import datetime
from decimal import Decimal, InvalidOperation
from types import SimpleNamespace

from fastapi import HTTPException
from sqlalchemy import select
from starlette.concurrency import run_in_threadpool

from app.db.models.platform import LegalEntity
from app.modules.audit_execution.models import EvidenceVersion
from app.modules.enterprise_ai.models import Chunk, Extraction
from app.modules.enterprise_ai.security import digest, scope, source
from app.modules.lineage.storage import ObjectStore, hash_file

FIELD_PATTERNS = {
    "reference": r"(?:Invoice|Reference|Agreement)\s*(?:number|no\.?|#)?\s*[:=]\s*([^\n]{1,80})",
    "amount": r"(?:Total|Principal|Amount|Closing balance)\s*[:=]\s*(?:AED|USD|EUR|GBP)?\s*([0-9,]+(?:\.[0-9]{1,8})?)",
    "currency": r"\b(AED|USD|EUR|GBP|SAR)\b",
    "date": r"(?:Date|Start date)\s*[:=]\s*(\d{4}-\d{2}-\d{2})",
    "counterparty": r"(?:Customer|Supplier|Borrower|Counterparty)\s*[:=]\s*([^\n]{1,100})",
    "tax_amount": r"(?:VAT|Tax)\s*[:=]\s*(?:AED|USD|EUR|GBP)?\s*([0-9,]+(?:\.[0-9]{1,8})?)",
}


def fields(pages):
    out = {}
    for number, text in pages:
        for name, pattern in FIELD_PATTERNS.items():
            match = re.search(pattern, text, flags=re.I)
            if match and name not in out:
                value = match.group(1).strip()
                try:
                    if name in {"amount", "tax_amount"}:
                        value = str(Decimal(value.replace(",", "")))
                    if name == "date":
                        datetime.strptime(value, "%Y-%m-%d")
                except (ValueError, InvalidOperation):
                    continue
                out[name] = {
                    "value": value,
                    "page": number,
                    "method": "explicit_label",
                    "quality": "unverified",
                }
    return out


def extract_bytes(data: bytes, media: str):
    if len(data) > 16 * 1024 * 1024:
        raise ValueError("Document extraction limited to 16 MiB")
    pages, tables, warnings = [], [], []
    if media == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted or len(reader.pages) > 100:
            raise ValueError("Unencrypted PDF with at most 100 pages required")
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                warnings.append(f"Page {i}: OCR_REQUIRED; no configured OCR engine")
            pages.append((i, text[:20000]))
        warnings.append("PDF table structure is not inferred from plain text")
    elif media == "text/csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        if len(rows) > 10000 or any(len(r) > 100 for r in rows):
            raise ValueError("CSV exceeds extraction bounds")
        tables = [{"page": 1, "section": "CSV", "rows": rows[:100]}]
        pages = [(1, "\n".join(" | ".join(r) for r in rows)[:100000])]
    elif "spreadsheetml" in media:
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if len(book.worksheets) > 20:
                raise ValueError("At most 20 worksheets supported")
            for i, sheet in enumerate(book.worksheets, 1):
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    if len(rows) >= 1000 or len(row) > 100:
                        raise ValueError("Worksheet extraction exceeds bounds")
                    rows.append([str(v) if v is not None else "" for v in row])
                tables.append({"page": i, "section": sheet.title, "rows": rows[:100]})
                pages.append((i, "\n".join(" | ".join(r) for r in rows)[:100000]))
        finally:
            book.close()
        warnings.append(
            "Sheet ordinal is a section reference, not a printed page; formula cached values may be absent"
        )
    elif "wordprocessingml" in media:
        from docx import Document

        doc = Document(io.BytesIO(data))
        pages = [(1, "\n".join(p.text for p in doc.paragraphs)[:100000])]
        tables = [
            {
                "page": 1,
                "section": f"Table {i}",
                "rows": [[c.text[:1000] for c in row.cells[:50]] for row in table.rows[:100]],
            }
            for i, table in enumerate(doc.tables[:20], 1)
        ]
        warnings.append("DOCX section reference only; printed pagination unavailable")
    else:
        raise ValueError("OCR_REQUIRED: image processing is not configured")
    if sum(len(text) for _, text in pages) > 250000:
        raise ValueError("Extracted text exceeds 250000 characters")
    return {
        "pages": [{"page": n, "text": t} for n, t in pages],
        "fields": fields(pages),
        "tables": tables,
        "missing_fields": sorted(set(FIELD_PATTERNS) - set(fields(pages))),
        "warnings": warnings,
        "extractor": "bounded-label-extractor/1",
        "confirmed": False,
    }


async def latest_evidence(db, ctx, eid, rid):
    obj = await source(db, ctx, eid, rid)
    if obj.kind != "evidence":
        if obj.kind not in {
            "working-papers",
            "ifrs-note",
            "grc-policy",
            "tax-rule",
            "vat-rule",
            "tp-rule",
        }:
            raise HTTPException(
                422, "Approved evidence, workpaper, policy or adopted technical rule required"
            )
        return obj, SimpleNamespace(
            id=uuid.uuid5(obj.id, str(obj.version)), number=obj.version, sha256=digest(obj.body)
        )
    version = await db.scalar(
        select(EvidenceVersion)
        .where(EvidenceVersion.tenant_id == ctx.tenant_id, EvidenceVersion.record_id == rid)
        .order_by(EvidenceVersion.number.desc())
        .limit(1)
    )
    if not version or version.scan_status != "CLEAN":
        raise HTTPException(409, "Current clean scanned evidence required")
    return obj, version


async def save(db, ctx, eid, obj, version, body, classification, confidentiality, parent=None):
    extraction = Extraction(
        tenant_id=ctx.tenant_id,
        created_by=ctx.user_id,
        engagement_id=eid,
        record_id=obj.id,
        evidence_version_id=version.id,
        record_version=obj.version,
        parent_id=parent,
        source_hash=version.sha256,
        classification=classification,
        confidentiality=confidentiality,
        body=body,
    )
    db.add(extraction)
    await db.flush()
    eng = await scope(db, ctx, eid)
    entity = await db.get(LegalEntity, eng.entity_id)
    for page in body["pages"]:
        for start in range(0, len(page["text"]), 1600):
            db.add(
                Chunk(
                    tenant_id=ctx.tenant_id,
                    created_by=ctx.user_id,
                    engagement_id=eid,
                    extraction_id=extraction.id,
                    page=page["page"],
                    section=f"{page['page']}:{start}",
                    text=page["text"][start : start + 1600],
                    metadata_details={
                        "tenant": str(ctx.tenant_id),
                        "client": str(entity.client_id),
                        "entity": str(eng.entity_id),
                        "period": str(eng.period_id),
                        "document": str(obj.id),
                        "document_version": version.number,
                        "record_version": obj.version,
                        "extraction_version": str(extraction.id),
                        "effective_date": None,
                        "classification": confidentiality,
                        "source_reference": obj.reference,
                        "source_hash": version.sha256,
                    },
                )
            )
    await db.flush()
    return extraction


async def extract(db, ctx, eid, data):
    await scope(db, ctx, eid, "ai.documents.index")
    obj, version = await latest_evidence(db, ctx, eid, data.record_id)
    if obj.kind != "evidence":
        text = json.dumps(obj.body, ensure_ascii=False, indent=2)
        if len(text) > 100000:
            raise HTTPException(413, "Approved source text exceeds indexing limit")
        body = {
            "pages": [{"page": 1, "text": text}],
            "fields": {},
            "tables": [],
            "missing_fields": [],
            "warnings": [
                "Approved workpaper section reference; not printed page",
                "Client/internal source; not independently authoritative legislation",
            ],
            "extractor": "approved-workpaper-section/1",
            "confirmed": False,
            "source_kind": obj.kind,
        }
        return await save(
            db, ctx, eid, obj, version, body, data.classification, data.confidentiality
        )

    def read():
        with ObjectStore().local(version.object_key, version.object_version) as path:
            if hash_file(path) != version.sha256:
                raise ValueError("Evidence hash differs from approved source")
            if path.stat().st_size > 16 * 1024 * 1024:
                raise ValueError("Document extraction limited to 16 MiB")
            return extract_bytes(path.read_bytes(), version.media_type)

    try:
        body = await run_in_threadpool(read)
    except (ValueError, OSError) as exc:
        raise HTTPException(422, str(exc)) from None
    return await save(db, ctx, eid, obj, version, body, data.classification, data.confidentiality)


async def authorized_extraction(db, ctx, eid, identifier):
    await scope(db, ctx, eid)
    obj = await db.scalar(
        select(Extraction).where(
            Extraction.tenant_id == ctx.tenant_id,
            Extraction.engagement_id == eid,
            Extraction.id == identifier,
        )
    )
    if not obj:
        raise HTTPException(404, "Resource unavailable")
    record, version = await latest_evidence(db, ctx, eid, obj.record_id)
    latest = await db.scalar(
        select(Extraction.id)
        .where(
            Extraction.tenant_id == ctx.tenant_id,
            Extraction.engagement_id == eid,
            Extraction.record_id == obj.record_id,
        )
        .order_by(Extraction.created_at.desc(), Extraction.id.desc())
        .limit(1)
    )
    if (
        obj.record_version != record.version
        or obj.evidence_version_id != version.id
        or obj.id != latest
    ):
        raise HTTPException(409, "Extraction superseded; reindex current evidence")
    return obj


async def authorized_chunk(db, ctx, eid, identifier):
    await scope(db, ctx, eid)
    chunk = await db.scalar(
        select(Chunk).where(
            Chunk.tenant_id == ctx.tenant_id, Chunk.engagement_id == eid, Chunk.id == identifier
        )
    )
    if not chunk:
        raise HTTPException(404, "Resource unavailable")
    ext = await authorized_extraction(db, ctx, eid, chunk.extraction_id)
    return {
        "id": str(chunk.id),
        "text": chunk.text,
        "page": chunk.page,
        "section": chunk.section,
        "record_version": ext.record_version,
        "metadata": chunk.metadata_details,
    }


async def search(db, ctx, eid, query, limit=10):
    await scope(db, ctx, eid, "audit.read")
    query = query.strip()
    if not query:
        return []
    # SQL text matching; source authorization occurs before every returned chunk.
    ids = (
        await db.scalars(
            select(Chunk.id)
            .where(
                Chunk.tenant_id == ctx.tenant_id,
                Chunk.engagement_id == eid,
                Chunk.text.icontains(query, autoescape=True),
            )
            .order_by(Chunk.created_at.desc(), Chunk.id)
            .limit(min(limit * 5, 100))
        )
    ).all()
    out = []
    for identifier in ids:
        try:
            out.append(await authorized_chunk(db, ctx, eid, identifier))
        except HTTPException as exc:
            if exc.status_code in {403, 404, 409}:
                continue
            raise
        if len(out) >= limit:
            break
    return out


def validate_corrections(fields):
    if not fields:
        raise HTTPException(422, "At least one corrected field is required")
    for name, value in fields.items():
        try:
            if name in {"amount", "tax_amount"}:
                number = Decimal(value)
                if (
                    not number.is_finite()
                    or abs(number) > Decimal("1e18")
                    or not isinstance(number.as_tuple().exponent, int)
                    or int(number.as_tuple().exponent) < -8
                ):
                    raise ValueError()
            elif name == "date":
                from datetime import date

                date.fromisoformat(value)
            elif name == "currency":
                import pycountry

                if not pycountry.currencies.get(alpha_3=value):
                    raise ValueError()
            if not value.strip():
                raise ValueError()
        except (ValueError, InvalidOperation):
            raise HTTPException(422, "Correction has an invalid typed field") from None


def match_signals(extracted, journal):
    def value(key):
        return extracted.get(key, {}).get("value")

    signals = {}
    amount = value("amount")
    if amount is not None:
        try:
            number = Decimal(amount)
            signals["amount"] = number.is_finite() and number == max(
                Decimal(str(journal["debit"])), Decimal(str(journal["credit"]))
            )
        except InvalidOperation:
            signals["amount"] = False
    for field, key in [
        ("currency", "currency"),
        ("date", "posting_date"),
        ("counterparty", "counterparty"),
    ]:
        if value(field):
            signals[field] = str(value(field)).casefold() == str(journal.get(key, "")).casefold()
    if value("reference"):
        signals["reference"] = (
            str(value("reference")).casefold() in str(journal.get("description", "")).casefold()
        )
    weights = {"amount": 4, "currency": 2, "date": 1, "counterparty": 1, "reference": 2}
    score = sum(weights[k] for k, ok in signals.items() if ok) / sum(weights.values())
    return {
        "signals": signals,
        "score": score,
        "confirmed": False,
        "score_type": "rule agreement; not calibrated probability",
        "mismatches": [k for k, ok in signals.items() if not ok],
    }
