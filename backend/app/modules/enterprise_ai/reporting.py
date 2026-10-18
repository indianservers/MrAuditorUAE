"""Draft exports from reauthorized receipts. No professional signature/publication."""

import io
from xml.sax.saxutils import escape


def report(run, events, kind):
    title = "AI investigation review draft"
    intro = "Prepared for the engagement reviewer. This source collection is a draft professional aid. Missing steps remain unresolved. No audit opinion, tax treatment or financial adjustment is approved by this document."
    sections = [("Investigation", f"Workflow {run.id}\nTask {run.task}\nState {run.state}")]
    ledger = []
    citations = []
    for event in events:
        if event.kind == "step":
            result = event.body["result"]
            sections.append(
                (result["tool"], result["status"] + ". " + " ".join(result["uncertainty"]))
            )
            citations.extend(result["sources"])
            if result["tool"] == "financial.get_general_ledger":
                ledger = result["data"].get("items", [])[:10]
        if event.kind == "narrative":
            output = event.body["output"]
            sections.append(
                (
                    "Unapproved narrative",
                    output["summary"] + "\n" + "\n".join(output["observations"]),
                )
            )
    sections.append(
        (
            "Source references",
            "\n".join(
                f"{c.get('type', 'record')} {c['id']} version {c.get('version', 'immutable snapshot')}"
                for c in citations
            )
            or "No source citations available",
        )
    )
    sections.append(
        (
            "Required professional action",
            "Inspect the original evidence, resolve missing and contradictory information, record an independent conclusion and follow the existing workpaper approval chain.",
        )
    )
    headers = ["Date", "Debit", "Credit", "Currency"]
    rows = [
        [str(x.get(k, "")) for k in ["posting_date", "debit", "credit", "currency"]] for x in ledger
    ]
    out = io.BytesIO()
    if kind == "docx":
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor

        doc = Document()
        doc.sections[0].top_margin = doc.sections[0].bottom_margin = Inches(0.75)
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(10)
        for name in ["Title", "Heading 1"]:
            doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
            # No template border decoration below a professional report title.
            element = doc.styles[name].element
            for border in list(element.xpath(".//w:pBdr")):
                border.getparent().remove(border)
        doc.add_paragraph(title, "Title")
        doc.add_paragraph(intro)
        if rows:
            doc.add_heading("Deterministic ledger excerpt", level=1)
            doc.add_paragraph(
                "Original accounting values from the authorized tool receipt. First ten lines only; no extrapolation."
            )
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, headers):
                cell.text = text
            for row in rows:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = text
        for heading, text in sections:
            doc.add_heading(heading.replace("_", " ").replace(".", " "), level=1)
            doc.add_paragraph(text)
        doc.save(out)
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.textColor = colors.black
        story = [
            Paragraph(title, styles["Title"]),
            Paragraph(escape(intro), styles["BodyText"]),
            Spacer(1, 12),
        ]
        if rows:
            story.append(Paragraph("Deterministic ledger excerpt", styles["Heading2"]))
            table = Table([headers, *rows], repeatRows=1, colWidths=[105, 115, 115, 80])
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                    ]
                )
            )
            story.append(table)
        for heading, text in sections:
            story.extend(
                [
                    Paragraph(
                        escape(heading.replace("_", " ").replace(".", " ")), styles["Heading2"]
                    ),
                    Paragraph(escape(text).replace("\n", "<br/>"), styles["BodyText"]),
                ]
            )
        SimpleDocTemplate(
            out, pagesize=A4, leftMargin=48, rightMargin=48, topMargin=42, bottomMargin=42
        ).build(story)
    return out.getvalue()
