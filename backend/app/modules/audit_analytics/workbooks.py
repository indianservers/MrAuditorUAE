"""Read-only bounded OOXML integrity scan; formulas/macros/links are never executed."""

import hashlib
import io
import re
import zipfile

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import range_boundaries


def inspect(data: bytes):
    if len(data) > 16 * 1024 * 1024:
        raise ValueError("Workbook exceeds 16 MiB")
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ValueError("A valid XLSX archive is required")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        parts = archive.infolist()
        if len(parts) > 2000 or sum(p.file_size for p in parts) > 64 * 1024 * 1024:
            raise ValueError("Workbook archive exceeds bounds")
        if any("vbaproject" in p.filename.lower() for p in parts):
            raise ValueError("Macro-enabled workbooks require a separate qualified scanner")
    book = load_workbook(io.BytesIO(data), data_only=False, read_only=False, keep_links=False)
    try:
        if len(book.worksheets) > 20:
            raise ValueError("At most 20 worksheets supported")
        findings: list[dict[str, str]] = []
        graph: dict[tuple[str, str], set[tuple[str, str]]] = {}
        formulas: dict[tuple[str, str], str] = {}
        count = 0
        unsupported = 0
        edges = 0

        def add(sheet, cell, code):
            if len(findings) < 2000:
                findings.append(dict(sheet=sheet, cell=cell, code=code))

        for sheet in book.worksheets:
            if sheet.max_row * sheet.max_column > 50000:
                raise ValueError("At most 50000 cells per worksheet")
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    count += 1
                    if count > 100000:
                        raise ValueError("At most 100000 populated cells supported")
                    key = (sheet.title, cell.coordinate)
                    if cell.data_type == "e":
                        add(*key, "error_value")
                    if cell.data_type != "f":
                        if isinstance(cell.value, (int, float)):
                            if (
                                sheet.sheet_state != "visible"
                                or sheet.row_dimensions[cell.row].hidden
                                or sheet.column_dimensions[cell.column_letter].hidden
                            ):
                                add(*key, "hidden_numeric_input")
                            left = sheet.cell(cell.row, max(1, cell.column - 1))
                            right = sheet.cell(cell.row, cell.column + 1)
                            if left.data_type == "f" and right.data_type == "f":
                                add(*key, "possible_hardcoded_override")
                        continue
                    formula = str(cell.value)
                    formulas[key] = formula
                    graph[key] = set()
                    if len(formulas) > 5000:
                        raise ValueError("At most 5000 formula cells supported")
                    if "#REF!" in formula:
                        add(*key, "broken_reference")
                    if "[" in formula:
                        add(*key, "external_or_structured_reference_review")
                    try:
                        tokens = Tokenizer(formula).items
                    except Exception:
                        add(*key, "unparsed_formula")
                        unsupported += 1
                        continue
                    for token in tokens:
                        if token.type != "OPERAND" or token.subtype != "RANGE":
                            continue
                        value = token.value
                        target_sheet = sheet.title
                        if "!" in value:
                            target_sheet, value = value.rsplit("!", 1)
                            target_sheet = target_sheet.strip("'").replace("''", "'")
                        if target_sheet not in book.sheetnames:
                            add(*key, "missing_sheet_reference")
                            continue
                        if not re.fullmatch(
                            r"\$?[A-Z]{1,3}\$?\d{1,7}(?::\$?[A-Z]{1,3}\$?\d{1,7})?", value
                        ):
                            unsupported += 1
                            continue
                        minc, minr, maxc, maxr = range_boundaries(value)
                        if (maxc - minc + 1) * (maxr - minr + 1) > 10000:
                            unsupported += 1
                            continue
                        edges += (maxc - minc + 1) * (maxr - minr + 1)
                        if edges > 100000:
                            raise ValueError("Formula dependency population exceeds 100000 edges")
                        target = book[target_sheet]
                        for cells in target.iter_rows(
                            min_row=minr, max_row=maxr, min_col=minc, max_col=maxc
                        ):
                            for dependency in cells:
                                graph[key].add((target_sheet, dependency.coordinate))
                    if cell.row > 1:
                        above = sheet.cell(cell.row - 1, cell.column)
                        if above.data_type == "f":
                            try:
                                if (
                                    Translator(
                                        str(above.value), origin=above.coordinate
                                    ).translate_formula(cell.coordinate)
                                    != formula
                                ):
                                    add(*key, "adjacent_formula_difference_review")
                            except Exception:
                                unsupported += 1
        # Iterative DFS avoids recursive stack exhaustion and detects direct/indirect formula cycles.
        colors: dict[tuple[str, str], int] = {}
        cycles: set[tuple[str, str]] = set()
        for root in graph:
            if colors.get(root):
                continue
            stack = [(root, False)]
            while stack:
                node, done = stack.pop()
                if done:
                    colors[node] = 2
                    continue
                if colors.get(node) == 1:
                    cycles.add(node)
                    continue
                if colors.get(node) == 2:
                    continue
                colors[node] = 1
                stack.append((node, True))
                for dependency in graph[node]:
                    if dependency in graph:
                        if colors.get(dependency) == 1:
                            cycles.add(dependency)
                            cycles.add(node)
                        elif not colors.get(dependency):
                            stack.append((dependency, False))
        for key in sorted(cycles):
            add(*key, "circular_formula_reference")
        return dict(
            engine_version="workbook-integrity/1",
            source_sha256=hashlib.sha256(data).hexdigest(),
            original_preserved=True,
            populated_cells=count,
            formula_cells=len(formulas),
            findings=findings,
            unsupported_reference_tokens=unsupported,
            findings_capped=len(findings) >= 2000,
            limitations=[
                "Static formula inspection; no Excel recalculation or macro execution",
                "Named/dynamic/structured references are reported as untraversed; cycle completeness is not established",
                "Adjacent formula differences and hidden numeric inputs are review candidates, not proven errors",
                "No determination that hardcoded assumptions are economically justified",
            ],
            professional_conclusion=None,
            publication="NOT_MODIFIED_NOT_SUBMITTED",
        )
    finally:
        book.close()
