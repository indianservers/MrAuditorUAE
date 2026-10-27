"""Tax workpapers reuse the verified export renderer; no live recalculation."""

import json
from decimal import ROUND_HALF_UP, Decimal

from app.modules.ifrs.exports import render as render_document


def amount(value):
    return format(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP), ",.2f")


def render(run, format, released=False, reviewers=()):
    data = run.results
    rows = [
        ["AUDITOS UAE", "Corporate Tax computation / specialist workpaper"],
        ["Status", "APPROVED" if released else "DRAFT - PROFESSIONAL REVIEW REQUIRED"],
        ["Entity", str(data.get("entity", "Scoped tax assessment"))],
        [
            "Tax period",
            str(data.get("period_start", "")) + " to " + str(data.get("period_end", "")),
        ],
        ["Run / methodology", str(run.id) + " / " + run.methodology_version],
        ["Prepared by", str(run.created_by)],
        ["Prepared UTC", str(run.created_at)],
        ["Reviewed by", ", ".join(reviewers) if released else "Pending"],
        ["SHA-256", run.sha256],
        ["Currency", "AED"],
        ["COMPUTATION BRIDGE"],
    ]
    for key in [
        "accounting_profit",
        "net_adjustments",
        "before_loss",
        "loss_utilised",
        "taxable_income",
        "gross_tax",
        "foreign_tax_credit",
        "current_tax",
        "current_tax_loss",
    ]:
        if key in data:
            rows.append([key.replace("_", " ").title(), amount(data[key])])
    rows.append(["Adjustment / source", "Amount (AED)", "Rule / source", "Status"])
    for item in data.get("lines", []):
        rows.append(
            [
                str(item.get("label", item.get("category", "Adjustment"))),
                amount(item.get("approved_adjustment", item.get("proposed_adjustment", "0"))),
                str(item.get("rule_id", "")),
                str(item.get("status", "Approved input")),
            ]
        )
    for section in [
        "exceptions",
        "audit_statement_obligation",
        "interest",
        "sbr",
        "qfzp",
        "loss_allocations",
        "members",
        "conditions",
        "segments",
    ]:
        content = data.get(section)
        if not content:
            continue
        rows.append([section.replace("_", " ").upper()])
        for row_key, value in (
            content.items() if isinstance(content, dict) else enumerate(content, 1)
        ):
            rows.append(
                [
                    str(row_key),
                    json.dumps(value, ensure_ascii=False)
                    if isinstance(value, (dict, list))
                    else str(value),
                ]
            )
    rows.append(["APPLICABLE LEGISLATION / VERSIONED SOURCES"])
    for category, snapshot in data.get("rules", {}).items():
        body = snapshot["body"]
        rows.extend(
            [
                [category.title(), body["decision_number"] + " / Article " + body["article"]],
                ["Rule version / record", str(body["rule_version"]) + " / " + snapshot["id"]],
                [
                    "Applicability / legal knowledge",
                    body["effective_from"]
                    + " to "
                    + str(body.get("effective_to") or "open")
                    + " / known "
                    + body["known_from"],
                ],
                ["Source", snapshot["source"]["url"]],
                ["Professional scope", body["conditions"]],
            ]
        )
    if run.kind not in {"corporate-tax", "tax-group"}:
        rows.append(["SPECIALIST RESULTS"])
        for key, value in data.items():
            if key not in {"rules", "period_start", "period_end", "currency"} and not isinstance(
                value, (dict, list)
            ):
                rows.append([key.replace("_", " "), str(value)])
    rows.append(["Scope", data.get("scope", "Preparation workpaper only")])
    rows.append(["Filing", "No direct FTA submission or professional signature is generated."])
    document, mime = render_document(
        run,
        format,
        released,
        reviewers,
        rows_override=rows,
        document_title="Corporate Tax workpaper",
    )
    if format == "xlsx":
        import io

        from openpyxl import load_workbook
        from openpyxl.styles import Alignment

        workbook = load_workbook(io.BytesIO(document))
        workbook.active.title = "Corporate Tax package"
        for name, value in [
            ("Adjustment schedule", data.get("lines", [])),
            ("Interest schedule", data.get("interest")),
            ("Loss utilization", data.get("loss_allocations", [])),
            ("QFZP assessment", data.get("qfzp")),
            ("CT audit obligation", data.get("audit_statement_obligation")),
            ("Group members", data.get("members", [])),
            ("Legal rule snapshots", data.get("rules", {})),
        ]:
            if not value:
                continue
            sheet = workbook.create_sheet(name)
            for index, text in enumerate(
                json.dumps(value, indent=2, ensure_ascii=False).splitlines(), 1
            ):
                cell = sheet.cell(index, 1, text)
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.column_dimensions["A"].width = 120
            sheet.freeze_panes = "A2"
        output = io.BytesIO()
        workbook.save(output)
        document = output.getvalue()
    return document, mime
