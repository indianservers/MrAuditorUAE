"""Explicitly published, whitelisted governance narratives, separate from working papers."""

import io
import uuid
from datetime import date
from html import escape

from fastapi import HTTPException
from sqlalchemy import select

from app.modules.audit_execution import service as audit
from app.modules.audit_execution.access import engagement_permissions
from app.modules.audit_execution.models import AuditRecord
from app.modules.grc import service
from app.modules.grc.models import GRCRun

REPORT_FIELDS = (
    "executive_summary",
    "objective",
    "scope",
    "methodology",
    "overall_assessment",
    "recommendations",
    "management_response",
    "appendices",
)


async def publish(db, ctx, eid, rid):
    await audit.engagement(db, ctx, eid, "governance.report.prepare", True)
    package = await service.record(db, ctx, eid, rid, "board")
    if package.body["valid_until"] < str(date.today()):
        raise HTTPException(409, "Package validity expired")
    deps = [service.dependency(package)]
    reports = []
    synthetic = bool(package.body["synthetic"])
    for identifier in package.body["report_ids"]:
        report = await service.record(db, ctx, eid, uuid.UUID(identifier), "report")
        deps.append(service.dependency(report))
        synthetic = synthetic or bool(report.body["synthetic"])
        findings = []
        for did in report.body["deficiency_ids"]:
            finding = await service.record(db, ctx, eid, uuid.UUID(did), "deficiency")
            deps.append(service.dependency(finding))
            findings.append(
                {
                    key: finding.body[key]
                    for key in (
                        "description",
                        "severity",
                        "root_cause",
                        "consequence",
                        "severity_methodology",
                    )
                }
            )
        reports.append(
            {
                "title": report.title,
                "reference": report.reference,
                **{key: report.body[key] for key in REPORT_FIELDS},
                "findings": findings,
            }
        )
    return await service.run(
        db,
        ctx,
        eid,
        "board-package",
        {"package_id": str(rid)},
        {
            "title": package.title,
            "executive_briefing": package.body["executive_briefing"],
            "authorized_audience": package.body["authorized_audience"],
            "valid_until": package.body["valid_until"],
            "synthetic": synthetic,
            "reports": reports,
        },
        deps,
        rid,
    )


async def published(db, ctx, eid, identifier):
    _, roles = await engagement_permissions(db, ctx, eid, "governance.board.read")
    obj = await db.scalar(
        select(GRCRun).where(
            GRCRun.tenant_id == ctx.tenant_id,
            GRCRun.engagement_id == eid,
            GRCRun.id == identifier,
            GRCRun.kind == "board-package",
        )
    )
    if not obj:
        raise HTTPException(404, "Published package not found")
    if obj.results["valid_until"] < str(date.today()):
        raise HTTPException(410, "Published package expired")
    audience = obj.results["authorized_audience"]
    professional = bool(roles - {"Audit Committee Member", "Board Viewer"})
    if (
        not professional
        and audience != "both"
        and not (
            audience == "board"
            and "Board Viewer" in roles
            or audience == "audit_committee"
            and "Audit Committee Member" in roles
        )
    ):
        raise HTTPException(403, "Package audience does not authorize this role")
    # Reopening or withdrawing an approved source revokes publication until republished.
    for dep in obj.dependencies:
        current = await db.scalar(
            select(AuditRecord).where(
                AuditRecord.tenant_id == ctx.tenant_id,
                AuditRecord.engagement_id == eid,
                AuditRecord.id == uuid.UUID(dep["id"]),
            )
        )
        if (
            not current
            or current.version != dep["version"]
            or current.status not in service.APPROVED
        ):
            raise HTTPException(
                409, "Published source changed; professional republication required"
            )
    return obj


def render(run, format):
    if format not in {"pdf", "docx"}:
        raise HTTPException(422, "Supported exports: pdf, docx")
    rows = [
        ["AUDITOS UAE", "Internal audit governance report"],
        ["Status", "Published from independently approved sources"],
        ["Package", str(run.id)],
        ["SHA-256", run.sha256],
    ]
    if run.results.get("synthetic"):
        rows.append(["SYNTHETIC SOFTWARE TEST DATA", "No real professional assurance conclusion"])
    rows.append(["Executive briefing", run.results["executive_briefing"]])
    for report in run.results["reports"]:
        rows.append([report["reference"], report["title"]])
        rows.extend([[key.replace("_", " ").title(), report[key]] for key in REPORT_FIELDS])
        for finding in report["findings"]:
            rows.append(["Finding / " + finding["severity"], finding["description"]])
            rows.append(["Root cause", finding["root_cause"]])
            rows.append(["Consequence", finding["consequence"]])
    if format == "docx":
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor

        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.8)
        for style in ["Normal", "Title", "Heading 1", "Heading 2"]:
            document.styles[style].font.name = "Calibri"
            document.styles[style].font.color.rgb = RGBColor(0, 0, 0)
        for style in document.styles:
            for border in style.element.xpath(".//w:pBdr"):
                border.getparent().remove(border)
        document.styles["Normal"].font.size = Pt(11)
        document.add_heading("Internal audit governance report", 0)
        for key, value in rows:
            document.add_heading(key, 2)
            document.add_paragraph(str(value))
        output = io.BytesIO()
        document.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    styles["BodyText"].fontSize = 9.5
    styles["BodyText"].leading = 13
    styles["Heading2"].fontSize = 11
    styles["Heading2"].spaceBefore = 8
    styles["Heading2"].spaceAfter = 4
    story = [Paragraph("Internal audit governance report", styles["Title"])]
    for key, value in rows:
        story.extend(
            [
                Paragraph(escape(str(key)), styles["Heading2"]),
                Paragraph(escape(str(value)).replace("\n", "<br/>"), styles["BodyText"]),
                Spacer(1, 2),
            ]
        )

    def footer(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(
            44, 25, f"AuditOS UAE | Published package {str(run.id)[:8]} | Page {document.page}"
        )

    output = io.BytesIO()
    SimpleDocTemplate(
        output, pagesize=A4, leftMargin=44, rightMargin=44, topMargin=40, bottomMargin=44
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue(), "application/pdf"
