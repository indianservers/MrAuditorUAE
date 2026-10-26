"""Bounded views and lossless scalar workbook rows from immutable SoD snapshots."""

import base64
import hashlib
import io
import uuid
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.modules.audit_execution import service as audit
from app.modules.grc import service
from app.modules.grc.models import GRCRun
from app.modules.grc.sod_sources import implementation_hash
from app.modules.intelligence.xlsx_values import ExactTextWriter
from app.modules.lineage.working_papers import require_current

COLLECTIONS = (
    "conflicts",
    "conflict_intervals",
    "assignment_pair_links",
    "assignments",
    "scope_mappings",
    "permission_pairs",
    "retained_source_receipts",
    "dependencies",
)
RESULT_COLLECTIONS = set(COLLECTIONS) - {"assignments", "dependencies"}


def rows(run, name):
    if name == "assignments":
        return [
            {"assignment_index": index, **item}
            for index, item in enumerate(run.inputs.get(name, []), 1)
        ]
    if name == "dependencies":
        return run.dependencies
    if name == "conflicts":
        return [
            {
                **{
                    key: value
                    for key, value in item.items()
                    if key not in {"intervals", "source_pairs"}
                },
                "interval_count": len(item.get("intervals", [])),
                "assignment_pair_count": len(item.get("source_pairs", [])),
            }
            for item in run.results.get(name, [])
        ]
    return run.results.get(name, [])


def compact(run):
    result = {key: value for key, value in run.results.items() if key not in RESULT_COLLECTIONS}
    result["collection_totals"] = {name: len(rows(run, name)) for name in COLLECTIONS}
    result["policy_pairs_retained"] = "permission_pairs" in run.results
    return result


async def read(db, ctx, eid, identifier):
    await audit.engagement(db, ctx, eid, "grc.read")
    if await service.owner_scoped(db, ctx, eid):
        raise HTTPException(403, "Professional read model permission required")
    run = await db.scalar(
        select(GRCRun).where(
            GRCRun.tenant_id == ctx.tenant_id,
            GRCRun.engagement_id == eid,
            GRCRun.id == identifier,
            GRCRun.kind == "sod-access",
        )
    )
    if run is None:
        raise HTTPException(404, "SoD execution snapshot not found")
    if run.results.get("sod_contract") != "permission-overlap-intervals/1":
        raise HTTPException(
            409, "Legacy SoD snapshot: freeze a new reviewed calculation for this view"
        )
    current = True
    for dependency in run.dependencies:
        try:
            source = await audit.get_record(db, ctx, eid, uuid.UUID(dependency["id"]), "grc.read")
            if source.status not in service.APPROVED or source.version != dependency["version"]:
                current = False
            await require_current(db, ctx, source)
        except HTTPException as error:
            if error.status_code not in {404, 409}:
                raise
            current = False
    state = dict(
        record_versions_status="CURRENT_RECORD_VERSIONS"
        if current
        else "STALE_HISTORICAL_SNAPSHOT",
        implementation_status="CURRENT_IMPLEMENTATION"
        if run.results.get("sod_implementation_sha256") == implementation_hash()
        else "HISTORICAL_IMPLEMENTATION",
        source_bytes_reverified_on_read=False,
    )
    return run, state


def summary(run, state):
    return dict(
        id=str(run.id),
        record_id=str(run.record_id),
        engagement_id=str(run.engagement_id),
        frozen_at=run.created_at.isoformat(),
        snapshot_sha256=run.sha256,
        results=compact(run),
        **state,
    )


def workbook(run, state):
    book = Workbook(write_only=True)
    encoded_values: list[dict[str, str]] = []
    writer = ExactTextWriter(
        book,
        lambda value: (
            "" if value is None else str(value).lower() if isinstance(value, bool) else str(value)
        ),
    )

    def sheet(name, values, headers=None):
        page = book.create_sheet(name)
        page.freeze_panes = "A2"
        fields = headers or list(dict.fromkeys(key for value in values for key in value))
        if not fields:
            fields = ["No retained rows"]
        writer.append(page, fields)
        for row_number, value in enumerate(values, 2):
            cells = []
            for column, field in enumerate(fields, 1):
                item = value.get(field)
                if isinstance(item, str) and ILLEGAL_CHARACTERS_RE.search(item):
                    identifier = f"EV{len(encoded_values) + 1:06d}"
                    raw = item.encode("utf-8")
                    encoded_values.append(
                        dict(
                            value_id=identifier,
                            source_sheet=name,
                            source_cell=f"{get_column_letter(column)}{row_number}",
                            utf8_sha256=hashlib.sha256(raw).hexdigest(),
                            utf8_base64=base64.b64encode(raw).decode(),
                        )
                    )
                    item = f"Exact UTF-8 value in Encoded values: {identifier}"
                if isinstance(item, (dict, list)):
                    raise ValueError("SoD workbook rows must be flattened without cell truncation")
                cells.append(item)
            writer.append(
                page,
                cells,
                integer_columns=frozenset(
                    i for i, item in enumerate(cells, 1) if type(item) is int and item >= 0
                ),
            )

    metadata = summary(run, state)
    scalar: list[dict[str, Any]] = []

    def flatten(section, path, value):
        if isinstance(value, dict) and value:
            for key, item in value.items():
                flatten(section, f"{path}.{key}" if path else key, item)
        elif isinstance(value, list) and value:
            for index, item in enumerate(value, 1):
                flatten(section, f"{path}[{index}]", item)
        else:
            scalar.append(
                dict(
                    section=section,
                    field=path,
                    value="{}" if value == {} else "[]" if value == [] else value,
                )
            )

    for section, values in (
        ("snapshot", {k: v for k, v in metadata.items() if k != "results"}),
        ("input", {k: v for k, v in run.inputs.items() if k != "assignments"}),
        ("result", compact(run)),
    ):
        flatten(section, "", values)
    scalar.append(
        dict(
            section="interpretation",
            field="Source verification",
            value="Source checks are frozen historical receipts; reading/exporting does not reverify current original-file bytes or establish completeness.",
        )
    )
    sheet("Scope", scalar, ["section", "field", "value"])
    for name, title in zip(
        COLLECTIONS,
        (
            "Conflicts",
            "Conflict intervals",
            "Assignment links",
            "Assignments",
            "Scope mappings",
            "Permission pairs",
            "Source receipts",
            "Dependencies",
        ),
        strict=True,
    ):
        sheet(title, rows(run, name))
    if encoded_values:
        sheet("Encoded values", encoded_values)
    writer.finish()
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
