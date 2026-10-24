"""Thin read models for the frontend; calculations remain in Phase 2 services."""

import uuid
from decimal import Decimal
from typing import Annotated, Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.encoders import jsonable_encoder
from sqlalchemy import case, func, select

from app.api.v1.accounting import serialize
from app.core.dependencies import DB, Context, require
from app.db.models.platform import AuditEvent, Client, Engagement, FinancialPeriod, LegalEntity
from app.modules.accounting import models as m
from app.modules.accounting.service import scoped

router = APIRouter(prefix="/workspace", tags=["Frontend read models"])
Read = Annotated[Context, Depends(require("clients.read"))]
Finance = Annotated[Context, Depends(require("accounting.read"))]
Analytics = Annotated[Context, Depends(require("analytics.read"))]


@router.post("/inspect-file")
async def inspect_file(
    ctx: Annotated[Context, Depends(require("accounting.import"))],
    file: Annotated[UploadFile, File()],
    sheet: Annotated[str | None, Form()] = None,
    delimiter: Annotated[str, Form()] = ",",
    encoding: Annotated[str, Form()] = "utf-8-sig",
):
    """Bounded, non-persisting source preview before mapping and upload."""
    import csv
    import io
    import zipfile
    from itertools import islice

    from openpyxl import load_workbook

    if file.size is None or file.size > 64 * 1024 * 1024:
        raise HTTPException(413, "File exceeds 64 MiB")
    if delimiter not in {",", ";", "\t"} or encoding not in {"utf-8-sig", "cp1252", "utf-16"}:
        raise HTTPException(422, "Unsupported source encoding or delimiter")
    try:
        sheets = []
        if (file.filename or "").lower().endswith(".xlsx"):
            with zipfile.ZipFile(file.file) as archive:
                entries = archive.infolist()
                if len(entries) > 5000 or sum(i.file_size for i in entries) > 512 * 1024 * 1024:
                    raise ValueError("Excel archive exceeds decompression limit")
                if any(
                    i.file_size > 1000000 and i.file_size / max(i.compress_size, 1) > 200
                    for i in entries
                ):
                    raise ValueError("Suspicious Excel compression ratio")
                if any(
                    "vbaproject" in i.filename.lower() or "externallinks/" in i.filename.lower()
                    for i in entries
                ):
                    raise ValueError("Macros and external links are prohibited")
            file.file.seek(0)
            book = load_workbook(file.file, read_only=True, data_only=False, keep_links=False)
            try:
                sheets = book.sheetnames
                if sheet and sheet not in sheets:
                    raise ValueError("Unknown worksheet")
                source = book[sheet or sheets[0]]
                if source.max_column is not None and source.max_column > 100:
                    raise ValueError("Too many columns")
                cells = list(islice(source.iter_rows(), 21))
                if any(c.data_type == "f" for row in cells for c in row):
                    raise ValueError("Formula cells prohibited; export values")
                values = [[str(c.value or "") for c in row] for row in cells]
            finally:
                book.close()
        elif (file.filename or "").lower().endswith(".csv"):
            wrapper = io.TextIOWrapper(file.file, encoding=encoding, newline="")
            try:
                values = list(islice(csv.reader(wrapper, delimiter=delimiter), 21))
            finally:
                wrapper.detach()
        else:
            raise ValueError("Preview supports CSV and XLSX. Tally XML/package uses direct import.")
        if not values or any(len(row) > 100 for row in values):
            raise ValueError("Empty source or too many columns")
        if any(len(cell) > 32000 for row in values for cell in row):
            raise ValueError("Cell exceeds preview limit")
        return {
            "sheets": sheets,
            "selected_sheet": sheet or (sheets[0] if sheets else None),
            "headers": values[0],
            "rows": values[1:],
            "preview_only": True,
        }
    except (ValueError, OSError, zipfile.BadZipFile, UnicodeError, csv.Error):
        raise HTTPException(
            422, "Unable to preview source; check file format, values, encoding and worksheet"
        ) from None


@router.get("/access")
async def access(ctx: Annotated[Context, Depends(require(""))]):
    return {"tenant_id": str(ctx.tenant_id), "permissions": sorted(ctx.permissions)}


@router.get("/directory")
async def directory(
    db: DB,
    ctx: Read,
    client_id: uuid.UUID | None = None,
    entity_id: uuid.UUID | None = None,
    country: str | None = Query(None, pattern=r"^[A-Z]{2}$"),
    offset: int = Query(0, ge=0),
    limit: int = Query(100, ge=1, le=100),
):
    if "entities.read" not in ctx.permissions:
        raise HTTPException(403, "Permission denied")
    if country:
        from app.modules.jurisdictions.packages import package

        package(country)
    if client_id:
        await scoped(db, ctx, Client, client_id)
    if entity_id:
        entity = await scoped(db, ctx, LegalEntity, entity_id)
        if country and entity.country != country:
            raise HTTPException(409, "Entity belongs to another working country")
        if client_id and entity.client_id != client_id:
            raise HTTPException(422, "Entity does not belong to selected client")
        model: Any = FinancialPeriod
        conditions = [model.tenant_id == ctx.tenant_id, model.entity_id == entity_id]
    else:
        model = LegalEntity
        conditions = [model.tenant_id == ctx.tenant_id]
        if country:
            conditions.append(LegalEntity.country == country)
        if client_id:
            conditions.append(model.client_id == client_id)
    total = await db.scalar(select(func.count()).select_from(model).where(*conditions))
    items = await db.scalars(
        select(model).where(*conditions).order_by(model.id).offset(offset).limit(limit)
    )
    return {
        "items": [serialize(r) for r in items],
        "total": total,
        "offset": offset,
        "limit": limit,
    }


@router.get("/dashboard")
async def dashboard(
    db: DB,
    ctx: Read,
    client_id: uuid.UUID | None = None,
    period_id: uuid.UUID | None = None,
    status: str | None = None,
):
    if client_id:
        await scoped(db, ctx, Client, client_id)
    if period_id:
        period = await scoped(db, ctx, FinancialPeriod, period_id)
        entity = await scoped(db, ctx, LegalEntity, period.entity_id)
        if client_id and entity.client_id != client_id:
            raise HTTPException(422, "Period does not belong to selected client")
    clients = select(func.count()).select_from(Client).where(Client.tenant_id == ctx.tenant_id)
    if client_id:
        clients = clients.where(Client.id == client_id)
    conditions = [Engagement.tenant_id == ctx.tenant_id]
    from app.modules.audit_execution.access import visible_engagements

    visible = await visible_engagements(db, ctx)
    if visible is not None:
        conditions.append(Engagement.id.in_(visible))
    if client_id:
        conditions.append(
            Engagement.entity_id.in_(
                select(LegalEntity.id).where(LegalEntity.client_id == client_id)
            )
        )
    if period_id:
        conditions.append(Engagement.period_id == period_id)
    if status:
        conditions.append(Engagement.status == status)
    engagements = []
    active = None
    if "engagements.read" in ctx.permissions:
        active = await db.scalar(
            select(func.count())
            .select_from(Engagement)
            .where(*conditions, Engagement.status == "active")
        )
        engagements = [
            serialize(r)
            for r in await db.scalars(
                select(Engagement)
                .where(*conditions)
                .order_by(Engagement.updated_at.desc())
                .limit(8)
            )
        ]
    findings = None
    if "analytics.read" in ctx.permissions:
        filters = [m.Finding.tenant_id == ctx.tenant_id]
        if client_id:
            filters.append(m.Finding.client_id == client_id)
        if period_id:
            filters.append(m.Finding.period_id == period_id)
        findings = await db.scalar(select(func.count()).select_from(m.Finding).where(*filters))
    # Activity is tenant metadata only; omit arbitrary event payloads.
    activity = []
    if "members.manage" in ctx.permissions and not client_id and not period_id:
        activity = [
            {"action": r.action, "created_at": r.created_at}
            for r in await db.scalars(
                select(AuditEvent)
                .where(AuditEvent.tenant_id == ctx.tenant_id)
                .order_by(AuditEvent.created_at.desc())
                .limit(8)
            )
        ]
    return {
        "clients": await db.scalar(clients),
        "active_engagements": active,
        "findings": findings,
        "engagements": engagements,
        "activity": activity,
        "deadlines": None,
        "compliance": "not_configured",
    }


@router.get("/ledger")
async def ledger(
    db: DB,
    ctx: Finance,
    entity_id: uuid.UUID,
    period_id: uuid.UUID,
    account_id: uuid.UUID | None = None,
    snapshot_id: uuid.UUID | None = None,
    search: str = Query("", max_length=200),
    date_from: str | None = None,
    date_to: str | None = None,
    minimum: Decimal | None = None,
    maximum: Decimal | None = None,
    journal_type: str | None = Query(None, max_length=100),
    offset: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=100),
    sort: str = "-posting_date",
):
    from datetime import date

    await scoped(db, ctx, LegalEntity, entity_id)
    period = await scoped(db, ctx, FinancialPeriod, period_id)
    if period.entity_id != entity_id:
        raise HTTPException(422, "Period belongs to another entity")
    conditions = [
        m.JournalLine.tenant_id == ctx.tenant_id,
        m.JournalLine.entity_id == entity_id,
        m.JournalLine.period_id == period_id,
    ]
    if snapshot_id:
        snapshot = await scoped(db, ctx, m.ImportSnapshot, snapshot_id)
        if snapshot.entity_id != entity_id or snapshot.period_id != period_id:
            raise HTTPException(422, "Snapshot context mismatch")
        conditions.append(m.JournalLine.snapshot_id == snapshot_id)
    else:
        conditions.append(m.ImportBatch.status == "active")
    if account_id:
        conditions.append(m.JournalLine.account_id == account_id)
    if search:
        escaped = search.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        conditions.append(m.JournalLine.description.ilike("%" + escaped + "%", escape="\\"))
    for bound, value, op in [("from", date_from, "ge"), ("to", date_to, "le")]:
        if value:
            try:
                parsed = date.fromisoformat(value)
            except ValueError:
                raise HTTPException(422, "Invalid " + bound + " date") from None
            conditions.append(
                m.JournalLine.posting_date >= parsed
                if op == "ge"
                else m.JournalLine.posting_date <= parsed
            )
    amount = func.greatest(m.JournalLine.base_debit, m.JournalLine.base_credit)
    if minimum is not None:
        conditions.append(amount >= minimum)
    if maximum is not None:
        conditions.append(amount <= maximum)
    if journal_type:
        conditions.append(m.JournalLine.attributes["journal_type"].astext == journal_type)
    orders = {"posting_date": m.JournalLine.posting_date, "amount": amount, "code": m.Account.code}
    field = sort.lstrip("-")
    if field not in orders:
        raise HTTPException(422, "Unsupported sort")
    query = (
        select(
            m.JournalLine,
            m.Account.code,
            m.Account.name,
            m.Journal.reference,
            m.JournalLine.attributes["journal_type"].astext,
            m.ImportBatch.source_system,
        )
        .join(m.Account, m.Account.id == m.JournalLine.account_id)
        .join(m.Journal, m.Journal.id == m.JournalLine.journal_id)
        .join(m.ImportBatch, m.ImportBatch.id == m.JournalLine.batch_id)
        .where(*conditions)
    )
    total = await db.scalar(select(func.count()).select_from(query.subquery()))
    column = orders[field]
    rows = (
        await db.execute(
            query.order_by(column.desc() if sort.startswith("-") else column, m.JournalLine.id)
            .offset(offset)
            .limit(limit)
        )
    ).all()
    return {
        "items": [
            {
                **serialize(r[0]),
                "account_code": r[1],
                "account_name": r[2],
                "voucher_number": r[3],
                "journal_type": r[4],
                "source_system": r[5],
            }
            for r in rows
        ],
        "total": total,
        "offset": offset,
        "limit": limit,
    }


@router.get("/analytics/{identifier}/summary")
async def analytics_summary(
    identifier: uuid.UUID,
    db: DB,
    ctx: Analytics,
    rule_id: str | None = None,
    band: str | None = None,
):
    run = await scoped(db, ctx, m.AnalyticsRun, identifier)
    selected_lines = select(m.Finding.line_id).where(
        m.Finding.run_id == run.id, m.Finding.tenant_id == ctx.tenant_id
    )
    filters = [m.Finding.run_id == run.id, m.Finding.tenant_id == ctx.tenant_id]
    if rule_id:
        selected_lines = selected_lines.where(m.Finding.rule_id == rule_id)
        filters.append(m.Finding.rule_id == rule_id)
    if band:
        band_lines = select(m.RiskScore.line_id).where(
            m.RiskScore.run_id == run.id,
            m.RiskScore.tenant_id == ctx.tenant_id,
            m.RiskScore.band == band,
        )
        selected_lines = selected_lines.where(m.Finding.line_id.in_(band_lines))
        filters.append(m.Finding.line_id.in_(band_lines))
    risk = (
        await db.execute(
            select(m.RiskScore.band, func.count())
            .where(
                m.RiskScore.run_id == run.id,
                m.RiskScore.tenant_id == ctx.tenant_id,
                m.RiskScore.line_id.in_(selected_lines),
            )
            .group_by(m.RiskScore.band)
        )
    ).all()
    rules = (
        await db.execute(
            select(m.Finding.rule_id, func.count()).where(*filters).group_by(m.Finding.rule_id)
        )
    ).all()
    groups = (
        await db.execute(
            select(m.Account.classification, m.Finding.rule_id, func.count())
            .select_from(m.Finding)
            .join(m.JournalLine, m.JournalLine.id == m.Finding.line_id)
            .join(m.Account, m.Account.id == m.JournalLine.account_id)
            .where(*filters)
            .group_by(m.Account.classification, m.Finding.rule_id)
        )
    ).all()
    timeline = (
        await db.execute(
            select(m.JournalLine.posting_date, func.count())
            .select_from(m.Finding)
            .join(m.JournalLine, m.JournalLine.id == m.Finding.line_id)
            .where(*filters)
            .group_by(m.JournalLine.posting_date)
            .order_by(m.JournalLine.posting_date)
        )
    ).all()
    amount = func.greatest(m.JournalLine.base_debit, m.JournalLine.base_credit)
    bucket = case(
        (amount < 10000, "0–10K"),
        (amount < 50000, "10–50K"),
        (amount < 100000, "50–100K"),
        (amount < 500000, "100–500K"),
        (amount < 1000000, "500K–1M"),
        else_=">1M",
    )
    distribution = (
        await db.execute(
            select(bucket, func.count())
            .where(m.JournalLine.tenant_id == ctx.tenant_id, m.JournalLine.id.in_(selected_lines))
            .group_by(bucket)
        )
    ).all()
    concentration = (
        await db.execute(
            select(m.Account.classification, func.count())
            .select_from(m.Finding)
            .join(m.JournalLine, m.JournalLine.id == m.Finding.line_id)
            .join(m.Account, m.Account.id == m.JournalLine.account_id)
            .where(*filters)
            .group_by(m.Account.classification)
        )
    ).all()
    return jsonable_encoder(
        {
            "risk": [{"name": r[0], "value": r[1]} for r in risk],
            "rules": [{"name": r[0], "value": r[1]} for r in rules],
            "groups": [{"group": r[0], "rule": r[1], "value": r[2]} for r in groups],
            "timeline": [{"name": str(r[0]), "value": r[1]} for r in timeline],
            "distribution": [{"name": r[0], "value": r[1]} for r in distribution],
            "concentration": [{"name": r[0], "value": r[1]} for r in concentration],
        }
    )


@router.get("/findings")
async def findings(
    db: DB,
    ctx: Analytics,
    run_id: uuid.UUID,
    rule_id: str | None = None,
    band: str | None = None,
    offset: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=100),
):
    await scoped(db, ctx, m.AnalyticsRun, run_id)
    conditions = [m.Finding.tenant_id == ctx.tenant_id, m.Finding.run_id == run_id]
    if rule_id:
        conditions.append(m.Finding.rule_id == rule_id)
    if band:
        conditions.append(m.RiskScore.band == band)
    query = (
        select(
            m.Finding,
            m.RiskScore.score,
            m.RiskScore.band,
            m.RiskScore.factors,
            m.Journal.reference,
            m.JournalLine.posting_date,
            m.Account.name,
            m.Account.code,
            m.JournalLine.base_debit,
            m.JournalLine.base_credit,
            m.JournalLine.description,
            m.JournalLine.counterparty,
        )
        .outerjoin(m.JournalLine, m.JournalLine.id == m.Finding.line_id)
        .outerjoin(m.Journal, m.Journal.id == m.JournalLine.journal_id)
        .outerjoin(m.Account, m.Account.id == m.JournalLine.account_id)
        .outerjoin(
            m.RiskScore, (m.RiskScore.line_id == m.Finding.line_id) & (m.RiskScore.run_id == run_id)
        )
        .where(*conditions)
    )
    total = await db.scalar(select(func.count()).select_from(query.subquery()))
    rows = (
        await db.execute(
            query.order_by(m.RiskScore.score.desc().nullslast(), m.Finding.id)
            .offset(offset)
            .limit(limit)
        )
    ).all()
    keys = [
        "score",
        "band",
        "factors",
        "voucher_number",
        "posting_date",
        "account_name",
        "account_code",
        "base_debit",
        "base_credit",
        "description",
        "counterparty",
    ]
    return jsonable_encoder(
        {
            "items": [{**serialize(r[0]), **dict(zip(keys, r[1:], strict=True))} for r in rows],
            "total": total,
            "offset": offset,
            "limit": limit,
        },
        custom_encoder={Decimal: str},
    )
