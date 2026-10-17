"""Read-model integration tests against the disposable Phase 1/2 database."""

import io
import os
import uuid

import pytest
from openpyxl import Workbook

from tests.integration.test_platform import environment as environment
from tests.phase2.test_pipeline import CSV, committed, scope, upload

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Disposable database required"
)


async def test_context_dashboard_and_ledger_filters(environment):
    client = environment
    headers, options = await scope(client)
    _, snapshot = await committed(client, headers, options)
    access = await client.get("/api/v1/workspace/access", headers=headers)
    assert "accounting.read" in access.json()["permissions"]
    directory = await client.get(
        "/api/v1/workspace/directory",
        headers=headers,
        params={"client_id": options["client_id"], "entity_id": options["entity_id"]},
    )
    assert (
        directory.status_code == 200 and directory.json()["items"][0]["id"] == options["period_id"]
    )
    dashboard = await client.get(
        "/api/v1/workspace/dashboard",
        headers=headers,
        params={"client_id": options["client_id"], "period_id": options["period_id"]},
    )
    assert dashboard.status_code == 200, dashboard.text
    assert dashboard.json()["clients"] == 1 and dashboard.json()["deadlines"] is None
    params = {
        "entity_id": options["entity_id"],
        "period_id": options["period_id"],
        "search": "Adjustment",
        "minimum": "100",
        "limit": 1,
    }
    ledger = await client.get("/api/v1/workspace/ledger", headers=headers, params=params)
    assert ledger.status_code == 200, ledger.text
    assert ledger.json()["total"] == 2 and len(ledger.json()["items"]) == 1
    assert ledger.json()["items"][0]["voucher_number"] == "1"
    for injected in ["%", "_", "' OR 1=1 --"]:
        result = await client.get(
            "/api/v1/workspace/ledger", headers=headers, params={**params, "search": injected}
        )
        assert result.status_code == 200 and result.json()["total"] == 0
    balance = await client.get(
        "/api/v1/accounting/trial-balance", headers=headers, params={"snapshot_id": snapshot}
    )
    assert balance.json()["total_debit"] == "100.20000000"
    assert balance.json()["total_credit"] == "100.20000000"
    assert balance.json()["mapped_accounts"] == 0 and balance.json()["unmapped_accounts"] == 3


async def test_analytics_aggregates_filter_and_scope(environment):
    client = environment
    headers, options = await scope(client)
    _, snapshot = await committed(client, headers, options)
    run = await client.post(
        "/api/v1/analytics/runs",
        headers=headers,
        json={"snapshot_id": snapshot, "materiality": "50"},
    )
    assert run.status_code == 201
    run_id = run.json()["id"]
    summary = await client.get(f"/api/v1/workspace/analytics/{run_id}/summary", headers=headers)
    assert summary.status_code == 200, summary.text
    assert summary.json()["rules"] and summary.json()["distribution"]
    findings = await client.get(
        "/api/v1/workspace/findings",
        headers=headers,
        params={"run_id": run_id, "rule_id": "J09", "limit": 1},
    )
    assert findings.status_code == 200, findings.text
    assert findings.json()["total"] == 2
    assert findings.json()["items"][0]["rule_id"] == "J09"
    filtered = await client.get(
        f"/api/v1/workspace/analytics/{run_id}/summary", headers=headers, params={"rule_id": "J09"}
    )
    assert filtered.json()["rules"] == [{"name": "J09", "value": 2}]
    empty_band = await client.get(
        f"/api/v1/workspace/analytics/{run_id}/summary",
        headers=headers,
        params={"band": "unavailable-band"},
    )
    assert empty_band.status_code == 200 and empty_band.json()["rules"] == []
    assert empty_band.json()["risk"] == []
    foreign = {**headers, "X-Tenant-ID": str(uuid.uuid4())}
    for path in [
        "access",
        "dashboard",
        f"analytics/{run_id}/summary",
        "findings?run_id=" + run_id,
        "ledger?entity_id=" + options["entity_id"] + "&period_id=" + options["period_id"],
        "directory",
    ]:
        response = await client.get("/api/v1/workspace/" + path, headers=foreign)
        assert response.status_code == 403, response.text


async def test_source_preview_sheet_selection_and_format_guards(environment):
    client = environment
    headers, _ = await scope(client)
    preview = await client.post(
        "/api/v1/workspace/inspect-file",
        headers=headers,
        files={"file": ("synthetic.csv", CSV, "text/csv")},
    )
    assert preview.status_code == 200 and len(preview.json()["rows"]) == 4
    book = Workbook()
    book.active.title = "First"
    book.active.append(["Account Code", "Account Name"])
    book.active.append(["100", "Cash"])
    other = book.create_sheet("Second")
    other.append(["Custom Code", "Custom Name"])
    other.append(["200", "Bank"])
    buffer = io.BytesIO()
    book.save(buffer)
    selected = await client.post(
        "/api/v1/workspace/inspect-file",
        headers=headers,
        data={"sheet": "Second"},
        files={"file": ("synthetic.xlsx", buffer.getvalue())},
    )
    assert selected.status_code == 200
    assert selected.json()["sheets"] == ["First", "Second"]
    assert selected.json()["headers"] == ["Custom Code", "Custom Name"]
    other["A2"] = "=1+1"
    buffer = io.BytesIO()
    book.save(buffer)
    formula = await client.post(
        "/api/v1/workspace/inspect-file",
        headers=headers,
        data={"sheet": "Second"},
        files={"file": ("synthetic.xlsx", buffer.getvalue())},
    )
    assert formula.status_code == 422


async def test_preserved_xlsx_validates_and_commits_from_opaque_storage(environment):
    import csv

    client = environment
    headers, options = await scope(client)
    book = Workbook()
    book.active.title = "Journal Data"
    for row in csv.reader(io.StringIO(CSV.decode())):
        book.active.append(row)
    stream = io.BytesIO()
    book.save(stream)
    batch = await upload(
        client,
        headers,
        {**options, "sheet": "Journal Data"},
        stream.getvalue(),
        "synthetic-preserved.xlsx",
    )
    validated = await client.post(f"/api/v1/accounting/imports/{batch}/validate", headers=headers)
    assert validated.status_code == 200, validated.text
    assert validated.json()["status"] == "validated", validated.text
    activation = await client.post(f"/api/v1/accounting/imports/{batch}/commit", headers=headers)
    assert activation.status_code == 200 and activation.json()["row_count"] == 4
