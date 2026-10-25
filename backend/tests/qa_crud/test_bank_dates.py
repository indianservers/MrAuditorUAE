"""Real API/DB date fields, paging, review blockers and full XLSX evidence.

Scanner and broker are explicit test doubles; the actual job processor executes.
"""

import io
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant, setup
from tests.qa_crud.test_country_workflows import (
    bank_payload,
    decision,
    double_boundaries,
    execute,
    source,
)
from tests.qa_crud.test_india_ecr_source import india_context

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable database required"
)


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
async def test_date_evidence_pages_blocks_fallback_and_preserves_full_export(
    environment, monkeypatch, country, currency
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "IN":
        h, eid, base = await india_context(c)
    else:
        h, _, eid, base = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "bank-date-reviewer")
    await grant(c, h, base, reviewer)
    rows = [
        dict(
            id=f"T-{i}",
            account="BANK-1",
            reference=f"SYN-{i}",
            date="2026-06-01",
            currency=currency,
            amount="1",
            bank_booking_date="2026-06-01",
            bank_value_date="2026-06-30",
            bank_date_basis="BOOKING_DATE",
        )
        for i in range(101)
    ]
    books = await source(
        c, h, eid, rows, "Synthetic explicit booking dates", totals={currency: "101"}
    )
    reviewed = await decision(c, reviewer[0], eid, books)
    assert reviewed.status_code == 201, reviewed.text
    books = reviewed.json()
    uncertain = [
        {
            **row,
            "bank_booking_date": "",
            "bank_value_date": "2026-06-01",
            "bank_date_basis": "VALUE_DATE_FALLBACK",
        }
        for row in rows
    ]
    statement = await source(
        c, h, eid, uncertain, "Synthetic value-date-only statement", totals={currency: "101"}
    )
    reviewed = await decision(c, reviewer[0], eid, statement)
    assert reviewed.status_code == 201, reviewed.text
    result = await execute(
        c,
        h,
        eid,
        "country-workflow",
        bank_payload(
            books, reviewed.json(), currency=currency, book_closing="101", statement_closing="101"
        ),
    )
    assert result["body"]["blocked"]
    assert result["body"]["booking_date_review_count"] == 101
    assert result["collection_totals"]["transaction_dates"] == 202
    assert len(result["body"]["transaction_dates"]) == 100
    assert (await decision(c, reviewer[0], eid, result)).status_code == 409
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    page = await c.get(prefix + "?date_offset=200&date_limit=5&limit=1", headers=h)
    assert page.status_code == 200, page.text
    assert len(page.json()["body"]["transaction_dates"]) == 2
    assert all(
        row["source_role"] == "statement" for row in page.json()["body"]["transaction_dates"]
    )
    assert (await c.get(prefix + "?date_offset=-1", headers=h)).status_code == 422
    assert (await c.get(prefix + "?date_limit=501", headers=h)).status_code == 422
    exported = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    sheet = load_workbook(io.BytesIO(exported.content))["Bank transaction dates"]
    assert sheet.max_row == 203
    assert sheet.cell(203, 7).value == "True"
    fixed = await source(
        c,
        h,
        eid,
        [{**row, "id": "S-" + row["id"]} for row in rows],
        "Synthetic corrected explicit statement",
        totals={currency: "101"},
    )
    reviewed = await decision(c, reviewer[0], eid, fixed)
    assert reviewed.status_code == 201, reviewed.text
    corrected = await execute(
        c,
        h,
        eid,
        "country-workflow",
        bank_payload(
            books, reviewed.json(), currency=currency, book_closing="101", statement_closing="101"
        ),
    )
    assert not corrected["body"]["blocked"] and corrected["body"]["booking_date_review_count"] == 0
    approval = await decision(c, reviewer[0], eid, corrected)
    assert approval.status_code == 201, approval.text
