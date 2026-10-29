"""Real API/DB/storage/processor with explicit scanner and dispatch doubles."""

import io
import os
from copy import deepcopy
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, create, grant
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_msme_tax import annual_population

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable migrated database required"
)


@pytest.mark.parametrize(
    "withdrawal",
    ["basis-document", "settlement-review", "other-settlement-document", "advance-document"],
)
async def test_annual_tax_source_review_yearend_observation_exports_and_staleness(
    environment, monkeypatch, withdrawal
):
    from app.modules.country_workflows.msme_tax_schemas import METRICS, ROLES
    from app.modules.country_workflows.schemas import ROW_TYPES

    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await india_context(c)
    checker = await actor(h, "Firm Administrator", "annual-tax-independent-checker")
    author = await actor(h, "Firm Administrator", "annual-tax-support-author")
    for member in (checker, author):
        await grant(c, h, base, member)
    prefix = f"/api/v1/intelligence/engagements/{eid}"
    catalogue = (await c.get(prefix + "/country-workflows", headers=h)).json()
    entry = next(row for row in catalogue["items"] if row["id"] == "india-msme-annual-tax")
    assert entry["available"] and len(entry["roles"]) == 8
    other = withdrawal == "other-settlement-document"
    advance = withdrawal == "advance-document"
    documents = {}
    for purpose in ["facts", "basis", "pnl", *(["other"] if other or advance else [])]:
        record = await create(
            c, author[0], base, "evidence", {"category": "Other", "retain_until": "2036-12-31"}
        )
        response = await c.post(
            base + "/evidence-files/" + record["id"],
            headers=author[0],
            files={
                "file": (
                    "SYNTHETIC-tax-" + purpose + ".txt",
                    (
                        "SYNTHETIC retained "
                        + purpose
                        + " only. Independent tax interpretation, not client or operative-law qualification."
                    ).encode(),
                    "text/plain",
                )
            },
        )
        assert response.status_code == 201, response.text
        documents[purpose] = (record, response.json())
    data, tables = annual_population(other)
    if advance:
        from tests.unit.test_msme_tax import other_settlement

        tables["bank"][0]["date"] = "2026-03-15"
        tables["tax_bases"][0]["history_from"] = "2026-03-01"
        tables["tax_settlements"][0] = {
            **other_settlement(tables),
            "id": "APPLIED",
            "mode": "ASSESSED_BANK_ADVANCE_APPLICATION",
            "amount": "100",
            "settlement_kind": "ADVANCE_APPROPRIATION",
            "principal_allocation_id": "LINK",
            "day": "2027-03-31",
        }
        tables["interest_pnl"][0]["section23_applies"] = False
        tables["reported"][0]["amount"] = "0"
        tables["reported"][1]["amount"] = "0"
    retained = {}
    for role, table_type in ROLES.items():
        records = [
            ROW_TYPES[table_type].model_validate(row).model_dump(mode="json")
            for row in tables[role]
        ]
        options = {}
        if role in {"assessments", "invoices"}:
            options.update(
                msme_context_evidence_version_id=documents["facts"][1]["id"],
                msme_context_assessment="Independent retained supplier, invoice acceptance and written-term scope",
            )
        if role in {"tax_bases", "interest_pnl"}:
            for row in records:
                row.pop("source_evidence_version_id")
                row.pop("source_sha256")
            options["rule_document_version_id"] = documents[
                "basis" if role == "tax_bases" else "pnl"
            ][1]["id"]
        if (other or advance) and role == "tax_settlements":
            for row in records:
                if row["mode"] != "BANK_ALLOCATION_PORTION":
                    row.update(
                        source_evidence_version_id=documents["other"][1]["id"],
                        source_sha256=documents["other"][1]["sha256"],
                    )
        totals = (
            {"INR": str(sum((Decimal(row["amount"]) for row in records), Decimal(0)))}
            if "amount" in records[0]
            else {}
        )
        retained[role] = await source(
            c, h, eid, records, "SYNTHETIC annual tax " + role, table_type, totals, **options
        )
    assert (await decision(c, author[0], eid, retained["tax_bases"])).status_code == 403
    parameters = dict(
        workflow="india-msme-annual-tax",
        source_ids={role: row["id"] for role, row in retained.items()},
        title="SYNTHETIC annual MSE tax bridge",
        rationale="Independent P&L interest, current tax expense, prior disallowed balances and post-year settlement observation",
        **{
            key: value.isoformat() if hasattr(value, "isoformat") else value
            for key, value in vars(data).items()
        },
    )
    await execute(c, h, eid, "country-workflow", parameters, failed="Independently review all")
    for role, item in retained.items():
        review = await decision(c, checker[0], eid, item)
        assert review.status_code == 201, review.text
        retained[role] = review.json()
    parameters["source_ids"] = {role: item["id"] for role, item in retained.items()}
    result = await execute(c, h, eid, "country-workflow", parameters)
    assert result["body"]["net_expected_tax_adjustment"] == (
        "-135" if advance else "-55" if other else "-35"
    )
    assert not result["body"]["blocked"] and result["body"]["exception_count"] == 0
    assert result["body"]["tax_totals"][METRICS[0]]["expected_amount"] == "15"
    assert result["body"]["tax_totals"][METRICS[1]]["expected_amount"] == (
        "0" if advance else "80" if other else "100"
    )
    assert result["body"]["tax_totals"][METRICS[2]]["expected_amount"] == "150"
    assert result["body"]["tax_observed_through"] == "2027-04-05"
    assert result["body"]["as_of"] == "2027-03-31"
    assert result["body"]["msme_document_binding"] == "RETAINED_DOCUMENT_BYTES_VERIFIED"
    assert result["body"]["msme_tax_document_binding"] == "RETAINED_DOCUMENT_BYTES_VERIFIED"
    assert not result["body"]["statutory_calculation_qualified"]
    assert (await decision(c, h, eid, result)).status_code == 403
    assert (await decision(c, author[0], eid, result)).status_code == 403
    approved = await decision(c, checker[0], eid, result)
    assert approved.status_code == 201, approved.text
    result = approved.json()
    page = (
        await c.get(
            prefix + f"/events/{result['id']}?limit=1&tax_invoice_offset=1&tax_invoice_limit=1",
            headers=h,
        )
    ).json()
    assert len(page["body"]["rows"]) == 1 and len(page["body"]["invoice_tax_details"]) == 1
    assert page["body"]["invoice_tax_details"][0]["id"] == "PRIOR"
    assert (
        page["collection_totals"]["rows"] == 6
        and page["collection_totals"]["invoice_tax_details"] == 2
    )
    exported = await c.get(prefix + f"/events/{result['id']}/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text[:200] if exported.status_code != 200 else ""
    workbook = load_workbook(io.BytesIO(exported.content))
    for title, count in {
        "Rows": 7,
        "MSE annual tax totals": 4,
        "MSE tax invoice scope": 3,
        "MSE tax bases": 3,
        "MSE tax PNL interest": 3,
        "MSE tax settlements": 5 if other else 4,
        "MSE tax principal links": 4,
        "MSE tax invoice facts": 3,
        "MSE tax supplier facts": 3,
        "MSE tax bank residuals": 4,
    }.items():
        assert workbook[title].max_row == count, title
    financial_body = deepcopy(result["body"])
    if withdrawal == "settlement-review":
        response = await decision(c, checker[0], eid, retained["tax_settlements"], "reject")
    else:
        purpose = "other" if other or advance else "basis"
        response = await c.post(
            base + "/evidence-files/" + documents[purpose][0]["id"],
            headers=author[0],
            files={
                "file": (
                    "SYNTHETIC-tax-support-replaced.txt",
                    b"SYNTHETIC changed assessed tax support",
                    "text/plain",
                )
            },
        )
    assert response.status_code == 201, response.text
    stale = (await c.get(prefix + "/events/" + result["id"], headers=h)).json()
    assert stale["freshness"] == "STALE" and stale["body"] == financial_body
    assert (await decision(c, checker[0], eid, result)).status_code == 409
