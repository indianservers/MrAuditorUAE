"""Professional SoD paging/export through retained, independently reviewed populations."""

import csv
import io
import json
import os
import uuid

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import action, actor, approve, create, grant, setup
from tests.qa_crud.test_country_workflows import double_boundaries
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_sod_intervals import assignment, mapping, policy, population

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required"
)


async def prepare(c, h, eid, audit, reviewer, country):
    base = f"/api/v1/grc/engagements/{eid}"
    rows = [
        assignment(
            permission,
            scope="A" if permission == "create" else "B",
            principal=f"SYN-{country}-user-{index:04d}",
            end="2026-06-30" if permission == "approve" else None,
        )
        for index in range(121)
        for permission in ("create", "approve")
    ]
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=list(rows[0]), lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    content = output.getvalue().encode()

    async def retain(name, content):
        record = await create(
            c, h, audit, "evidence", dict(category="Other", retain_until="2036-12-31")
        )
        response = await c.post(
            audit + "/evidence-files/" + record["id"],
            headers=h,
            files={"file": (name, content, "text/csv" if name.endswith(".csv") else "text/plain")},
        )
        assert response.status_code == 201, response.text
        record["version"] += 1
        return await approve(c, h, reviewer[0], audit, record)

    source = await retain("SYNTHETIC-access.csv", content)
    definition = await retain(
        "SYNTHETIC-policy.txt",
        b"Fictional 2026 policy: create/approve incompatible. Appendix 2: A/B directly overlap.\n",
    )
    control = await retain(
        "SYNTHETIC-controls.txt",
        b"Fictional independently assessed full 2026 history coverage, 242 assignments, 121 principals.\n",
    )
    rule = policy(scope_mode="reviewed-overlap", scope_overlaps=[mapping()]).model_dump(mode="json")
    rule["definition_evidence_id"] = definition["id"]
    response = await c.post(
        base + "/records/sod-rule",
        headers=h,
        json=dict(
            reference=f"SYN-{country}-SOD-POLICY", title="Synthetic reviewed policy", body=rule
        ),
    )
    assert response.status_code == 201, response.text
    rule = await approve(c, h, reviewer[0], base, response.json())
    body = population(rows).model_dump(mode="json")
    del body["assignments"]
    del body["source_hash"]
    body.update(
        rule_id=rule["id"],
        source_evidence_id=source["id"],
        population_control_evidence_id=control["id"],
    )
    response = await c.post(
        base + "/sod/import",
        headers=h,
        data={
            "metadata": json.dumps(
                dict(
                    reference=f"SYN-{country}-SOD-121",
                    title="Synthetic 121-principal historical population",
                    body=body,
                )
            )
        },
        files={"file": ("SYNTHETIC-access.csv", content, "text/csv")},
    )
    assert response.status_code == 201, response.text
    access = await approve(c, h, reviewer[0], base, response.json())
    response = await c.post(base + f"/records/{access['id']}/calculate?compact_sod=true", headers=h)
    assert response.status_code == 201, response.text
    return dict(
        engagement=eid,
        access=access,
        source=source,
        run=response.json(),
        reference=f"SYN-{country}-SOD-121",
        reviewer=reviewer,
    )


@pytest.mark.parametrize("country", ["AE", "IN"])
async def test_complete_reporting_paging_permissions_and_stale_snapshot(
    environment, monkeypatch, country
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, audit = await setup(c)
    else:
        h, eid, audit = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "sod-report-reviewer")
    await grant(c, h, audit, reviewer)
    case = await prepare(c, h, eid, audit, reviewer, country)
    run = case["run"]
    base = f"/api/v1/grc/engagements/{eid}"
    detail = base + f"/runs/{run['id']}/sod"
    assert "assignments" not in run["inputs"]
    assert "conflicts" not in run["results"]
    assert run["results"]["collection_totals"]["conflicts"] == 121
    summary = (await c.get(detail, headers=h)).json()
    assert summary["record_versions_status"] == "CURRENT_RECORD_VERSIONS"
    assert summary["source_bytes_reverified_on_read"] is False
    for collection, total in (
        ("conflicts", 121),
        ("conflict_intervals", 121),
        ("assignment_pair_links", 121),
        ("assignments", 242),
    ):
        page = await c.get(detail + f"/rows/{collection}?offset=100&limit=21", headers=h)
        assert page.status_code == 200, page.text
        assert page.json()["total"] == total and len(page.json()["items"]) == 21
        assert page.json()["items"][0].get("principal", "").startswith(f"SYN-{country}-user-")
    page = await c.get(detail + "/rows/assignments?offset=241&limit=1", headers=h)
    assert page.json()["items"][0]["assignment_index"] == 242
    canonical = await c.get(detail + "/rows/assignment-pair-links?offset=100&limit=21", headers=h)
    assert canonical.status_code == 200 and len(canonical.json()["items"]) == 21
    assert (await c.get(detail + "/rows/conflicts?offset=121", headers=h)).json()["items"] == []
    assert (await c.get(detail + "/rows/conflicts?limit=501", headers=h)).status_code == 422
    assert (await c.get(detail + "/rows/unknown", headers=h)).status_code == 404
    assert (await c.get(base + f"/runs/{uuid.uuid4()}/sod", headers=h)).status_code == 404
    listing = (
        await c.get(
            base + f"/runs?kind=sod-access&record_id={case['access']['id']}&compact_sod=true",
            headers=h,
        )
    ).json()
    assert listing["total"] == 1 and "assignments" not in listing["items"][0]["inputs"]
    exported = await c.get(detail + "/export", headers=h)
    assert exported.status_code == 200, exported.text if exported.status_code != 200 else ""
    book = load_workbook(io.BytesIO(exported.content))
    assert book["Conflicts"].max_row == book["Assignment links"].max_row == 122
    assert book["Assignments"].max_row == 243
    assert book["Permission pairs"].max_row == 2
    assert book["Assignments"].cell(243, 1).value == 242
    for role in ("Risk Owner", "Board Viewer", "Internal Auditor"):
        person = await actor(h, role, "report-" + role.replace(" ", "-"))
        if role != "Internal Auditor":
            await grant(c, h, audit, person)
        for suffix in ("", "/export", "/rows/assignments"):
            assert (await c.get(detail + suffix, headers=person[0])).status_code == 403
    assert (await c.get(detail, headers={**h, "X-Tenant-ID": str(uuid.uuid4())})).status_code == 403
    await action(c, reviewer[0], audit, case["source"], "reopen")
    stale = await c.get(detail, headers=h)
    assert stale.status_code == 200, stale.text
    assert stale.json()["record_versions_status"] == "STALE_HISTORICAL_SNAPSHOT"
    assert stale.json()["snapshot_sha256"] == summary["snapshot_sha256"]
    assert stale.json()["results"] == summary["results"]
    stale_export = await c.get(detail + "/export", headers=h)
    assert stale_export.status_code == 200
    scope = list(load_workbook(io.BytesIO(stale_export.content))["Scope"].values)
    assert ("snapshot", "record_versions_status", "STALE_HISTORICAL_SNAPSHOT") in scope
