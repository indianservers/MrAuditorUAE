"""Synthetic offsite custody documents through actual retained sources and review."""

import hashlib
import io
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, create, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_inventory_custody import controls, followup, pending, population

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required"
)


async def prepare(c, h, eid, reviewer, case="received", count=1, reviewed=True):
    base = f"/api/v1/engagements/{eid}/audit"
    rows = dict(positions=[], responses=[])
    for index in range(count):
        template = population()
        position = template["positions"][0]
        letter = template["responses"][0]
        position.update(id=f"POS-{index:04d}", sku=f"SKU-{index:04d}")
        letter.update(
            id=f"LETTER-{index:04d}",
            position_id=position["id"],
            sku=position["sku"],
            request_reference=f"SYN-REQUEST-LINE-{index:04d}",
            response_reference=f"SYN-RESPONSE-LINE-{index:04d}",
        )
        if case == "zero":
            position["quantity"] = "0"
            letter.update(
                requested_quantity="0",
                response_quantity="0",
                condition="not-applicable",
                held_for_entity="none-held",
            )
        if case == "pending":
            letter = followup(pending(letter))
            letter["followup_reference"] = f"SYN-FOLLOWUP-LINE-{index:04d}"
        for role in ("request", "response", "followup"):
            letter[role + "_evidence_version_id"] = ""
            letter[role + "_source_sha256"] = ""
        rows["positions"].append(position)
        rows["responses"].append(letter)
    documents = {}
    for role in ("request", "followup") if case == "pending" else ("request", "response"):
        text = (
            "SYNTHETIC ONLY. Fictional offsite quantity and condition "
            + role
            + "; no authenticated custodian or title.\n"
        )
        text += "\n".join(
            f"SYN-{role.upper()}-LINE-{index:04d}|POS-{index:04d}|WAREHOUSE-1|OFFSITE-1|SKU-{index:04d}|LOT-1|KG|2026-06-30|{0 if case == 'zero' else 10}|"
            + (
                "ALTERNATIVE PROCEDURES ASSESSED; NO REPLY"
                if role == "followup"
                else "QUANTITY AND CONDITION REQUEST"
                if role == "request"
                else "NOT APPLICABLE|NONE HELD FOR ENTITY"
                if case == "zero"
                else "GOOD|HELD FOR ENTITY YES"
            )
            for index in range(count)
        )
        record = await create(
            c, h, base, "evidence", dict(category="Other", retain_until="2036-12-31")
        )
        response = await c.post(
            base + "/evidence-files/" + record["id"],
            headers=h,
            files={"file": (f"SYNTHETIC-custody-{role}.txt", text.encode(), "text/plain")},
        )
        assert response.status_code == 201, response.text
        documents[role] = {**response.json(), "record": record["id"], "expected_text": text}
    events = {}
    for role, kind in (
        ("positions", "inventory-custody-positions"),
        ("responses", "inventory-custody-confirmations"),
    ):
        options = dict(custody_quantity_controls=controls(rows[role], role == "responses"))
        if role == "responses":
            options.update(
                {
                    "confirmation_" + role + "_document_version_id": document["id"]
                    for role, document in documents.items()
                }
            )
        event = await source(
            c, h, eid, rows[role], "SYNTHETIC custody " + case + " " + role, kind, {}, **options
        )
        if reviewed:
            response = await decision(c, reviewer[0], eid, event)
            assert response.status_code == 201, response.text
            event = response.json()
        events[role] = event
    return dict(
        engagement=eid,
        documents=documents,
        events=events,
        sources={role: event["id"] for role, event in events.items()},
        payload=dict(
            workflow="inventory-custody-confirmations",
            source_ids={role: event["id"] for role, event in events.items()},
            as_of="2026-06-30",
            review_through="2026-07-15",
            title="SYNTHETIC " + case + " inventory custody",
            rationale="Independent scoped offsite quantities, condition assessment and retained correspondence; no title assurance",
        ),
    )


@pytest.mark.parametrize("country", ["AE", "IN"])
@pytest.mark.parametrize("case", ["received", "pending", "zero"])
async def test_custody_review_retained_download_approval_export_and_staleness(
    environment, monkeypatch, country, case
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "custody-source-reviewer")
    final = await actor(h, "Firm Administrator", "custody-final-reviewer")
    for person in (reviewer, final):
        await grant(c, h, base, person)
    fixture = await prepare(c, h, eid, reviewer, case, reviewed=False)
    await execute(
        c, h, eid, "country-workflow", fixture["payload"], failed="Independently review custody"
    )
    for role, event in fixture["events"].items():
        response = await decision(c, reviewer[0], eid, event)
        assert response.status_code == 201, response.text
        fixture["events"][role] = response.json()
        fixture["payload"]["source_ids"][role] = response.json()["id"]
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    row = result["body"]["rows"][0]
    assert not result["body"]["blocked"]
    assert row["state"] == (
        "MATCHED" if case != "pending" else "CUSTODY_ALTERNATIVE_PROCEDURES_REQUIRE_REVIEW"
    )
    assert row["quantity_difference"] == ("0" if case != "pending" else None)
    if case == "pending":
        assert (await decision(c, final[0], eid, result)).status_code == 409
        accepted = await decision(c, h, eid, result, action="accept", item_id=row["id"])
        assert accepted.status_code == 201, accepted.text
        result = accepted.json()
    approved = await decision(c, final[0], eid, result)
    assert approved.status_code == 201, approved.text
    result = approved.json()
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    for role, document in fixture["documents"].items():
        response = await c.get(
            base + f"/evidence-files/{document['record']}/{document['id']}/download", headers=h
        )
        assert (
            response.status_code == 200 and response.content == document["expected_text"].encode()
        )
        assert (
            hashlib.sha256(response.content).hexdigest()
            == row["confirmation_evidence"][role + "_source_sha256"]
        )
        assert any(
            receipt.get("id") == document["id"] and receipt.get("record_id") == document["record"]
            for receipt in result["dependencies"]
        )
    response = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    assert len(list(book["Custody quantity controls"].values)) == 3
    assert len(list(book["Custody documents"].values)) == 3
    book.close()
    document = fixture["documents"]["request"]
    revised = await c.post(
        base + "/evidence-files/" + document["record"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-revised-request.txt",
                b"SYNTHETIC revised request needs independent re-assessment",
                "text/plain",
            )
        },
    )
    assert revised.status_code == 201, revised.text
    assert (await c.get(prefix, headers=h)).json()["freshness"] == "STALE"


async def test_custody_controls_and_all_correspondence_export_beyond_first_hundred(
    environment, monkeypatch
):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, base = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "custody-pagination-reviewer")
    await grant(c, h, base, reviewer)
    fixture = await prepare(c, h, eid, reviewer, count=121)
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    page = await c.get(
        prefix + "?offset=120&limit=1&custody_control_offset=241&custody_control_limit=1", headers=h
    )
    assert page.status_code == 200, page.text
    assert page.json()["body"]["rows"][0]["position_source_id"] == "POS-0120"
    assert page.json()["body"]["custody_quantity_controls"][0]["position_id"] == "POS-0120"
    assert page.json()["collection_totals"]["custody_quantity_controls"] == 242
    exported = await c.get(prefix + "/export?format=xlsx", headers=h)
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert len(list(book["Custody documents"].values)) == 243
    assert len(list(book["Custody quantity controls"].values)) == 243
    book.close()
