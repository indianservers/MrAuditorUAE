"""Synthetic scope workpaper integration; not a qualified India client pilot."""

import io
import os
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.modules.country_workflows.india_company_scope import PINNED, ROLES
from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant
from tests.qa_crud.test_country_workflows import decision, double_boundaries
from tests.qa_crud.test_india_ecr_source import india_context
from tests.qa_crud.test_uae_company_accounts import context as uae_context
from tests.qa_crud.test_uae_ct_audit_duty import upload
from tests.unit.test_india_company_scope_core import facts

pytestmark = pytest.mark.skipif(os.getenv("INTEGRATION_TESTS") != "1", reason="Owned DB required")
FOLDER = (
    Path(__file__).resolve().parents[3]
    / "docs/country-workflows/authority-references/india-company-reporting-2026-10-31"
)
FILES = {
    "small_company_amendment": "small-company-gsr880-2025.pdf",
    "caro_commencement": "caro2020-commencement-amendment.pdf",
    "ifc_notification": "ifc-private-gsr583-2017-original.pdf",
}


async def documents(c, h, base):
    sources, documents = {}, []
    for role in sorted(ROLES):
        # Unpinned legal files deliberately remain SYNTHETIC, not authority-qualified.
        content = (
            (FOLDER / FILES[role]).read_bytes()
            if role in PINNED
            else (
                f"%PDF-1.4\nSYNTHETIC {role} {uuid4()} no professional or original legal qualification"
            ).encode()
        )
        sources[role] = await upload(c, h, base, role + ".pdf", content, "application/pdf")
        version = sources[role][1]
        documents.append(
            {"role": role, "evidence_version_id": version["id"], "sha256": version["sha256"]}
        )
    return sources, {
        "profile": "india-company-reporting-scope/1",
        "label": "SYNTHETIC India reporting scope",
        "facts": facts(),
        "adopted_from": "2025-04-01",
        "adopted_through": "2026-06-30",
        "source_chain_status": "independently_assessed",
        "other_operative_changes": "assessed_none",
        "legal_adoption_assessment": "SYNTHETIC adopted test candidates; no real legal qualification",
        "documents": documents,
    }


@pytest.mark.parametrize("changed_role", sorted(ROLES))
async def test_scope_review_export_and_every_original_stales(
    environment, monkeypatch, changed_role
):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await india_context(c, year=2025)
    reviewer = await actor(h, "Firm Administrator", "india-company-reviewer")
    await grant(c, h, base, reviewer)
    sources, data = await documents(c, h, base)
    prefix = f"/api/v1/intelligence/engagements/{eid}"
    response = await c.post(prefix + "/india-company-reporting-scope", headers=h, json=data)
    assert response.status_code == 201, response.text
    result = response.json()
    row = result["body"]["rows"][0]
    assert row["companies_act_section139_audit"] is True
    assert row["small_company"] is True and row["caro2020_applies"] is False
    assert result["body"]["ready_for_statutory_issue"] is False
    assert (await decision(c, h, eid, result)).status_code == 403
    response = await decision(c, reviewer[0], eid, result)
    assert response.status_code == 201, response.text
    reviewed = response.json()
    report = None
    if changed_role == "company_facts":
        from app.modules.regulatory.report_packs import manifest

        preset = manifest("india-company-scope:standalone/1")
        template = await c.post(
            f"/api/v1/regulatory/engagements/{eid}/report-templates",
            headers=h,
            json={
                "title": "SYNTHETIC retained India company applicability summary",
                "category": "india-statutory",
                "requirements_pack": preset["id"],
                "requirements": [{**r, "disposition": "REQUIRED"} for r in preset["requirements"]],
                "requirements_assessed": True,
                "requirements_basis": "SYNTHETIC six explicit decisions from one current reviewed workpaper",
                "blocks": [
                    {
                        "key": "scope",
                        "title": "Retained company decisions",
                        "kind": "regulatory_checklist",
                        "source_kind": "intelligence",
                        "source_id": reviewed["id"],
                        "requirement_keys": [r["key"] for r in preset["requirements"]],
                    }
                ],
                "rationale": "SYNTHETIC applicability summary, not a full statutory report",
            },
        )
        assert template.status_code == 201, template.text
        approved_template = await decision(c, reviewer[0], eid, template.json())
        assert approved_template.status_code == 201, approved_template.text
        generated = await c.post(
            f"/api/v1/regulatory/engagements/{eid}/reports",
            headers=h,
            json={
                "template_id": approved_template.json()["id"],
                "rationale": "SYNTHETIC exact reviewed scope and preserved false duty decisions",
            },
        )
        assert generated.status_code == 201, generated.text
        report = generated.json()
        assert (
            report["body"]["blocks"][0]["data"]["rows"][0]["companies_act_section139_audit"] is True
        )
        assert report["body"]["blocks"][0]["data"]["rows"][0]["caro2020_applies"] is False
        assert report["body"]["requirements_coverage"]["complete"] is True
        assert report["body"]["requirements_coverage"]["statutory_completeness_qualified"] is False
        exported = await c.get(
            prefix + "/events/" + report["id"] + "/export?format=xlsx", headers=h
        )
        assert exported.status_code == 200, exported.text
        exported_book = load_workbook(io.BytesIO(exported.content), read_only=True)
        cells = {
            str(cell)
            for sheet in exported_book.worksheets
            for row in sheet.values
            for cell in row
            if cell is not None
        }
        assert all(requirement["key"] in cells for requirement in preset["requirements"])
        exported_book.close()
        for export_format in ("docx", "pdf"):
            rendered = await c.get(
                prefix + "/events/" + report["id"] + "/export?format=" + export_format, headers=h
            )
            assert rendered.status_code == 200, rendered.text[:300]
            assert len(rendered.content) > 500
    path = prefix + "/events/" + reviewed["id"]
    export = await c.get(path + "/export?format=xlsx", headers=h)
    assert export.status_code == 200, export.text
    book = load_workbook(io.BytesIO(export.content), read_only=True)
    rows = list(book["Rows"].values)
    assert str(rows[1][rows[0].index("companies_act_section139_audit")]).lower() == "true"
    book.close()
    replacement = await c.post(
        base + "/evidence-files/" + sources[changed_role][0]["id"],
        headers=h,
        files={"file": ("replaced.txt", b"SYNTHETIC source changed; reassess scope", "text/plain")},
    )
    assert replacement.status_code == 201, replacement.text
    stale = await c.get(path, headers=h)
    assert stale.json()["freshness"] == "STALE"
    assert stale.json()["body"] == reviewed["body"]
    assert (await decision(c, reviewer[0], eid, reviewed)).status_code == 409
    if report:
        retained_report = await c.get(prefix + "/events/" + report["id"], headers=h)
        assert retained_report.json()["freshness"] == "STALE"
        assert retained_report.json()["body"] == report["body"]


@pytest.mark.parametrize(
    "failure", ["country", "period", "hash", "pinned", "duplicate", "integrity"]
)
async def test_scope_guards(environment, monkeypatch, failure):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = (
        await uae_context(c) if failure == "country" else await india_context(c, year=2025)
    )
    _, data = await documents(c, h, base)
    if failure == "period":
        data["facts"]["financial_period_end"] = "2026-03-30"
    if failure in {"hash", "pinned"}:
        source = next(
            d
            for d in data["documents"]
            if d["role"] == ("company_facts" if failure == "hash" else "ifc_notification")
        )
        source["sha256"] = "f" * 64
    if failure == "duplicate":
        data["documents"][1] = data["documents"][0]
    if failure == "integrity":
        from app.modules.country_workflows import service

        monkeypatch.setattr(service, "hash_file", lambda _: "0" * 64)
    response = await c.post(
        f"/api/v1/intelligence/engagements/{eid}/india-company-reporting-scope",
        headers=h,
        json=data,
    )
    assert response.status_code == 422, response.text


@pytest.mark.parametrize("unresolved", ["adoption", "facts"])
async def test_unresolved_scope_cannot_be_approved(environment, monkeypatch, unresolved):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await india_context(c, year=2025)
    reviewer = await actor(h, "Firm Administrator", "india-pending-reviewer")
    await grant(c, h, base, reviewer)
    _, data = await documents(c, h, base)
    if unresolved == "adoption":
        data["source_chain_status"] = "unresolved"
    else:
        data["facts"]["other_special_exemption_claimed"] = None
    response = await c.post(
        f"/api/v1/intelligence/engagements/{eid}/india-company-reporting-scope",
        headers=h,
        json=data,
    )
    assert response.status_code == 201, response.text
    assert response.json()["body"]["rows"][0]["companies_act_section139_audit"] is None
    rejected = await decision(c, reviewer[0], eid, response.json())
    assert rejected.status_code == 409, rejected.text
