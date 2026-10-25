"""Real retained imports, independent reviews, workflow receipts and export."""

import io
import os
import uuid
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.db.session import Session, tenant_context
from app.modules.audit_execution.models import EvidenceVersion
from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.unit.test_import_vat import line
from tests.unit.test_vat_box_scope import full_filed_boxes

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Disposable DB required"
)


@pytest.mark.parametrize("review_sources", [True, False])
async def test_import_vat_retained_sources_review_export_and_replacement(
    environment, monkeypatch, review_sources
):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, base = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "import-vat-source-reviewer")
    final = await actor(h, "Firm Administrator", "import-vat-final-reviewer")
    bridge_reviewer = await actor(h, "Firm Administrator", "import-vat-bridge-reviewer")
    for person in (reviewer, final, bridge_reviewer):
        await grant(c, h, base, person)
    selected, originals = {}, {}
    for role in ("customs", "portal", "books"):
        rows = [
            line(),
            line(
                id="ADJ-1",
                entry_kind="adjustment",
                adjustment_reference="ADJ-1",
                net="-100",
                tax="-5",
            ),
        ]
        for row in rows:
            row["evidence_reference"] = f"SYNTHETIC-{role.upper()}-{row['id']}"
        prepared = await source(
            c,
            h,
            eid,
            rows,
            f"SYNTHETIC {role} import VAT",
            table_type="import-vat-lines",
            totals={"AED": "45"},
            import_taxable_base_control="900",
        )
        originals[role] = prepared
        denied = await decision(c, h, eid, prepared)
        assert denied.status_code == 403
        if review_sources:
            approved = await decision(c, reviewer[0], eid, prepared)
            assert approved.status_code == 201, approved.text
            selected[role] = approved.json()["id"]
        else:
            selected[role] = prepared["id"]
    output = await execute(
        c,
        h,
        eid,
        "country-workflow",
        {
            "workflow": "uae-import-vat",
            "source_ids": selected,
            "as_of": "2026-06-30",
            "currency": "AED",
            "title": "SYNTHETIC UAE import VAT declaration comparison",
            "rationale": "Independent declarations, assessed portal details and book movements compared without legal entitlement inference",
        },
    )
    assert output["body"]["exception_count"] == 0 and not output["body"]["blocked"]
    assert output["body"]["role_totals"]["customs"]["declaration"]["tax"] == "50"
    assert output["body"]["role_totals"]["customs"]["adjustment"]["tax"] == "-5"
    denied = await decision(c, h, eid, output)
    assert denied.status_code == 403
    prefix = f"/api/v1/intelligence/engagements/{eid}"
    bridge_input = {
        "source_role": "books",
        "label": "SYNTHETIC assessed import VAT boxes",
        "assessment": "Independent book role and declaration/adjustment controls assessed for Box 6 and Box 7 only",
        "boxes": {
            "6": {"expected_rows": 1, "taxable_amount": "1000", "tax": "50"},
            "7": {"expected_rows": 1, "taxable_amount": "-100", "tax": "-5"},
        },
    }
    premature = await c.post(
        prefix + f"/events/{output['id']}/import-vat-box-source", headers=h, json=bridge_input
    )
    assert premature.status_code == 409
    reviewed = await decision(c, final[0], eid, output)
    if not review_sources:
        assert reviewed.status_code == 409 and "independent review" in reviewed.text
        return
    assert reviewed.status_code == 201, reviewed.text
    bridge = await c.post(
        prefix + f"/events/{reviewed.json()['id']}/import-vat-box-source",
        headers=h,
        json=bridge_input,
    )
    assert bridge.status_code == 201, bridge.text
    bridge = bridge.json()
    assert bridge["body"]["review_status"] == "UNVERIFIED" and bridge["body"]["row_count"] == 2
    assert bridge["body"]["native_metadata"]["source_role"] == "books"
    assert (await decision(c, h, eid, bridge)).status_code == 403
    # A prior review is not data preparation; the existing maker/checker rule applies.
    bridge_reviewed = await decision(c, final[0], eid, bridge)
    assert bridge_reviewed.status_code == 201, bridge_reviewed.text
    filed_rows = full_filed_boxes()
    filed = await source(
        c,
        h,
        eid,
        filed_rows,
        "SYNTHETIC independently retained filed import boxes",
        table_type="return-boxes",
        totals={"AED": "230"},
    )
    filed_reviewed = await decision(c, reviewer[0], eid, filed)
    assert filed_reviewed.status_code == 201, filed_reviewed.text
    vat_payload = {
        "workflow": "uae-vat-return",
        "source_ids": {
            "books": bridge_reviewed.json()["id"],
            "return": filed_reviewed.json()["id"],
        },
        "as_of": "2026-06-30",
        "required_boxes": ["6", "7"],
        "restrict_to_required_boxes": True,
        "box_scope_assessment": "Only the assessed import boxes 6 and 7 are compared; all other filed boxes remain outside this review",
        "title": "SYNTHETIC selected import boxes versus independent filed evidence",
        "rationale": "Compare only assessed selected import boxes with separately retained filed VAT201 evidence",
    }
    default_result = await execute(
        c,
        h,
        eid,
        "country-workflow",
        {**vat_payload, "restrict_to_required_boxes": False, "box_scope_assessment": ""},
    )
    assert (
        default_result["body"]["exception_count"] == 12
        and default_result["body"]["excluded_box_rows"] == []
    )
    vat_result = await execute(c, h, eid, "country-workflow", vat_payload)
    assert vat_result["body"]["exception_count"] == 0
    assert len(vat_result["body"]["excluded_box_rows"]) == 12
    assert (
        vat_result["body"]["box_comparison_scope"]["whole_return_completeness"] == "NOT_DETERMINED"
    )
    excluded_page = await c.get(
        prefix + f"/events/{vat_result['id']}?excluded_offset=2&excluded_limit=2", headers=h
    )
    assert excluded_page.status_code == 200
    assert (
        len(excluded_page.json()["body"]["rows"]) == 2
        and len(excluded_page.json()["body"]["excluded_box_rows"]) == 2
    )
    assert excluded_page.json()["collection_totals"]["excluded_box_rows"] == 12
    vat_reviewed = await decision(c, bridge_reviewer[0], eid, vat_result)
    assert vat_reviewed.status_code == 201, vat_reviewed.text
    scoped_export = await c.get(
        prefix + f"/events/{vat_reviewed.json()['id']}/export?format=xlsx", headers=h
    )
    assert scoped_export.status_code == 200
    scoped_workbook = load_workbook(io.BytesIO(scoped_export.content))
    assert (
        scoped_workbook["Rows"].max_row == 3 and scoped_workbook["Excluded VAT boxes"].max_row == 13
    )
    scoped_paper = await c.post(
        prefix + f"/events/{vat_reviewed.json()['id']}/working-paper",
        headers=h,
        json={
            "reference": "SYN-VAT-IMPORT-SCOPE",
            "title": "Synthetic selected import VAT scope",
            "objective": "Document only the assessed import boxes while preserving other filed boxes as excluded evidence",
        },
    )
    assert scoped_paper.status_code == 201, scoped_paper.text
    tables = scoped_paper.json()["body"]["financial_tables"]
    assert tables[0]["values"]["box_comparison_scope"]["mode"] == "SELECTED_REQUIRED_BOXES"
    assert tables[-1]["total_rows"] == 12
    assert any(
        section["heading"] == "Selected VAT201 comparison coverage"
        and "Whole-return completeness is not determined" in section["text"]
        for section in scoped_paper.json()["body"]["sections"]
    )
    await execute(
        c,
        h,
        eid,
        "country-workflow",
        {
            **vat_payload,
            "source_ids": {
                "books": filed_reviewed.json()["id"],
                "return": bridge_reviewed.json()["id"],
            },
        },
        failed="books role",
    )
    exported = await c.get(
        prefix + f"/events/{reviewed.json()['id']}/export?format=xlsx", headers=h
    )
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content))
    values = [str(cell.value) for sheet in workbook for row in sheet for cell in row]
    assert "declaration" in values and "adjustment" in values
    assert any("IMPORT-1" in value for value in values)
    evidence_id = originals["customs"]["body"]["evidence_version_id"]
    async with Session() as db:
        await tenant_context(db, h["X-Tenant-ID"])
        version = await db.get(EvidenceVersion, uuid.UUID(evidence_id))
        record_id = str(version.record_id)
    raw = await c.get(base + f"/evidence-files/{record_id}/{evidence_id}/download", headers=h)
    assert raw.status_code == 200
    assert Decimal(originals["customs"]["body"]["control_totals"]["AED"]) == 45
    paper = await c.post(
        prefix + f"/events/{reviewed.json()['id']}/working-paper",
        headers=h,
        json={
            "reference": "SYN-IMPORT-VAT",
            "title": "Synthetic declaration VAT review",
            "objective": "Retain declaration, adjustment and independent portal/book evidence comparison",
        },
    )
    assert paper.status_code == 201, paper.text
    replacement = await c.post(
        base + f"/evidence-files/{record_id}",
        headers=h,
        files={"file": ("SYNTHETIC-customs-v2.csv", raw.content + b"\n", "text/csv")},
    )
    assert replacement.status_code == 201, replacement.text
    stale = await c.get(prefix + f"/events/{reviewed.json()['id']}", headers=h)
    assert stale.status_code == 200 and stale.json()["freshness"] == "STALE"
    for result_id in (bridge_reviewed.json()["id"], vat_reviewed.json()["id"]):
        stale = await c.get(prefix + f"/events/{result_id}", headers=h)
        assert stale.status_code == 200 and stale.json()["freshness"] == "STALE"
