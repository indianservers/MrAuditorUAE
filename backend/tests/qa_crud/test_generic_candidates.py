"""Committed ledgers and retained supporting CSVs, with scanner/broker doubles."""

import io
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase2.test_pipeline import committed
from tests.phase4.test_workflows import actor, create, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute

pytestmark = pytest.mark.skipif(os.getenv("INTEGRATION_TESTS") != "1", reason="Owned database")


async def context(c, country, existing=None):
    h, options, eid, audit = existing or await setup(c)
    options = dict(options)
    if country == "IN":
        entity = await c.post(
            "/api/v1/entities",
            headers=h,
            json=dict(
                client_id=options["client_id"],
                name="SYNTHETIC India generic candidates",
                country="IN",
                jurisdiction="Synthetic India",
                functional_currency="INR",
                presentation_currency="INR",
            ),
        )
        assert entity.status_code == 201, entity.text
        period = await c.post(
            "/api/v1/financial-periods",
            headers=h,
            json=dict(
                entity_id=entity.json()["id"], start_date="2026-04-01", end_date="2027-03-31"
            ),
        )
        assert period.status_code == 201, period.text
        engagement = await c.post(
            "/api/v1/engagements",
            headers=h,
            json=dict(
                entity_id=entity.json()["id"],
                period_id=period.json()["id"],
                name="SYNTHETIC India generic reconciliation",
                engagement_type="External Audit",
            ),
        )
        assert engagement.status_code == 201, engagement.text
        eid = engagement.json()["id"]
        audit = f"/api/v1/engagements/{eid}/audit"
        options.update(entity_id=entity.json()["id"], period_id=period.json()["id"])
    return h, options, eid, audit


async def prepare(c, h, options, eid, audit, country, count=121):
    currency = "INR" if country == "IN" else "AED"
    # One balanced journal, 121 indistinguishable cash postings and one equity line.
    header = "Account Code,Account Name,Journal Number,Posting Date,Debit,Credit,Source ID,Opening,Manual,Narration\n"
    lines = [
        f"100,Cash,SYN-SHARED,2026-06-01,1,0,L{i:03},false,false,SYNTHETIC duplicate amount"
        for i in range(count)
    ]
    lines += [
        f"300,Equity,SYN-SHARED,2026-06-01,0,{count},OFFSET,false,false,SYNTHETIC balancing line"
    ]
    _, snapshot = await committed(
        c,
        h,
        {
            **options,
            "expected_rows": count + 1,
            "currency": currency,
            "expected_debit": str(count),
            "expected_credit": str(count),
        },
        (header + "\n".join(lines) + "\n").encode(),
    )
    evidence = await create(
        c, h, audit, "evidence", dict(category="Other", retain_until="2036-12-31")
    )
    party = "SYNTHETIC شركة ألف" if currency == "AED" else "SYNTHETIC भारत आपूर्तिकर्ता"
    csv = "reference,date,currency,amount,counterparty\n" + "".join(
        f"SYN-SHARED,2026-06-01,{currency},1,{party}\n" for _ in range(count)
    )
    uploaded = await c.post(
        audit + "/evidence-files/" + evidence["id"],
        headers=h,
        files={"file": ("SYNTHETIC-generic-bank.csv", csv.encode(), "text/csv")},
    )
    assert uploaded.status_code == 201, uploaded.text
    table = await execute(
        c,
        h,
        eid,
        "source-table",
        dict(evidence_version_id=uploaded.json()["id"], table_kind="bank"),
    )
    reviewer = await actor(h, "Firm Administrator", f"generic-candidate-reviewer-{country}")
    await grant(c, h, audit, reviewer)
    approved = await decision(c, reviewer[0], eid, table)
    assert approved.status_code == 201, approved.text
    prefix = f"/api/v1/intelligence/engagements/{eid}"
    accounts = (
        await c.get(prefix + "/accounts", headers=h, params={"snapshot_id": snapshot})
    ).json()["items"]
    account = next(a for a in accounts if a["code"] == "100")
    return dict(
        snapshot_id=snapshot, account_ids=[account["id"]], statement_source_id=approved.json()["id"]
    ), reviewer


@pytest.mark.parametrize("country", ["AE", "IN"])
async def test_generic_result_retains_every_original_candidate_and_export(
    environment, monkeypatch, country
):
    c = environment
    double_boundaries(monkeypatch)
    h, options, eid, audit = await context(c, country)
    payload, reviewer = await prepare(c, h, options, eid, audit, country)
    result = await execute(c, h, eid, "reconciliation", payload)
    body = result["body"]
    assert not body["matches"] and body["ambiguity_link_count"] == 14641
    assert result["collection_totals"]["ambiguous"] == 121
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    last_groups = (
        await c.get(prefix + "?ambiguity_offset=100&ambiguity_limit=25", headers=h)
    ).json()["body"]["ambiguous"]
    assert len(last_groups) == 21
    last = last_groups[-1]
    page = await c.get(
        prefix + "/bank-candidates",
        headers=h,
        params=dict(left_id=last["left_id"], offset=100, limit=25),
    )
    assert page.status_code == 200, page.text
    assert page.json()["total"] == 121 and len(page.json()["items"]) == 21
    assert page.json()["left"]["reference"] == "SYN-SHARED"
    assert last["candidate_preview_truncated"] is True
    assert page.json()["items"][-1]["id"] == "row-122"
    assert page.json()["items"][-1]["currency"] == ("AED" if country == "AE" else "INR")
    assert page.json()["items"][-1]["machine_match"] is None
    assert (await decision(c, reviewer[0], eid, result)).status_code == 409
    exported = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    links = list(book["Matching candidates"].values)
    assert len(links) == 14642
    assert links[-1][:4] == (last["left_id"], "row-122", "121", "COMPLETE_RETAINED_CANDIDATES")


@pytest.mark.parametrize("country", ["AE", "IN"])
async def test_acceptance_and_both_sides_of_unmatched_disposition_are_required(
    environment, monkeypatch, country
):
    c = environment
    double_boundaries(monkeypatch)
    h, options, eid, audit = await context(c, country)
    payload, _ = await prepare(c, h, options, eid, audit, country, count=2)
    final = await actor(h, "Firm Administrator", "generic-final-reviewer")
    await grant(c, h, audit, final)
    obj = await execute(c, h, eid, "reconciliation", payload)
    items = [*obj["body"]["unmatched_left"], *obj["body"]["unmatched_right"]]
    assert len(items) == 4 and not obj["body"]["matches"]
    requested = await decision(c, h, eid, obj, "request-evidence", item_id=items[0]["id"])
    assert requested.status_code == 201, requested.text
    obj = requested.json()
    assert (await decision(c, final[0], eid, obj)).status_code == 409
    for item in items[:-1]:
        response = await decision(c, h, eid, obj, "finding", item_id=item["id"])
        assert response.status_code == 201, response.text
        obj = response.json()
    assert (await decision(c, final[0], eid, obj)).status_code == 409
    response = await decision(c, h, eid, obj, "finding", item_id=items[-1]["id"])
    assert response.status_code == 201, response.text
    obj = response.json()
    assert (await decision(c, h, eid, obj)).status_code == 403
    approved = await decision(c, final[0], eid, obj)
    assert approved.status_code == 201, approved.text
    path = f"/api/v1/intelligence/engagements/{eid}/events/{approved.json()['id']}"
    full = (await c.get(path, headers=h)).json()
    summary = full["body"]["review_decision_summary"]
    assert summary["unmatched_transactions"] == summary["unmatched_with_findings"] == 4
    assert full["body"]["unmatched_left"] and full["body"]["unmatched_right"]
    exported = await c.get(path + "/export?format=xlsx", headers=h)
    workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
    rows = list(workbook["Reconciliation decisions"].values)
    assert sum(row[2] == "findings" for row in rows if len(row) > 2) == 4
    assert sum(row[2] == "requests" for row in rows if len(row) > 2) == 1


async def test_proposed_match_cannot_be_approved_until_preparer_records_acceptance(
    environment, monkeypatch
):
    c = environment
    double_boundaries(monkeypatch)
    h, options, eid, audit = await context(c, "AE")
    payload, _ = await prepare(c, h, options, eid, audit, "AE", count=1)
    final = await actor(h, "Firm Administrator", "generic-final-match-reviewer")
    await grant(c, h, audit, final)
    obj = await execute(c, h, eid, "reconciliation", payload)
    assert len(obj["body"]["matches"]) == 1
    assert (await decision(c, final[0], eid, obj)).status_code == 409
    accepted = await decision(
        c, h, eid, obj, "accept", item_id=obj["body"]["matches"][0]["left_ids"][0]
    )
    assert accepted.status_code == 201, accepted.text
    approved = await decision(c, final[0], eid, accepted.json())
    assert approved.status_code == 201, approved.text
