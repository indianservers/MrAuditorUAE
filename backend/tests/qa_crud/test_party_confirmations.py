"""Synthetic positive balance evidence through persisted sources, reviews and exports."""

import hashlib
import io
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, create, grant, setup
from tests.qa_crud.test_country_workflows import (
    csv_bytes,
    decision,
    double_boundaries,
    execute,
    source,
)
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_party_confirmations import followup, nonresponse, population

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required"
)


async def prepare(
    c,
    h,
    eid,
    reviewer,
    currency="AED",
    positions=1,
    status="received",
    with_followup=False,
    side="receivable",
    request_basis="stated-balance",
    include_coverage=False,
):
    base = f"/api/v1/engagements/{eid}/audit"
    template = population(currency)
    account = "AP" if side == "payable" else "AR"
    for records in template.values():
        for row in records:
            row.update(account=account, side=side)
    parties = [f"PARTY-{index:04d}" for index in range(positions)]
    documents = {}
    written = status in {"received", "disputed", "unreliable", "partially-received"}
    for prefix in (
        ("request", "response") if written else () if status == "not-sent" else ("request",)
    ) + (("followup",) if with_followup else ()):
        text = (
            "SYNTHETIC ONLY. Fictional positive balance "
            + prefix
            + " batch; no real client or authenticated sender.\n"
        )
        document_amount = (
            "BLANK"
            if prefix == "request" and request_basis == "blank-positive"
            else "100"
            if prefix == "response" and status == "partially-received"
            else "150"
        )
        text += "\n".join(
            f"SYN-{prefix.upper()}-LINE-{index:04d}|{party}|{account}|{side}|{currency}|2026-06-30|{document_amount}"
            for index, party in enumerate(parties)
        )
        record = await create(
            c, h, base, "evidence", {"category": "Other", "retain_until": "2036-12-31"}
        )
        response = await c.post(
            base + "/evidence-files/" + record["id"],
            headers=h,
            files={"file": (f"SYNTHETIC-party-{prefix}.txt", text.encode(), "text/plain")},
        )
        assert response.status_code == 201, response.text
        documents[prefix] = {**response.json(), "record": record["id"], "expected_text": text}
    items, letters = [], []
    for index, party in enumerate(parties):
        items.extend(
            [
                {**item, "id": item["id"] + f"-{index:04d}", "counterparty": party}
                for item in template["items"]
            ]
        )
        letter = {
            **template["responses"][0],
            "id": f"SYN-LETTER-{index:04d}",
            "counterparty": party,
            "request_reference": f"SYN-REQUEST-LINE-{index:04d}",
            "response_reference": f"SYN-RESPONSE-LINE-{index:04d}",
        }
        if not written:
            letter = nonresponse(letter, status)
        else:
            letter["status"] = status
            if status == "unreliable":
                letter.update(contact_verified=False, independently_received=False)
        if request_basis == "blank-positive":
            letter.update(request_basis=request_basis, scope_balance="150", requested_amount=None)
        if status == "partially-received":
            letter["response_amount"] = "100"
        if with_followup:
            letter = followup(
                letter,
                "partial-response-procedures"
                if status == "partially-received"
                else "exception-investigation"
                if status in {"received", "disputed"}
                else "alternative-procedures",
            )
            letter["followup_reference"] = f"SYN-FOLLOWUP-LINE-{index:04d}"
        for prefix in ("request", "response", "followup"):
            letter[prefix + "_evidence_version_id"] = ""
            letter[prefix + "_source_sha256"] = ""
        letters.append(letter)
    controls = [{**template["controls"][0], "amount": str(150 * positions)}]
    selected, events = {}, {}
    tables = [
        ("items", "audit-open-items", items),
        ("controls", "balance-controls", controls),
        ("responses", "party-confirmations", letters),
    ]
    if include_coverage:
        tables.append(
            (
                "coverage",
                "confirmation-item-coverage",
                [
                    {
                        "id": f"COVERAGE-{index:04d}",
                        "confirmation_id": f"SYN-LETTER-{index:04d}",
                        "item_id": items[index * 3]["id"],
                        "as_of": "2026-06-30",
                        "evidence_reference": f"SYN-RESPONSE-LINE-{index:04d}",
                    }
                    for index in range(positions)
                ],
            )
        )
    for role, table, rows in tables:
        options = {}
        if role == "responses":
            options = {
                "confirmation_request_totals": {}
                if request_basis == "blank-positive"
                else {currency: str(150 * positions)},
                "confirmation_response_totals": {
                    currency: str((100 if status == "partially-received" else 150) * positions)
                }
                if written
                else {},
                **{
                    "confirmation_" + prefix + "_document_version_id": document["id"]
                    for prefix, document in documents.items()
                },
            }
        event = await source(
            c,
            h,
            eid,
            rows,
            "SYNTHETIC party confirmation " + role,
            table,
            {} if role == "coverage" else {currency: str(150 * positions)},
            **options,
        )
        reviewed = await decision(c, reviewer[0], eid, event)
        assert reviewed.status_code == 201, reviewed.text
        events[role] = reviewed.json()
        selected[role] = reviewed.json()["id"]
    return {
        "engagement": eid,
        "documents": documents,
        "events": events,
        "sources": selected,
        "payload": {
            "workflow": "party-balance-confirmations",
            "source_ids": selected,
            "as_of": "2026-06-30",
            "review_through": "2026-07-15",
            "title": "SYNTHETIC positive party balance evidence",
            "rationale": "Positive balances over many invoice/credit items with retained written correspondence and independent GL controls",
        },
    }


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
@pytest.mark.parametrize("side", ["receivable", "payable"])
async def test_party_letter_review_download_export_and_document_staleness(
    environment, monkeypatch, country, currency, side
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "party-source-reviewer")
    final = await actor(h, "Firm Administrator", "party-final-reviewer")
    for person in (reviewer, final):
        await grant(c, h, base, person)
    fixture = await prepare(c, h, eid, reviewer, currency, side=side)
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    party = result["body"]["rows"][-1]
    assert result["body"]["exception_count"] == 0 and party["open_item_count"] == 3
    assert (
        party["confirmation_evidence"]["response_source_sha256"]
        == fixture["documents"]["response"]["sha256"]
    )
    assert any(
        receipt.get("type") == "evidence"
        and receipt.get("id") == fixture["documents"]["response"]["id"]
        and receipt.get("record_id") == fixture["documents"]["response"]["record"]
        for receipt in result["dependencies"]
    )
    assert (await decision(c, h, eid, result)).status_code == 403
    approved = await decision(c, final[0], eid, result)
    assert approved.status_code == 201, approved.text
    result = approved.json()
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/"
    for view, kind in (("parties", "party-confirmation"), ("gl-controls", "control")):
        page = await c.get(prefix + result["id"] + f"?confirmation_view={view}&limit=1", headers=h)
        assert (
            page.status_code == 200
            and page.json()["collection_totals"]["rows"] == 1
            and page.json()["body"]["rows"][0]["kind"] == kind
        )
    document = fixture["documents"]["response"]
    original = await c.get(
        base + f"/evidence-files/{document['record']}/{document['id']}/download", headers=h
    )
    assert original.status_code == 200 and original.content == document["expected_text"].encode()
    assert (
        hashlib.sha256(original.content).hexdigest()
        == party["confirmation_evidence"]["response_source_sha256"]
    )
    exported = await c.get(prefix + result["id"] + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    values = [value for row in book["Rows"].values for value in row]
    assert "PARTY-0000" in values and any("SYN-CREDIT-1-0000" in str(value) for value in values)
    book.close()
    replaced = await c.post(
        base + "/evidence-files/" + document["record"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-party-response-revised.txt",
                b"SYNTHETIC revised response needs fresh scope and review",
                "text/plain",
            )
        },
    )
    assert replaced.status_code == 201, replaced.text
    assert (await c.get(prefix + result["id"], headers=h)).json()["freshness"] == "STALE"
    await execute(c, h, eid, "country-workflow", fixture["payload"], failed="source is stale")


@pytest.mark.parametrize("status", ["pending", "unreliable"])
async def test_nonresponse_requires_bound_alternatives_and_independent_disposition(
    environment, monkeypatch, status
):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, base = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "party-alternative-source-reviewer")
    final = await actor(h, "Firm Administrator", "party-alternative-final-reviewer")
    for person in (reviewer, final):
        await grant(c, h, base, person)
    fixture = await prepare(c, h, eid, reviewer, status=status)
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    assert (
        result["body"]["blocked"]
        and result["body"]["rows"][-1]["state"] == "ALTERNATIVE_PROCEDURES_MISSING"
    )
    accepted = await decision(c, h, eid, result, action="accept", item_id="party-confirmation:1")
    assert accepted.status_code == 201, accepted.text
    assert (await decision(c, final[0], eid, accepted.json())).status_code == 409
    fixture = await prepare(c, h, eid, reviewer, status=status, with_followup=True)
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    assert (
        not result["body"]["blocked"]
        and result["body"]["rows"][-1]["state"] == "ALTERNATIVE_PROCEDURES_REQUIRE_REVIEW"
    )
    assert result["body"]["rows"][-1]["confirmation_status"] == status
    assert result["body"]["rows"][-1]["response_amount"] == (
        "150" if status == "unreliable" else None
    )
    assert (await decision(c, final[0], eid, result)).status_code == 409
    accepted = await decision(c, h, eid, result, action="accept", item_id="party-confirmation:1")
    assert accepted.status_code == 201, accepted.text
    assert (await decision(c, final[0], eid, accepted.json())).status_code == 201


async def test_gl_view_and_all_parties_remain_available_after_first_hundred(
    environment, monkeypatch
):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, base = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "party-pagination-reviewer")
    await grant(c, h, base, reviewer)
    fixture = await prepare(c, h, eid, reviewer, positions=121)
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    last = await c.get(prefix + "?confirmation_view=parties&offset=100&limit=100", headers=h)
    assert last.status_code == 200 and last.json()["collection_totals"]["rows"] == 121
    assert (
        len(last.json()["body"]["rows"]) == 21
        and last.json()["body"]["rows"][-1]["counterparty"] == "PARTY-0120"
    )
    gl = await c.get(prefix + "?confirmation_view=gl-controls", headers=h)
    assert (
        gl.status_code == 200
        and gl.json()["collection_totals"]["rows"] == 1
        and gl.json()["body"]["rows"][0]["subledger_amount"] == "18150"
    )
    exported = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    rows = list(book["Rows"].values)
    records = [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]
    assert len(records) == 122 and records[-1]["counterparty"] == "PARTY-0120"
    links = list(book["Confirmation item links"].values)
    documents = list(book["Confirmation documents"].values)
    assert len(links) == 364 and links[-1][-1] == "SYN-CREDIT-1-0120"
    assert len(documents) == 243 and documents[-1][-1] == "RETAINED_IDENTITY_ONLY"
    book.close()


async def test_unreviewed_assessments_bad_hashes_and_self_evidence_fail_closed(
    environment, monkeypatch
):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, base = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "party-negative-source-reviewer")
    await grant(c, h, base, reviewer)
    fixture = await prepare(c, h, eid, reviewer)
    rows = fixture["events"]["responses"]["body"]["rows"]
    unknown_source = await source(
        c,
        h,
        eid,
        [{**rows[0], "counterparty": "UNKNOWN-PARTY"}],
        "SYNTHETIC unknown party assessment",
        "party-confirmations",
        {"AED": "150"},
        confirmation_request_totals={"AED": "150"},
        confirmation_response_totals={"AED": "150"},
    )
    reviewed_unknown = await decision(c, reviewer[0], eid, unknown_source)
    assert reviewed_unknown.status_code == 201, reviewed_unknown.text
    unknown_result = await execute(
        c,
        h,
        eid,
        "country-workflow",
        {
            **fixture["payload"],
            "source_ids": {**fixture["sources"], "responses": reviewed_unknown.json()["id"]},
        },
    )
    unknown = next(
        row for row in unknown_result["body"]["rows"] if row["state"] == "UNKNOWN_PARTY_POPULATION"
    )
    assert (
        unknown_result["body"]["blocked"]
        and unknown["book_balance"] is None
        and unknown["difference"] is None
        and unknown["response_amount"] == "150"
    )
    event = await source(
        c,
        h,
        eid,
        rows,
        "SYNTHETIC unreviewed assessment",
        "party-confirmations",
        {"AED": "150"},
        confirmation_request_totals={"AED": "150"},
        confirmation_response_totals={"AED": "150"},
    )
    await execute(
        c,
        h,
        eid,
        "country-workflow",
        {**fixture["payload"], "source_ids": {**fixture["sources"], "responses": event["id"]}},
        failed="Independently review open items",
    )
    await source(
        c,
        h,
        eid,
        [{**rows[0], "response_source_sha256": "a" * 64}],
        "SYNTHETIC incorrect response hash",
        "party-confirmations",
        {"AED": "150"},
        confirmation_request_totals={"AED": "150"},
        confirmation_response_totals={"AED": "150"},
        failed="document hash differs",
    )
    await source(
        c,
        h,
        eid,
        rows,
        "SYNTHETIC incorrect response total",
        "party-confirmations",
        {"AED": "150"},
        confirmation_request_totals={"AED": "150"},
        confirmation_response_totals={"AED": "149"},
        failed="written-response currency totals do not agree",
    )
    record = await create(
        c, h, base, "evidence", {"category": "Other", "retain_until": "2036-12-31"}
    )
    uploaded = await c.post(
        base + "/evidence-files/" + record["id"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-self-response-assessment.csv",
                csv_bytes(
                    [{**rows[0], "response_evidence_version_id": "", "response_source_sha256": ""}]
                ),
                "text/csv",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    await execute(
        c,
        h,
        eid,
        "workflow-source",
        {
            "evidence_version_id": uploaded.json()["id"],
            "table_type": "party-confirmations",
            "label": "SYNTHETIC self evidence guard",
            "expected_rows": 1,
            "control_totals": {"AED": "150"},
            "confirmation_request_totals": {"AED": "150"},
            "confirmation_response_totals": {"AED": "150"},
            "confirmation_response_document_version_id": uploaded.json()["id"],
        },
        failed="distinct retained request/response/follow-up documents",
    )


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
@pytest.mark.parametrize("partial_reply", [False, True])
async def test_blank_requests_and_explicit_partial_coverage_persist_and_export(
    environment, monkeypatch, country, currency, partial_reply
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "extension-source-reviewer")
    await grant(c, h, base, reviewer)
    final = await actor(h, "Firm Administrator", "extension-final-reviewer")
    await grant(c, h, base, final)
    fixture = await prepare(
        c,
        h,
        eid,
        reviewer,
        currency,
        status="partially-received" if partial_reply else "received",
        with_followup=partial_reply,
        request_basis="blank-positive",
        include_coverage=partial_reply,
    )
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    body = result["body"]
    party = body["rows"][-1]
    assert not body["blocked"] and party["requested_amount"] is None
    assert party["scope_balance"] == "150"
    assert "|BLANK" in fixture["documents"]["request"]["expected_text"]
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    if partial_reply:
        assert party["difference"] is None
        assert party["covered_book_balance"] == "100"
        assert party["uncovered_book_balance"] == "50"
        assert party["covered_response_difference"] == "0"
        assert len(party["uncovered_item_source_ids"]) == 2
        assert (await decision(c, reviewer[0], eid, result)).status_code == 409
        page = await c.get(
            prefix + "?confirmation_link_offset=1&confirmation_link_limit=1", headers=h
        )
        assert page.status_code == 200, page.text
        assert page.json()["body"]["confirmation_coverage_links"] == []
        assert page.json()["collection_totals"]["confirmation_coverage_links"] == 1
        accepted = await decision(
            c, h, eid, result, action="accept", item_id="party-confirmation:1"
        )
        assert accepted.status_code == 201, accepted.text
        approved = await decision(c, final[0], eid, accepted.json())
        assert approved.status_code == 201, approved.text
    else:
        assert party["difference"] == "0"
    response = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    mappings = list(book["Confirmation coverage links"].values)
    assert len(mappings) == (2 if partial_reply else 1)
    if partial_reply:
        assert mappings[-1][2] == party["covered_item_source_ids"][0]
    book.close()
