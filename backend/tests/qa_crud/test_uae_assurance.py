import os, uuid, json
from pathlib import Path
import pytest
from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import setup, actor, grant, create, approve, action
from tests.unit.test_uae_assurance import fixture
from app.modules.jurisdictions.ae.advanced import Schedule, calculate
from app.modules.enterprise_ai.security import serial

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Disposable database only"
)


async def test_all_assurance_cases_approved_execution_and_staleness(
    environment,
):
    c = environment
    h, options, eid, base = await setup(c)
    reviewer = await actor(h, "Audit Manager", "advanced-reviewer")
    await grant(c, h, base, reviewer)
    partner = await actor(h, "Audit Partner", "advanced-partner")
    await grant(c, h, base, partner)
    # Synthetic reviewers are explicitly qualified for all specialist read scopes.
    from app.db.session import Session, tenant_context
    from app.db.models.platform import RolePermission, Permission
    from sqlalchemy import select

    async with Session() as db, db.begin():
        await tenant_context(db, h["X-Tenant-ID"])
        for rid in [reviewer[2], partner[2]]:
            for code in [
                "accounting.read",
                "vat.read",
                "tax.profile.read",
                "tp.read",
                "dmtt.read",
                "ifrs.read",
            ]:
                pid = await db.scalar(select(Permission.id).where(Permission.code == code))
                present = await db.scalar(
                    select(RolePermission.role_id).where(
                        RolePermission.role_id == uuid.UUID(rid),
                        RolePermission.permission_id == pid,
                    )
                )
                if not present:
                    db.add(RolePermission(role_id=uuid.UUID(rid), permission_id=pid))
    proof = await create(
        c,
        h,
        base,
        "working-papers",
        dict(
            folder="SYNTHETIC",
            sections=[
                dict(
                    heading="Evidence",
                    text="SYNTHETIC approved policy and population; test fixture only",
                )
            ],
            conclusion="Reviewed synthetic inputs only",
        ),
    )
    proof = await approve(c, h, reviewer[0], base, proof)
    proof = await action(c, partner[0], base, proof, "approve")
    route = f"/api/v1/ai/engagements/{eid}/use-cases"
    listing = await c.get(route, headers=h)
    assert listing.status_code == 200
    assert len(listing.json()["items"]) == 200
    newproof = await create(
        c,
        h,
        base,
        "working-papers",
        dict(
            folder="SYNTHETIC",
            sections=[dict(heading="New rule", text="Synthetic separately preserved new source")],
            conclusion="Reviewed",
        ),
    )
    newproof = await approve(c, h, reviewer[0], base, newproof)
    newproof = await action(c, partner[0], base, newproof, "approve")
    receipts = []
    first = None
    for i in range(151, 201):
        data = fixture(i)
        data["policy"]["source_id"] = proof["id"]
        for row in data["rows"]:
            row["evidence_id"] = proof["id"]
            for key in row:
                if key.endswith("_source_id"):
                    row[key] = newproof["id"] if key == "new_source_id" else proof["id"]
            if i == 193:
                row["supporting_sources"] = [proof["id"]]
        body = dict(
            folder="SYNTHETIC",
            sections=[dict(heading="Schedule", text=f"Synthetic case {i}")],
            conclusion="Awaiting review",
            sources=[dict(kind="record", id=proof["id"])]
            + ([dict(kind="record", id=newproof["id"])] if i == 151 else []),
            advanced_input=data,
        )
        paper = await create(c, h, base, "working-papers", body)
        command = dict(
            action="advanced-review",
            record_id=paper["id"],
            idempotency_key=f"synthetic-assurance-{i}",
        )
        denied = await c.post(f"{route}/{i}/actions", headers=h, json=command)
        assert denied.status_code == 409, (i, denied.text)
        if i == 151:
            listing = await c.get(route + "/151/inputs", headers=h)
            assert next(r for r in listing.json()["items"] if r["id"] == paper["id"])[
                "draft_preview_id"
            ]
        submitted = await action(c, h, base, paper, "submit")
        await action(c, h, base, submitted, "approve", 403)
        paper = await action(c, reviewer[0], base, submitted, "approve")
        paper = await action(c, partner[0], base, paper, "approve")
        result = await c.post(f"{route}/{i}/actions", headers=h, json=command)
        assert result.status_code == 201, (i, result.text)
        value = result.json()
        assert value["body"]["output"]["results"] == serial(
            calculate(Schedule.model_validate(data))
        )
        repeated = await c.post(f"{route}/{i}/actions", headers=h, json=command)
        assert repeated.status_code == 201 and repeated.json()["id"] == value["id"]
        if i == 151:
            first = value
            impact = await c.post(
                f"{route}/151/actions",
                headers=h,
                json=dict(
                    action="rule-impact", record_id=paper["id"], idempotency_key="rule-impact-test"
                ),
            )
            assert impact.status_code == 201, impact.text
            assert any(
                n["id"] == paper["id"] for n in impact.json()["body"]["output"]["affected_items"]
            )
        if i == 190:
            incomplete = {
                **body,
                "advanced_input": {**data, "rows": [{**data["rows"][0], "completed": False}]},
            }
            blocked = await create(c, h, base, "working-papers", incomplete)
            completion = await c.get(base + "/completion-checks", headers=h)
            assert completion.status_code == 200
            assert (
                next(
                    check
                    for check in completion.json()["checks"]
                    if check["code"] == "quality_consultation_" + blocked["id"]
                )["passed"]
                is False
            )
        if i == 198:
            assert paper["body"]["factual_findings_only"] is True
            assert "no audit opinion" in paper["body"]["conclusion"]
        receipts.append(
            dict(id=i, status="approved execution and replay verified", event_id=value["id"])
        )
    proof = await action(c, reviewer[0], base, proof, "reopen")
    response = await c.get(
        f"/api/v1/intelligence/engagements/{eid}/events/{first['id']}", headers=h
    )
    assert response.status_code == 200 and response.json()["freshness"] == "STALE"
    destination = (
        Path(__file__).resolve().parents[3]
        / "docs/mr-auditor-business/evidence/assurance-case-results.json"
    )
    destination.write_text(json.dumps(receipts, indent=2), encoding="utf-8")


async def test_original_provenance_respects_quarantine_and_source_replacement(
    environment, monkeypatch, tmp_path
):
    from tests.unit.test_workbook_integrity import workbook
    from app.core.config import settings
    from app.db.session import Session, tenant_context
    from app.modules.audit_execution.models import EvidenceVersion
    import hashlib

    c = environment
    h, _, eid, base = await setup(c)
    monkeypatch.setattr(settings(), "source_storage_path", str(tmp_path))
    reviewer = await actor(h, "Audit Manager", "workbook-reviewer")
    await grant(c, h, base, reviewer)
    proof = await create(
        c, h, base, "evidence", dict(category="Workbook", retain_until="2036-12-31")
    )
    data = workbook()
    uploaded = await c.post(
        base + "/evidence-files/" + proof["id"],
        headers=h,
        files={
            "file": (
                "synthetic.xlsx",
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    route = f"/api/v1/ai/engagements/{eid}/use-cases/192/actions"
    command = dict(
        action="inspect-provenance",
        record_id=proof["id"],
        idempotency_key="synthetic-workbook-scan",
    )
    assert (await c.post(route, headers=h, json=command)).status_code == 409
    # Explicit disposable scanner boundary: real versions remain immutable.
    import io
    from openpyxl import load_workbook

    monkeypatch.setattr("app.api.v1.audit_evidence.scan", lambda stream: "CLEAN")
    book = load_workbook(io.BytesIO(data))
    book.active["A1"] = 12
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    data = stream.getvalue()
    clean = await c.post(
        base + "/evidence-files/" + proof["id"],
        headers=h,
        files={
            "file": (
                "synthetic-clean.xlsx",
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert clean.status_code == 201 and clean.json()["scan_status"] == "CLEAN"
    proof = (await c.get(base + "/records/" + proof["id"], headers=h)).json()
    proof = await approve(c, h, reviewer[0], base, proof)
    result = await c.post(route, headers=h, json=command)
    assert result.status_code == 201, result.text
    out = result.json()["body"]["output"]
    assert out["observed_sha256"] == hashlib.sha256(data).hexdigest()
    assert out["matches_registered_hash"] is True and out["forgery_conclusion"] is None
    assert out["analysis_mode"] == "rules_baseline"
    duplicate = await create(
        c, h, base, "evidence", dict(category="Duplicate control", retain_until="2036-12-31")
    )
    uploaded = await c.post(
        base + "/evidence-files/" + duplicate["id"],
        headers=h,
        files={
            "file": (
                "duplicate.xlsx",
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 201
    duplicate = (await c.get(base + "/records/" + duplicate["id"], headers=h)).json()
    duplicate = await approve(c, h, reviewer[0], base, duplicate)
    duplicate_check = await c.post(
        route, headers=h, json={**command, "idempotency_key": "provenance-duplicate-corpus"}
    )
    assert duplicate_check.status_code == 201, duplicate_check.text
    assert any(
        r["record_id"] == duplicate["id"]
        for r in duplicate_check.json()["body"]["output"]["duplicate_candidates"]
    )
    # A distinct new version is quarantined and invalidates the retained approved-file scan.
    import io
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data))
    book.active["A1"] = 999
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    proof = await action(c, reviewer[0], base, proof, "reopen")
    monkeypatch.setattr("app.api.v1.audit_evidence.scan", lambda stream: "QUARANTINED")
    replacement = await c.post(
        base + "/evidence-files/" + proof["id"],
        headers=h,
        files={
            "file": (
                "synthetic-v2.xlsx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert replacement.status_code == 201, replacement.text
    detail = await c.get(
        f"/api/v1/intelligence/engagements/{eid}/events/{result.json()['id']}", headers=h
    )
    assert detail.status_code == 200 and detail.json()["freshness"] == "STALE"
