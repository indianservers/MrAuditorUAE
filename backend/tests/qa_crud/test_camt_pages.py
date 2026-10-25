"""Retained multipart statements through API, independent review and BRS.

Scanner and broker use explicit doubles; jobs execute through the real processor.
"""

import io
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, create, grant, setup
from tests.qa_crud.test_country_workflows import (
    bank_payload,
    decision,
    double_boundaries,
    execute,
    source,
)
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_camt_pages import packet

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable database required"
)


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
async def test_retained_multipart_statement_review_reconciliation_and_export(
    environment, monkeypatch, country, currency
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "IN":
        h, eid, base = await india_context(c)
    else:
        h, _, eid, base = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "multipart-statement-reviewer")
    await grant(c, h, base, reviewer)
    evidence = await create(
        c, h, base, "evidence", {"category": "Other", "retain_until": "2036-12-31"}
    )
    content = packet(currency, middle=True)
    upload = await c.post(
        base + "/evidence-files/" + evidence["id"],
        headers=h,
        files={"file": ("SYNTHETIC-three-page.xml", content, "application/xml")},
    )
    assert upload.status_code == 201, upload.text
    native = await execute(
        c,
        h,
        eid,
        "workflow-source",
        {
            "evidence_version_id": upload.json()["id"],
            "file_format": "camt053",
            "table_type": "transactions",
            "label": "SYNTHETIC complete bank message chain",
            "expected_rows": 2,
            "control_totals": {currency: "95"},
            "account_identifier": "BANK-1",
            "statement_identifier": "SYNTHETIC-STMT",
        },
    )
    assert native["body"]["parser_version"] == "camt053-2"
    assert native["body"]["native_metadata"]["page_count"] == 3
    assert [r["source_statement_page"] for r in native["body"]["rows"]] == [1, 3]
    original = await c.get(
        base + "/evidence-files/" + evidence["id"] + "/" + upload.json()["id"] + "/download",
        headers=h,
    )
    assert original.status_code == 200 and original.content == content
    assert (await decision(c, h, eid, native)).status_code == 403
    reviewed = await decision(c, reviewer[0], eid, native)
    assert reviewed.status_code == 201, reviewed.text
    book = await source(
        c,
        h,
        eid,
        [
            {
                "id": "B1",
                "date": "2026-06-01",
                "reference": "PAY-1",
                "amount": "100",
                "currency": currency,
                "account": "BANK-1",
                "counterparty": "",
            },
            {
                "id": "B2",
                "date": "2026-06-02",
                "reference": "FEE",
                "amount": "-5",
                "currency": currency,
                "account": "BANK-1",
                "counterparty": "",
            },
        ],
        "SYNTHETIC independent cashbook",
        totals={currency: "95"},
    )
    book_review = await decision(c, reviewer[0], eid, book)
    assert book_review.status_code == 201, book_review.text
    result = await execute(
        c,
        h,
        eid,
        "country-workflow",
        bank_payload(
            book_review.json(),
            reviewed.json(),
            currency=currency,
            book_closing="95",
            statement_closing="95",
        ),
    )
    assert result["body"]["exception_count"] == 0
    approved = await decision(c, reviewer[0], eid, result)
    assert approved.status_code == 201, approved.text
    exported = await c.get(
        f"/api/v1/intelligence/engagements/{eid}/events/{reviewed.json()['id']}/export?format=xlsx",
        headers=h,
    )
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content))
    assert workbook["CAMT page controls"].max_row == 4
    replacement = await c.post(
        base + "/evidence-files/" + evidence["id"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-three-page.xml",
                content.replace(b"Customer", b"Revised customer"),
                "application/xml",
            )
        },
    )
    assert replacement.status_code == 201, replacement.text
    for event in (reviewed.json(), approved.json()):
        current = await c.get(
            f"/api/v1/intelligence/engagements/{eid}/events/{event['id']}",
            headers=h,
        )
        assert current.status_code == 200 and current.json()["freshness"] == "STALE"
