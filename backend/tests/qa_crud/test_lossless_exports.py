"""Large reviewed UAE/India allocations retain complete evidence in downloads."""

import hashlib
import io
import json
import os
import uuid

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.db.models.platform import AuditEvent
from app.db.session import Session, tenant_context
from app.modules.intelligence import xlsx_values
from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant, setup
from tests.qa_crud.test_country_workflows import (
    csv_bytes,
    decision,
    double_boundaries,
    execute,
    source,
)
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.country_workflow_cases import obligation, transaction
from tests.unit.test_dues_allocations import parameters
from tests.unit.test_xlsx_values import originals

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Disposable DB required"
)


async def retained_allocations(c, h, eid, reviewer, country):
    currency = "AED" if country == "AE" else "INR"
    expected = obligation(
        id="E1",
        amount="500",
        currency=currency,
        due_date="2026-06-30",
        period="2026-05",
        component="SALARY" if country == "AE" else "PF:EMPLOYEE",
        payment_date=None,
        subject="SYNTHETIC-EXPORT-SUBJECT",
    )
    bank = [
        transaction(
            id=f"B{i:04}",
            reference=f"SYNTHETIC-PAYMENT-{i:04}",
            currency=currency,
            amount="-1",
            date="2026-06-01",
        )
        for i in range(500)
    ]
    allocations = [
        dict(
            id=f"A{i:04}",
            expected_id="E1",
            reported_id="R1",
            transaction_id=row["id"],
            currency=currency,
            amount="1",
            evidence_reference="SYNTHETIC-ALLOCATION-SCHEDULE",
            rationale="Fictional independently assessed split payment",
        )
        for i, row in enumerate(bank)
    ]
    values = {
        "expected": [expected],
        "reported": [{**expected, "id": "R1"}],
        "bank": bank,
        "allocations": allocations,
    }
    ids, retained = {}, {}
    for role, table, total in [
        ("expected", "obligations", "500"),
        ("reported", "obligations", "500"),
        ("bank", "transactions", "-500"),
        ("allocations", "obligation-allocations", "500"),
    ]:
        prepared = await source(
            c, h, eid, values[role], "SYNTHETIC large export " + role, table, {currency: total}
        )
        retained[role] = prepared
        reviewed = await decision(c, reviewer[0], eid, prepared)
        assert reviewed.status_code == 201, reviewed.text
        ids[role] = reviewed.json()["id"]
    result = await execute(
        c,
        h,
        eid,
        "country-workflow",
        {
            **parameters(),
            "source_ids": ids,
            "currency": currency,
            "obligation_month_end": "2026-05",
        },
    )
    assert result["body"]["exception_count"] == 0
    approved = await decision(c, reviewer[0], eid, result)
    assert approved.status_code == 201, approved.text
    return approved.json(), retained, values


@pytest.mark.parametrize("country", ["AE", "IN"])
async def test_complete_large_export_stale_history_and_explicit_limits(
    environment, monkeypatch, country
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "lossless-export-reviewer")
    await grant(c, h, base, reviewer)
    result, retained, values = await retained_allocations(c, h, eid, reviewer, country)
    path = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}/export?format=xlsx"

    async def download(expected_freshness):
        response = await c.get(path, headers=h)
        assert response.status_code == 200, (
            response.text[:200] if response.status_code != 200 else ""
        )
        assert response.headers["cache-control"] == "no-store"
        assert response.headers["x-export-sha256"] == hashlib.sha256(response.content).hexdigest()
        async with Session() as db:
            await tenant_context(db, h["X-Tenant-ID"])
            receipt = await db.scalar(
                select(AuditEvent).where(
                    AuditEvent.tenant_id == uuid.UUID(h["X-Tenant-ID"]),
                    AuditEvent.resource_id == uuid.UUID(result["id"]),
                    AuditEvent.action == "intelligence.exported",
                    AuditEvent.details["content_sha256"].as_string()
                    == response.headers["x-export-sha256"],
                )
            )
            assert receipt is not None and receipt.details["content_bytes"] == len(response.content)
        book = load_workbook(io.BytesIO(response.content), read_only=True)
        metadata = dict(book["Source and review"].values)
        assert metadata["Source freshness at export"] == expected_freshness
        restored = originals(book)
        header = next(
            i
            for i, row in enumerate(book["Source and review"].values, 1)
            if row[0] == "bank_consumption_manifest"
        )
        assert (
            json.loads(restored[("Source and review", f"B{header}")])
            == result["body"]["bank_consumption_manifest"]
        )
        assert len(list(book["Bank cash claims"].values)) == 501
        rows = list(book["Rows"].values)
        assert len(rows) == 503  # Two obligations plus 500 bank movements and header.
        link_column = rows[0].index("payment_allocations")
        first = dict(zip(rows[0], rows[1], strict=True))
        assert first["payment_allocations"].startswith("Full value in Long values:")
        from openpyxl.utils import get_column_letter

        links = json.loads(restored[("Rows", f"{get_column_letter(link_column + 1)}2")])
        assert len(links) == 500 and sum(int(row["amount"]) for row in links) == 500

    await download("CURRENT")
    with monkeypatch.context() as bounded:
        bounded.setattr(xlsx_values, "MAX_LONG_BYTES", 10)
        rejected = await c.get(path, headers=h)
        assert rejected.status_code == 422 and "continuation limit" in rejected.text
        assert "content-disposition" not in rejected.headers
    receipt = next(
        item for item in retained["expected"]["dependencies"] if item["type"] == "evidence"
    )
    changed = [{**values["expected"][0], "amount": "501"}]
    replaced = await c.post(
        base + "/evidence-files/" + receipt["record_id"],
        headers=h,
        files={"file": ("SYNTHETIC changed large export.csv", csv_bytes(changed), "text/csv")},
    )
    assert replaced.status_code == 201, replaced.text
    await download("STALE")
