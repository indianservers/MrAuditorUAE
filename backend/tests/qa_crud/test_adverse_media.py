"""Real retained files, HTTP jobs, source review and stale-media dependencies."""

import csv
import hashlib
import io
import os

import pytest
from openpyxl import load_workbook

from app.modules.country_workflows.adverse_media import ROLES, ROW_TYPES
from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, create, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_adverse_media import packet

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable database required"
)


async def prepare(c, h, eid, reviewer, country="AE", case="collision", *, reviewed=True, count=1):
    tables = packet(country, case)
    audit = f"/api/v1/engagements/{eid}/audit"
    documents = {}
    for role in ("subject", "query", "article", "policy"):
        content = (
            f"SYNTHETIC ONLY {country} {case}: retained {role} assessment. Fictional subject V1, query Q1, {count} article observations. No real person, provider, registry verification or media clearance.\n"
        ).encode()
        record = await create(
            c, h, audit, "evidence", dict(category="Other", retain_until="2036-12-31")
        )
        response = await c.post(
            audit + "/evidence-files/" + record["id"],
            headers=h,
            files={"file": (f"SYNTHETIC-media-{role}.txt", content, "text/plain")},
        )
        assert response.status_code == 201, response.text
        documents[role] = dict(response.json(), record=record["id"], bytes=content)
    for role, doc in dict(
        subjects="subject", searches="query", mentions="article", policy="policy"
    ).items():
        for row in tables[role]:
            row.update(
                source_evidence_version_id=documents[doc]["id"],
                source_sha256=documents[doc]["sha256"],
            )
            if row.get("identity_evidence_version_id"):
                row.update(
                    identity_evidence_version_id=documents["subject"]["id"],
                    identity_source_sha256=documents["subject"]["sha256"],
                )
    if count > 1:
        tables["mentions"] = [
            dict(
                tables["mentions"][0],
                id=f"M{i:04}",
                article_id=f"ARTICLE-{i:04}",
                url=f"https://example.invalid/synthetic/article-{i:04}",
            )
            for i in range(count)
        ]
        tables["searches"][0]["expected_mentions"] = count
    events = {}
    for role, kind in ROLES.items():
        rows = tables[role]
        if rows:
            event = await source(
                c, h, eid, rows, f"SYNTHETIC {country} media {case} {role}", kind, {}
            )
        else:
            stream = io.StringIO(newline="")
            csv.writer(stream).writerow(list(ROW_TYPES[kind].model_fields))
            record = await create(
                c, h, audit, "evidence", dict(category="Other", retain_until="2036-12-31")
            )
            uploaded = await c.post(
                audit + "/evidence-files/" + record["id"],
                headers=h,
                files={
                    "file": (
                        "SYNTHETIC-media-no-mentions.csv",
                        stream.getvalue().encode(),
                        "text/csv",
                    )
                },
            )
            assert uploaded.status_code == 201, uploaded.text
            event = await execute(
                c,
                h,
                eid,
                "workflow-source",
                dict(
                    evidence_version_id=uploaded.json()["id"],
                    table_type=kind,
                    label=f"SYNTHETIC {case} mention population",
                    expected_rows=0,
                    control_totals={},
                    nil_reason="Synthetic query receipt declares no returned mentions; status is separately assessed",
                ),
            )
        if reviewed:
            response = await decision(c, reviewer[0], eid, event)
            assert response.status_code == 201, response.text
            event = response.json()
        events[role] = event
    return dict(
        events=events,
        documents=documents,
        payload=dict(
            workflow="adverse-media-review",
            source_ids={role: event["id"] for role, event in events.items()},
            as_of="2026-06-30",
            title=f"SYNTHETIC {country} media {case}",
            rationale="Independent declared search coverage and retained human identity assessment",
        ),
    )


@pytest.mark.parametrize("country", ["AE", "IN"])
@pytest.mark.parametrize("case", ["collision", "nil", "same-subject", "unresolved", "unavailable"])
async def test_media_review_evidence_scope_and_approval(environment, monkeypatch, country, case):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, audit = await setup(c)
    else:
        h, eid, audit = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "media-source-reviewer")
    final = await actor(h, "Firm Administrator", "media-final-reviewer")
    for person in (reviewer, final):
        await grant(c, h, audit, person)
    fixture = await prepare(c, h, eid, reviewer, country, case, reviewed=False)
    await execute(
        c, h, eid, "country-workflow", fixture["payload"], failed="Independently review all media"
    )
    for role, event in fixture["events"].items():
        assert (await decision(c, h, eid, event)).status_code == 403
        response = await decision(c, reviewer[0], eid, event)
        assert response.status_code == 201, response.text
        fixture["payload"]["source_ids"][role] = response.json()["id"]
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    assert result["body"]["media_document_binding"] == "RETAINED_DOCUMENT_BYTES_VERIFIED"
    assert result["body"]["blocked"] == (case in {"unresolved", "unavailable"})
    assert result["body"]["automatic_rejection"] is False
    assert (await decision(c, h, eid, result)).status_code == 403
    if case in {"same-subject", "unresolved", "unavailable"}:
        assert (await decision(c, final[0], eid, result)).status_code == 409
        for row in result["body"]["rows"]:
            if row["state"] != "MATCHED":
                response = await decision(c, h, eid, result, action="accept", item_id=row["id"])
                assert response.status_code == 201, response.text
                result = response.json()
    approved = await decision(c, final[0], eid, result)
    assert approved.status_code == (409 if result["body"]["blocked"] else 201), approved.text
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    export = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert export.status_code == 200, export.text if export.status_code != 200 else ""
    workbook = load_workbook(io.BytesIO(export.content), read_only=True)
    receipts = result["body"]["retained_media_documents"]
    assert len(list(workbook["Media documents"].values)) == len(receipts) + 1
    for receipt in receipts:
        response = await c.get(
            audit + f"/evidence-files/{receipt['record_id']}/{receipt['id']}/download", headers=h
        )
        assert response.status_code == 200
        assert hashlib.sha256(response.content).hexdigest() == receipt["sha256"]
    query = fixture["documents"]["query"]
    changed = await c.post(
        audit + "/evidence-files/" + query["record"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-revised-search.txt",
                b"SYNTHETIC changed search coverage; reassessment required",
                "text/plain",
            )
        },
    )
    assert changed.status_code == 201, changed.text
    assert (await c.get(prefix, headers=h)).json()["freshness"] != "CURRENT"


async def test_media_paging_preserves_every_article_and_rationale_in_export(
    environment, monkeypatch
):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, audit = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "media-page-reviewer")
    await grant(c, h, audit, reviewer)
    fixture = await prepare(c, h, eid, reviewer, count=121)
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    assert result["collection_totals"]["rows"] == 124
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    second = await c.get(prefix + "?offset=100&limit=100", headers=h)
    assert second.status_code == 200, second.text
    assert len(second.json()["body"]["rows"]) == 24
    export = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert export.status_code == 200
    workbook = load_workbook(io.BytesIO(export.content), read_only=True)
    rows = list(workbook["Rows"].values)
    assert len(rows) == 125
    article = rows[0].index("article_id")
    assert any(row[article] == "ARTICLE-0120" for row in rows[1:])
