"""Original native quarters, retained independent mapping, complete export and stale lineage."""

import io
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute
from tests.qa_crud.test_india_ecr_source import india_context
from tests.qa_crud.test_india_withholding_native import upload
from tests.unit.test_india_24q_native import controls24, options24, original24
from tests.unit.test_india_salary_annual import options_q4, q4_fixture
from tests.unit.test_india_withholding_native import edit

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required"
)


async def native(c, h, reviewer, eid, base, quarter):
    if quarter == "Q4":
        raw, controls, annual = q4_fixture()
        total = "1000"
    else:
        month = {"Q1": "04", "Q2": "07", "Q3": "10"}[quarter]
        total = "10000" if quarter == "Q1" else "500"
        raw, controls = original24(), controls24()
        for line, field, value in [
            (1, 5, "01042025"),
            (2, 16, "202526"),
            (2, 17, "202425"),
            (2, 18, quarter),
            (2, 47, total + ".00"),
            (3, 22, total + ".00"),
            (3, 27, total + ".00"),
            (3, 29, total + ".00"),
            (3, 30, total + ".00"),
            (3, 33, total + ".00"),
            (3, 18, "14" + month + "2024"),
            (4, 14, total + ".00"),
            (4, 17, total + ".00"),
            (4, 19, total + ".00"),
            (4, 23, "16" + month + "2024"),
            (4, 24, "16" + month + "2024"),
            (4, 25, "14" + month + "2024"),
        ]:
            raw = edit(raw, line, field, value)
        controls.update(
            financial_year="202425",
            quarter=quarter,
            income_tax=total,
            allocated_deposit=total,
            challan_deposit=total,
        )
        controls["groups"][0].update(
            period="2024-" + month, transaction_date="2024-" + month + "-16"
        )
    _, file = await upload(c, h, base, f"SYNTHETIC-{quarter}-salary.txt", raw)
    record, document = await upload(
        c,
        h,
        base,
        f"SYNTHETIC-{quarter}-crosswalk.txt",
        f"SYNTHETIC separately assessed {quarter} native party crosswalk".encode(),
    )
    controls["crosswalk_document_version_id"] = document["id"]
    payload = (
        options_q4(controls, annual, evidence_version_id=file["id"])
        if quarter == "Q4"
        else options24(
            evidence_version_id=file["id"],
            native_withholding_controls=controls,
            control_totals={"INR": total},
        )
    ).model_dump(mode="json")
    source = await execute(c, h, eid, "workflow-source", payload)
    reviewed = await decision(c, reviewer[0], eid, source)
    assert reviewed.status_code == 201, reviewed.text
    return reviewed.json(), record, document


def payload(sources, document):
    return {
        "label": "SYNTHETIC four-quarter current-employer comparison",
        "quarter_source_ids": {
            q: sources[q]["id"] if q in sources else None for q in ("Q1", "Q2", "Q3", "Q4")
        },
        "nil_quarter_assessments": {
            q: "SYNTHETIC independently assessed nil employer party population, not filing evidence"
            for q in ("Q1", "Q2", "Q3")
            if q not in sources
        },
        "crosswalk_document_version_id": document["id"],
        "crosswalk_sha256": document["sha256"],
        "assessment": "SYNTHETIC independent full employee crosswalk and all-quarter controls",
        "employees": [
            {
                "annual_key": "SD/1/1",
                "parties": [{"quarter": q, "party_key": "1/1/1"} for q in sources],
                "identity_status": "assessed_same_employee",
                "assessment": "SYNTHETIC independently assessed employee identity across original source records",
            }
        ],
    }


@pytest.mark.parametrize("replacement", ["annual-crosswalk", "quarter-crosswalk"])
async def test_complete_native_four_quarter_review_export_and_dependency_replacement(
    environment, monkeypatch, replacement
):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await india_context(c, year=2024)
    reviewer = await actor(h, "Firm Administrator", "salary-annual-comparison-reviewer")
    await grant(c, h, base, reviewer)
    sources, quarter_records = {}, {}
    for q in ("Q1", "Q2", "Q3", "Q4"):
        sources[q], quarter_records[q], _ = await native(c, h, reviewer, eid, base, q)
    record, document = await upload(
        c,
        h,
        base,
        "SYNTHETIC-independent-annual-mapping.txt",
        b"SYNTHETIC independently specified four-quarter employee identity mapping and complete annual population",
    )
    prefix = f"/api/v1/intelligence/engagements/{eid}"
    response = await c.post(
        prefix + "/salary-annual-comparison", headers=h, json=payload(sources, document)
    )
    assert response.status_code == 201, response.text
    result = response.json()
    assert result["body"]["rows"][0]["comparison_state"] == "AGREES"
    assert result["body"]["rows"][0]["quarterly_tax"] == "12000.00"
    assert result["body"]["exception_count"] == 0
    assert (await decision(c, h, eid, result)).status_code == 403
    response = await decision(c, reviewer[0], eid, result)
    assert response.status_code == 201, response.text
    reviewed = response.json()
    path = prefix + "/events/" + reviewed["id"]
    paged = await c.get(path + "?salary_link_offset=3&salary_link_limit=1", headers=h)
    assert paged.status_code == 200, paged.text
    assert paged.json()["collection_totals"]["salary_quarterly_links"] == 4
    assert paged.json()["body"]["salary_quarterly_links"][0]["quarter"] == "Q4"
    exported = await c.get(path + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
    scope = dict(list(workbook["Source and review"].values))
    assert scope["tan"] == "MUMN77777A" and scope["financial_year"] == "202425"
    links = list(workbook["Salary quarter links"].values)
    assert len(links) == 5 and "identity_assessment" in links[0]
    assert "Unassigned salary parties" in workbook.sheetnames
    workbook.close()
    changed = record if replacement == "annual-crosswalk" else quarter_records["Q2"]
    response = await c.post(
        base + "/evidence-files/" + changed["id"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-revised-crosswalk.txt",
                b"SYNTHETIC changed employee identity mapping requires independent reassessment",
                "text/plain",
            )
        },
    )
    assert response.status_code == 201, response.text
    stale = await c.get(path, headers=h)
    assert stale.status_code == 200 and stale.json()["freshness"] == "STALE"
    assert stale.json()["body"] == reviewed["body"]


@pytest.mark.parametrize(
    "failure", ["same-crosswalk", "cloned-crosswalk", "quarantined-crosswalk", "short-period"]
)
async def test_annual_mapping_and_period_cannot_be_substituted(environment, monkeypatch, failure):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await india_context(c, year=2026 if failure == "short-period" else 2024)
    reviewer = await actor(h, "Firm Administrator", "salary-quarter-reviewer")
    await grant(c, h, base, reviewer)
    source, _, original_doc = await native(c, h, reviewer, eid, base, "Q4")
    document = original_doc
    if failure != "same-crosswalk":
        if failure == "quarantined-crosswalk":
            double_boundaries(monkeypatch, scan="QUARANTINED")
        _, document = await upload(
            c,
            h,
            base,
            "SYNTHETIC-annual-mapping.txt",
            b"SYNTHETIC separately assessed Q4 native party crosswalk"
            if failure == "cloned-crosswalk"
            else b"SYNTHETIC independent annual employee and nil-quarter mapping",
        )
    response = await c.post(
        f"/api/v1/intelligence/engagements/{eid}/salary-annual-comparison",
        headers=h,
        json=payload({"Q4": source}, document),
    )
    assert response.status_code in {409, 422}, response.text
    expected = {
        "same-crosswalk": "separate",
        "cloned-crosswalk": "cloned",
        "quarantined-crosswalk": "quarantined",
        "short-period": "complete assessed financial year",
    }[failure]
    assert expected in response.text.lower()
