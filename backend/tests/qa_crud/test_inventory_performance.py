"""Actual committed current/prior ledgers, controls, review, export and stale history."""

import io
import os
from copy import deepcopy
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase2.test_pipeline import committed
from tests.phase4.test_workflows import actor, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries
from tests.qa_crud.test_india_ecr_source import india_context

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required"
)


def ledger(year, inventory, revenue, costs):
    lines = [
        "Account Code,Account Name,Journal Number,Posting Date,Debit,Credit,Source ID,Opening,Manual,Narration"
    ]
    entries = [
        ("100", "Inventory", "OPEN", inventory, 0, True),
        ("300", "Equity", "OPEN", 0, inventory, True),
        ("400", "Revenue", "OPEN", 0, 50, True),
        ("300", "Equity", "OPEN", 50, 0, True),
        ("500", "Cost of sales", "OPEN", 100, 0, True),
        ("300", "Equity", "OPEN", 0, 100, True),
        ("400", "Revenue", "SALE", 0, revenue, False),
        ("200", "Cash", "SALE", revenue, 0, False),
        ("500", "Cost of sales", "COST", costs, 0, False),
        ("200", "Cash", "COST", 0, costs, False),
        ("100", "Inventory", "STOCK", 20, 0, False),
        ("300", "Equity", "STOCK", 0, 20, False),
    ]
    for index, (code, name, ref, debit, credit, opening) in enumerate(entries):
        lines.append(
            f"{code},{name},{ref},{year}-06-01,{debit},{credit},SYN-{index},{str(opening).lower()},false,SYNTHETIC {ref}"
        )
    return ("\n".join(lines) + "\n").encode()


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
async def test_committed_stock_gp_comparison_controls_review_export_and_prior_supersession(
    environment, monkeypatch, country, currency
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, options, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
        eng = (await c.get(f"/api/v1/engagements/{eid}", headers=h)).json()
        entity = (await c.get(f"/api/v1/entities/{eng['entity_id']}", headers=h)).json()
        options = {
            "entity_id": eng["entity_id"],
            "period_id": eng["period_id"],
            "client_id": entity["client_id"],
        }
    prior_period = await c.post(
        "/api/v1/financial-periods",
        headers=h,
        json={
            "entity_id": options["entity_id"],
            "start_date": "2025-01-01" if country == "AE" else "2025-04-01",
            "end_date": "2025-12-31" if country == "AE" else "2026-03-31",
        },
    )
    assert prior_period.status_code == 201, prior_period.text
    current_options = {
        **options,
        "expected_rows": 12,
        "expected_debit": "2070",
        "expected_credit": "2070",
        "source_complete": True,
        "currency": currency,
    }
    _, current = await committed(c, h, current_options, ledger(2026, 100, 1000, 800))
    prior_options = {
        **options,
        "period_id": prior_period.json()["id"],
        "expected_rows": 12,
        "expected_debit": "1650",
        "expected_credit": "1650",
        "source_complete": True,
        "currency": currency,
    }
    prior_batch, prior = await committed(c, h, prior_options, ledger(2025, 80, 800, 600))
    prefix = f"/api/v1/intelligence/engagements/{eid}"
    choices = await c.get(prefix + "/inventory-performance-options", headers=h)
    assert choices.status_code == 200, choices.text
    assert current in {row["id"] for row in choices.json()["current"]}
    assert prior in {row["id"] for row in choices.json()["prior"]}
    payload = {
        "currency": currency,
        "periods_comparable": True,
        "comparability_assessment": "SYNTHETIC complete annual periods and consistently assessed accounting scope",
        "stock_change_threshold_percent": "20",
        "gp_change_threshold_pp": "5",
        "rounding_places": 2,
        "rounding_method": "HALF_UP",
        "label": "SYNTHETIC committed stock and gross profit",
        "rationale": "Independently controlled current and prior ledger selections",
    }
    for side, snapshot, amounts in (
        ("current", current, ("120", "1000", "800")),
        ("prior", prior, ("100", "800", "600")),
    ):
        response = await c.get(
            prefix + "/accounts",
            headers=h,
            params={"snapshot_id": snapshot, "comparative": side == "prior"},
        )
        assert response.status_code == 200, response.text
        accounts = {row["code"]: row["id"] for row in response.json()["items"]}
        payload[side] = {"snapshot_id": snapshot}
        for metric, code, amount in zip(
            ("inventory", "revenue", "cost_of_sales"), ("100", "400", "500"), amounts, strict=True
        ):
            payload[side][metric] = {
                "account_ids": [accounts[code]],
                "expected_accounts": 1,
                "expected_amount": amount,
                "assessment": "SYNTHETIC account classification, signs and complete selected scope",
            }
    for change in ("amount", "missing_account", "currency"):
        invalid = deepcopy(payload)
        if change == "amount":
            invalid["prior"]["inventory"]["expected_amount"] = "101"
        elif change == "missing_account":
            invalid["current"]["revenue"]["account_ids"] = [str(uuid4())]
        else:
            invalid["currency"] = "USD"
        response = await c.post(prefix + "/inventory-performance", headers=h, json=invalid)
        assert response.status_code == 422, response.text
    response = await c.post(prefix + "/inventory-performance", headers=h, json=payload)
    assert response.status_code == 201, response.text
    result = response.json()
    assert result["body"]["exception_count"] == 0 and not result["body"]["blocked"]
    gp = result["body"]["rows"][-1]
    assert (gp["current_display"], gp["prior_display"], gp["change_display"]) == (
        "20.00",
        "25.00",
        "-5.00",
    )
    assert result["body"]["rows"][0]["current_display"] == "120.00"
    unreviewed_paper = await c.post(
        prefix + "/events/" + result["id"] + "/working-paper",
        headers=h,
        json={
            "reference": "UNREVIEWED-STOCK-GP",
            "title": "Unreviewed synthetic stock comparison",
            "objective": "Attempt generation before independent stock comparison review",
        },
    )
    assert unreviewed_paper.status_code == 409, unreviewed_paper.text
    assert (await decision(c, h, eid, result)).status_code == 403
    reviewer = await actor(h, "Firm Administrator", "stock-gp-independent-reviewer")
    await grant(c, h, base, reviewer)
    reviewed = await decision(c, reviewer[0], eid, result)
    assert reviewed.status_code == 201, reviewed.text
    result = reviewed.json()
    path = prefix + "/events/" + result["id"]
    paper_input = {
        "reference": "STOCK-GP-SYNTHETIC",
        "title": "Synthetic stock and gross profit working paper",
        "objective": "Assess selected stock and gross profit movements and source controls",
    }
    paper = await c.post(path + "/working-paper", headers=h, json=paper_input)
    assert paper.status_code == 201, paper.text
    paper_body = paper.json()["body"]
    assert paper_body["sources"] == [{"kind": "intelligence", "id": result["id"]}]
    calculations = paper_body["financial_tables"][0]
    assert calculations["source_event_id"] == result["id"]
    assert calculations["source_sha256"] == result["sha256"]
    assert calculations["input_version"] == result["version"]
    for key in (
        "inventory_performance_contract",
        "inventory_performance_binding",
        "inventory_performance_scope",
        "inventory_performance_controls",
        "gp_arithmetic_observation",
        "population_completeness",
        "statutory_calculation_qualified",
        "limitations",
    ):
        assert calculations["values"][key] == result["body"][key]
    assert len(calculations["values"]["inventory_performance_controls"]) == 6
    assert paper_body["financial_tables"][1]["rows"] == result["body"]["rows"]
    assert paper_body["conclusion"].startswith("Draft only.")
    page = await c.get(
        path + "?inventory_performance_offset=4&inventory_performance_limit=2", headers=h
    )
    assert page.status_code == 200, page.text
    assert page.json()["collection_totals"]["inventory_performance_links"] == 6
    assert [row["metric"] for row in page.json()["body"]["inventory_performance_links"]] == [
        "revenue",
        "cost_of_sales",
    ]
    workbook_response = await c.get(path + "/export?format=xlsx", headers=h)
    assert workbook_response.status_code == 200, workbook_response.text
    workbook = load_workbook(io.BytesIO(workbook_response.content))
    assert workbook["Stock GP account lineage"].max_row == 7
    assert workbook["Stock GP source controls"].max_row == 7
    retained = (await c.get(path, headers=h)).json()
    replacement = {
        **prior_options,
        "supersedes_id": prior_batch,
        "expected_debit": "1651",
        "expected_credit": "1651",
    }
    await committed(c, h, replacement, ledger(2025, 80, 801, 600))
    historical = (await c.get(path, headers=h)).json()
    assert historical["freshness"] == "STALE"
    assert historical["body"] == retained["body"] and historical["version"] == retained["version"]
    assert (await decision(c, reviewer[0], eid, result)).status_code == 409
    stale_paper = await c.post(path + "/working-paper", headers=h, json=paper_input)
    assert stale_paper.status_code == 409, stale_paper.text
