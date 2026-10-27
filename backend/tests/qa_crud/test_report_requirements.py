"""Declared requirement coverage through real UAE/India sources and independent review."""

import io
import os
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, create, grant, setup
from tests.qa_crud.report_requirement_fixtures import declared_requirements
from tests.qa_crud.test_country_workflows import decision, double_boundaries
from tests.qa_crud.test_india_ecr_source import india_context
from tests.qa_crud.test_lossless_exports import retained_allocations

pytestmark = pytest.mark.skipif(os.getenv("INTEGRATION_TESTS") != "1", reason="Dedicated test DB")


@pytest.mark.parametrize("country,category", [("AE", "uae-ifrs"), ("IN", "india-statutory")])
async def test_required_items_need_assessment_mapping_and_reviewed_sources(
    environment, monkeypatch, country, category
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "requirements-reviewer")
    await grant(c, h, base, reviewer)
    source, _, _ = await retained_allocations(c, h, eid, reviewer, country)
    prefix = f"/api/v1/regulatory/engagements/{eid}"
    data = {
        "title": "SYNTHETIC requirements-bound report",
        "category": category,
        "rationale": "Synthetic declared source scope, not legal qualification",
        **declared_requirements("source"),
        "blocks": [
            {
                "key": "source",
                "title": "Reviewed source output",
                "kind": "financial_table",
                "source_kind": "intelligence",
                "source_id": source["id"],
                "requirement_keys": ["source"],
            }
        ],
    }

    async def save(body):
        result = await c.post(prefix + "/report-templates", headers=h, json=body)
        assert result.status_code == 201, result.text
        return result.json()

    for fault in [
        "unassessed-scope",
        "unassessed-item",
        "missing-coverage",
        "not-applicable-without-source",
    ]:
        raw = deepcopy(data)
        if fault == "unassessed-scope":
            raw["requirements_assessed"] = False
        elif fault == "unassessed-item":
            raw["requirements"][0]["disposition"] = "UNASSESSED"
        elif fault == "missing-coverage":
            raw["blocks"][0]["requirement_keys"] = []
        else:
            raw["requirements"][0].update(disposition="NOT_APPLICABLE", source_required=False)
            raw["blocks"] = [
                dict(
                    key="note",
                    title="Narrative",
                    kind="narrative",
                    text="Synthetic authored explanation cannot replace reviewed sources",
                    requirement_keys=["source"],
                )
            ]
        draft = await save(raw)
        assert not draft["body"]["requirements_coverage"]["complete"]
        assert (await decision(c, reviewer[0], eid, draft)).status_code == 409
    paper = await create(
        c,
        h,
        base,
        "working-papers",
        dict(
            folder="SYNTHETIC pending source",
            sections=[dict(heading="Pending", text="Source is not yet reviewed")],
            conclusion="Pending independent review",
        ),
    )
    raw = deepcopy(data)
    raw["blocks"][0].update(source_kind="record", source_id=paper["id"])
    unreviewed = await save(raw)
    assert (await decision(c, reviewer[0], eid, unreviewed)).status_code == 409
    template = await save(data)
    assert (await decision(c, h, eid, template)).status_code == 403
    reviewed = await decision(c, reviewer[0], eid, template)
    assert reviewed.status_code == 201, reviewed.text
    assert (
        reviewed.json()["body"]["requirements_coverage"]["source_validation"]
        == "CURRENT_REVIEWED_SOURCES_RECHECKED"
    )
    generated = await c.post(
        prefix + "/reports",
        headers=h,
        json={
            "template_id": reviewed.json()["id"],
            "rationale": "Synthetic declared scope and sources rechecked",
        },
    )
    assert generated.status_code == 201, generated.text
    saved = generated.json()
    assert saved["body"]["requirements_coverage"]["complete"]
    assert saved["body"]["requirements_coverage"]["statutory_completeness_qualified"] is False
    assert saved["body"]["blocks"][0]["requirement_keys"] == ["source"]

    exported = await c.get(
        f"/api/v1/intelligence/engagements/{eid}/events/{saved['id']}/export?format=xlsx", headers=h
    )
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert "Requirements coverage" in workbook.sheetnames
    rows = list(workbook["Requirements coverage"].values)
    assert rows[1][0] == "source" and rows[1][3] == "COVERED"
    assert "not a complete statutory manifest" in rows[1][6]
    workbook.close()
