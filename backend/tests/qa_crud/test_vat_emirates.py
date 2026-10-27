import io
import os
import uuid

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_vat_ecommerce import split_rows
from tests.unit.test_vat_emirates import detailed_rows

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable database required"
)


async def prepare_sources(c, h, eid, reviewer, layout="emirate-20"):
    rows = split_rows("BOOK") if layout == "ecommerce-20" else detailed_rows("BOOK")
    components = []
    for label, records, tax in (
        ("emirate book allocations", rows[:7], "10"),
        ("other book boxes", rows[7:], "220"),
    ):
        imported = await source(
            c, h, eid, records, "SYNTHETIC " + label, "return-boxes", {"AED": tax}
        )
        reviewed = await decision(c, reviewer[0], eid, imported)
        assert reviewed.status_code == 201, reviewed.text
        components.append(reviewed.json()["id"])
    filed = (
        split_rows("INDEPENDENT-FILED")
        if layout == "ecommerce-20"
        else detailed_rows("INDEPENDENT-FILED")
    )
    filed[0].update(taxable_amount="120", tax="6")
    filed[1].update(taxable_amount="80", tax="4")
    if layout == "ecommerce-20":
        for row in filed[:2]:
            row.update(ecommerce_taxable_amount=row["taxable_amount"], ecommerce_tax=row["tax"])
    imported = await source(
        c,
        h,
        eid,
        filed,
        "SYNTHETIC retained filed emirate allocation",
        "return-boxes",
        {"AED": "230"},
    )
    reviewed = await decision(c, reviewer[0], eid, imported)
    assert reviewed.status_code == 201, reviewed.text
    return components, reviewed.json()["id"]


@pytest.mark.parametrize("layout", ["emirate-20", "ecommerce-20"])
async def test_emirate_composition_review_mismatch_missing_zero_export_and_staleness(
    environment, monkeypatch, layout
):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, base = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "vat-emirate-reviewer")
    await grant(c, h, base, reviewer)
    components, filed_id = await prepare_sources(c, h, eid, reviewer, layout)
    path = f"/api/v1/intelligence/engagements/{eid}"
    config = await c.get(path + "/country-workflows", headers=h)
    boxes = config.json()["vat201_layouts"][layout]["boxes"]
    assert len(boxes) == 20 and config.json()["vat201_emirates"]["1g"] == "Fujairah"
    payload = {
        "source_ids": components,
        "layout": layout,
        "taxable_amount": "2400",
        "tax": "230",
        "adjustment": "0",
        "label": "SYNTHETIC composed emirate books",
        "assessment": "All seven geographic allocations and explicit nil emirates independently assessed against retained supporting sources",
    }
    response = await c.post(path + "/vat201-composition-source", headers=h, json=payload)
    assert response.status_code == 201, response.text
    composed = response.json()
    assert (
        composed["body"]["row_count"] == 20
        and composed["body"]["rows"][0]["emirate"] == "Abu Dhabi"
    )
    compare = {
        "workflow": "uae-vat-return",
        "vat201_layout": layout,
        "source_ids": {"books": composed["id"], "return": filed_id},
        "as_of": "2026-06-30",
        "currency": "AED",
        "required_boxes": boxes,
        "title": "SYNTHETIC emirate comparison",
        "rationale": "Independent filed regional totals compared with all assessed geographic book allocations",
    }
    await execute(
        c, h, eid, "country-workflow", compare, failed="Independently review emirate allocation"
    )
    assert (await decision(c, h, eid, composed)).status_code == 403
    approved = await decision(c, reviewer[0], eid, composed)
    assert approved.status_code == 201, approved.text
    composed = approved.json()
    compare["source_ids"]["books"] = composed["id"]
    result = await execute(c, h, eid, "country-workflow", compare)
    assert result["body"]["exception_count"] == 2 and not result["body"]["blocked"]
    assert [row["emirate"] for row in result["body"]["rows"][:2]] == ["Abu Dhabi", "Dubai"]
    assert (await decision(c, reviewer[0], eid, result)).status_code == 409
    exported = await c.get(path + f"/events/{result['id']}/export?format=xlsx", headers=h)
    assert exported.status_code == 200
    book = load_workbook(io.BytesIO(exported.content))
    rows = list(book["Rows"].values)
    assert len(rows) == 21 and "emirate" in rows[0]
    await execute(
        c,
        h,
        eid,
        "country-workflow",
        {**compare, "vat201_layout": "custom"},
        failed="explicit twenty-row",
    )
    missing = await source(
        c,
        h,
        eid,
        [
            row
            for row in (
                split_rows() if layout == "ecommerce-20" else detailed_rows("MISSING-FILED")
            )
            if row["box"] != "1g"
        ],
        "SYNTHETIC missing zero emirate",
        "return-boxes",
        {"AED": "230"},
    )
    reviewed = await decision(c, reviewer[0], eid, missing)
    assert reviewed.status_code == 201, reviewed.text
    incomplete = await execute(
        c,
        h,
        eid,
        "country-workflow",
        {**compare, "source_ids": {**compare["source_ids"], "return": reviewed.json()["id"]}},
    )
    assert incomplete["body"]["blocked"] is True
    assert (await decision(c, reviewer[0], eid, incomplete)).status_code == 409
    # The same wrong summary in independent books and filed populations must block.
    bad_ids = {}
    for role in ("books", "return"):
        wrong = (
            split_rows("ARITHMETIC-" + role)
            if layout == "ecommerce-20"
            else detailed_rows("ARITHMETIC-" + role)
        )
        next(row for row in wrong if row["box"] == "8")["tax"] = "55.01"
        imported = await source(
            c,
            h,
            eid,
            wrong,
            "SYNTHETIC wrong " + role + " totals",
            "return-boxes",
            {"AED": "230.01"},
        )
        checked = await decision(c, reviewer[0], eid, imported)
        assert checked.status_code == 201, checked.text
        bad_ids[role] = checked.json()["id"]
    invalid = await execute(c, h, eid, "country-workflow", {**compare, "source_ids": bad_ids})
    assert all(row["state"] == "MATCHED" for row in invalid["body"]["rows"])
    assert invalid["body"]["blocked"]
    assert invalid["body"]["vat201_arithmetic_error_count"] == 4
    assert (await decision(c, reviewer[0], eid, invalid)).status_code == 409
    exported = await c.get(path + f"/events/{invalid['id']}/export?format=xlsx", headers=h)
    assert exported.status_code == 200
    arithmetic = list(load_workbook(io.BytesIO(exported.content))["VAT201 arithmetic"].values)
    assert len(arithmetic) == (73 if layout == "ecommerce-20" else 31)
    assert any("ARITHMETIC_ERROR" in row for row in arithmetic)
    original = await c.get(path + "/events/" + components[0], headers=h)
    rejected = await decision(c, reviewer[0], eid, original.json(), "reject")
    assert rejected.status_code == 201, rejected.text
    for event in (composed, result, incomplete):
        stale = await c.get(path + "/events/" + event["id"], headers=h)
        assert stale.json()["freshness"] == "STALE"


@pytest.mark.parametrize("layout", ["emirate-20", "ecommerce-20"])
async def test_detailed_uae_layout_is_not_enabled_in_india(environment, monkeypatch, layout):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, _ = await india_context(c)
    path = f"/api/v1/intelligence/engagements/{eid}"
    config = await c.get(path + "/country-workflows", headers=h)
    assert config.json()["vat201_layouts"] == config.json()["vat201_emirates"] == {}
    response = await c.post(
        path + "/vat201-composition-source",
        headers=h,
        json={
            "source_ids": [str(uuid.uuid4()), str(uuid.uuid4())],
            "layout": layout,
            "taxable_amount": "0",
            "tax": "0",
            "adjustment": "0",
            "label": "SYNTHETIC wrong country",
            "assessment": "A UAE domestic layout must not activate for an Indian engagement",
        },
    )
    assert response.status_code == 422 and "UAE engagement" in response.text
