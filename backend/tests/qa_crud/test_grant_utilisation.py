"""Retained grant mandates, source review, allocation controls and complete exports."""

import csv
import hashlib
import io
import os
import uuid
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.modules.country_workflows.grant_utilisation import ROW_TYPES
from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, create, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_grant_utilisation import packet

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required"
)


async def prepare(c, h, eid, reviewer, country="AE", case="matched", count=1, reviewed=True):
    tables = packet(country)
    currency = "AED" if country == "AE" else "INR"
    start = "2026-01-01" if country == "AE" else "2026-04-01"
    opening = "2025-12-31" if country == "AE" else "2026-03-31"
    tables["controls"][0].update(period_start=start, period_end="2026-06-30")
    tables["controls"][1].update(period_start=opening, period_end=opening)
    audit = f"/api/v1/engagements/{eid}/audit"
    documents = {}
    for role in ("mandate", "credit") if case == "credit" else ("mandate",):
        text = (
            (
                "SYNTHETIC ONLY. Fictional grant mandate: F1 / training / P1 / D1; effective 2025-01-01 to 2027-12-31; category ceiling "
                + ("400" if case == "over-budget" else "10000" if count > 1 else "1000")
                + "; whole fund ceiling "
                + ("15000" if count > 1 else "1500")
                + " "
                + currency
                + ". No legal eligibility or donor authentication.\n"
            )
            if role == "mandate"
            else "SYNTHETIC ONLY. Assessed credit recovery -50 against SYN-INV-1, grant F1/training. Original charge and donor permission still require professional assessment.\n"
        )
        record = await create(
            c, h, audit, "evidence", dict(category="Other", retain_until="2036-12-31")
        )
        response = await c.post(
            audit + "/evidence-files/" + record["id"],
            headers=h,
            files={"file": (f"SYNTHETIC-grant-{role}.txt", text.encode(), "text/plain")},
        )
        assert response.status_code == 201, response.text
        documents[role] = {**response.json(), "record": record["id"], "bytes": text.encode()}
    tables["budgets"][0].update(
        source_evidence_version_id=documents["mandate"]["id"],
        source_sha256=documents["mandate"]["sha256"],
    )
    if case == "over-budget":
        tables["budgets"][0]["category_ceiling"] = "400"
    if count > 1:
        expense = tables["expenses"][0]
        allocation = tables["allocations"][0]
        tables["expenses"] = [
            dict(expense, id=f"E-{i:04d}", reference=f"SYN-INV-{i:04d}", amount="1")
            for i in range(count)
        ]
        tables["allocations"] = [
            dict(allocation, id=f"A-{i:04d}", expense_id=f"E-{i:04d}", amount="1")
            for i in range(count)
        ]
        tables["controls"][0].update(
            amount=str(count), positive_amount=str(count), expected_rows=count
        )
        tables["budgets"][0].update(category_ceiling="10000", fund_ceiling="15000")
    if case == "credit":
        tables["expenses"].append(
            dict(
                tables["expenses"][0],
                id="CR1",
                reference="SYN-CREDIT-1",
                kind="credit",
                amount="-50",
                original_reference="SYN-INV-1",
            )
        )
        tables["allocations"].append(
            dict(
                tables["allocations"][0],
                id="CA1",
                expense_id="CR1",
                amount="-50",
                credit_original_reference="SYN-INV-1",
                credit_evidence_version_id=documents["credit"]["id"],
                credit_source_sha256=documents["credit"]["sha256"],
            )
        )
        tables["controls"][0].update(
            amount="250", positive_amount="300", negative_amount="-50", expected_rows=2
        )
    if case == "nil":
        tables["expenses"] = []
        tables["allocations"] = []
        tables["controls"][0].update(
            amount="0", positive_amount="0", negative_amount="0", expected_rows=0
        )
    if case == "unallocated":
        tables["allocations"][0]["amount"] = "250"
    events = {}
    for role, kind in dict(
        budgets="grant-budgets",
        expenses="grant-expenses",
        allocations="grant-allocations",
        controls="grant-utilisation-controls",
    ).items():
        rows = tables[role]
        total = sum(
            (
                Decimal(row["category_ceiling"] if role == "budgets" else row["amount"])
                for row in rows
            ),
            Decimal(0),
        )
        options = (
            {"rule_document_version_id": documents["mandate"]["id"]} if role == "budgets" else {}
        )
        if rows:
            event = await source(
                c,
                h,
                eid,
                rows,
                f"SYNTHETIC {country} grant {case} {role}",
                kind,
                {currency: str(total)},
                **options,
            )
        else:
            output = io.StringIO(newline="")
            csv.writer(output).writerow(list(ROW_TYPES[kind].model_fields))
            record = await create(
                c, h, audit, "evidence", dict(category="Other", retain_until="2036-12-31")
            )
            uploaded = await c.post(
                audit + "/evidence-files/" + record["id"],
                headers=h,
                files={
                    "file": (f"SYNTHETIC-nil-{role}.csv", output.getvalue().encode(), "text/csv")
                },
            )
            assert uploaded.status_code == 201, uploaded.text
            event = await execute(
                c,
                h,
                eid,
                "workflow-source",
                dict(
                    evidence_version_id=uploaded.json()["id"],
                    table_type=kind,
                    label=f"SYNTHETIC nil {role}",
                    expected_rows=0,
                    control_totals={currency: "0"},
                    nil_reason="Independent synthetic review found no expenditure/allocations in the declared source scope",
                ),
            )
        if reviewed:
            response = await decision(c, reviewer[0], eid, event)
            assert response.status_code == 201, response.text
            event = response.json()
        events[role] = event
    return dict(
        events=events,
        documents=documents,
        payload=dict(
            workflow="grant-utilisation",
            source_ids={role: event["id"] for role, event in events.items()},
            period_start=start,
            as_of="2026-06-30",
            title=f"SYNTHETIC {country} grant {case}",
            rationale="Independent synthetic mandate, opening utilisation, expense and allocation review",
            currency=currency,
        ),
    )


@pytest.mark.parametrize("country", ["AE", "IN"])
@pytest.mark.parametrize("case", ["matched", "credit", "nil", "over-budget"])
async def test_grant_source_review_dispositions_download_export_and_stale(
    environment, monkeypatch, country, case
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, audit = await setup(c)
    else:
        h, eid, audit = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "grant-source-reviewer")
    final = await actor(h, "Firm Administrator", "grant-final-reviewer")
    for person in (reviewer, final):
        await grant(c, h, audit, person)
    fixture = await prepare(c, h, eid, reviewer, country, case, reviewed=False)
    await execute(
        c, h, eid, "country-workflow", fixture["payload"], failed="Independently review all grant"
    )
    assert (await decision(c, h, eid, fixture["events"]["budgets"])).status_code == 403
    for role, event in fixture["events"].items():
        reviewed = await decision(c, reviewer[0], eid, event)
        assert reviewed.status_code == 201, reviewed.text
        fixture["events"][role] = reviewed.json()
        fixture["payload"]["source_ids"][role] = reviewed.json()["id"]
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    assert result["body"]["grant_document_binding"] == "RETAINED_DOCUMENT_BYTES_VERIFIED"
    assert not result["body"]["blocked"]
    assert result["body"]["certificate_issued"] is False
    assert (await decision(c, h, eid, result)).status_code == 403
    if case in {"credit", "over-budget"}:
        assert (await decision(c, final[0], eid, result)).status_code == 409
        for row in result["body"]["rows"]:
            if row["state"] != "MATCHED":
                accepted = await decision(c, h, eid, result, action="accept", item_id=row["id"])
                assert accepted.status_code == 201, accepted.text
                result = accepted.json()
    response = await decision(c, final[0], eid, result)
    assert response.status_code == 201, response.text
    result = response.json()
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    exported = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text if exported.status_code != 200 else ""
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert len(list(book["Grant documents"].values)) == len(fixture["documents"]) + 1
    assert len(list(book["Rows"].values)) == len(result["body"]["rows"]) + 1
    for document in fixture["documents"].values():
        download = await c.get(
            audit + f"/evidence-files/{document['record']}/{document['id']}/download", headers=h
        )
        assert (
            download.status_code == 200
            and hashlib.sha256(download.content).hexdigest() == document["sha256"]
        )
    mandate = fixture["documents"]["mandate"]
    superseded = await c.post(
        audit + "/evidence-files/" + mandate["record"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-amended-mandate.txt",
                b"Synthetic amended grant terms; must independently reassess existing sources.",
                "text/plain",
            )
        },
    )
    assert superseded.status_code == 201, superseded.text
    assert (await c.get(prefix, headers=h)).json()["freshness"] != "CURRENT"


async def test_unallocated_expense_cannot_be_overridden_to_approve(environment, monkeypatch):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, audit = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "grant-blocked-reviewer")
    await grant(c, h, audit, reviewer)
    fixture = await prepare(c, h, eid, reviewer, case="unallocated")
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    assert result["body"]["blocked"]
    for row in result["body"]["rows"]:
        if row["state"] != "MATCHED":
            accepted = await decision(c, h, eid, result, action="accept", item_id=row["id"])
            assert accepted.status_code == 201, accepted.text
            result = accepted.json()
    assert (await decision(c, reviewer[0], eid, result)).status_code == 409


async def test_grant_paging_and_full_export_keep_all_original_allocations(environment, monkeypatch):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, audit = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "grant-page-reviewer")
    await grant(c, h, audit, reviewer)
    fixture = await prepare(c, h, eid, reviewer, count=121)
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    assert result["collection_totals"]["rows"] == 245
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    page = await c.get(prefix + "?offset=244&limit=1", headers=h)
    assert page.status_code == 200 and page.json()["body"]["rows"][0]["kind"] == "fund-budget"
    assert page.json()["body"]["rows"][0]["closing_utilisation"] == "321"
    book = load_workbook(
        io.BytesIO((await c.get(prefix + "/export?format=xlsx", headers=h)).content), read_only=True
    )
    assert len(list(book["Rows"].values)) == 246


async def test_grant_calculation_rechecks_actual_mandate_bytes(environment, monkeypatch):
    from app.db.session import Session, tenant_context
    from app.modules.audit_execution.models import EvidenceVersion
    from app.modules.lineage.storage import ObjectStore

    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, audit = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "grant-integrity-reviewer")
    await grant(c, h, audit, reviewer)
    fixture = await prepare(c, h, eid, reviewer)
    document = fixture["documents"]["mandate"]
    async with Session() as db, db.begin():
        await tenant_context(db, h["X-Tenant-ID"])
        version = await db.get(EvidenceVersion, uuid.UUID(document["id"]))
        with ObjectStore().local(version.object_key, version.object_version) as path:
            original = path.read_bytes()
            try:
                path.write_bytes(
                    b"Synthetic altered mandate must not be accepted against its old stored hash"
                )
                await execute(c, h, eid, "country-workflow", fixture["payload"], failed="integrity")
            finally:
                path.write_bytes(original)
