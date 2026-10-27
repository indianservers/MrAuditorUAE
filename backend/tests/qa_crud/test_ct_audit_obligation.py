"""Synthetic UAE CT audit obligation and independently linked support regressions."""

import io
import os
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import action, actor, create, grant
from tests.phase6.test_workflows import compute, prepared

pytestmark = pytest.mark.skipif(os.getenv("INTEGRATION_TESTS") != "1", reason="Dedicated test DB")


@pytest.mark.parametrize(
    "classification,scale,required,reason",
    [
        ("resident_juridical", 1000001, True, "REVENUE_EXCEEDS_50M"),
        ("qfzp", 40000, True, "QFZP_REGARDLESS_OF_REVENUE"),
        ("nonresident_juridical", 40000, None, "NONRESIDENT_REVENUE_SCOPE_UNASSESSED"),
    ],
)
async def test_obligation_cannot_be_suppressed_by_false_profile_flag(
    environment, classification, scale, required, reason
):
    c = environment
    h, eid, base, manager, partner, original, adj, note, rules, ready, financial = await prepared(
        c, scale=scale
    )
    body = {key: value for key, value in original["body"].items() if not key.startswith("review_")}
    body.update(
        classification=classification,
        resident=classification != "nonresident_juridical",
        free_zone=classification == "qfzp",
    )
    profile = await ready("profile", body)
    if classification == "qfzp":
        from app.modules.tax.calculations import CORE_QFZP, QFZP_REQUIRED

        segment = dict(
            source_line_code="400",
            counterparty="Synthetic buyer",
            counterparty_type="non_free_zone",
            activity="Synthetic manufacturing",
            jurisdiction="AE",
            contract_evidence_id=note["id"],
            rationale="Synthetic independently assessed segment",
        )
        await ready(
            "qfzp",
            {
                "profile_id": profile["id"],
                "conditions": [
                    {
                        "requirement": key,
                        "assessment": "PASS" if key in CORE_QFZP else "NOT_APPLICABLE",
                        "rationale": "Synthetic reviewed condition",
                        "evidence_record_ids": [note["id"]] if key in CORE_QFZP else [],
                    }
                    for key in sorted(QFZP_REQUIRED)
                ],
                "segments": [
                    {
                        **segment,
                        "reference": "Q",
                        "revenue": "1900000",
                        "taxable_profit": "1140000",
                        "classification": "qualifying",
                    },
                    {
                        **segment,
                        "reference": "N",
                        "revenue": "100000",
                        "taxable_profit": "60000",
                        "classification": "nonqualifying",
                    },
                ],
                "procedures_applicability_basis": "Synthetic no distribution activity",
            },
        )
    result = await compute(c, h, eid, profile)
    guard = result["results"]["audit_statement_obligation"]
    assert guard["required_for_ct"] is required
    assert reason in guard["reasons"] and guard["issues"]
    assert not result["results"]["validation"]["ready"]
    if classification == "qfzp":
        assert result["results"]["qfzp"]["assessment"] == "PASS"
        assert result["results"]["exceptions"] == guard["issues"]

    if classification == "nonresident_juridical":
        assert guard["threshold_revenue"] is None


async def test_audit_support_is_a_direct_dependency_and_reopening_blocks_reuse(environment):
    c = environment
    h, eid, base, manager, partner, original, adj, note, rules, ready, financial = await prepared(c)
    audit_manager = await actor(h, "Audit Manager", "ct-audit-support-manager")
    audit_partner = await actor(h, "Audit Partner", "ct-audit-support-partner")
    await grant(c, h, base, audit_manager)
    await grant(c, h, base, audit_partner)
    support = await create(
        c,
        h,
        base,
        "working-papers",
        {
            "folder": "SYNTHETIC CT documentation",
            "sections": [
                {
                    "heading": "Synthetic independent audit support",
                    "text": "Synthetic retained evidence assessment only; no real audit opinion",
                }
            ],
            "conclusion": "Synthetic supporting record for dependency controls",
        },
    )
    support = await action(c, h, base, support, "submit")
    support = await action(c, audit_manager[0], base, support, "approve")
    support = await action(c, audit_partner[0], base, support, "approve")
    body = {key: value for key, value in original["body"].items() if not key.startswith("review_")}
    body.update(audited_financials_required=True, audited_financials_evidence_id=support["id"])
    profile = await ready("profile", body)
    result = await compute(c, h, eid, profile)
    endpoint = f"/api/v1/tax/engagements/{eid}/runs/{result['id']}"
    assert result["results"]["validation"]["ready"]
    assert any(item["id"] == support["id"] for item in result["dependencies"])
    assert (await c.get(endpoint, headers=h)).json()["dependencies_current"]
    exported = await c.get(endpoint + "/export/xlsx", headers=h)
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert "CT audit obligation" in workbook.sheetnames
    assert "MD84-2025" in " ".join(str(row[0]) for row in workbook["CT audit obligation"].values)
    workbook.close()
    retained = deepcopy(result["results"])
    await action(c, audit_partner[0], base, support, "reopen")
    source_financial = await c.get(
        f"/api/v1/ifrs/engagements/{eid}/runs/{financial['id']}", headers=h
    )
    assert source_financial.status_code == 200 and source_financial.json()["dependencies_current"]
    historical = await c.get(endpoint, headers=h)
    assert historical.status_code == 200
    assert not historical.json()["dependencies_current"]
    assert historical.json()["results"] == retained
    assert (
        await c.post(f"/api/v1/tax/engagements/{eid}/computations/{profile['id']}", headers=h)
    ).status_code == 409
