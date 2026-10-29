"""Real HTTP/DB/storage/processor; scanner and dispatch are explicit doubles."""

import io
import os
from copy import deepcopy
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, create, grant
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_msme_interest import settled_fixture

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable migrated database required"
)


@pytest.mark.parametrize(
    "withdrawal", ["rate-document", "fact-context", "principal-review", "noncash-document"]
)
async def test_msme_interest_original_sources_buckets_bank_pool_exports_and_staleness(
    environment, monkeypatch, withdrawal
):
    from app.modules.country_workflows.msme_interest_schemas import ROLES
    from app.modules.country_workflows.schemas import ROW_TYPES

    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await india_context(c)
    checker = await actor(h, "Firm Administrator", "interest-independent-checker")
    author = await actor(h, "Firm Administrator", "interest-support-author")
    for member in (checker, author):
        await grant(c, h, base, member)
    prefix = f"/api/v1/intelligence/engagements/{eid}"
    catalogue = (await c.get(prefix + "/country-workflows", headers=h)).json()
    available = next(row for row in catalogue["items"] if row["id"] == "india-msme-interest")
    assert available["available"] and len(available["roles"]) == 9
    documents = {}
    for purpose in [
        "facts",
        "policy",
        "rate",
        *(["credit"] if withdrawal == "noncash-document" else []),
    ]:
        record = await create(
            c, author[0], base, "evidence", {"category": "Other", "retain_until": "2036-12-31"}
        )
        uploaded = await c.post(
            base + "/evidence-files/" + record["id"],
            headers=author[0],
            files={
                "file": (
                    "SYNTHETIC-interest-" + purpose + ".txt",
                    (
                        "SYNTHETIC retained "
                        + purpose
                        + " assessment only; no client or operative-law qualification"
                    ).encode(),
                    "text/plain",
                )
            },
        )
        assert uploaded.status_code == 201, uploaded.text
        documents[purpose] = (record, uploaded.json())
    data, tables = settled_fixture()
    if withdrawal == "noncash-document":
        policy = tables["interest_policies"][0]
        tables["principal_changes"].append(
            dict(
                id="CREDIT",
                **{
                    key: policy[key]
                    for key in (
                        "buyer_id",
                        "supplier_id",
                        "invoice_id",
                        "currency",
                        "source_url",
                        "source_version",
                        "applicability_review",
                    )
                },
                mode="DOCUMENTED_PRINCIPAL_REDUCTION",
                amount="500",
                day="2026-06-16",
                source_evidence_version_id=documents["credit"][1]["id"],
                source_sha256=documents["credit"][1]["sha256"],
                evidence_reference="SYNTHETIC separately assessed credit note",
                rationale="Independent dated non-cash principal reduction",
            )
        )
        tables["interest_allocations"][0]["amount"] = "50"
        tables["reported"][0]["amount"] = "131.37"
    retained = {}
    for role, table_type in ROLES.items():
        records = [
            ROW_TYPES[table_type].model_validate(row).model_dump(mode="json")
            for row in tables[role]
        ]
        options = {}
        if role in {"assessments", "invoices"}:
            options.update(
                msme_context_evidence_version_id=documents["facts"][1]["id"],
                msme_context_assessment="Independent retained supplier, acceptance, original invoice and terms facts",
            )
        if role in {"interest_policies", "bank_rates"}:
            for row in records:
                row.pop("source_evidence_version_id")
                row.pop("source_sha256")
            options["rule_document_version_id"] = documents[
                "policy" if role == "interest_policies" else "rate"
            ][1]["id"]
        totals = (
            {"INR": str(sum((Decimal(row["amount"]) for row in records), Decimal(0)))}
            if "amount" in records[0]
            else {}
        )
        retained[role] = await source(
            c, h, eid, records, "SYNTHETIC interest " + role, table_type, totals, **options
        )
    assert (await decision(c, h, eid, retained["bank_rates"])).status_code == 403
    assert (await decision(c, author[0], eid, retained["bank_rates"])).status_code == 403
    parameters = dict(
        workflow="india-msme-interest",
        source_ids={role: row["id"] for role, row in retained.items()},
        **{
            key: value.isoformat() if hasattr(value, "isoformat") else value
            for key, value in vars(data).items()
        },
        title="SYNTHETIC original MSE interest and bank reconciliation",
        rationale="Independent invoice principal, source-assessed conventions, bank rates and interest payment controls",
    )
    await execute(c, h, eid, "country-workflow", parameters, failed="Independently review all")
    for role, row in retained.items():
        reviewed = await decision(c, checker[0], eid, row)
        assert reviewed.status_code == 201, reviewed.text
        retained[role] = reviewed.json()
    parameters["source_ids"] = {role: row["id"] for role, row in retained.items()}
    result = await execute(c, h, eid, "country-workflow", parameters)
    expected = "131.37" if withdrawal == "noncash-document" else "93.23"
    assert not result["body"]["blocked"] and result["body"]["expected_total"] == expected
    assert result["body"]["principal_timing_exception_count"] == 1
    assert result["body"]["msme_document_binding"] == "RETAINED_DOCUMENT_BYTES_VERIFIED"
    assert result["body"]["msme_interest_document_binding"] == "RETAINED_DOCUMENT_BYTES_VERIFIED"
    assert not result["body"]["statutory_calculation_qualified"]
    assert (await decision(c, h, eid, result)).status_code == 403
    assert (await decision(c, author[0], eid, result)).status_code == 403
    reviewed = await decision(c, checker[0], eid, result)
    assert reviewed.status_code == 201, reviewed.text
    result = reviewed.json()
    exported = await c.get(prefix + f"/events/{result['id']}/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content))
    for title in [
        "MSE interest intervals",
        "MSE interest events",
        "MSE monthly rests",
        "MSE gross principal links",
        "MSE interest rate sources",
        "MSE interest bank controls",
        "MSE interest policies",
    ]:
        assert title in workbook.sheetnames
    rests = list(workbook["MSE monthly rests"].values)
    assert len(rests) == 3
    bank = list(workbook["MSE interest bank controls"].values)
    assert bank[1][bank[0].index("gross_principal_allocated")] == "900"
    assert bank[1][bank[0].index("interest_allocated")] == (
        "50" if withdrawal == "noncash-document" else "90"
    )
    assert bank[1][bank[0].index("unallocated_outflow")] == (
        "50" if withdrawal == "noncash-document" else "10"
    )
    original = deepcopy(result["body"])
    if withdrawal == "principal-review":
        response = await decision(c, checker[0], eid, retained["principal_changes"], "reject")
    else:
        purpose = {"rate-document": "rate", "fact-context": "facts", "noncash-document": "credit"}[
            withdrawal
        ]
        response = await c.post(
            base + "/evidence-files/" + documents[purpose][0]["id"],
            headers=author[0],
            files={
                "file": (
                    "SYNTHETIC-replaced-interest-support.txt",
                    b"SYNTHETIC changed retained supporting facts",
                    "text/plain",
                )
            },
        )
    assert response.status_code == 201, response.text
    stale = (await c.get(prefix + "/events/" + result["id"], headers=h)).json()
    assert stale["freshness"] == "STALE" and stale["body"] == original
    assert (await decision(c, checker[0], eid, result)).status_code == 409
