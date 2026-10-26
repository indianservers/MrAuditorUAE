"""Synthetic retained cutoff schedules through source review, lineage and export."""

import io
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_cutoff_schedules import population

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required"
)


async def prepare(
    c, h, eid, reviewer, currency="AED", case="split_reversal", count=1, reviewed=True
):
    rows = dict(documents=[], recognitions=[], postings=[])
    for index in range(count):
        template = population(currency)
        if case == "late_credit":
            template["documents"][0].update(document_kind="credit-note", amount="-1200")
            for allocation in template["recognitions"]:
                allocation["amount"] = "-" + allocation["amount"]
            template["postings"][0].update(amount="-100", posting_date="2026-07-10")
        else:
            posting = template["postings"][0]
            template["postings"] = [
                {**posting, "id": f"POST-{offset}", "amount": amount}
                for offset, amount in enumerate(("80", "40", "-20"))
            ]
        for role, records in template.items():
            for row in records:
                row["id"] += f"-{index:04d}"
                if "document_id" in row:
                    row["document_id"] += f"-{index:04d}"
                if "allocation_id" in row:
                    row["allocation_id"] += f"-{index:04d}"
                if role == "documents":
                    row["reference"] += f"-{index:04d}"
                rows[role].append(row)
    events = {}
    for role, kind in (
        ("documents", "cutoff-documents"),
        ("recognitions", "cutoff-recognition-lines"),
        ("postings", "cutoff-posting-lines"),
    ):
        amount = (
            (-1200 if role != "postings" else -100)
            if case == "late_credit"
            else 1200
            if role != "postings"
            else 100
        )
        event = await source(
            c,
            h,
            eid,
            rows[role],
            "SYNTHETIC " + case + " cutoff " + role,
            kind,
            {currency: str(amount * count)},
        )
        if reviewed:
            response = await decision(c, reviewer[0], eid, event)
            assert response.status_code == 201, response.text
            event = response.json()
        events[role] = event
    return dict(
        engagement=eid,
        events=events,
        sources={role: event["id"] for role, event in events.items()},
        payload=dict(
            workflow="expense-cutoff",
            source_ids={role: event["id"] for role, event in events.items()},
            period_start="2026-04-01" if currency == "INR" else "2026-01-01",
            as_of="2026-06-30",
            review_through="2026-07-31",
            title="SYNTHETIC " + case + " expense cutoff",
            rationale="Independently reviewed invoice allocation and exact dated ledger links",
        ),
    )


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
@pytest.mark.parametrize("case", ["split_reversal", "late_credit"])
async def test_cutoff_sources_review_period_amounts_export_and_staleness(
    environment, monkeypatch, country, currency, case
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "cutoff-source-reviewer")
    final = await actor(h, "Firm Administrator", "cutoff-final-reviewer")
    for person in (reviewer, final):
        await grant(c, h, base, person)
    fixture = await prepare(c, h, eid, reviewer, currency, case, reviewed=False)
    refused = await execute(
        c, h, eid, "country-workflow", fixture["payload"], failed="Independently review"
    )
    assert "Independently review" in str(refused)
    for role, event in fixture["events"].items():
        reviewed = await decision(c, reviewer[0], eid, event)
        assert reviewed.status_code == 201, reviewed.text
        fixture["events"][role] = reviewed.json()
        fixture["sources"][role] = reviewed.json()["id"]
        fixture["payload"]["source_ids"][role] = reviewed.json()["id"]
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    assert not result["body"]["blocked"]
    current, future = result["body"]["rows"]
    assert future["expected_period_amount"] == "0"
    assert current["posted_period_amount"] == ("100" if case == "split_reversal" else "0")
    assert current["state"] == ("MATCHED" if case == "split_reversal" else "LATE_RECOGNITION")
    if case == "late_credit":
        assert current["subsequent_posted_amount"] == "-100" and current["difference"] == "100"
        assert (await decision(c, final[0], eid, result)).status_code == 409
        accepted = await decision(c, h, eid, result, action="accept", item_id=current["id"])
        assert accepted.status_code == 201, accepted.text
        result = accepted.json()
    approved = await decision(c, final[0], eid, result)
    assert approved.status_code == 201, approved.text
    result = approved.json()
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    exported = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert len(list(book["Cutoff document controls"].values)) == 2
    assert len(list(book["Cutoff posting links"].values)) == (4 if case == "split_reversal" else 2)
    book.close()
    rereview = await decision(c, reviewer[0], eid, fixture["events"]["recognitions"])
    assert rereview.status_code == 201, rereview.text
    stale = await c.get(prefix, headers=h)
    assert stale.json()["freshness"] == "STALE"


async def test_cutoff_collections_page_independently_and_export_every_link(
    environment, monkeypatch
):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, base = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "cutoff-pagination-reviewer")
    await grant(c, h, base, reviewer)
    fixture = await prepare(c, h, eid, reviewer, count=121)
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    page = await c.get(
        prefix
        + "?offset=240&limit=2&cutoff_control_offset=120&cutoff_control_limit=1&cutoff_posting_offset=362&cutoff_posting_limit=1",
        headers=h,
    )
    assert page.status_code == 200, page.text
    body = page.json()["body"]
    assert len(body["rows"]) == 2 and body["rows"][-1]["document_source_id"] == "DOC-1-0120"
    assert body["cutoff_document_controls"][0]["document_source_id"] == "DOC-1-0120"
    assert body["cutoff_posting_links"][0]["posting_source_id"] == "POST-2-0120"
    assert page.json()["collection_totals"]["cutoff_posting_links"] == 363
    exported = await c.get(prefix + "/export?format=xlsx", headers=h)
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert len(list(book["Rows"].values)) == 243
    links = list(book["Cutoff posting links"].values)
    assert len(links) == 364 and links[-1][0] == "POST-2-0120"
    book.close()
