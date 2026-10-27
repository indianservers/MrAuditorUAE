import io
import os
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_payroll_ledger import population

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required"
)


async def prepare_sources(c, h, eid, currency, reviewer, populations=None):
    selected, originals = {}, {}
    for role, rows in (populations or population(currency)).items():
        imported = await source(
            c,
            h,
            eid,
            rows,
            "SYNTHETIC payroll ledger " + role,
            table_type="payroll-account-balances"
            if role in {"opening", "closing"}
            else "payroll-ledger-lines",
            totals={currency: str(sum((Decimal(row["amount"]) for row in rows), Decimal(0)))},
        )
        originals[role] = imported
        response = await decision(c, reviewer[0], eid, imported)
        assert response.status_code == 201, response.text
        selected[role] = response.json()["id"]
    return selected, originals


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
async def test_controlled_payroll_accounting_review_export_paging_and_withdrawal(
    environment, monkeypatch, country, currency
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "payroll-ledger-reviewer")
    await grant(c, h, base, reviewer)
    selected, _ = await prepare_sources(c, h, eid, currency, reviewer)
    payload = {
        "workflow": "payroll-ledger-rollforward",
        "source_ids": selected,
        "period_start": "2026-06-01",
        "as_of": "2026-06-30",
        "title": "SYNTHETIC payroll accounting close",
        "rationale": "Independent payroll allocation, posting date and normal-side basis assessed against retained GL balances",
    }
    expected_rows = population(currency)["payroll"]
    unreviewed = await source(
        c,
        h,
        eid,
        expected_rows,
        "SYNTHETIC unreviewed payroll",
        table_type="payroll-ledger-lines",
        totals={currency: str(sum((Decimal(row["amount"]) for row in expected_rows), Decimal(0)))},
    )
    await execute(
        c,
        h,
        eid,
        "country-workflow",
        {**payload, "source_ids": {**selected, "payroll": unreviewed["id"]}},
        failed="Independently review payroll accounting",
    )
    await execute(
        c,
        h,
        eid,
        "country-workflow",
        {**payload, "period_start": "2025-12-01"},
        failed="within the engagement period",
    )
    result = await execute(c, h, eid, "country-workflow", payload)
    assert result["body"]["exception_count"] == 0 and not result["body"]["blocked"]
    assert len(result["body"]["rows"]) == 11
    assert result["collection_totals"]["payroll_posting_links"] == 12
    assert (await decision(c, h, eid, result)).status_code == 403
    response = await decision(c, reviewer[0], eid, result)
    assert response.status_code == 201, response.text
    result = response.json()
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    page = await c.get(prefix + "?limit=1&posting_offset=2&posting_limit=1", headers=h)
    assert page.status_code == 200, page.text
    assert len(page.json()["body"]["rows"]) == 1
    assert len(page.json()["body"]["payroll_posting_links"]) == 1
    assert page.json()["body"]["payroll_posting_links"][0]["source_row_id"] == "PAY-2"
    for view, count, kind in (
        ("accounts", 5, "account-rollforward"),
        ("postings", 6, "posting-comparison"),
    ):
        filtered = await c.get(prefix + f"?payroll_view={view}", headers=h)
        assert filtered.status_code == 200, filtered.text
        assert filtered.json()["collection_totals"]["rows"] == count
        assert all(row["kind"] == kind for row in filtered.json()["body"]["rows"])
    exported = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content))
    assert book["Payroll posting links"].max_row == 13
    closing_event = await c.get(
        f"/api/v1/intelligence/engagements/{eid}/events/{selected['closing']}", headers=h
    )
    withdrawn = await decision(c, reviewer[0], eid, closing_event.json(), "reject")
    assert withdrawn.status_code == 201, withdrawn.text
    stale = await c.get(prefix, headers=h)
    assert stale.json()["freshness"] == "STALE"
    stale_export = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert stale_export.status_code == 200
    stale_book = load_workbook(io.BytesIO(stale_export.content))
    cover = dict(stale_book["Source and review"].values)
    assert cover["Source freshness at export"] == "STALE"
    assert cover["Export status"] == "STALE_HISTORICAL_RESULT_REQUIRES_RECALCULATION"


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
async def test_subject_substitution_survives_review_paging_and_export(
    environment, monkeypatch, country, currency
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "subject-allocation-reviewer")
    await grant(c, h, base, reviewer)
    populations = population(currency)
    populations["posted"][0]["subject"] = "WRONG-EMPLOYEE"
    selected, _ = await prepare_sources(c, h, eid, currency, reviewer, populations)
    result = await execute(
        c,
        h,
        eid,
        "country-workflow",
        {
            "workflow": "payroll-ledger-rollforward",
            "source_ids": selected,
            "period_start": "2026-06-01",
            "as_of": "2026-06-30",
            "title": "SYNTHETIC employee substitution",
            "rationale": "Independently retained payroll and GL have equal totals but different employee attribution",
        },
    )
    rows = result["body"]["rows"]
    assert result["body"]["exception_count"] == 2
    assert all(row["state"] == "MATCHED" for row in rows if row["kind"] == "account-rollforward")
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    response = await c.get(prefix + "?payroll_view=postings&limit=100", headers=h)
    assert response.status_code == 200, response.text
    comparisons = response.json()["body"]["rows"]
    exceptions = {row["subject"]: row["state"] for row in comparisons if row["state"] != "MATCHED"}
    assert exceptions == {
        "EMP-1": "MISSING_GL_PAYROLL_MOVEMENT",
        "WRONG-EMPLOYEE": "UNEXPECTED_GL_PAYROLL_MOVEMENT",
    }
    exported = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content))
    # Full result rows and original source links both retain the misallocation.
    values = [value for sheet in book for row in sheet.values for value in row]
    assert "WRONG-EMPLOYEE" in values and "UNEXPECTED_GL_PAYROLL_MOVEMENT" in values
