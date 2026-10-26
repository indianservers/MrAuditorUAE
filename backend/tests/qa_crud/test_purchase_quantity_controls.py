"""UAE then India purchase-unit controls through persisted jobs, review and XLSX."""

import io
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_practice_audit_workflows import purchase

pytestmark = pytest.mark.skipif(os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required")


async def prepare(c, h, eid, reviewer, currency="AED", positions=1, receipt_unit="PCS"):
    selected, events = {}, {}
    for role, kind in (("orders", "order"), ("receipts", "receipt"), ("invoices", "invoice"), ("payments", "payment")):
        uom = receipt_unit if role == "receipts" else "PCS"
        rows = [purchase(kind, id=f"SYN-{kind}-{index:04d}", line_reference=f"LINE-{index:04d}", currency=currency, uom=uom) for index in range(positions)]
        controls = [{"order_id": "PO-1", "line_reference": f"LINE-{index:04d}", "vendor": "V1", "currency": currency, "uom": uom, "document_type": kind, "expected_rows": 1, "quantity": "2"} for index in range(positions)] if role != "payments" else []
        event = await source(c, h, eid, rows, "SYN controlled purchase " + role, "purchase-lines", {currency: str(100 * positions)}, purchase_quantity_controls=controls)
        reviewed = await decision(c, reviewer[0], eid, event)
        assert reviewed.status_code == 201, reviewed.text
        events[role] = reviewed.json()
        selected[role] = reviewed.json()["id"]
    return {"engagement": eid, "sources": selected, "events": events, "payload": {"workflow": "procurement-match", "source_ids": selected, "as_of": "2026-06-30", "title": "SYN controlled purchase quantity review", "rationale": "Independent order receipt invoice and cash controls; exact units, no inferred conversion"}}


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
async def test_reviewed_purchase_quantities_and_unit_mismatch_cannot_be_overridden(environment, monkeypatch, country, currency):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "purchase-source-reviewer")
    final = await actor(h, "Firm Administrator", "purchase-final-reviewer")
    for person in (reviewer, final):
        await grant(c, h, base, person)
    fixture = await prepare(c, h, eid, reviewer, currency)
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    assert result["body"]["exception_count"] == 0 and not result["body"]["blocked"]
    assert result["body"]["rows"][0]["uom"] == "PCS"
    assert (await decision(c, h, eid, result)).status_code == 403
    approved = await decision(c, final[0], eid, result)
    assert approved.status_code == 201, approved.text
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/"
    exported = await c.get(prefix + result["id"] + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert "PCS" in [value for values in book["Rows"].values for value in values]
    book.close()
    bad = await prepare(c, h, eid, reviewer, currency, receipt_unit="KG")
    result = await execute(c, h, eid, "country-workflow", bad["payload"])
    row = result["body"]["rows"][0]
    assert result["body"]["blocked"] and row["quantities"] == {} and row["quantity_comparison_status"] == "NOT_COMPARABLE"
    accepted = await decision(c, h, eid, result, action="accept", item_id=row["id"])
    assert accepted.status_code == 201, accepted.text
    denied = await decision(c, final[0], eid, accepted.json())
    assert denied.status_code == 409 and "Resolve missing evidence" in denied.text


async def test_purchase_source_controls_are_separately_paged_and_exported(environment, monkeypatch):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, base = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "purchase-page-reviewer")
    await grant(c, h, base, reviewer)
    fixture = await prepare(c, h, eid, reviewer, positions=121)
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/"
    identifier = fixture["sources"]["orders"]
    page = await c.get(prefix + identifier + "?limit=1&purchase_offset=100", headers=h)
    assert page.status_code == 200, page.text
    body = page.json()
    assert len(body["body"]["rows"]) == 1 and body["collection_totals"]["purchase_quantity_controls"] == 121
    assert len(body["body"]["purchase_quantity_controls"]) == 21 and body["body"]["purchase_quantity_controls"][-1]["line_reference"] == "LINE-0120"
    exported = await c.get(prefix + identifier + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    rows = list(book["Purchase quantity controls"].values)
    assert len(rows) == 122 and rows[-1][1] == "LINE-0120" and rows[-1][4] == "PCS"
    book.close()


async def test_unreviewed_purchase_source_is_not_usable(environment, monkeypatch):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, base = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "purchase-unreviewed-source-reviewer")
    await grant(c, h, base, reviewer)
    fixture = await prepare(c, h, eid, reviewer)
    event = await source(c, h, eid, [purchase("order")], "SYN unreviewed purchase source", "purchase-lines", {"AED": "100"}, purchase_quantity_controls=[{"order_id": "PO-1", "line_reference": "LINE-1", "vendor": "V1", "currency": "AED", "uom": "PCS", "document_type": "order", "expected_rows": 1, "quantity": "2"}])
    await execute(c, h, eid, "country-workflow", {**fixture["payload"], "source_ids": {**fixture["sources"], "orders": event["id"]}}, failed="Independently review all procurement sources")
