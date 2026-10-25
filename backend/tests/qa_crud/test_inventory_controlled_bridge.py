"""Original inventory evidence, independent controls/review and bilateral date bridge."""

import io
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant, setup
from tests.qa_crud.test_country_workflows import (
    csv_bytes,
    decision,
    double_boundaries,
    execute,
    source,
)
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_inventory_controlled_bridge import fixture as inventory_fixture

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable database required"
)


async def retain(c, h, eid, currency):
    rows, controls = inventory_fixture(currency)
    sources = {}
    for role, population in rows.items():
        sources[role] = await source(
            c,
            h,
            eid,
            population,
            "Synthetic controlled inventory " + role,
            "stock-movements" if role == "movements" else "stock",
            {},
            inventory_quantity_controls=controls[role],
        )
    return sources, rows, controls


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
async def test_controlled_inventory_review_export_and_original_replacement(
    environment, monkeypatch, country, currency
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "inventory-source-reviewer")
    await grant(c, h, base, reviewer)
    originals, rows, controls = await retain(c, h, eid, currency)
    payload = {
        "workflow": "inventory-count-bridge",
        "source_ids": {role: item["id"] for role, item in originals.items()},
        "as_of": "2026-06-30",
        "currency": currency,
        "review_through": "2026-07-03",
        "stock_bridge_mode": "bidirectional-controlled",
        "stock_cutoff_assessment": "Independent exclusive earlier day-end and inclusive later day-end lot movement populations assessed",
        "title": "Synthetic reviewed inventory forward/backward bridge",
        "rationale": "Separate original book, count and signed movement evidence controlled and reviewed by lot",
    }
    await execute(
        c, h, eid, "country-workflow", payload, failed="review every controlled inventory source"
    )
    for role, item in originals.items():
        assert item["body"]["inventory_controls_verified"] is True
        assert item["body"]["inventory_quantity_controls"] == controls[role]
        assert (await decision(c, h, eid, item)).status_code == 403
        approved = await decision(c, reviewer[0], eid, item)
        assert approved.status_code == 201, approved.text
        payload["source_ids"][role] = approved.json()["id"]
    result = await execute(c, h, eid, "country-workflow", payload)
    assert result["body"]["exception_count"] == 0 and len(result["body"]["rows"]) == 2
    assert result["body"]["count_bridges"][1]["direction"] == "ROLL_BACK"
    assert result["body"]["count_bridges"][1]["reporting_date_quantity"] == "20"
    approved = await decision(c, reviewer[0], eid, result)
    assert approved.status_code == 201, approved.text
    path = f"/api/v1/intelligence/engagements/{eid}"
    exported = await c.get(path + f"/events/{result['id']}/export?format=xlsx", headers=h)
    assert exported.status_code == 200 and exported.content[:2] == b"PK"
    workbook = load_workbook(io.BytesIO(exported.content))
    assert workbook["Count date bridges"].max_row == 3
    assert workbook["Inventory movement links"].max_row == 4
    bridges = await c.get(
        path + f"/events/{result['id']}?bridge_offset=1&bridge_limit=1", headers=h
    )
    assert bridges.status_code == 200 and len(bridges.json()["body"]["rows"]) == 2
    assert bridges.json()["collection_totals"]["count_bridges"] == 2
    assert bridges.json()["body"]["count_bridges"][0]["lot"] == "B"
    source_path = path + "/events/" + payload["source_ids"]["books"]
    controls_page = await c.get(source_path + "?inventory_offset=1&inventory_limit=1", headers=h)
    assert controls_page.status_code == 200 and len(controls_page.json()["body"]["rows"]) == 2
    assert controls_page.json()["body"]["inventory_quantity_controls"][0]["lot"] == "B"
    controlled_export = await c.get(source_path + "/export?format=xlsx", headers=h)
    assert controlled_export.status_code == 200
    assert (
        load_workbook(io.BytesIO(controlled_export.content))["Inventory quantity controls"].max_row
        == 3
    )
    original = next(
        item for item in originals["movements"]["dependencies"] if item["type"] == "evidence"
    )
    changed = [{**row, "quantity": "2"} if row["id"] == "M1" else row for row in rows["movements"]]
    replaced = await c.post(
        base + "/evidence-files/" + original["record_id"],
        headers=h,
        files={"file": ("SYNTHETIC-changed-inventory-moves.csv", csv_bytes(changed), "text/csv")},
    )
    assert replaced.status_code == 201
    history = await c.get(path + f"/events/{approved.json()['id']}", headers=h)
    assert history.status_code == 200 and history.json()["freshness"] == "STALE"
    assert history.json()["body"]["count_bridges"][1]["reporting_date_quantity"] == "20"
