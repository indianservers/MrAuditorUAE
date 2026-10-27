import json
import os
import uuid
from pathlib import Path

import pytest

from app.modules.enterprise_ai.security import serial
from app.modules.jurisdictions.ae.advanced import Schedule, calculate
from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import action, actor, approve, create, grant, setup
from tests.unit.test_uae_advanced import fixture

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Disposable database only"
)


async def test_all_new_cases_have_real_approved_engine_actions_and_lineage(environment):
    c = environment
    h, options, eid, base = await setup(c)
    reviewer = await actor(h, "Audit Manager", "advanced-reviewer")
    await grant(c, h, base, reviewer)
    partner = await actor(h, "Audit Partner", "advanced-partner")
    await grant(c, h, base, partner)
    # Synthetic reviewers are explicitly qualified for all specialist read scopes.
    from sqlalchemy import select

    from app.db.models.platform import Permission, RolePermission
    from app.db.session import Session, tenant_context

    async with Session() as db, db.begin():
        await tenant_context(db, h["X-Tenant-ID"])
        for rid in [reviewer[2], partner[2]]:
            for code in [
                "accounting.read",
                "vat.read",
                "tax.profile.read",
                "tp.read",
                "dmtt.read",
                "ifrs.read",
            ]:
                pid = await db.scalar(select(Permission.id).where(Permission.code == code))
                present = await db.scalar(
                    select(RolePermission.role_id).where(
                        RolePermission.role_id == uuid.UUID(rid),
                        RolePermission.permission_id == pid,
                    )
                )
                if not present:
                    db.add(RolePermission(role_id=uuid.UUID(rid), permission_id=pid))
    proof = await create(
        c,
        h,
        base,
        "working-papers",
        dict(
            folder="SYNTHETIC",
            sections=[
                dict(
                    heading="Evidence",
                    text="SYNTHETIC approved policy and population; test fixture only",
                )
            ],
            conclusion="Reviewed synthetic inputs only",
        ),
    )
    proof = await approve(c, h, reviewer[0], base, proof)
    proof = await action(c, partner[0], base, proof, "approve")
    route = f"/api/v1/ai/engagements/{eid}/use-cases"
    listing = await c.get(route, headers=h)
    assert listing.status_code == 200
    assert len(listing.json()["items"]) == 200
    receipts = []
    first = None
    payload = None
    for i in range(51, 101):
        data = fixture(i)
        data["policy"]["source_id"] = proof["id"]
        for row in data["rows"]:
            row["evidence_id"] = proof["id"]
        body = dict(
            folder="SYNTHETIC",
            sections=[dict(heading="Schedule", text=f"SYNTHETIC independent input case {i}")],
            conclusion="Input awaiting independent review",
            sources=[dict(kind="record", id=proof["id"])],
            advanced_input=data,
        )
        paper = await create(c, h, base, "working-papers", body)
        command = dict(
            action="advanced-review",
            record_id=paper["id"],
            idempotency_key=f"synthetic-advanced-{i}",
        )
        denied = await c.post(f"{route}/{i}/actions", headers=h, json=command)
        assert denied.status_code == 409, (i, denied.text)
        submitted = await action(c, h, base, paper, "submit")
        await action(c, h, base, submitted, "approve", 403)
        paper = await action(c, reviewer[0], base, submitted, "approve")
        paper = await action(c, partner[0], base, paper, "approve")
        result = await c.post(f"{route}/{i}/actions", headers=h, json=command)
        assert result.status_code == 201, (i, result.text)
        value = result.json()
        out = value["body"]["output"]
        assert out["results"] == serial(calculate(Schedule.model_validate(data)))
        assert out["entity_id"] == options["entity_id"]
        assert len(value["dependencies"]) >= 3
        repeated = await c.post(f"{route}/{i}/actions", headers=h, json=command)
        assert repeated.status_code == 201 and repeated.json()["id"] == value["id"]
        command["payload"] = {"rows": []}
        conflict = await c.post(f"{route}/{i}/actions", headers=h, json=command)
        assert conflict.status_code == 409
        if i == 51:
            first = value
            payload = {**command, "payload": {}}
        receipts.append(
            dict(
                use_case_id=i,
                unapproved=409,
                self_approval=403,
                engine_http=201,
                exact_result_match=True,
                idempotent=True,
            )
        )
    # Actual file readiness reads persisted population; a new draft invalidates the prior snapshot.
    file_check = await c.post(
        route + "/100/actions",
        headers=h,
        json=dict(action="file-readiness", idempotency_key="synthetic-file-readiness"),
    )
    assert file_check.status_code == 201, file_check.text
    assert file_check.json()["body"]["output"]["audit_opinion"] is None
    file_id = file_check.json()["id"]
    await create(
        c,
        h,
        base,
        "working-papers",
        dict(
            folder="SYN",
            sections=[dict(heading="New", text="SYN")],
            conclusion="Unreviewed new paper",
        ),
    )
    changed = await c.get(f"/api/v1/intelligence/engagements/{eid}/events/{file_id}", headers=h)
    assert changed.status_code == 200 and changed.json()["freshness"] == "STALE"
    # Draft paper metadata must not be promoted to approved source receipts in the assistant plan.
    command = await c.post(
        "/api/v1/ai/command",
        headers=h,
        json={
            "engagement_id": eid,
            "task": "uae-52",
            "prompt": "SYNTHETIC collect approved allocation evidence with drafts present",
            "idempotency_key": "synthetic-advanced-collection-with-drafts",
            "inputs": {},
        },
    )
    assert command.status_code == 201, command.text
    final = None
    for _ in command.json()["plan"]:
        final = await c.post(
            f"/api/v1/ai/engagements/{eid}/investigations/{command.json()['id']}/step", headers=h
        )
        assert final.status_code == 200, final.text
    assert final.json()["state"] in {"COMPLETED", "PARTIALLY_COMPLETED"}
    # Source reopen must invalidate the reviewed input and every dependent retained action.
    await action(c, reviewer[0], base, proof, "reopen")
    stale = await c.get(f"/api/v1/intelligence/engagements/{eid}/events/{first['id']}", headers=h)
    assert stale.status_code == 200 and stale.json()["freshness"] == "STALE", stale.text
    inbox = await c.get(route + "/51/alerts", headers=reviewer[0])
    assert inbox.status_code == 200 and len(inbox.json()["items"]) == 1
    assert inbox.json()["items"][0]["previous_approved_version_preserved"]
    assert inbox.json()["email_or_push_delivered"] is False
    repeat = await c.post(f"{route}/51/actions", headers=h, json=payload)
    assert repeat.status_code == 409, repeat.text
    wrong_tenant = await c.get(route, headers={**h, "X-Tenant-ID": str(uuid.uuid4())})
    assert wrong_tenant.status_code in {403, 404}
    target = (
        Path(__file__).resolve().parents[3]
        / "docs/mr-auditor-business/evidence/advanced-case-results.json"
    )
    target.write_text(
        json.dumps(
            dict(
                synthetic=True,
                cases=receipts,
                source_reopen_invalidates=True,
                tenant_isolation=True,
            ),
            indent=2,
        )
    )


async def test_original_xlsx_scanner_respects_quarantine_and_source_replacement(
    environment, monkeypatch, tmp_path
):
    import hashlib

    from app.core.config import settings
    from tests.unit.test_workbook_integrity import workbook

    c = environment
    h, _, eid, base = await setup(c)
    monkeypatch.setattr(settings(), "source_storage_path", str(tmp_path))
    reviewer = await actor(h, "Audit Manager", "workbook-reviewer")
    await grant(c, h, base, reviewer)
    proof = await create(
        c, h, base, "evidence", dict(category="Workbook", retain_until="2036-12-31")
    )
    data = workbook()
    uploaded = await c.post(
        base + "/evidence-files/" + proof["id"],
        headers=h,
        files={
            "file": (
                "synthetic.xlsx",
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    route = f"/api/v1/ai/engagements/{eid}/use-cases/97/actions"
    command = dict(
        action="workbook-integrity",
        record_id=proof["id"],
        idempotency_key="synthetic-workbook-scan",
    )
    assert (await c.post(route, headers=h, json=command)).status_code == 409
    # Explicit disposable scanner boundary: real versions remain immutable.
    import io

    from openpyxl import load_workbook

    monkeypatch.setattr("app.api.v1.audit_evidence.scan", lambda stream: "CLEAN")
    book = load_workbook(io.BytesIO(data))
    book.active["A1"] = 12
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    data = stream.getvalue()
    clean = await c.post(
        base + "/evidence-files/" + proof["id"],
        headers=h,
        files={
            "file": (
                "synthetic-clean.xlsx",
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert clean.status_code == 201 and clean.json()["scan_status"] == "CLEAN"
    proof = (await c.get(base + "/records/" + proof["id"], headers=h)).json()
    proof = await approve(c, h, reviewer[0], base, proof)
    result = await c.post(route, headers=h, json=command)
    assert result.status_code == 201, result.text
    out = result.json()["body"]["output"]
    assert out["source_sha256"] == hashlib.sha256(data).hexdigest()
    assert any(f["code"] == "circular_formula_reference" for f in out["findings"])
    # A distinct new version is quarantined and invalidates the retained approved-file scan.
    import io

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data))
    book.active["A1"] = 999
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    proof = await action(c, reviewer[0], base, proof, "reopen")
    monkeypatch.setattr("app.api.v1.audit_evidence.scan", lambda stream: "QUARANTINED")
    replacement = await c.post(
        base + "/evidence-files/" + proof["id"],
        headers=h,
        files={
            "file": (
                "synthetic-v2.xlsx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert replacement.status_code == 201, replacement.text
    detail = await c.get(
        f"/api/v1/intelligence/engagements/{eid}/events/{result.json()['id']}", headers=h
    )
    assert detail.status_code == 200 and detail.json()["freshness"] == "STALE"
