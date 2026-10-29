"""Synthetic tax-period/form guards through retained facts and independent review."""

import io
import os
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import action, actor, approve, create, grant, setup
from tests.qa_crud.regulatory_source_fixtures import retain
from tests.qa_crud.report_requirement_fixtures import declared_requirements
from tests.qa_crud.test_country_workflows import decision, double_boundaries
from tests.unit.test_regulatory_evaluation import fixture

pytestmark = pytest.mark.skipif(os.getenv("INTEGRATION_TESTS") != "1", reason="Dedicated test DB")


@pytest.mark.parametrize(
    "modern,other", [(False, True), (False, False), (True, True), (True, False)]
)
async def test_period_form_review_export_and_source_correction(
    environment, monkeypatch, modern, other
):
    c = environment
    double_boundaries(monkeypatch)
    h, opts, _, _ = await setup(c)
    start, end = ("2026-04-01", "2027-03-31") if modern else ("2025-04-01", "2026-03-31")
    form = "26" if modern else "3CA-3CD" if other else "3CB-3CD"
    part = ("A" if other else "B") if modern else None
    facts = {
        "tax_period_start": start,
        "tax_period_end": end,
        "tax_audit_form": form,
        "audited_under_other_law": other,
        **({"audit_report_part": part} if part else {}),
    }
    entity = await c.post(
        "/api/v1/entities",
        headers=h,
        json={
            "client_id": opts["client_id"],
            "name": "SYNTHETIC India tax-period entity",
            "country": "IN",
            "jurisdiction": "SYNTHETIC India",
            "functional_currency": "INR",
            "presentation_currency": "INR",
        },
    )
    assert entity.status_code == 201, entity.text
    period = await c.post(
        "/api/v1/financial-periods",
        headers=h,
        json={"entity_id": entity.json()["id"], "start_date": start, "end_date": end},
    )
    assert period.status_code == 201, period.text
    engagement = await c.post(
        "/api/v1/engagements",
        headers=h,
        json={
            "entity_id": entity.json()["id"],
            "period_id": period.json()["id"],
            "name": "SYNTHETIC tax period engagement",
            "engagement_type": "External Audit",
        },
    )
    assert engagement.status_code == 201, engagement.text
    eid = engagement.json()["id"]
    base = f"/api/v1/engagements/{eid}/audit"
    reviewer = await actor(h, "Firm Administrator", "tax-scope-reviewer")
    manager = await actor(h, "Audit Manager", "tax-scope-manager")
    partner = await actor(h, "Audit Partner", "tax-scope-partner")
    await grant(c, h, base, reviewer)
    await grant(c, h, base, manager)
    await grant(c, h, base, partner)
    context = f"/api/v1/jurisdictions/engagements/{eid}/context"
    proposed = await c.post(
        context,
        headers=h,
        json={
            "country": "IN",
            "framework": "Ind AS",
            "entity_type": "company",
            "as_of": end,
            "base_version": 0,
            "reason": "Synthetic independently assessed reporting jurisdiction",
        },
    )
    assert proposed.status_code == 201, proposed.text
    approved_context = await c.post(
        context + "/decision",
        headers=reviewer[0],
        json={
            "request_id": proposed.json()["id"],
            "approve": True,
            "comments": "Synthetic independent jurisdiction assessment",
        },
    )
    assert approved_context.status_code == 201, approved_context.text
    evidence = await create(
        c, h, base, "evidence", {"category": "Other", "retain_until": "2036-12-31"}
    )
    uploaded = await c.post(
        base + "/evidence-files/" + evidence["id"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-scope.txt",
                b"SYNTHETIC period and other-law basis; not a legally qualified audit report.",
                "text/plain",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    evidence = (await c.get(base + "/records/" + evidence["id"], headers=h)).json()
    evidence = await approve(c, h, reviewer[0], base, evidence)
    from app.modules.regulatory.report_packs import manifest

    source_scope = None
    if modern:
        item_pack = manifest(f"india-form26-index:{part}/1")
        source_scope = {
            "preset_id": item_pack["id"],
            "period_start": start,
            "period_end": end,
            "items": [
                {
                    "key": r["key"],
                    "disposition": "NOT_APPLICABLE"
                    if r["key"] == "form26.schedule.quantities"
                    else "REQUIRED",
                    "section_index": 0,
                    "supporting_sources": [{"kind": "record", "id": evidence["id"]}],
                    "assessment": "SYNTHETIC mapping control mechanics only; no substantive legal evidence qualification",
                }
                for r in item_pack["requirements"]
            ],
        }
    paper = await create(
        c,
        h,
        base,
        "working-papers",
        {
            "folder": "SYNTHETIC tax scope",
            "report_requirement_scope": source_scope,
            "sections": [
                {
                    "heading": "Assessed source facts",
                    "text": "Synthetic reviewed period and other-law audit facts",
                }
            ],
            "sources": [{"kind": "record", "id": evidence["id"]}],
            "financial_tables": [{"title": "SYNTHETIC retained scope facts", "values": facts}],
            "conclusion": "Synthetic selection only; no legal form or audit opinion qualification",
        },
    )
    paper = await action(c, h, base, paper, "submit")
    paper = await action(c, manager[0], base, paper, "approve")
    paper = await action(c, partner[0], base, paper, "approve")
    rule = fixture(
        country="IN",
        rule_key="IN-SYNTHETIC-TAX-SCOPE",
        framework="Ind AS",
        effective_from=start,
        effective_to=end,
        applicability=[],
        input_requirements=list(facts),
        validation_status="VALIDATED",
        professional_review_status="REVIEWED",
        last_reviewed="2026-10-27",
    ).model_dump(mode="json")
    rule["citations"][0]["url"] = "https://incometax.gov.in/"
    await retain(c, h, rule)
    created = await c.post("/api/v1/regulatory/rules", headers=h, json=rule)
    assert created.status_code == 201, created.text
    rule_id = created.json()["id"]
    vector = {
        "id": "SYNTHETIC-APPLICABLE",
        "kind": "checklist",
        "expected_reference": "Independent synthetic scope selection inputs",
        "period_start": start,
        "period_end": end,
        "facts": facts,
        "expected_applicability": "APPLICABLE",
    }
    for value, who, sequence in [("approve", reviewer[0], 1), ("publish", h, 2)]:
        response = await c.post(
            f"/api/v1/regulatory/rules/{rule_id}/decisions",
            headers=who,
            json={
                "action": value,
                "expected_sequence": sequence,
                "comments": "Synthetic mechanics only; not adopted Indian legal content",
                **({"validation_cases": [vector]} if value == "approve" else {}),
            },
        )
        assert response.status_code == 201, response.text
    prefix = f"/api/v1/regulatory/engagements/{eid}"
    case = await c.post(
        prefix + "/checklists",
        headers=h,
        json={
            "area": "india-tax-audit",
            "title": "SYNTHETIC assessed tax scope",
            "rationale": "Synthetic independently bound form and tax period",
            "items": [
                {
                    "clause": "SYNTHETIC-FORM-SELECTION",
                    "rule_id": rule_id,
                    "state": "PREPARED",
                    "facts": facts,
                    "fact_sources": {
                        key: {
                            "kind": "record",
                            "id": paper["id"],
                            "path": ["financial_tables", 0, "values", key],
                            "comparison": "exact",
                        }
                        for key in facts
                    },
                    "procedure": "Synthetic retained source facts independently compared",
                    "conclusion": "Synthetic chosen period and other-law requirement agree",
                    "evidence_ids": [evidence["id"]],
                }
            ],
        },
    )
    assert case.status_code == 201, case.text
    scope = {
        "applicability_id": case.json()["id"],
        "tax_period_start": start,
        "tax_period_end": end,
        "audited_under_other_law": other,
        "form": form,
        "audit_report_part": part,
    }
    data = {
        "title": "SYNTHETIC source-linked tax working draft",
        "category": "india-tax-audit",
        **declared_requirements("scope"),
        "tax_audit_scope": scope,
        "blocks": [
            {
                "key": "scope",
                "requirement_keys": ["scope"],
                "title": "Retained scope assessment",
                "kind": "regulatory_checklist",
                "source_kind": "intelligence",
                "source_id": case.json()["id"],
            }
        ],
        "rationale": "Synthetic reviewed period selection; not an official form",
    }
    denied = await c.post(prefix + "/report-templates", headers=h, json=data)
    assert denied.status_code == 409  # independent case review is mandatory
    reviewed_case = await decision(c, reviewer[0], eid, case.json())
    assert reviewed_case.status_code == 201, reviewed_case.text
    data["tax_audit_scope"]["applicability_id"] = reviewed_case.json()["id"]
    data["blocks"][0]["source_id"] = reviewed_case.json()["id"]
    if modern:
        outdated = deepcopy(data)
        outdated["tax_audit_scope"]["audit_report_part"] = "C" if other else "D"
        assert (
            await c.post(prefix + "/report-templates", headers=h, json=outdated)
        ).status_code == 422
    wrong = deepcopy(data)
    wrong["tax_audit_scope"]["form"] = "3CB-3CD" if modern else "26"
    assert (await c.post(prefix + "/report-templates", headers=h, json=wrong)).status_code == 422
    outside = deepcopy(data)
    outside["tax_audit_scope"]["tax_period_start"] = "2027-04-01" if modern else "2024-04-01"
    outside["tax_audit_scope"]["tax_period_end"] = "2028-03-31" if modern else "2025-03-31"
    assert (await c.post(prefix + "/report-templates", headers=h, json=outside)).status_code == 422
    unbound_period = deepcopy(data)
    unbound_period["tax_audit_scope"]["tax_period_start"] = "2026-04-02" if modern else "2025-04-02"
    assert (
        await c.post(prefix + "/report-templates", headers=h, json=unbound_period)
    ).status_code == 409
    template = await c.post(prefix + "/report-templates", headers=h, json=data)
    assert template.status_code == 201, template.text
    reviewed_template = await decision(c, reviewer[0], eid, template.json())
    assert reviewed_template.status_code == 201, reviewed_template.text
    generated = await c.post(
        prefix + "/reports",
        headers=h,
        json={
            "template_id": reviewed_template.json()["id"],
            "rationale": "Synthetic independently reviewed tax-period report composition",
        },
    )
    assert generated.status_code == 201, generated.text
    saved = generated.json()
    assert saved["body"]["tax_audit_scope"] == data["tax_audit_scope"]
    assert saved["body"]["tax_form_content_qualified"] is False
    assert saved["body"]["tax_form_scope_contract"] == "assessed-tax-period-form/2"
    if modern:
        assert (
            "https://www.incometaxindia.gov.in/w/rule-47-5"
            in saved["body"]["tax_form_selection_sources"]
        )
    endpoint = f"/api/v1/intelligence/engagements/{eid}/events/{saved['id']}"
    exported = await c.get(endpoint + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    metadata = dict(book["Report and review"].values)
    assert metadata["Selected tax-audit form"] == form
    assert metadata["Tax period start"] == start and metadata["Tax period end"] == end
    book.close()
    indexed_report = None
    if modern:
        catalogue = await c.get(prefix + "/report-presets", headers=h)
        assert catalogue.status_code == 200, catalogue.text
        pack = next(p for p in catalogue.json()["items"] if p["report_part"] == part)
        assert len(pack["requirements"]) == 65
        invalid_scope = deepcopy(source_scope)
        invalid_scope["items"][0]["section_index"] = 5
        invalid_paper = await c.post(
            base + "/working-papers",
            headers=h,
            json={
                "reference": "SYNTHETIC-invalid-item-scope",
                "title": "SYNTHETIC invalid section declaration",
                "body": {
                    "folder": "SYNTHETIC scope validation",
                    "sections": [
                        {
                            "heading": "One section",
                            "text": "Synthetic substantive text and source assessment",
                        }
                    ],
                    "sources": [{"kind": "record", "id": evidence["id"]}],
                    "report_requirement_scope": invalid_scope,
                    "conclusion": "Synthetic invalid item section rejection fixture",
                },
            },
        )
        assert invalid_paper.status_code == 422, invalid_paper.text
        # Consolidated synthetic workpaper mappings exercise controls, not substantive legal coverage.
        indexed = {
            **data,
            "requirements_pack": pack["id"],
            "requirements": [{**r, "disposition": "REQUIRED"} for r in pack["requirements"]],
            "requirements_basis": "Synthetic numbered-item mapping mechanics only; not substantive form completion",
            "blocks": [
                {
                    "key": "index",
                    "title": "Synthetic consolidated workpaper mapping",
                    "kind": "working_paper_reference",
                    "source_kind": "record",
                    "source_id": paper["id"],
                    "requirement_keys": [r["key"] for r in pack["requirements"]],
                }
            ],
        }
        next(r for r in indexed["requirements"] if r["key"] == "form26.schedule.quantities")[
            "disposition"
        ] = "NOT_APPLICABLE"
        omitted = await c.post(
            prefix + "/report-templates",
            headers=h,
            json={**indexed, "requirements": indexed["requirements"][:-1]},
        )
        assert omitted.status_code == 422
        generic = await c.post(
            prefix + "/report-templates",
            headers=h,
            json={
                **indexed,
                "blocks": [
                    {
                        **indexed["blocks"][0],
                        "kind": "regulatory_checklist",
                        "source_kind": "intelligence",
                        "source_id": reviewed_case.json()["id"],
                    }
                ],
            },
        )
        assert generic.status_code == 201, generic.text
        assert (await decision(c, reviewer[0], eid, generic.json())).status_code == 409
        unscoped = await create(
            c,
            h,
            base,
            "working-papers",
            {
                "folder": "SYNTHETIC no item declaration",
                "sections": [
                    {
                        "heading": "General procedures",
                        "text": "Synthetic approved paper without report item declarations",
                    }
                ],
                "sources": [{"kind": "record", "id": evidence["id"]}],
                "conclusion": "Synthetic source without declared clause coverage",
            },
        )
        unscoped = await action(c, h, base, unscoped, "submit")
        unscoped = await action(c, manager[0], base, unscoped, "approve")
        unscoped = await action(c, partner[0], base, unscoped, "approve")
        undeclared = await c.post(
            prefix + "/report-templates",
            headers=h,
            json={**indexed, "blocks": [{**indexed["blocks"][0], "source_id": unscoped["id"]}]},
        )
        assert undeclared.status_code == 201, undeclared.text
        rejected = await decision(c, reviewer[0], eid, undeclared.json())
        assert rejected.status_code == 409 and "Declare item coverage" in rejected.text
        prepared = await c.post(prefix + "/report-templates", headers=h, json=indexed)
        assert prepared.status_code == 201, prepared.text
        assert (await decision(c, h, eid, prepared.json())).status_code == 403
        reviewed = await decision(c, reviewer[0], eid, prepared.json())
        assert reviewed.status_code == 201, reviewed.text
        output = await c.post(
            prefix + "/reports",
            headers=h,
            json={
                "template_id": reviewed.json()["id"],
                "rationale": "Synthetic numbered-item evidence index mechanics only",
            },
        )
        assert output.status_code == 201, output.text
        indexed_report = output.json()
        assert len(indexed_report["body"]["requirements_coverage"]["items"]) == 65
        assert (
            indexed_report["body"]["requirements_coverage"]["requirements_pack"]["sha256"]
            == pack["sha256"]
        )
        assert not indexed_report["body"]["tax_form_content_qualified"]
        export = await c.get(
            f"/api/v1/intelligence/engagements/{eid}/events/{indexed_report['id']}/export?format=xlsx",
            headers=h,
        )
        assert export.status_code == 200
        index_book = load_workbook(io.BytesIO(export.content), read_only=True)
        assert len(list(index_book["Requirements coverage"].values)) == 66
        index_book.close()
    if modern:
        await expanded_field_report(
            c, h, eid, base, reviewer, manager, partner, evidence, data, part, start, end
        )
    historical = deepcopy(saved["body"])
    await action(c, partner[0], base, paper, "reopen")
    stale = await c.get(endpoint, headers=h)
    assert stale.status_code == 200 and stale.json()["freshness"] == "STALE"
    assert stale.json()["body"] == historical
    if indexed_report:
        stale_index = await c.get(
            f"/api/v1/intelligence/engagements/{eid}/events/{indexed_report['id']}", headers=h
        )
        assert (
            stale_index.json()["freshness"] == "STALE"
            and stale_index.json()["body"] == indexed_report["body"]
        )


async def expanded_field_report(
    c, h, eid, base, reviewer, manager, partner, evidence, data, part, start, end
):
    """Full retained field population across bounded papers; synthetic control evidence only."""
    from app.modules.regulatory.report_packs import manifest

    pack = manifest(f"india-form26-index:{part}/2")
    blocks, papers = [], []
    for offset in range(0, len(pack["requirements"]), 100):
        group = pack["requirements"][offset : offset + 100]
        paper = await create(
            c,
            h,
            base,
            "working-papers",
            {
                "folder": "SYNTHETIC expanded field mapping controls",
                "sections": [
                    {
                        "heading": "Synthetic source and field scope",
                        "text": "SYNTHETIC mapping mechanics; no substantive professional field sufficiency or legal coverage is established",
                    }
                ],
                "sources": [{"kind": "record", "id": evidence["id"]}],
                "conclusion": "Synthetic exact field mapping control test only",
                "report_requirement_scope": {
                    "preset_id": pack["id"],
                    "period_start": start,
                    "period_end": end,
                    "items": [
                        {
                            "key": r["key"],
                            "disposition": "REQUIRED",
                            "section_index": 0,
                            "supporting_sources": [{"kind": "record", "id": evidence["id"]}],
                            "assessment": "Synthetic declaration plumbing only, no substantive client or statutory qualification",
                        }
                        for r in group
                    ],
                },
            },
        )
        paper = await action(c, h, base, paper, "submit")
        paper = await action(c, manager[0], base, paper, "approve")
        paper = await action(c, partner[0], base, paper, "approve")
        papers.append(paper)
        blocks.append(
            {
                "key": f"fields-{offset}",
                "title": "SYNTHETIC field coverage chunk",
                "kind": "working_paper_reference",
                "source_kind": "record",
                "source_id": paper["id"],
                "requirement_keys": [r["key"] for r in group],
            }
        )
    raw = {
        **data,
        "requirements_pack": pack["id"],
        "requirements": [{**r, "disposition": "REQUIRED"} for r in pack["requirements"]],
        "requirements_basis": "Synthetic full inventory plumbing, not professional assessment of statutory content",
        "blocks": blocks,
    }
    prefix = f"/api/v1/regulatory/engagements/{eid}"
    invalid = deepcopy(raw)
    invalid["requirements"].pop()
    assert (await c.post(prefix + "/report-templates", headers=h, json=invalid)).status_code == 422
    response = await c.post(prefix + "/report-templates", headers=h, json=raw)
    assert response.status_code == 201, response.text
    assert (await decision(c, h, eid, response.json())).status_code == 403
    reviewed = await decision(c, reviewer[0], eid, response.json())
    assert reviewed.status_code == 201, reviewed.text
    output = await c.post(
        prefix + "/reports",
        headers=h,
        json={
            "template_id": reviewed.json()["id"],
            "rationale": "Synthetic complete published inventory, not statutory calculations",
        },
    )
    assert output.status_code == 201, output.text
    saved = output.json()
    body = deepcopy(saved["body"])
    coverage = body["requirements_coverage"]
    assert len(coverage["items"]) == 877 and coverage["complete"]
    assert coverage["requirements_pack"]["sha256"] == pack["sha256"]
    assert (
        not coverage["statutory_completeness_qualified"] and not body["tax_form_content_qualified"]
    )
    endpoint = f"/api/v1/intelligence/engagements/{eid}/events/{saved['id']}"
    export = await c.get(endpoint + "/export?format=xlsx", headers=h)
    assert export.status_code == 200, export.text
    book = load_workbook(io.BytesIO(export.content), read_only=True)
    rows = list(book["Requirements coverage"].values)
    assert len(rows) == 878
    assert {row[0] for row in rows[1:]} == {r["key"] for r in pack["requirements"]}
    book.close()
    await action(c, partner[0], base, papers[-1], "reopen")
    stale = await c.get(endpoint, headers=h)
    assert stale.status_code == 200 and stale.json()["freshness"] == "STALE"
    assert stale.json()["body"] == body


async def test_large_template_allowance_does_not_relax_other_routes_or_authentication(environment):
    c = environment
    path = "/api/v1/regulatory/engagements/00000000-0000-0000-0000-000000000001/report-templates"
    ordinary = await c.post(
        "/api/v1/entities", content=b" " * 65537, headers={"Content-Type": "application/json"}
    )
    assert ordinary.status_code == 413
    oversized = await c.post(
        path, content=b" " * (8 * 1024 * 1024 + 1), headers={"Content-Type": "application/json"}
    )
    assert oversized.status_code == 413
    unauthorized = await c.post(path, json={"title": "No authentication"})
    assert unauthorized.status_code == 401
