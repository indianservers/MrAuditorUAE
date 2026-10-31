import csv
import io
import json
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant
from tests.qa_crud.test_country_workflows import decision, double_boundaries
from tests.qa_crud.test_uae_company_accounts import context
from tests.qa_crud.test_uae_ct_audit_duty import CLARIFICATION, upload
from tests.unit.test_uae_tax_group_statements import existing_group, line
from tests.unit.test_uae_tax_group_statements import sample as paper_sample
from tests.unit.test_uae_tax_group_workpaper import sample

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required"
)
FTA7 = CLARIFICATION.parent / "fta-decision7-2025-unofficial-english.pdf"


async def documents(
    c,
    h,
    base,
    *,
    mapping_headers=None,
    comparative=False,
    wrong_original=False,
    broken_mapping=False,
    quarantine_member=False,
    monkeypatch=None,
):
    files = {}
    for role, path in (("fta7", FTA7), ("ctp007", CLARIFICATION)):
        files[role] = await upload(
            c,
            h,
            base,
            path.name,
            path.read_bytes() + (b"\n" if wrong_original and role == "fta7" else b""),
            "application/pdf",
        )
    files["adoption"] = await upload(
        c,
        h,
        base,
        "SYNTHETIC-tax-group-adoption.txt",
        b"SYNTHETIC independently declared adoption of FTA7/CTP007 for the full completed period; no legal qualification",
        "text/plain",
    )
    paper = existing_group() if comparative else paper_sample()
    for period_name, source_period in (
        ("current", paper["current"]),
        ("comparative", paper["comparative"]),
    ):
        if source_period is None:
            continue
        for index, member in enumerate(source_period["members"]):
            if period_name == "current" and index == 1 and quarantine_member:
                double_boundaries(monkeypatch, scan="QUARANTINED")
            member["lines"] += [
                line(f"ZERO-{index}-{i}", "profit_loss", "expense", 0) for i in range(55)
            ]
            text = io.StringIO()
            writer = csv.writer(text)
            writer.writerow(
                ["member", "period_start", "period_end", "reference", "statement", "kind", "amount"]
            )
            writer.writerows(
                [
                    member["member"],
                    source_period["period_start"],
                    source_period["period_end"],
                    *[row[k] for k in ("reference", "statement", "kind", "amount")],
                ]
                for row in member["lines"]
            )
            role = member["member"] if period_name == "current" else "prior:" + member["member"]
            files[role] = await upload(
                c,
                h,
                base,
                f"SYNTHETIC-{period_name}-member-{member['member']}.csv",
                text.getvalue().encode(),
                "text/csv",
            )
            version = files[role][1]
            member.update(source_evidence_version_id=version["id"], source_sha256=version["sha256"])
    double_boundaries(monkeypatch)
    raw = b'{"contract":"bad"}' if broken_mapping else json.dumps(paper).encode()
    files["mapping"] = await upload(
        c,
        mapping_headers or h,
        base,
        "SYNTHETIC-independent-tax-group-mapping.json",
        raw,
        "application/json",
    )
    data = sample(
        rule_evidence_version_id=files["fta7"][1]["id"],
        clarification_evidence_version_id=files["ctp007"][1]["id"],
        adoption_evidence_version_id=files["adoption"][1]["id"],
        adoption_sha256=files["adoption"][1]["sha256"],
        mapping_evidence_version_id=files["mapping"][1]["id"],
        mapping_sha256=files["mapping"][1]["sha256"],
    )
    if comparative:
        data["adopted_from"] = "2024-01-01"
    return data, files


@pytest.mark.parametrize("replaced", ["fta7", "ctp007", "adoption", "mapping", "A", "B"])
async def test_original_binding_separate_review_complete_export_paging_and_staleness(
    environment, monkeypatch, replaced
):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await context(c)
    mapping_preparer = await actor(h, "Firm Administrator", "tax-group-mapping-preparer")
    reviewer = await actor(h, "Firm Administrator", "tax-group-independent-reviewer")
    await grant(c, h, base, mapping_preparer)
    await grant(c, h, base, reviewer)
    data, files = await documents(
        c, h, base, mapping_headers=mapping_preparer[0], monkeypatch=monkeypatch
    )
    prefix = f"/api/v1/intelligence/engagements/{eid}"
    response = await c.post(prefix + "/uae-tax-group-statements", headers=h, json=data)
    assert response.status_code == 201, response.text
    result = response.json()
    body = result["body"]
    assert body["tax_group_totals"][0]["assets"] == "220"
    assert (
        body["source_binding_established"]
        and not body["validation"]["ready"]
        and not body["statutory_issuance"]
    )
    assert result["collection_totals"]["tax_group_source_lines"] == 131
    assert len(body["tax_group_source_lines"]) == 100
    assert (await decision(c, h, eid, result)).status_code == 403
    assert (await decision(c, mapping_preparer[0], eid, result)).status_code == 403
    approved = await decision(c, reviewer[0], eid, result)
    assert approved.status_code == 201, approved.text
    reviewed = approved.json()
    path = prefix + "/events/" + reviewed["id"]
    page = await c.get(
        path + "/tax-group-collections?collection_name=sources&offset=125&limit=25", headers=h
    )
    assert (
        page.status_code == 200 and page.json()["total"] == 131 and len(page.json()["items"]) == 6
    )
    notes = await c.get(
        path + "/tax-group-collections?collection_name=notes&period=current", headers=h
    )
    assert len(notes.json()["items"]) == 4
    exported = await c.get(path + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
    source_rows = list(workbook["Tax Group sources"].values)
    assert len(source_rows) == 132
    assert (
        sum(row[source_rows[0].index("reference")] == "ZERO-1-54" for row in source_rows[1:]) == 1
    )
    assert len(list(workbook["Tax Group documents"].values)) == 7
    assert len(list(workbook["Tax Group notes"].values)) == 5
    workbook.close()
    changed = await c.post(
        base + "/evidence-files/" + files[replaced][0]["id"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-revised.txt",
                b"SYNTHETIC revised retained Tax Group source requires reassessment",
                "text/plain",
            )
        },
    )
    assert changed.status_code == 201, changed.text
    stale = await c.get(path, headers=h)
    assert stale.status_code == 200 and stale.json()["freshness"] == "STALE"
    assert stale.json()["body"]["tax_group_totals"] == reviewed["body"]["tax_group_totals"]


@pytest.mark.parametrize(
    "failure",
    [
        "wrong-original",
        "broken-mapping",
        "quarantined-member",
        "hash",
        "wrong-period",
        "unresolved",
    ],
)
async def test_invalid_sources_and_unresolved_adoption_cannot_be_approved(
    environment, monkeypatch, failure
):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await context(c)
    reviewer = await actor(h, "Firm Administrator", "tax-group-negative-reviewer")
    await grant(c, h, base, reviewer)
    data, _ = await documents(
        c,
        h,
        base,
        wrong_original=failure == "wrong-original",
        broken_mapping=failure == "broken-mapping",
        quarantine_member=failure == "quarantined-member",
        monkeypatch=monkeypatch,
    )
    if failure == "hash":
        data["mapping_sha256"] = "0" * 64
    if failure == "wrong-period":
        data["assessed_on"] = "2025-07-16"
    if failure == "unresolved":
        data["operative_changes"] = "unresolved"
    prefix = f"/api/v1/intelligence/engagements/{eid}"
    response = await c.post(prefix + "/uae-tax-group-statements", headers=h, json=data)
    if failure == "unresolved":
        assert response.status_code == 201, response.text
        assert response.json()["body"]["blocked"]
        assert (await decision(c, reviewer[0], eid, response.json())).status_code == 409
    else:
        assert response.status_code in {409, 422}, response.text


async def test_evidence_choices_are_current_clean_paginated_and_hash_filtered(
    environment, monkeypatch
):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await context(c)
    versions = []
    for i in range(27):
        _, version = await upload(
            c,
            h,
            base,
            f"SYNTHETIC-choice-{i}.txt",
            f"SYNTHETIC distinct retained evidence choice {i}".encode(),
            "text/plain",
        )
        versions.append(version)
    prefix = f"/api/v1/intelligence/engagements/{eid}"
    first = await c.get(prefix + "/evidence-choices?limit=25", headers=h)
    last = await c.get(prefix + "/evidence-choices?offset=25&limit=25", headers=h)
    assert first.status_code == 200 and first.json()["total"] == 27
    assert len(last.json()["items"]) == 2
    selected = await c.get(prefix + "/evidence-choices?sha256=" + versions[0]["sha256"], headers=h)
    assert [v["id"] for v in selected.json()["items"]] == [versions[0]["id"]]
    double_boundaries(monkeypatch, scan="QUARANTINED")
    revised = await c.post(
        base + "/evidence-files/" + versions[0]["record_id"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-quarantined-revision.txt",
                b"SYNTHETIC quarantined newer original",
                "text/plain",
            )
        },
    )
    assert revised.status_code == 201, revised.text
    old = await c.get(prefix + "/evidence-choices?sha256=" + versions[0]["sha256"], headers=h)
    assert old.status_code == 200 and old.json()["total"] == 0
    schema = await c.get(prefix + "/tax-group-mapping-schema", headers=h)
    assert schema.status_code == 200 and "PeriodStatements" in schema.json()["$defs"]


async def test_another_engagement_original_cannot_be_used_as_a_group_owned_source(
    environment, monkeypatch
):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await context(c)
    data, files = await documents(c, h, base, monkeypatch=monkeypatch)
    root = (await c.get(f"/api/v1/engagements/{eid}", headers=h)).json()
    other = await c.post(
        "/api/v1/engagements",
        headers=h,
        json={
            "entity_id": root["entity_id"],
            "period_id": root["period_id"],
            "name": "SYNTHETIC separate engagement",
            "engagement_type": "External Audit",
        },
    )
    assert other.status_code == 201, other.text
    _, foreign = await upload(
        c,
        h,
        f"/api/v1/engagements/{other.json()['id']}/audit",
        "SYNTHETIC-other-file.txt",
        b"SYNTHETIC other engagement standalone source",
        "text/plain",
    )
    paper = paper_sample()
    paper["current"]["members"][0].update(
        source_evidence_version_id=foreign["id"], source_sha256=foreign["sha256"]
    )
    paper["current"]["members"][1].update(
        source_evidence_version_id=files["B"][1]["id"], source_sha256=files["B"][1]["sha256"]
    )
    _, mapping = await upload(
        c,
        h,
        base,
        "SYNTHETIC-foreign-reference.json",
        json.dumps(paper).encode(),
        "application/json",
    )
    data.update(mapping_evidence_version_id=mapping["id"], mapping_sha256=mapping["sha256"])
    response = await c.post(
        f"/api/v1/intelligence/engagements/{eid}/uae-tax-group-statements", headers=h, json=data
    )
    assert response.status_code == 404, response.text


async def test_comparative_period_uses_its_own_sources_and_prior_replacement_stales_review(
    environment, monkeypatch
):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await context(c)
    reviewer = await actor(h, "Firm Administrator", "tax-group-comparative-reviewer")
    await grant(c, h, base, reviewer)
    data, files = await documents(c, h, base, comparative=True, monkeypatch=monkeypatch)
    prefix = f"/api/v1/intelligence/engagements/{eid}"
    created = await c.post(prefix + "/uae-tax-group-statements", headers=h, json=data)
    assert created.status_code == 201, created.text
    result = created.json()
    assert result["body"]["tax_group_totals"][0]["opening_equity"] == "150"
    assert result["body"]["tax_group_totals"][1]["closing_equity"] == "150"
    assert not result["body"]["blocked"]
    approved = await decision(c, reviewer[0], eid, result)
    assert approved.status_code == 201, approved.text
    path = prefix + "/events/" + approved.json()["id"]
    prior = await c.get(
        path
        + "/tax-group-collections?collection_name=sources&period=comparative&offset=100&limit=100",
        headers=h,
    )
    assert (
        prior.status_code == 200
        and prior.json()["total"] == 131
        and len(prior.json()["items"]) == 31
    )
    export = await c.get(path + "/export?format=xlsx", headers=h)
    assert export.status_code == 200, export.text
    book = load_workbook(io.BytesIO(export.content), read_only=True)
    assert len(list(book["Tax Group sources"].values)) == 263
    assert len(list(book["Tax Group notes"].values)) == 9
    book.close()
    changed = await c.post(
        base + "/evidence-files/" + files["prior:A"][0]["id"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-revised-prior.txt",
                b"SYNTHETIC revised prior period standalone facts",
                "text/plain",
            )
        },
    )
    assert changed.status_code == 201, changed.text
    stale = await c.get(path, headers=h)
    assert stale.status_code == 200 and stale.json()["freshness"] == "STALE"
    assert stale.json()["body"]["tax_group_totals"] == approved.json()["body"]["tax_group_totals"]


async def test_india_engagement_cannot_prepare_a_uae_tax_group_workpaper(environment, monkeypatch):
    from tests.qa_crud.test_india_ecr_source import india_context

    double_boundaries(monkeypatch)
    h, eid, _ = await india_context(environment)
    response = await environment.post(
        f"/api/v1/intelligence/engagements/{eid}/uae-tax-group-statements", headers=h, json=sample()
    )
    assert response.status_code == 422 and "UAE engagement" in response.text
