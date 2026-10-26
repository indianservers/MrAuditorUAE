"""Retained ownership source review, exact paths, independent approval and exports."""

import io
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_diligence_audit_workflows import edge, policy
from tests.unit.test_ownership_paths import population

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required"
)


async def prepare(c, h, eid, reviewer, country="AE", case="complete", reviewed=True, count=None):
    rows = population(country)
    if case == "incomplete":
        rows["ownership"] = [edge("A", "P1", "0.5"), edge("A", "MISSING", "0.5", "company")]
    if count:
        rows["ownership"] = [
            edge("A", "B", "1", "company"),
            *[edge("B", f"P{index:04d}", "0.0078125") for index in range(count)],
        ]
    events = {}
    for role, kind in (("ownership", "ownership-edges"), ("policy", "ownership-policies")):
        event = await source(
            c, h, eid, rows[role], "SYNTHETIC ownership " + case + " " + role, kind, {}
        )
        if reviewed:
            response = await decision(c, reviewer[0], eid, event)
            assert response.status_code == 201, response.text
            event = response.json()
        events[role] = event
    return dict(
        engagement=eid,
        events=events,
        sources={role: event["id"] for role, event in events.items()},
        payload=dict(
            workflow="ownership-lookthrough",
            source_ids={role: event["id"] for role, event in events.items()},
            subjects=["A"],
            as_of="2026-06-30",
            title="SYNTHETIC " + case + " economic ownership paths",
            rationale="Retained supplied economic graph and independently reviewed fictional country policy; no legal threshold qualification",
        ),
    )


@pytest.mark.parametrize("country", ["AE", "IN"])
@pytest.mark.parametrize("case", ["complete", "incomplete"])
async def test_ownership_review_paths_independent_approval_export_and_staleness(
    environment, monkeypatch, country, case
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "ownership-source-reviewer")
    final = await actor(h, "Firm Administrator", "ownership-final-reviewer")
    for person in (reviewer, final):
        await grant(c, h, base, person)
    fixture = await prepare(c, h, eid, reviewer, country, case, reviewed=False)
    await execute(
        c, h, eid, "country-workflow", fixture["payload"], failed="Independently review ownership"
    )
    for role, event in fixture["events"].items():
        response = await decision(c, reviewer[0], eid, event)
        assert response.status_code == 201, response.text
        fixture["events"][role] = response.json()
        fixture["payload"]["source_ids"][role] = response.json()["id"]
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    assert result["body"]["ownership_path_contract"] == "economic-ownership-paths/1"
    assert result["body"]["ownership_scope"]["country"] == country
    for row in list(result["body"]["rows"]):
        if row["state"] != "MATCHED":
            accepted = await decision(c, h, eid, result, action="accept", item_id=row["id"])
            assert accepted.status_code == 201, accepted.text
            result = accepted.json()
    approved = await decision(c, final[0], eid, result)
    if case == "incomplete":
        assert approved.status_code == 409
        person = next(row for row in result["body"]["rows"] if row.get("person_id"))
        assert person["threshold_candidate"] is None
        assert any(
            path["gap_reason"] == "CORPORATE_OWNER_EVIDENCE_MISSING"
            for path in result["body"]["ownership_paths"]
        )
    else:
        assert approved.status_code == 201, approved.text
        result = approved.json()
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    exported = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert len(list(book["Ownership paths"].values)) == len(result["body"]["ownership_paths"]) + 1
    assert (
        len(list(book["Ownership path steps"].values))
        == len(result["body"]["ownership_path_steps"]) + 1
    )
    book.close()
    changed = await decision(c, reviewer[0], eid, fixture["events"]["policy"])
    assert changed.status_code == 201, changed.text
    stale = await c.get(prefix, headers=h)
    assert stale.json()["freshness"] == "STALE"


async def test_ownership_path_and_step_pagers_are_independent_and_export_all_records(
    environment, monkeypatch
):
    c = environment
    double_boundaries(monkeypatch)
    h, _, eid, base = await setup(c)
    reviewer = await actor(h, "Firm Administrator", "ownership-pagination-reviewer")
    await grant(c, h, base, reviewer)
    fixture = await prepare(c, h, eid, reviewer, count=128)
    result = await execute(c, h, eid, "country-workflow", fixture["payload"])
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    page = await c.get(
        prefix
        + "?limit=1&ownership_path_offset=127&ownership_path_limit=1&ownership_step_offset=255&ownership_step_limit=1",
        headers=h,
    )
    assert page.status_code == 200, page.text
    body = page.json()["body"]
    assert body["ownership_paths"][0]["person_id"] == "P0127"
    assert body["ownership_path_steps"][0]["edge_source_id"] == "B:P0127"
    assert page.json()["collection_totals"]["ownership_paths"] == 128
    assert page.json()["collection_totals"]["ownership_path_steps"] == 256
    exported = await c.get(prefix + "/export?format=xlsx", headers=h)
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert len(list(book["Ownership paths"].values)) == 129
    assert len(list(book["Ownership path steps"].values)) == 257
    book.close()
