"""Synthetic FIFO through real DB, HTTP, storage and review; scanner/broker doubled."""

import io
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, create, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_inventory_fifo import book, control, event

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable database required"
)


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
async def test_fifo_review_complete_export_paging_and_original_support_staleness(
    environment, monkeypatch, country, currency
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "fifo-independent-source-reviewer")
    await grant(c, h, base, reviewer)
    original = await create(
        c, h, base, "evidence", {"category": "Other", "retain_until": "2036-12-31"}
    )
    uploaded = await c.post(
        base + "/evidence-files/" + original["id"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-cost-policy.txt",
                b"SYNTHETIC assessed FIFO opening layers 1 unit at 10 each; no real client evidence",
                "text/plain",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    document = uploaded.json()
    common = {
        "currency": currency,
        "source_evidence_version_id": document["id"],
        "source_sha256": document["sha256"],
    }
    # 25 layers exceed the default 20-row detail projection; one issue consumes all.
    rows = [event(f"O{i}", "opening", "1", "10", sequence=i, **common) for i in range(25)]
    rows.append(event("I", "issue", "25", day="2026-06-02", **common))
    controls = [
        control(
            currency=currency,
            expected_rows=26,
            opening_quantity="25",
            receipt_quantity="0",
            issue_quantity="25",
            opening_cost="250",
            receipt_cost="0",
        )
    ]
    await source(
        c,
        h,
        eid,
        [{**row, "source_sha256": "b" * 64} for row in rows],
        "SYNTHETIC mismatched original hash",
        "inventory-cost-movements",
        {},
        inventory_fifo_controls=controls,
        failed="hash differs",
    )
    movements = await source(
        c,
        h,
        eid,
        rows,
        "SYNTHETIC FIFO movements",
        "inventory-cost-movements",
        {},
        inventory_fifo_controls=controls,
    )
    closing = book(currency=currency, quantity="0", unit_cost="10")
    quantity = [
        {
            **{key: closing[key] for key in ("site", "sku", "lot", "uom", "currency")},
            "expected_rows": 1,
            "positive_quantity": "0",
            "negative_quantity": "0",
        }
    ]
    books = await source(
        c,
        h,
        eid,
        [closing],
        "SYNTHETIC FIFO closing books",
        "stock",
        {},
        inventory_quantity_controls=quantity,
    )
    payload = {
        "workflow": "inventory-fifo-cost",
        "source_ids": {"movements": movements["id"], "books": books["id"]},
        "period_start": "2026-06-01",
        "as_of": "2026-06-30",
        "tolerance": "0",
        "date_tolerance": 0,
        "inventory_scope_assessment": "SYNTHETIC independently assessed FIFO opening order, gross cost and rounding policy",
        "title": "SYNTHETIC FIFO review",
        "rationale": "Separate original cost support and independently controlled closing quantities",
        "currency": currency,
    }
    await execute(c, h, eid, "country-workflow", payload, failed="review")
    for imported in (movements, books):
        reviewed = await decision(c, reviewer[0], eid, imported)
        assert reviewed.status_code == 201, reviewed.text
        payload["source_ids"]["movements" if imported == movements else "books"] = reviewed.json()[
            "id"
        ]
    result = await execute(c, h, eid, "country-workflow", payload)
    assert not result["body"]["blocked"]
    assert result["body"]["rows"][0]["covered_issue_cost"] == "250"
    assert result["collection_totals"]["fifo_allocations"] == 25
    assert len(result["body"]["fifo_allocations"]) == 20
    assert (await decision(c, h, eid, result)).status_code == 403
    final = await actor(h, "Firm Administrator", "fifo-independent-final-reviewer")
    await grant(c, h, base, final)
    approved = await decision(c, final[0], eid, result)
    assert approved.status_code == 201, approved.text
    result = approved.json()
    path = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    paged = await c.get(path + "?fifo_offset=20&fifo_limit=20", headers=h)
    assert paged.status_code == 200, paged.text
    assert len(paged.json()["body"]["fifo_allocations"]) == 5
    assert paged.json()["body"]["fifo_allocations"][-1]["layer_source_row_id"] == "O24"
    exported = await c.get(path + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content))
    assert workbook["FIFO allocations"].max_row == 26
    assert workbook["FIFO events"].max_row == 27
    assert workbook["FIFO source controls"].max_row == 2
    retained = (await c.get(path, headers=h)).json()
    updated = await c.post(
        base + "/evidence-files/" + original["id"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-revised-policy.txt",
                b"SYNTHETIC corrected original FIFO cost support requiring reassessment",
                "text/plain",
            )
        },
    )
    assert updated.status_code == 201, updated.text
    stale = (await c.get(path, headers=h)).json()
    assert stale["freshness"] == "STALE"
    assert stale["body"] == retained["body"] and stale["version"] == retained["version"]
