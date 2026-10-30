"""Retained missing-PAN assessments, independent review, full export and staleness."""

import io
import json
import os
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, create, grant
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_withholding_pan_status import treatment
from tests.unit.test_withholding_workflows import population

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required"
)


@pytest.mark.parametrize("role", ["deductor", "collector", "deductee", "collectee"])
async def test_missing_pan_stays_pending_until_scoped_retained_treatment_is_reviewed(
    environment, monkeypatch, role
):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "pan-source-reviewer")
    final = await actor(h, "Firm Administrator", "pan-result-reviewer")
    for person in (reviewer, final):
        await grant(c, h, base, person)
    liability = role in {"deductor", "collector"}
    assessed = "expected" if liability else "certificate"
    document = await create(
        c, h, base, "evidence", {"category": "Other", "retain_until": "2036-12-31"}
    )
    content = b"SYNTHETIC independently assessed missing-PAN identity and treatment facts; no registry or statutory qualification"
    uploaded = await c.post(
        base + "/evidence-files/" + document["id"],
        headers=h,
        files={"file": ("SYNTHETIC-pan-treatment.txt", content, "text/plain")},
    )
    assert uploaded.status_code == 201, uploaded.text
    evidence = uploaded.json()
    records = population(role)
    for rows in records.values():
        for row in rows:
            row["counterparty"] = "PANNOTAVBL"
    records[assessed][0].update(
        assessment_evidence_version_id=evidence["id"], assessment_sha256=evidence["sha256"]
    )
    prepared = {}

    async def prepare(name, rows):
        created = await source(
            c,
            h,
            eid,
            rows,
            "SYNTHETIC PAN " + name,
            "withholding-records",
            {"INR": "100"},
            **({"rule_document_version_id": evidence["id"]} if name.startswith(assessed) else {}),
        )
        response = await decision(c, reviewer[0], eid, created)
        assert response.status_code == 201, response.text
        return response.json()

    for name, rows in records.items():
        prepared[name] = await prepare(name, rows)
    payload = {
        "workflow": "india-withholding-liability" if liability else "india-withholding-credit",
        "withholding_role": role,
        "source_ids": {name: item["id"] for name, item in prepared.items()},
        "as_of": "2026-06-30",
        "title": "SYNTHETIC missing-PAN assessment review",
        "rationale": "Equal amounts preserve the source marker and require separately scoped retained treatment",
    }
    pending = await execute(c, h, eid, "country-workflow", payload)
    assert (
        pending["body"]["blocked"] and pending["body"]["rows"][0]["comparison_state"] == "MATCHED"
    )
    assert pending["body"]["withholding_pan_counts"] == {
        "rows_with_marker": 1,
        "pending_rows": 1,
        "documented_rows": 0,
    }
    original = deepcopy(pending["body"])
    rejected = await decision(c, final[0], eid, pending)
    assert rejected.status_code == 409 and "missing-PAN" in rejected.text
    noted = await decision(c, h, eid, pending, "accept", item_id=pending["body"]["rows"][0]["id"])
    assert noted.status_code == 201, noted.text
    assert (await decision(c, final[0], eid, noted.json())).status_code == 409
    assessed_rows = deepcopy(records[assessed])
    assessed_rows[0]["pan_treatment"] = json.dumps(treatment())
    prepared[assessed] = await prepare(assessed + "-with-treatment", assessed_rows)
    supported = await execute(
        c,
        h,
        eid,
        "country-workflow",
        {**payload, "source_ids": {name: item["id"] for name, item in prepared.items()}},
    )
    row = supported["body"]["rows"][0]
    assert not supported["body"]["blocked"] and row["state"] == "MATCHED"
    assert row["amounts"] == original["rows"][0]["amounts"] and row["counterparty"] == "PANNOTAVBL"
    assert (
        row["pan_treatment_state"] == "DOCUMENTED_ON_RETAINED_ASSESSMENT"
        and row["pan_authentication"] == "NOT_PERFORMED"
    )
    assert row["assessment_provenance"]["assessment_sha256"] == evidence["sha256"]
    approved = await decision(c, final[0], eid, supported)
    assert approved.status_code == 201, approved.text
    approved = approved.json()
    prefix = f"/api/v1/intelligence/engagements/{eid}/events"
    exported = await c.get(f"{prefix}/{approved['id']}/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    values = list(book["Withholding PAN treatment"].values)
    assert len(values) == 2
    record = dict(zip(values[0], values[1], strict=True))
    assert (
        record["pan_status"] == "PANNOTAVBL"
        and record["party_reference"] == treatment()["party_reference"]
    )
    assert (
        record["pan_authentication"] == "NOT_PERFORMED"
        and record["assessment_sha256"] == evidence["sha256"]
    )
    book.close()
    paper = await c.post(
        f"{prefix}/{approved['id']}/working-paper",
        headers=h,
        json={
            "reference": "SYN-PAN",
            "title": "SYNTHETIC documented missing-PAN treatment",
            "objective": "Retain uploaded marker, assessed treatment and original evidence without authenticating PAN",
        },
    )
    assert paper.status_code == 201, paper.text
    replacement = await c.post(
        base + "/evidence-files/" + document["id"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-pan-treatment-v2.txt",
                content + b"; reassessed original",
                "text/plain",
            )
        },
    )
    assert replacement.status_code == 201, replacement.text
    historical = (await c.get(f"{prefix}/{approved['id']}", headers=h)).json()
    assert historical["freshness"] == "STALE" and historical["body"] == approved["body"]


async def test_missing_pan_beyond_first_page_still_blocks_approval_and_exports_complete_schedule(
    environment, monkeypatch
):
    c = environment
    double_boundaries(monkeypatch)
    h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "pan-large-reviewer")
    await grant(c, h, base, reviewer)
    doc = await create(c, h, base, "evidence", {"category": "Other", "retain_until": "2036-12-31"})
    file = await c.post(
        base + "/evidence-files/" + doc["id"],
        headers=h,
        files={
            "file": (
                "SYNTHETIC-complete-rule.txt",
                b"SYNTHETIC independent complete expected population and rule assessment",
                "text/plain",
            )
        },
    )
    assert file.status_code == 201, file.text
    evidence = file.json()
    prepared = {}
    for name, example in population("collector").items():
        rows = []
        for i in range(121):
            row = {
                **example[0],
                "id": f"SYN-{name}-{i}",
                "transaction_reference": f"SYN-{i:03}",
                "counterparty": "PANINVALID" if i == 120 else "AAA-SYN-ORDINARY-PARTY",
            }
            if name == "expected":
                row.update(
                    assessment_evidence_version_id=evidence["id"],
                    assessment_sha256=evidence["sha256"],
                )
            rows.append(row)
        result = await source(
            c,
            h,
            eid,
            rows,
            "SYNTHETIC 121 " + name,
            "withholding-records",
            {"INR": "12100"},
            **({"rule_document_version_id": evidence["id"]} if name == "expected" else {}),
        )
        response = await decision(c, reviewer[0], eid, result)
        assert response.status_code == 201, response.text
        prepared[name] = response.json()["id"]
    result = await execute(
        c,
        h,
        eid,
        "country-workflow",
        {
            "workflow": "india-withholding-liability",
            "withholding_role": "collector",
            "source_ids": prepared,
            "as_of": "2026-06-30",
            "title": "SYNTHETIC complete PAN population",
            "rationale": "Complete population retains a missing-PAN marker beyond the first visible page",
        },
    )
    assert result["body"]["blocked"] and result["body"]["withholding_pan_counts"] == {
        "rows_with_marker": 1,
        "pending_rows": 1,
        "documented_rows": 0,
    }
    assert len(result["body"]["rows"]) == 100 and all(
        row["counterparty"] != "PANINVALID" for row in result["body"]["rows"]
    )
    assert (await decision(c, reviewer[0], eid, result)).status_code == 409
    exported = await c.get(
        f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}/export?format=xlsx",
        headers=h,
    )
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert len(list(book["Rows"].values)) == 122
    values = list(book["Withholding PAN treatment"].values)
    assert len(values) == 2
    row = dict(zip(values[0], values[1], strict=True))
    assert row["transaction_reference"] == "SYN-120" and row["pan_status"] == "PANINVALID"
    assert "PAN_TREATMENT_ASSESSMENT_MISSING" in row["pan_issues"]
    book.close()
