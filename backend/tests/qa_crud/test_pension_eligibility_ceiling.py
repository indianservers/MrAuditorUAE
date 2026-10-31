"""Published synthetic eligibility parameters retain exact basis/source lineage."""

import io
import os
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.qa_crud.test_country_workflows import decision, execute, source
from tests.qa_crud.test_published_pension_source import prepared, publish
from tests.unit.test_pension_eligibility_ceiling import boundary_vectors
from tests.unit.test_pension_workflow import pension_fixture

pytestmark = pytest.mark.skipif(os.getenv("INTEGRATION_TESTS") != "1", reason="Owned DB")


async def test_published_ceiling_uses_original_salary_at_boundary_then_exports_and_stales(
    environment, monkeypatch
):
    c = environment
    h, eid, _, reviewer, final, identifiers, _, _, _, _ = await prepared(c, monkeypatch)
    new_government, _ = await publish(
        c,
        h,
        reviewer[0],
        "GOVERNMENT",
        sequence=3,
        pension_overrides={"exclusive_eligibility_ceiling": "20000"},
        validation_fixtures=[case.model_dump(mode="json") for case in boundary_vectors()],
    )
    identifiers[2] = new_government
    bases, observed = [], []
    for i, salary in enumerate(["19999.99", "20000", "20000.01"]):
        case_bases, _, case_reported = pension_fixture()
        for row in case_bases:
            row.update(id=f"{i}-{row['id']}", subject=f"PERSON-{i}", amount=salary)
        for row in case_reported:
            row.update(id=f"{i}-{row['id']}", subject=f"PERSON-{i}")
            if i and row["component"] == "PENSION:GOVERNMENT":
                row["amount"] = "0"
            if i and row["component"] == "PENSION:EMPLOYER":
                row["amount"] = "1800"
        bases.extend(case_bases)
        observed.extend(case_reported)
    selected = {}
    for role, table, rows, total in [
        ("bases", "pension-bases", bases, "180000"),
        ("reported", "pension-reported", observed, "7500"),
    ]:
        raw = await source(c, h, eid, rows, f"SYNTHETIC ceiling {role}", table, {"AED": total})
        reviewed = await decision(c, reviewer[0], eid, raw)
        assert reviewed.status_code == 201, reviewed.text
        selected[role] = reviewed.json()["id"]
    created = await c.post(
        f"/api/v1/intelligence/engagements/{eid}/pension-rule-source",
        headers=h,
        json={
            "basis_source_id": selected["bases"],
            "rule_version_ids": identifiers,
            "expected_basis_rows": 9,
            "label": "Synthetic ceiling-bound pension rules",
            "assessment": "Independent contribution salary and exclusive eligibility boundary review",
        },
    )
    assert created.status_code == 201, created.text
    government = next(
        row for row in created.json()["body"]["rows"] if row["component"] == "PENSION:GOVERNMENT"
    )
    assert government["exclusive_eligibility_ceiling"] == "20000"
    approved = await decision(c, reviewer[0], eid, created.json())
    assert approved.status_code == 201, approved.text
    selected["rules"] = approved.json()["id"]
    result = await execute(
        c,
        h,
        eid,
        "country-workflow",
        {
            "workflow": "uae-pension-review",
            "source_ids": selected,
            "currency": "AED",
            "as_of": "2026-06-30",
            "title": "Synthetic reviewed eligibility ceiling",
            "rationale": "Independent boundary salary, capped basis and government offsets assessed",
        },
    )
    body = result["body"]
    assert not body["blocked"] and body["exception_count"] == 0
    assert body["expected_total"] == "7500.00"
    rows = {(row["subject"], row["component"]): row for row in body["rows"]}
    for i in [1, 2]:
        assert rows[(f"PERSON-{i}", "PENSION:GOVERNMENT")]["expected_amount"] == "0"
        assert rows[(f"PERSON-{i}", "PENSION:EMPLOYER")]["expected_amount"] == "1800.00"
    checked = await decision(c, final[0], eid, result)
    assert checked.status_code == 201, checked.text
    saved = checked.json()
    endpoint = f"/api/v1/intelligence/engagements/{eid}/events/{saved['id']}"
    exported = await c.get(endpoint + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200
    book = load_workbook(io.BytesIO(exported.content), read_only=True)
    assert "exclusive_eligibility_ceiling" in next(book["Rows"].values)
    book.close()
    # Equal overall reported totals must not hide contradictory support salary.
    bad_bases, bad_reported = deepcopy(bases), deepcopy(observed)
    for row in bad_bases:
        if row["subject"] == "PERSON-0" and row["component"] == "PENSION:GOVERNMENT":
            row["amount"] = "20000"
    for row in bad_reported:
        if row["subject"] == "PERSON-0":
            if row["component"] == "PENSION:GOVERNMENT":
                row["amount"] = "0"
            elif row["component"] == "PENSION:EMPLOYER":
                row["amount"] = "1800"
    bad_selected = {}
    for role, table, values, total in [
        ("bases", "pension-bases", bad_bases, "180000.01"),
        ("reported", "pension-reported", bad_reported, "7500"),
    ]:
        raw = await source(
            c, h, eid, values, f"SYNTHETIC contradictory {role}", table, {"AED": total}
        )
        reviewed = await decision(c, reviewer[0], eid, raw)
        assert reviewed.status_code == 201, reviewed.text
        bad_selected[role] = reviewed.json()["id"]
    bad_rules = await c.post(
        f"/api/v1/intelligence/engagements/{eid}/pension-rule-source",
        headers=h,
        json={
            "basis_source_id": bad_selected["bases"],
            "rule_version_ids": identifiers,
            "expected_basis_rows": 9,
            "label": "SYNTHETIC contradictory government salary",
            "assessment": "Synthetic retained contrary declarations must remain visible and block netting",
        },
    )
    assert bad_rules.status_code == 201, bad_rules.text
    reviewed = await decision(c, reviewer[0], eid, bad_rules.json())
    assert reviewed.status_code == 201, reviewed.text
    bad_selected["rules"] = reviewed.json()["id"]
    bad = await execute(
        c,
        h,
        eid,
        "country-workflow",
        {
            "workflow": "uae-pension-review",
            "source_ids": bad_selected,
            "currency": "AED",
            "as_of": "2026-06-30",
            "title": "SYNTHETIC contradictory support salary",
            "rationale": "Same combined amount with inconsistent component salary cannot establish a matched net employer",
        },
    )
    assert bad["body"]["blocked"] and not bad["body"]["expected_total_complete"]
    employer = next(
        row
        for row in bad["body"]["rows"]
        if row["subject"] == "PERSON-0" and row["component"] == "PENSION:EMPLOYER"
    )
    assert employer["state"] == "GOVERNMENT_SUPPORT_SALARY_BASIS_MISMATCH"
    assert employer["government_assessed_base"] == "20000"
    assert employer["expected_amount"] is None
    rejected = await decision(c, final[0], eid, bad)
    assert rejected.status_code == 409
    bad_export = await c.get(
        f"/api/v1/intelligence/engagements/{eid}/events/{bad['id']}/export?format=xlsx", headers=h
    )
    assert bad_export.status_code == 200
    book = load_workbook(io.BytesIO(bad_export.content), read_only=True)
    exported_rows = list(book["Rows"].values)
    headers = exported_rows[0]
    assert "government_assessed_base" in headers
    assert any(
        row[headers.index("state")] == "GOVERNMENT_SUPPORT_SALARY_BASIS_MISMATCH"
        for row in exported_rows[1:]
    )
    book.close()
    historical = deepcopy(saved["body"])
    withdrawn = await c.post(
        f"/api/v1/regulatory/rules/{new_government}/decisions",
        headers=h,
        json={
            "action": "withdraw",
            "expected_sequence": 6,
            "comments": "Synthetic reviewed ceiling version withdrawn for source reassessment",
        },
    )
    assert withdrawn.status_code == 201, withdrawn.text
    stale = await c.get(endpoint, headers=h)
    assert stale.status_code == 200 and stale.json()["freshness"] == "STALE"
    assert stale.json()["body"] == historical
