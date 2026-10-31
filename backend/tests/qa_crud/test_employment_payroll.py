import io
import os
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.modules.country_workflows.employment_payroll import ROLES
from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_employment_payroll import population

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable DB required"
)


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
async def test_employment_review_export_and_withdrawal(environment, monkeypatch, country, currency):
    c = environment
    double_boundaries(monkeypatch)
    if country == "AE":
        h, _, eid, base = await setup(c)
    else:
        h, eid, base = await india_context(c)
    reviewer = await actor(h, "Firm Administrator", "employment-payroll-reviewer")
    await grant(c, h, base, reviewer)
    selected, originals = {}, {}
    populations = population(currency)
    old = dict(populations["register"][0], id="R0", amount="1500", record_status="superseded")
    populations["register"][0]["revision"] = 2
    populations["register"].append(old)
    for role, rows in populations.items():
        totals = (
            {}
            if role == "employment"
            else {currency: str(sum((Decimal(row["amount"]) for row in rows), Decimal(0)))}
        )
        imported = await source(
            c, h, eid, rows, "SYNTHETIC employment " + role, table_type=ROLES[role], totals=totals
        )
        originals[role] = imported
        assert (await decision(c, h, eid, imported)).status_code == 403
        response = await decision(c, reviewer[0], eid, imported)
        assert response.status_code == 201, response.text
        selected[role] = response.json()["id"]
    payload = dict(
        workflow="employment-payroll-review",
        source_ids=selected,
        period_start="2026-06-01",
        as_of="2026-06-30",
        title="SYNTHETIC employee month review",
        rationale="Independently assessed contract, attendance and payroll expectation",
    )
    await execute(
        c,
        h,
        eid,
        "country-workflow",
        {**payload, "source_ids": {**selected, "controls": originals["controls"]["id"]}},
        failed="select the reviewed current version",
    )
    result = await execute(c, h, eid, "country-workflow", payload)
    assert not result["body"]["blocked"] and result["body"]["exception_count"] == 0
    assert len(result["body"]["rows"]) == 3
    approved = await decision(c, reviewer[0], eid, result)
    assert approved.status_code == 201, approved.text
    result = approved.json()
    prefix = f"/api/v1/intelligence/engagements/{eid}/events/{result['id']}"
    for view, count, kinds in (
        ("assessment", 1, {"entry-comparison", "population-coverage"}),
        ("history", 2, {"revision-history"}),
        ("all", 3, {"entry-comparison", "population-coverage", "revision-history"}),
    ):
        filtered = await c.get(prefix + f"?employment_view={view}&limit=1", headers=h)
        assert filtered.status_code == 200, filtered.text
        assert filtered.json()["collection_totals"]["rows"] == count
        assert len(filtered.json()["body"]["rows"]) == 1
        assert filtered.json()["body"]["rows"][0]["kind"] in kinds
    history_page = await c.get(prefix + "?employment_view=history&offset=1&limit=1", headers=h)
    assert history_page.status_code == 200, history_page.text
    assert history_page.json()["body"]["rows"][0]["retained_amount"] == "1500"
    assert (await c.get(prefix + "?employment_view=unknown", headers=h)).status_code == 422
    preparer = await actor(h, "Senior Auditor", "employment-restricted-preparer")
    await grant(c, h, base, preparer)
    for restricted in (
        prefix,
        prefix + "/export?format=xlsx",
        prefix + "?employment_view=history",
        f"/api/v1/intelligence/engagements/{eid}/events/{selected['register']}",
    ):
        denied = await c.get(restricted, headers=preparer[0])
        assert denied.status_code == 403, denied.text
    exported = await c.get(prefix + "/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    book = load_workbook(io.BytesIO(exported.content))
    assert any(
        sheet.max_row == 4 and "revision-history" in str(list(sheet.values)) for sheet in book
    )
    incomplete = await execute(
        c,
        h,
        eid,
        "country-workflow",
        {**payload, "as_of": "2026-07-31", "title": "SYNTHETIC missing July payroll controls"},
    )
    missing_prefix = f"/api/v1/intelligence/engagements/{eid}/events/{incomplete['id']}"
    missing_view = await c.get(missing_prefix + "?employment_view=assessment", headers=h)
    assert missing_view.status_code == 200, missing_view.text
    assert missing_view.json()["collection_totals"]["rows"] == 2
    missing = next(
        row for row in missing_view.json()["body"]["rows"] if row["kind"] == "population-coverage"
    )
    assert missing["payroll_month"] == "2026-07" and missing["expected_amount"] is None
    assert missing_view.json()["body"]["blocked"] is True
    assert (await decision(c, reviewer[0], eid, incomplete)).status_code == 409
    assert (await c.get(prefix, headers=h)).json()["freshness"] == "CURRENT"
    original = await c.get(
        f"/api/v1/intelligence/engagements/{eid}/events/{selected['employment']}", headers=h
    )
    withdrawn = await decision(c, reviewer[0], eid, original.json(), "reject")
    assert withdrawn.status_code == 201, withdrawn.text
    stale = await c.get(prefix, headers=h)
    assert stale.json()["freshness"] == "STALE"
    assert stale.json()["body"]["rows"] == result["body"]["rows"]
