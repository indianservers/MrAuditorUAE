"""Actual API/DB evidence and review; scanner and dispatch are explicit doubles."""

import io
import os

import pytest
from openpyxl import load_workbook

from tests.integration.test_platform import environment as environment
from tests.phase4.test_workflows import actor, grant, setup
from tests.qa_crud.test_country_workflows import decision, double_boundaries, execute, source
from tests.qa_crud.test_india_ecr_source import india_context
from tests.unit.test_bank_accounts import parameters, population

pytestmark = pytest.mark.skipif(
    os.getenv("INTEGRATION_TESTS") != "1", reason="Owned disposable database required"
)


@pytest.mark.parametrize("country", ["AE", "IN"])
async def test_independent_account_evidence_entity_binding_export_and_staleness(
    environment, monkeypatch, country
):
    c = environment
    double_boundaries(monkeypatch)
    if country == "IN":
        h, eid, base = await india_context(c)
        entity_name = "SYNTHETIC ECR India entity"
    else:
        h, _, eid, base = await setup(c)
        entity_name = "FICTIONAL synthetic entity"
    reviewer = await actor(h, "Firm Administrator", "bank-account-reviewer")
    await grant(c, h, base, reviewer)
    records = population(country)
    prepared = {}
    for role, rows in records.items():
        rows[0]["owner_name"] = entity_name
        prepared[role] = await source(
            c,
            h,
            eid,
            rows,
            "SYNTHETIC " + role + " bank account evidence",
            "bank-account-evidence",
            {},
        )
    payload = parameters(country, bank_owner_name=entity_name).model_dump(mode="json")
    payload["source_ids"] = {role: item["id"] for role, item in prepared.items()}
    await execute(
        c,
        h,
        eid,
        "country-workflow",
        {**payload, "bank_owner_name": "Unrelated declared company"},
        failed="engagement legal entity name",
    )
    await execute(c, h, eid, "country-workflow", payload, failed="Independently review all")
    for role, item in prepared.items():
        response = await decision(c, reviewer[0], eid, item)
        assert response.status_code == 201, response.text
        prepared[role] = response.json()
    payload["source_ids"] = {role: item["id"] for role, item in prepared.items()}
    result = await execute(c, h, eid, "country-workflow", payload)
    assert not result["body"]["blocked"]
    assert result["body"]["bank_authentication_status"] == "NOT_ESTABLISHED_BY_DOCUMENT_COMPARISON"
    assert (await decision(c, h, eid, result)).status_code == 403
    approval = await decision(c, reviewer[0], eid, result)
    assert approval.status_code == 201, approval.text
    approved = approval.json()
    prefix = f"/api/v1/intelligence/engagements/{eid}"
    exported = await c.get(prefix + f"/events/{approved['id']}/export?format=xlsx", headers=h)
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content))
    cover = {row[0].value: row[1].value for row in workbook["Source and review"]}
    assert cover["bank_authentication_status"] == "NOT_ESTABLISHED_BY_DOCUMENT_COMPARISON"
    assert workbook["Rows"].max_row == 4
    records["bank"][0]["owner_identifier"] = "UNRELATED-ID"
    wrong = await source(
        c, h, eid, records["bank"], "SYNTHETIC mismatched bank owner", "bank-account-evidence", {}
    )
    response = await decision(c, reviewer[0], eid, wrong)
    assert response.status_code == 201, response.text
    bad = await execute(
        c,
        h,
        eid,
        "country-workflow",
        {**payload, "source_ids": {**payload["source_ids"], "bank": response.json()["id"]}},
    )
    assert bad["body"]["blocked"]
    assert (await decision(c, reviewer[0], eid, bad)).status_code == 409
    engagement = (await c.get("/api/v1/engagements/" + eid, headers=h)).json()
    entity = (await c.get("/api/v1/entities/" + engagement["entity_id"], headers=h)).json()
    renamed = await c.patch(
        "/api/v1/entities/" + entity["id"],
        headers=h,
        json={"version": entity["version"], "name": entity_name + " renamed"},
    )
    assert renamed.status_code == 200, renamed.text
    stale = await c.get(prefix + f"/events/{approved['id']}", headers=h)
    assert stale.json()["freshness"] == "STALE"
