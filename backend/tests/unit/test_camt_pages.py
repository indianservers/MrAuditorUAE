"""Synthetic CGI statement pagination; no provider-authentication qualification."""

import io
from copy import deepcopy
from decimal import Decimal
from xml.etree import ElementTree as xml

import pytest
from openpyxl import load_workbook

from app.modules.country_workflows.native_sources import parse_native
from app.modules.intelligence.exports import render
from tests.unit.native_source_examples import CAMT
from tests.unit.test_native_workflow_sources import options

NS = "{urn:iso:std:iso:20022:tech:xsd:camt.053.001.08}"


def packet(currency="AED", *, middle=True, page_count=None):
    root = xml.fromstring(CAMT.replace(b"AED", currency.encode()))
    container = root.find(NS + "BkToCstmrStmt")
    original = container.find(NS + "Stmt")
    entries = original.findall(NS + "Ntry")
    container.remove(original)
    count = page_count or (3 if middle else 2)
    for index in range(1, count + 1):
        page = deepcopy(original)
        for entry in page.findall(NS + "Ntry"):
            page.remove(entry)
        xml.SubElement(page, NS + "ElctrncSeqNb").text = "27"
        pagination = xml.SubElement(page, NS + "StmtPgntn")
        xml.SubElement(pagination, NS + "PgNb").text = str(index)
        xml.SubElement(pagination, NS + "LastPgInd").text = "true" if index == count else "false"
        opening, closing = page.findall(NS + "Bal")
        opening.find(NS + "Amt").text = "0.00" if index == 1 else "100.00"
        closing.find(NS + "Amt").text = "95.00" if index == count else "100.00"
        if index > 1:
            opening.find(NS + "Dt/" + NS + "Dt").text = "2026-06-30"
            subtype = xml.SubElement(opening.find(NS + "Tp"), NS + "SubTp")
            xml.SubElement(subtype, NS + "Cd").text = "INTM"
        if index < count:
            subtype = xml.SubElement(closing.find(NS + "Tp"), NS + "SubTp")
            xml.SubElement(subtype, NS + "Cd").text = "INTM"
        summary = page.find(NS + "TxsSummry")
        if index > 1:
            page.remove(summary)
        else:
            total = summary.find(NS + "TtlNtries")
            xml.SubElement(total, NS + "Sum").text = "105.00"
            xml.SubElement(total, NS + "TtlNetNtryAmt").text = "95.00"
            xml.SubElement(total, NS + "CdtDbtInd").text = "CRDT"
            for group, amount in (("TtlCdtNtries", "100.00"), ("TtlDbtNtries", "5.00")):
                control = xml.SubElement(summary, NS + group)
                xml.SubElement(control, NS + "NbOfNtries").text = "1"
                xml.SubElement(control, NS + "Sum").text = amount
        if index == 1:
            page.append(deepcopy(entries[0]))
        if index == count:
            page.append(deepcopy(entries[1]))
        container.append(page)
    return xml.tostring(root)


def read(content, currency="AED", country="AE"):
    data = options().model_copy(update={"control_totals": {currency: Decimal("95")}})
    return parse_native(content, data, country, "SYNTHETIC-paginated-camt.xml")


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR"), ("AE", "USD")])
@pytest.mark.parametrize("middle", [False, True])
def test_complete_statement_pages_preserve_totals_and_positions(country, currency, middle):
    result = read(packet(currency, middle=middle), currency, country)
    assert result["parser_version"] == "camt053-3"
    assert [Decimal(row["amount"]) for row in result["rows"]] == [Decimal("100"), Decimal("-5")]
    pages = 3 if middle else 2
    assert [row["source_statement_page"] for row in result["rows"]] == [1, pages]
    assert [row["source_row"] for row in result["rows"]] == [1, 2]
    assert [row["source_page_row"] for row in result["rows"]] == [1, 1]
    assert len({row["id"] for row in result["rows"]}) == 2
    metadata = result["native_metadata"]
    assert metadata["page_count"] == pages
    assert metadata["opening"] == "0.00" and metadata["closing"] == "95.00"
    assert metadata["summary_scope"] == "WHOLE_STATEMENT_FIRST_PAGE"
    assert metadata["summary_controls"]["TtlNtries"]["Sum"] == "105.00"
    assert metadata["bank_identity_verified"] is False
    assert [item["entry_count"] for item in metadata["page_controls"]] == (
        [1, 0, 1] if middle else [1, 1]
    )


def mutate(path, value, *, page=1, remove=False, duplicate=False):
    root = xml.fromstring(packet())
    selected = root.findall(NS + "BkToCstmrStmt/" + NS + "Stmt")[page - 1]
    pieces = path.split("/")
    parent = selected
    for part in pieces[:-1]:
        parent = parent.find(NS + part)
    child = parent.find(NS + pieces[-1])
    if remove:
        parent.remove(child)
    elif duplicate:
        parent.append(deepcopy(child))
    else:
        child.text = value
    return xml.tostring(root)


@pytest.mark.parametrize(
    "path,value,page,error",
    [
        ("StmtPgntn/PgNb", "3", 2, "contiguous"),
        ("StmtPgntn/PgNb", "01", 1, "contiguous"),
        ("StmtPgntn/LastPgInd", "true", 1, "last-page"),
        ("StmtPgntn/LastPgInd", "false", 3, "incomplete"),
        ("StmtPgntn/LastPgInd", "TRUE", 3, "incomplete"),
        ("ElctrncSeqNb", "28", 2, "mix"),
        ("Acct/Ccy", "USD", 2, "mix"),
        ("Acct/Id/Othr/Id", "OTHER", 2, "unambiguous"),
        ("Bal/Amt", "101", 2, "carry forward"),
        ("Bal/Dt/Dt", "2026-06-29", 2, "carry forward"),
        ("Bal/Tp/SubTp/Cd", "OTHER", 2, "INTM"),
        ("TxsSummry/TtlNtries/NbOfNtries", "1", 1, "entry-count"),
        ("TxsSummry/TtlNtries/Sum", "95", 1, "gross"),
        ("TxsSummry/TtlNtries/TtlNetNtryAmt", "105", 1, "net"),
        ("TxsSummry/TtlDbtNtries/NbOfNtries", "2", 1, "entry-count"),
        ("TxsSummry/TtlDbtNtries/Sum", "6", 1, "gross"),
        ("Ntry/Sts/Cd", "PDNG", 3, "Only booked"),
    ],
)
def test_broken_pages_and_full_statement_controls_fail_closed(path, value, page, error):
    with pytest.raises(ValueError, match=error):
        read(mutate(path, value, page=page))


@pytest.mark.parametrize(
    "path",
    ["StmtPgntn", "StmtPgntn/PgNb", "ElctrncSeqNb", "Acct/Ccy", "TxsSummry/TtlNtries/NbOfNtries"],
)
def test_duplicate_controls_are_not_silently_first_wins(path):
    with pytest.raises(ValueError, match="Duplicate"):
        read(mutate(path, None, duplicate=True))


def test_missing_page_and_repeated_summary_do_not_pass_as_complete():
    root = xml.fromstring(packet())
    container = root.find(NS + "BkToCstmrStmt")
    pages = container.findall(NS + "Stmt")
    container.remove(pages[1])
    with pytest.raises(ValueError, match="contiguous"):
        read(xml.tostring(root))
    container.insert(2, pages[1])
    pages[1].find(NS + "StmtPgntn/" + NS + "PgNb").text = "3"
    # A complete chain with an illicit repeated summary must also be rejected.
    root = xml.fromstring(packet())
    pages = root.findall(NS + "BkToCstmrStmt/" + NS + "Stmt")
    pages[1].append(deepcopy(pages[0].find(NS + "TxsSummry")))
    with pytest.raises(ValueError, match="only on the first"):
        read(xml.tostring(root))


def test_single_page_interim_balance_is_not_a_final_statement():
    content = CAMT.replace(b"</CdOrPrtry>", b"</CdOrPrtry><SubTp><Cd>INTM</Cd></SubTp>", 1)
    with pytest.raises(ValueError, match="Interim"):
        read(content)


def test_xlsx_retains_all_page_controls_without_mutating_the_result():
    body = read(packet())
    original = deepcopy(body)
    event = {
        "id": "SYNTHETIC-CAMT-PAGES",
        "kind": "workflow-source",
        "version": 1,
        "sha256": "a" * 64,
        "created_at": "2026-06-30T00:00:00Z",
        "body": body,
        "dependencies": [],
    }
    content, _ = render(event, "xlsx")
    workbook = load_workbook(io.BytesIO(content))
    assert workbook["CAMT page controls"].max_row == 4
    cover = {row[0].value: row[1].value for row in workbook["Source and review"]}
    assert "3 page control records retained" in cover["native_metadata"]
    assert body == original


def test_single_page_final_pagination_keeps_the_existing_non_cgi_profile():
    content = CAMT.replace(
        b"</Stmt>", b"<StmtPgntn><PgNb>1</PgNb><LastPgInd>true</LastPgInd></StmtPgntn></Stmt>"
    )
    assert read(content)["row_count"] == 2


def test_maximum_chain_keeps_empty_pages_and_rejects_an_extra_page():
    body = read(packet(page_count=100))
    assert len(body["native_metadata"]["page_controls"]) == 100
    assert [row["source_statement_page"] for row in body["rows"]] == [1, 100]
    with pytest.raises(ValueError, match="file bound"):
        read(packet(page_count=101))


def test_explicit_statement_selection_does_not_merge_another_account():
    root = xml.fromstring(packet())
    container = root.find(NS + "BkToCstmrStmt")
    other = xml.fromstring(CAMT).find(NS + "BkToCstmrStmt/" + NS + "Stmt")
    other.find(NS + "Id").text = "OTHER-STMT"
    other.find(NS + "Acct/" + NS + "Id/" + NS + "Othr/" + NS + "Id").text = "BANK-2"
    container.insert(1, other)
    with pytest.raises(ValueError, match="unambiguous"):
        read(xml.tostring(root))
    data = options(statement_identifier="SYNTHETIC-STMT", account_identifier="BANK-1")
    body = parse_native(xml.tostring(root), data, "AE", "SYNTHETIC.xml")
    assert body["native_metadata"]["file_statement_count"] == 4
    assert [row["source_statement_position"] for row in body["rows"]] == [2, 4]
    assert body["row_count"] == 2
    other_data = options(statement_identifier="OTHER-STMT", account_identifier="BANK-2")
    other_body = parse_native(xml.tostring(root), other_data, "AE", "SYNTHETIC.xml")
    assert other_body["native_metadata"]["file_statement_count"] == 4
    assert other_body["rows"][0]["account"] == "BANK-2"
