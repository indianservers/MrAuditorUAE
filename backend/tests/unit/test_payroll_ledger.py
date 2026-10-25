import io
import uuid
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from app.modules.country_workflows.calculations import calculate
from app.modules.country_workflows.catalogue import WORKFLOWS
from app.modules.country_workflows.schemas import WorkflowInput
from app.modules.intelligence.exports import render
from tests.unit.test_shared_audit_workflows import source


def population(currency="AED"):
    common = {"registration": "SYNTHETIC-EMPLOYER", "currency": currency, "evidence_reference": "SYNTHETIC-ORIGINAL"}
    accounts = [("SALARY-EXP", "expense", "0", "1000"), ("NET-PAY", "liability", "200", "200"), ("DEDUCTIONS", "liability", "0", "100"), ("PENSION-EXP", "expense", "0", "50"), ("PENSION-PAY", "liability", "0", "50")]
    opening = [{**common, "id": f"OPEN-{index}", "account": account, "account_class": cls, "amount": amount, "as_of": "2026-05-31"} for index, (account, cls, amount, _) in enumerate(accounts)]
    closing = [{**row, "id": f"CLOSE-{index}", "amount": accounts[index][3], "as_of": "2026-06-30"} for index, row in enumerate(opening)]
    lines = [("SALARY-EXP", "expense", "SALARY:GROSS", "accrual", "1000"), ("NET-PAY", "liability", "SALARY:NET", "accrual", "900"), ("DEDUCTIONS", "liability", "EMPLOYEE_DEDUCTION", "accrual", "100"), ("NET-PAY", "liability", "SALARY:NET", "settlement", "-900"), ("PENSION-EXP", "expense", "PENSION:EMPLOYER", "accrual", "50"), ("PENSION-PAY", "liability", "PENSION:EMPLOYER", "accrual", "50")]
    payroll = [{**common, "id": f"PAY-{index}", "account": account, "account_class": cls, "amount": amount, "posting_date": "2026-06-20", "wage_month": "2026-06", "component": component, "category": category, "subject": "EMP-1", "reference": f"SYN-EXPECTED-{index}", "assessment": "Independently assessed fictional payroll allocation and posting date"} for index, (account, cls, component, category, amount) in enumerate(lines)]
    posted = [{**row, "id": f"GL-{index}", "reference": f"SYN-POSTED-{index}", "evidence_reference": "SYNTHETIC-GL-ORIGINAL"} for index, row in enumerate(payroll)]
    return {"opening": opening, "payroll": payroll, "posted": posted, "closing": closing}


def run(rows=None, country="AE", **changes):
    rows = rows or population()
    sources = {role: source(WORKFLOWS["payroll-ledger-rollforward"][2][role], records) for role, records in rows.items()}
    params = WorkflowInput.model_validate({"workflow": "payroll-ledger-rollforward", "source_ids": {role: uuid.uuid4() for role in sources}, "period_start": "2026-06-01", "as_of": "2026-06-30", "title": "SYNTHETIC payroll accounting", "rationale": "Independent expected payroll and posted GL evidence with both balance controls", **changes})
    return calculate(params, country, sources)


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
def test_separate_gross_net_deductions_pension_and_prior_period_balance(country, currency):
    p = population(currency)
    original = deepcopy(p)
    result = run(p, country)
    assert p == original and result["exception_count"] == 0 and not result["blocked"]
    assert len(result["rows"]) == 11 and len(result["payroll_posting_links"]) == 12
    net = next(row for row in result["rows"] if row["kind"] == "account-rollforward" and row["account"] == "NET-PAY")
    assert net["opening_amount"] == net["actual_closing"] == "200"
    assert net["payroll_movement"] == "0" and net["normal_sign"] == "CREDIT_POSITIVE"
    expense = next(row for row in result["rows"] if row["kind"] == "account-rollforward" and row["account"] == "SALARY-EXP")
    assert expense["expected_closing"] == "1000" and expense["normal_sign"] == "DEBIT_POSITIVE"
    assert result["payroll_ledger_scope"]["opening_date"] == "2026-05-31"
    assert all(link["comparison_row_id"] in {row["id"] for row in result["rows"]} for link in result["payroll_posting_links"])


def test_offsetting_extra_postings_cannot_hide_behind_equal_net_or_closing_balance():
    p = population()
    for index, amount in enumerate(("100", "-100")):
        p["posted"].append({**p["posted"][0], "id": f"EXTRA-{index}", "reference": f"EXTRA-{index}", "amount": amount})
    result = run(p)
    comparison = next(row for row in result["rows"] if row["kind"] == "posting-comparison" and row["account"] == "SALARY-EXP")
    assert comparison["positive_difference"] == "100" and comparison["negative_difference"] == "-100"
    assert comparison["state"] == "DIRECTIONAL_POSTING_VARIANCE"
    assert all(row["state"] == "MATCHED" for row in result["rows"] if row["kind"] == "account-rollforward")


@pytest.mark.parametrize("field,value", [("wage_month", "2026-05"), ("component", "SALARY:OTHER"), ("category", "adjustment"), ("registration", "OTHER-EMPLOYER"), ("currency", "USD"), ("account", "OTHER-ACCOUNT")])
def test_partition_changes_never_clear_other_payroll_buckets(field, value):
    p = population()
    p["posted"][0][field] = value
    result = run(p)
    assert any(row["state"] == "MISSING_GL_PAYROLL_MOVEMENT" for row in result["rows"])
    assert any(row["state"] == "UNEXPECTED_GL_PAYROLL_MOVEMENT" for row in result["rows"])


@pytest.mark.parametrize("role,field,value,error", [
    ("opening", "as_of", "2026-06-01", "balance dates"),
    ("closing", "as_of", "2026-07-01", "balance dates"),
    ("payroll", "posting_date", "2026-07-01", "outside"),
    ("posted", "account_class", "liability", "class conflicts"),
    ("payroll", "assessment", "          ", "assessment"),
    ("posted", "wage_month", "0000-06", "record"),
    ("payroll", "amount", "0", "record"),
])
def test_bad_scope_and_unassessed_rows_fail(role, field, value, error):
    p = population()
    p[role][0][field] = value
    with pytest.raises(ValueError, match=error):
        run(p)


def test_missing_balance_is_blocked_not_inferred_zero_and_duplicates_fail():
    p = population()
    p["opening"].pop()
    result = run(p)
    assert result["blocked"] and any(row["expected_closing"] is None for row in result["rows"] if row["kind"] == "account-rollforward")
    p = population()
    p["posted"].append({**p["posted"][0], "id": "COPY"})
    with pytest.raises(ValueError, match="Duplicate payroll accounting observation"):
        run(p)
    p = population()
    p["closing"].append({**p["closing"][0], "id": "COPY"})
    with pytest.raises(ValueError, match="Duplicate payroll registration"):
        run(p)


def test_wrong_country_domestic_component_and_missing_accounting_period_are_rejected():
    p = population()
    p["payroll"][0]["component"] = "PF:EMPLOYEE"
    with pytest.raises(ValueError, match="Indian domestic"):
        run(p)
    with pytest.raises(ValueError, match="accounting period"):
        run(period_start=None)
    with pytest.raises(ValueError, match="accounting period"):
        run(period_start="0001-01-01")


def test_posted_balance_error_is_distinct_from_payroll_allocation_error():
    p = population()
    p["closing"][0]["amount"] = "999"
    row = next(row for row in run(p)["rows"] if row["kind"] == "account-rollforward" and row["account"] == "SALARY-EXP")
    assert row["payroll_closing_difference"] == row["posted_closing_difference"] == "-1"
    assert row["state"] == "GL_BALANCE_ROLLFORWARD_VARIANCE"


def test_xlsx_preserves_individual_posting_lineage_without_large_json_cells():
    body = run()
    data, _ = render({"id": str(uuid.uuid4()), "kind": "country-workflow", "version": 1, "body": body, "dependencies": [], "sha256": "a" * 64, "created_at": "2026-06-30T00:00:00Z"}, "xlsx")
    book = load_workbook(io.BytesIO(data))
    rows = list(book["Payroll posting links"].values)
    assert len(rows) == 13
    assert "source_event_id" in rows[0] and "comparison_row_id" in rows[0]
    assert rows[-1][rows[0].index("source_row_id")] == "GL-5"


def test_maximum_posting_population_retains_all_individual_xlsx_links():
    p = population()
    p["opening"] = [p["opening"][0]]
    p["closing"] = [{**p["closing"][0], "amount": "5000"}]
    p["payroll"] = [{**p["payroll"][0], "id": f"PAY-{index}-" + "X" * 100, "subject": f"EMP-{index}", "amount": "1"} for index in range(5000)]
    p["posted"] = [{**row, "id": f"GL-{index}-" + "X" * 100, "reference": f"GL-{index}"} for index, row in enumerate(p["payroll"])]
    body = run(p)
    assert body["exception_count"] == 0 and len(body["payroll_posting_links"]) == 10000
    data, _ = render({"id": str(uuid.uuid4()), "kind": "country-workflow", "version": 1, "body": body, "dependencies": [], "sha256": "a" * 64, "created_at": "2026-06-30T00:00:00Z"}, "xlsx")
    book = load_workbook(io.BytesIO(data), read_only=True)
    links = book["Payroll posting links"]
    records = list(links.values)
    assert len(records) == 10001
    row_id = records[0].index("source_row_id")
    assert records[-1][row_id] == "GL-4999-" + "X" * 100


@pytest.mark.parametrize("format", ["xlsx", "docx", "pdf"])
def test_stale_export_status_remains_visible_separately_from_recorded_review(format):
    body = run()
    body["review_status"] = "REVIEWED"
    content, _ = render({"id": str(uuid.uuid4()), "kind": "country-workflow", "version": 1, "body": body, "dependencies": [], "sha256": "a" * 64, "created_at": "2026-06-30T00:00:00Z", "exported_at": "2026-07-01T00:00:00Z", "freshness": "STALE"}, format)
    if format == "xlsx":
        cover = dict(load_workbook(io.BytesIO(content))["Source and review"].values)
        assert cover["Source freshness at export"] == "STALE"
        assert cover["Review status"] == "REVIEWED"
        assert cover["Export timestamp"] == "2026-07-01T00:00:00Z"
    elif format == "docx":
        from zipfile import ZipFile
        with ZipFile(io.BytesIO(content)) as archive:
            text = archive.read("word/document.xml").decode("utf-8")
        assert "STALE_HISTORICAL_RESULT_REQUIRES_RECALCULATION" in text
        assert "Source freshness at export" in text
    else:
        from pypdf import PdfReader
        text = " ".join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)
        assert "Source freshness at export: STALE" in text
        assert "Review status: REVIEWED" in text
