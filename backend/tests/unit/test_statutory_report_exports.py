import io
import json
import zipfile

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.modules.intelligence.exports import render


def report_event():
    return {
        "id": "SYNTHETIC-REPORT",
        "engagement_id": "SYNTHETIC-ENGAGEMENT",
        "kind": "statutory-report",
        "version": 1,
        "sha256": "a" * 64,
        "created_at": "2026-10-24T00:00:00Z",
        "dependencies": [],
        "body": {
            "title": "SYNTHETIC source report",
            "country": "IN",
            "reporting_framework": "Ind AS",
            "template_version": 1,
            "review_status": "UNVERIFIED",
            "blocks": [
                {
                    "number": 1,
                    "title": "Actual source fixture",
                    "kind": "financial_table",
                    "text": "Synthetic draft only",
                    "source_receipt": {"id": "actual-source-fixture", "version": 2},
                    "data": {
                        "rows": [
                            {"account": "=FORMULA()", "amount": "12345678901234567890.12345678"}
                        ]
                    },
                }
            ],
        },
    }


def test_report_exports_keep_source_figures_provenance_and_draft_status():
    event = report_event()
    xlsx, _ = render(event, "xlsx")
    cells = [c for s in load_workbook(io.BytesIO(xlsx)) for row in s for c in row]
    assert any(c.value == "12345678901234567890.12345678" and c.data_type == "s" for c in cells)
    assert any(c.value == "=FORMULA()" and c.data_type == "s" for c in cells)
    word, _ = render(event, "docx")
    with zipfile.ZipFile(io.BytesIO(word)) as archive:
        assert b"12345678901234567890.12345678" in archive.read("word/document.xml")
        assert b"DRAFT" in archive.read("word/header1.xml")
    pdf, _ = render(event, "pdf")
    extracted = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "DRAFT" in extracted and "NOT SIGNED OR FILED" in extracted
    assert "actual-source-fixture" in extracted


@pytest.mark.parametrize("country,framework", [("AE", "IFRS"), ("IN", "Ind AS")])
@pytest.mark.parametrize("freshness", ["CURRENT", "STALE"])
def test_report_complete_text_chunks_and_freshness(country, framework, freshness):
    from tests.unit.test_xlsx_values import originals

    event = report_event()
    event.update(freshness=freshness, exported_at="2026-10-26T00:00:00Z")
    event["body"].update(country=country, reporting_framework=framework)
    block = event["body"]["blocks"][0]
    # Retained legacy fields and nested source data must not be abbreviated.
    block["text"] = "SYNTHETIC legacy narrative " + "ع भारत 😀" * 6000
    block["source_receipt"]["legacy_note"] = "x" * 40000
    block["data"]["emoji_note"] = "😀" * 40000
    block["data"]["rows"][0]["detail"] = [
        {"reference": f"=SOURCE-{i}", "amount": "0.00000001", "note": "ع भारत 😀" * 30}
        for i in range(500)
    ]
    content, _ = render(event, "xlsx")
    book = load_workbook(io.BytesIO(content), read_only=True)
    full = originals(book)
    cover = dict(book["Report and review"].values)
    assert cover["Source freshness at export"] == freshness
    assert cover["Status"] == "DRAFT — NOT SIGNED OR FILED"
    assert cover["Export timestamp"] == event["exported_at"]
    assert cover["Country"] == country
    assert full[("Ordered blocks", "D2")] == block["text"]
    assert json.loads(full[("Ordered blocks", "E2")]) == block["source_receipt"]
    raw = []
    for number, row in enumerate(list(book["Complete source chunks"].values)[1:], 2):
        assert row[1] == number - 1
        raw.append(full.get(("Complete source chunks", f"C{number}"), row[2]))
    assert json.loads("".join(raw)) == block["data"]
    table = list(book["Table 1"].values)
    column = table[1].index("detail")
    from openpyxl.utils import get_column_letter

    assert (
        json.loads(full[("Table 1", f"{get_column_letter(column + 1)}3")])
        == block["data"]["rows"][0]["detail"]
    )
    assert all(
        cell.data_type != "f"
        for sheet in book
        for row in sheet
        for cell in row
        if cell.value is not None
    )


def test_statutory_report_limit_never_returns_partial_file(monkeypatch):
    from app.modules.intelligence import xlsx_values

    event = report_event()
    event["body"]["blocks"][0]["text"] = "x" * 40000
    monkeypatch.setattr(xlsx_values, "MAX_LONG_BYTES", 10)
    with pytest.raises(xlsx_values.XlsxValueLimitError):
        render(event, "xlsx")


@pytest.mark.parametrize("format", ["pdf", "docx"])
def test_statutory_paper_exports_show_historical_freshness_and_draft_status(format):
    event = report_event()
    event.update(freshness="STALE", exported_at="2026-10-26T00:00:00Z")
    data, _ = render(event, format)
    if format == "pdf":
        text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)
    else:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            text = archive.read("word/document.xml").decode()
    assert "STALE" in text and "NOT SIGNED OR FILED" in text
    assert "2026-10-26T00:00:00Z" in text
