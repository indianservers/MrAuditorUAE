"""Explicit quantities, lots and day-end movement windows in both directions."""

import csv
import io
import uuid
from copy import deepcopy

import pytest
from pydantic import ValidationError

from app.modules.country_workflows.calculations import calculate
from app.modules.country_workflows.inventory_controls import verify_quantity_controls
from app.modules.country_workflows.schemas import SourceInput, WorkflowInput
from app.modules.country_workflows.sources import parse_csv
from tests.unit.country_workflow_cases import stock


def fixture(currency="AED"):
    books = [
        stock(
            id="BOOK-" + lot, lot=lot, currency=currency, quantity=qty, unit_cost="5", unit_nrv="5"
        )
        for lot, qty in (("A", "10"), ("B", "20"))
    ]
    counts = [
        {**books[0], "id": "COUNT-A", "quantity": "9", "count_date": "2026-06-29"},
        {**books[1], "id": "COUNT-B", "quantity": "23", "count_date": "2026-07-02"},
    ]
    moves = [
        dict(
            id=identifier,
            site="WH-1",
            sku="SKU-1",
            lot=lot,
            uom="EA",
            currency=currency,
            date=when,
            quantity=qty,
            evidence_reference="LINE-" + identifier,
        )
        for identifier, lot, when, qty in (
            ("M1", "A", "2026-06-30", "1"),
            ("M2", "B", "2026-07-01", "4"),
            ("M3", "B", "2026-07-02", "-1"),
        )
    ]

    def control(lot, rows, positive, negative):
        return dict(
            site="WH-1",
            sku="SKU-1",
            lot=lot,
            uom="EA",
            currency=currency,
            expected_rows=rows,
            positive_quantity=positive,
            negative_quantity=negative,
        )

    controls = {
        "books": [control("A", 1, "10", "0"), control("B", 1, "20", "0")],
        "counts": [control("A", 1, "9", "0"), control("B", 1, "23", "0")],
        "movements": [control("A", 1, "1", "0"), control("B", 2, "4", "-1")],
    }
    return {"books": books, "counts": counts, "movements": moves}, controls


def prepared(rows, controls, table_type):
    buffer = io.StringIO(newline="")
    fields = (
        list(rows[0])
        if rows
        else [
            "id",
            "site",
            "sku",
            "lot",
            "uom",
            "currency",
            "date",
            "quantity",
            "evidence_reference",
        ]
    )
    writer = csv.DictWriter(buffer, fieldnames=fields)
    writer.writeheader()
    writer.writerows(rows)
    data = SourceInput(
        evidence_version_id=uuid.uuid4(),
        table_type=table_type,
        expected_rows=len(rows),
        control_totals={},
        label="Synthetic controlled inventory",
        nil_reason="Independent retained nil movement population" if not rows else None,
        inventory_quantity_controls=controls,
    )
    return parse_csv(buffer.getvalue().encode(), data)


def request(**changes):
    return WorkflowInput.model_validate(
        dict(
            workflow="inventory-count-bridge",
            source_ids={role: uuid.uuid4() for role in ("books", "counts", "movements")},
            as_of="2026-06-30",
            review_through="2026-07-03",
            stock_bridge_mode="bidirectional-controlled",
            stock_cutoff_assessment="Independently reviewed exclusive earlier day-end and inclusive later day-end movement cutoffs",
            title="Synthetic bidirectional count bridge",
            rationale="Original lot counts and all independently controlled signed movements retained",
            **changes,
        )
    )


def run(rows=None, controls=None, currency="AED"):
    defaults, default_controls = fixture(currency)
    rows = rows or defaults
    controls = controls or default_controls
    sources = {
        role: prepared(
            population, controls[role], "stock-movements" if role == "movements" else "stock"
        )
        for role, population in rows.items()
    }
    return calculate(request(), "AE" if currency == "AED" else "IN", sources)


@pytest.mark.parametrize("currency", ["AED", "INR"])
def test_controlled_forward_and_backward_bridge_preserves_lots_and_original_counts(currency):
    rows, controls = fixture(currency)
    original = deepcopy(rows)
    result = run(rows, controls, currency)
    assert result["exception_count"] == 0 and result["coverage"]["book_value"] == "150"
    assert "calculated reporting-date quantities" in result["coverage"]["basis"]
    assert [row["lot"] for row in result["rows"]] == ["A", "B"]
    bridge = result["count_bridges"][1]
    assert bridge["direction"] == "ROLL_BACK" and bridge["movement_quantity"] == "3"
    assert bridge["applied_movement_quantity"] == "-3" and bridge["reporting_date_quantity"] == "20"
    assert bridge["movement_source_ids"] == ["M2", "M3"]
    assert bridge["reporting_date_condition"] == "NOT_ESTABLISHED_BY_QUANTITY_BRIDGE"
    assert rows == original


def test_opposing_lot_variances_cannot_cancel_at_sku_level():
    rows, controls = fixture()
    rows["counts"][0]["quantity"], rows["counts"][1]["quantity"] = "10", "22"
    controls["counts"][0]["positive_quantity"], controls["counts"][1]["positive_quantity"] = (
        "10",
        "22",
    )
    result = run(rows, controls)
    assert result["exception_count"] == 2
    assert {row["quantity_difference"] for row in result["rows"]} == {"1", "-1"}


@pytest.mark.parametrize("field", ["lot", "uom", "currency"])
def test_quantity_controls_never_net_incompatible_partitions(field):
    rows, controls = fixture()
    altered = deepcopy(rows["movements"])
    altered[0][field] = "DIFFERENT" if field != "currency" else "USD"
    with pytest.raises(ValueError, match="quantity controls do not agree"):
        verify_quantity_controls(altered, controls["movements"])


def test_directional_controls_detect_offsetting_missing_or_changed_movements():
    rows, controls = fixture()
    rows["movements"][1]["quantity"], rows["movements"][2]["quantity"] = "5", "-2"
    with pytest.raises(ValueError, match="quantity controls do not agree"):
        run(rows, controls)


def test_duplicate_line_observations_are_not_additional_movements():
    rows, controls = fixture()
    rows["movements"].append({**rows["movements"][1], "id": "COPY", "quantity": "4.00"})
    controls["movements"][1].update(expected_rows=3, positive_quantity="8")
    with pytest.raises(ValueError, match="Duplicate stock movement"):
        run(rows, controls)


@pytest.mark.parametrize("when", ["2026-06-30", "2026-07-03"])
def test_backward_window_excludes_reporting_day_end_and_dates_after_count(when):
    rows, controls = fixture()
    rows["movements"][1]["date"] = when
    with pytest.raises(ValueError, match="outside its count"):
        run(rows, controls)


def test_controlled_mode_requires_explicit_window_and_assessment():
    for changes in (
        {"review_through": None},
        {"stock_cutoff_assessment": " " * 20},
        {"review_through": "2026-06-29"},
    ):
        data = request().model_dump()
        data.update(changes)
        with pytest.raises(ValidationError):
            WorkflowInput.model_validate(data)


def test_every_source_needs_verified_controls_and_controls_are_rechecked():
    rows, controls = fixture()
    sources = {
        role: prepared(
            population, controls[role], "stock-movements" if role == "movements" else "stock"
        )
        for role, population in rows.items()
    }
    sources["counts"].pop("inventory_controls_verified")
    with pytest.raises(ValueError, match="every source"):
        calculate(request(), "AE", sources)
    sources["counts"]["inventory_controls_verified"] = True
    sources["counts"]["rows"][0]["quantity"] = "100"
    with pytest.raises(ValueError, match="quantity controls do not agree"):
        calculate(request(), "AE", sources)


def test_explicit_nil_movement_source_can_support_a_zero_movement_bridge():
    rows, controls = fixture()
    rows["counts"] = [
        {**row, "id": "COUNT-" + row["lot"], "count_date": "2026-06-29"} for row in rows["books"]
    ]
    rows["movements"], controls["movements"] = [], []
    controls["counts"] = deepcopy(controls["books"])
    result = run(rows, controls)
    assert result["exception_count"] == 0
    assert all(bridge["applied_movement_quantity"] == "0" for bridge in result["count_bridges"])


def test_duplicate_control_partitions_and_uncontrolled_partitions_are_refused():
    rows, controls = fixture()
    with pytest.raises(ValueError, match="Duplicate inventory"):
        verify_quantity_controls(rows["books"], controls["books"] * 2)
    with pytest.raises(ValueError, match="every site"):
        verify_quantity_controls(rows["books"], controls["books"][:1])


def test_xlsx_retains_five_thousand_long_movement_links_without_cell_truncation():
    from openpyxl import load_workbook

    from app.modules.intelligence.exports import render

    identifiers = [f"SYNTHETIC-{index:04}-" + "X" * 105 for index in range(5000)]
    event = {
        "id": str(uuid.uuid4()),
        "kind": "country-workflow",
        "version": 1,
        "sha256": "a" * 64,
        "created_at": "2026-06-30T00:00:00Z",
        "dependencies": [],
        "body": {
            "workflow": "inventory-count-bridge",
            "rows": [],
            "count_bridges": [
                {
                    "count_source_id": "COUNT-A",
                    "site": "WH-1",
                    "sku": "SKU-1",
                    "lot": "A",
                    "uom": "EA",
                    "currency": "AED",
                    "movement_source_ids": identifiers,
                }
            ],
        },
    }
    content, _ = render(event, "xlsx")
    book = load_workbook(io.BytesIO(content), read_only=True)
    links = list(book["Inventory movement links"].iter_rows(min_row=2, values_only=True))
    assert len(links) == 5000 and [row[1] for row in links] == identifiers
    assert len(next(book["Count date bridges"].iter_rows(values_only=True))) == 6
