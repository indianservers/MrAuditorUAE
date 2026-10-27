import io
import json
import zipfile

import pytest
from openpyxl import load_workbook

from app.modules.intelligence.exports import render


@pytest.mark.parametrize("format", ["xlsx", "docx", "pdf"])
def test_workbook_scope_is_retained_in_review_summary(format):
    scope = {
        "worksheet": "Invoice register",
        "unselected_worksheets": ["Archive", "Payroll"],
        "original_sha256": "a" * 64,
        "assessment": "Only the current-period invoice population was selected",
    }
    event = dict(
        id="SYNTHETIC-WORKBOOK",
        kind="workflow-source",
        version=1,
        sha256="b" * 64,
        created_at="2026-10-27T00:00:00Z",
        dependencies=[],
        body={"rows": [{"id": "I1"}], "workbook_metadata": scope},
    )
    content, _ = render(event, format)
    if format == "docx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            text = archive.read("word/document.xml").decode("utf-8")
        assert all(value in text for value in ["Invoice register", "Archive", "Payroll", "a" * 64])
        return
    if format == "pdf":
        from pypdf import PdfReader

        text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(content)).pages)
        assert all(
            value in text
            for value in ["Invoice register", "Archive", "Payroll", "workbook_metadata"]
        )
        return
    book = load_workbook(io.BytesIO(content), read_only=True)
    try:
        summary = {row[0]: row[1] for row in book["Source and review"].values if len(row) >= 2}
        assert "workbook_metadata" in summary
        assert json.loads(summary["workbook_metadata"]) == scope
    finally:
        book.close()


def test_exact_decimal_and_formula_injection_stay_text_in_exports():
    event = dict(
        id="SYNTHETIC-RESULT",
        kind="lead",
        version=1,
        sha256="a" * 64,
        created_at="2026-10-24T00:00:00Z",
        dependencies=[],
        body={
            "review_status": "UNVERIFIED",
            "rows": [{"name": "=HYPERLINK(unsafe)", "amount": "99999999999999999999.12345678"}],
        },
    )
    xlsx, media = render(event, "xlsx")
    assert "spreadsheet" in media
    workbook = load_workbook(io.BytesIO(xlsx))
    cells = [cell for sheet in workbook for row in sheet for cell in row]
    assert any(c.value == "99999999999999999999.12345678" and c.data_type == "s" for c in cells)
    assert any(c.value == "=HYPERLINK(unsafe)" and c.data_type == "s" for c in cells)
    pdf, _ = render(event, "pdf")
    assert pdf.startswith(b"%PDF-")
    docx, _ = render(event, "docx")
    with zipfile.ZipFile(io.BytesIO(docx)) as archive:
        assert b"UNVERIFIED" in archive.read("word/document.xml")
