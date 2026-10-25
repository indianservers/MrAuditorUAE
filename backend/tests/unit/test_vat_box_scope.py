import io
import uuid
import zipfile
from copy import deepcopy
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.modules.country_workflows.calculations import calculate
from app.modules.country_workflows.schemas import WorkflowInput
from app.modules.intelligence.exports import render
from app.modules.vat.catalogue import BOXES
from tests.unit.country_workflow_cases import return_box
from tests.unit.test_country_workflows import parsed


def full_filed_boxes():
    """Fictional reported values across every supported box; not a filing fixture."""
    values = {
        "1": ("200", "10"),
        "6": ("1000", "50"),
        "7": ("-100", "-5"),
        "8": ("1100", "55"),
        "9": ("100", "5"),
        "11": ("100", "5"),
        "12": ("0", "55"),
        "13": ("0", "5"),
        "14": ("0", "50"),
    }
    return [
        return_box(
            id="FILED-" + box,
            registration="100000000000001",
            box=box,
            taxable_amount=values.get(box, ("0", "0"))[0],
            tax=values.get(box, ("0", "0"))[1],
        )
        for box in BOXES
    ]


def compare(books=None, filed=None, selected=True):
    filed = full_filed_boxes() if filed is None else filed
    books = (
        [
            {**row, "id": row["id"].replace("FILED", "BOOK")}
            for row in full_filed_boxes()
            if row["box"] in {"6", "7"}
        ]
        if books is None
        else books
    )
    data = WorkflowInput(
        workflow="uae-vat-return",
        source_ids={"books": uuid.uuid4(), "return": uuid.uuid4()},
        as_of="2026-06-30",
        title="SYNTHETIC explicit VAT box scope",
        rationale="Independent import box book values against a separately retained wider filed source",
        required_boxes=["6", "7"],
        restrict_to_required_boxes=selected,
        box_scope_assessment="Only assessed import boxes 6 and 7; all other retained boxes remain outside this comparison"
        if selected
        else "",
    )
    return calculate(
        data,
        "AE",
        {
            role: parsed("return-boxes", rows)
            for role, rows in {"books": books, "return": filed}.items()
        },
    )


def test_selected_boxes_keep_the_full_filed_population_and_exact_partition_controls():
    original = full_filed_boxes()
    snapshot = deepcopy(original)
    output = compare(filed=original)
    assert output["exception_count"] == 0 and not output["blocked"]
    assert original == snapshot
    assert {row["box"] for row in output["rows"]} == {"6", "7"}
    assert len(output["excluded_box_rows"]) == 12
    assert output["box_comparison_scope"]["excluded_rows"] == {"books": 0, "return": 12}
    assert output["box_comparison_scope"]["whole_return_completeness"] == "NOT_DETERMINED"
    included = next(
        row for row in output["included_box_controls"] if row["source_role"] == "return"
    )
    excluded = output["excluded_box_controls"][0]
    assert included["row_count"] + excluded["row_count"] == len(original)
    for field in ("taxable_amount", "tax", "adjustment"):
        assert Decimal(included[field]) + Decimal(excluded[field]) == sum(
            Decimal(row[field]) for row in original
        )
    assert Decimal(included["tax"]) == 45 and Decimal(excluded["tax"]) == 185


def test_default_scope_still_compares_every_retained_box():
    output = compare(selected=False)
    assert len(output["rows"]) == 14 and output["exception_count"] == 12
    assert output["excluded_box_rows"] == []
    assert output["box_comparison_scope"]["mode"] == "ALL_RETAINED_BOXES"


def test_missing_selected_box_stays_missing_despite_other_retained_boxes():
    output = compare(filed=[row for row in full_filed_boxes() if row["box"] != "6"])
    row = next(row for row in output["rows"] if row["box"] == "6")
    assert row["state"] == "BOOKS_ONLY" and row["differences"]["tax"] is None
    unselected = [row for row in full_filed_boxes() if row["box"] not in {"6", "7"}]
    output = compare(books=unselected, filed=unselected)
    assert output["blocked"] and all(
        row["state"] == "MISSING_BOX_EVIDENCE" for row in output["rows"]
    )


def test_box_selection_cannot_hide_an_extra_registration_or_future_period():
    filed = full_filed_boxes()
    filed.append(return_box(id="OTHER-TRN", registration="100000000000002", box="1"))
    output = compare(filed=filed)
    assert output["blocked"] and any(
        row["state"] == "MISSING_BOX_EVIDENCE" for row in output["rows"]
    )
    filed[-1]["period_end"] = "2026-07-31"
    with pytest.raises(ValueError, match="period extends"):
        compare(filed=filed)


def test_unselected_duplicate_boxes_remain_in_provenance_without_claiming_they_were_checked():
    filed = full_filed_boxes()
    filed.append({**filed[0], "id": "DUPLICATE-OUTSIDE-SCOPE"})
    output = compare(filed=filed)
    assert output["exception_count"] == 0 and len(output["excluded_box_rows"]) == 13
    assert output["box_comparison_scope"]["mode"] == "SELECTED_REQUIRED_BOXES"
    with pytest.raises(ValueError, match="Duplicate domain key"):
        compare(filed=filed, selected=False)


def test_scoped_export_has_an_explicit_cover_and_full_excluded_row_sheet():
    body = compare()
    event = {
        "id": str(uuid.uuid4()),
        "kind": "country-workflow",
        "version": 1,
        "sha256": "a" * 64,
        "created_at": "2026-06-30T00:00:00Z",
        "body": body,
        "dependencies": [],
    }
    content, _ = render(event, "xlsx")
    workbook = load_workbook(io.BytesIO(content))
    assert workbook["Rows"].max_row == 3
    assert workbook["Excluded VAT boxes"].max_row == 13
    cover = {row[0].value: row[1].value for row in workbook["Source and review"]}
    assert "SELECTED_REQUIRED_BOXES" in cover["box_comparison_scope"]
    assert "12 outside the selected box scope" in cover["Excluded VAT box rows"]
    assert "185" in cover["excluded_box_controls"]
    assert "Whole-return completeness is not determined" in cover["Comparison scope"]
    pdf, _ = render(event, "pdf")
    pdf_text = " ".join(
        " ".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages).split()
    )
    assert "Selected VAT201 boxes only: 6, 7." in pdf_text
    assert "12 outside the selected box scope" in pdf_text
    docx, _ = render(event, "docx")
    with zipfile.ZipFile(io.BytesIO(docx)) as archive:
        assert b"Whole-return completeness is not determined" in archive.read("word/document.xml")


@pytest.mark.parametrize(
    "changes",
    [
        {"box_scope_assessment": " " * 20},
        {"workflow": "bank-brs"},
        {"required_boxes": [" "]},
        {"required_boxes": []},
    ],
)
def test_scope_selection_requires_valid_box_scope_and_an_assessment(changes):
    params = dict(
        workflow="uae-vat-return",
        source_ids={"books": uuid.uuid4(), "return": uuid.uuid4()},
        as_of="2026-06-30",
        title="SYNTHETIC scope",
        rationale="Selected import-only scope requires explicit assessment",
        required_boxes=["6", "7"],
        restrict_to_required_boxes=True,
        box_scope_assessment="Assessed selected import boxes only",
    )
    with pytest.raises(ValueError):
        WorkflowInput(**{**params, **changes})
