import base64
import copy
import hashlib
import io
import uuid
from datetime import UTC, datetime
from types import SimpleNamespace as NS

from openpyxl import load_workbook

from app.modules.grc.sod_intervals import calculate
from app.modules.grc.sod_reporting import COLLECTIONS, compact, rows, workbook
from tests.unit.test_sod_intervals import assignment, policy, population


def snapshot(count=121, principal="SYNTHETIC-user"):
    assignments = [
        assignment(permission, principal=f"{principal}-{i:04d}")
        for i in range(count)
        for permission in ("create", "approve")
    ]
    access = population(assignments)
    return NS(
        id=uuid.uuid4(),
        record_id=uuid.uuid4(),
        engagement_id=uuid.uuid4(),
        created_at=datetime.now(UTC),
        sha256="a" * 64,
        inputs=access.model_dump(mode="json"),
        results=calculate(policy(), access),
        dependencies=[],
    )


STATE = dict(
    record_versions_status="CURRENT_RECORD_VERSIONS",
    implementation_status="CURRENT_IMPLEMENTATION",
    source_bytes_reverified_on_read=False,
)


def test_compact_pages_do_not_mutate_the_frozen_snapshot():
    run = snapshot()
    original = copy.deepcopy(run.results)
    result = compact(run)
    assert result["collection_totals"]["conflicts"] == 121
    assert result["collection_totals"]["assignments"] == 242
    assert result["policy_pairs_retained"] is True
    assert all(name not in result for name in COLLECTIONS)
    assert rows(run, "assignments")[-1]["assignment_index"] == 242
    assert "source_pairs" not in rows(run, "conflicts")[0]
    assert rows(run, "conflicts")[0]["assignment_pair_count"] == 1
    assert run.results == original


def test_complete_workbook_retains_last_rows_links_definition_and_literal_values():
    run = snapshot(principal="=DANGEROUS()")
    run.results["retained_source_receipts"] = [dict(role="population", sha256="b" * 64)]
    book = load_workbook(io.BytesIO(workbook(run, STATE)))
    for title, count in (
        ("Conflicts", 121),
        ("Conflict intervals", 121),
        ("Assignment links", 121),
        ("Assignments", 242),
        ("Permission pairs", 1),
        ("Source receipts", 1),
    ):
        assert book[title].max_row == count + 1
    cells = {cell.value: cell.column for cell in book["Conflicts"][1]}
    principal = book["Conflicts"].cell(122, cells["principal"])
    assert principal.value == "=DANGEROUS()-0120" and principal.data_type == "s"
    assert book["Conflicts"].cell(122, cells["compensating_control_verified"]).value == "false"
    assert book["Conflicts"].cell(122, cells["scope_evidence_reference"]).value is None
    assert book["Assignments"].cell(243, 1).value == 242
    values = [tuple(cell.value for cell in row) for row in book["Scope"]]
    assert ("snapshot", "source_bytes_reverified_on_read", "false") in values


def test_export_preserves_text_beyond_excel_line_limits_and_xml_control_characters():
    run = snapshot(1)
    rationale = "=untrusted\n" * 300
    run.results["permission_pairs"][0]["rationale"] = rationale
    run.results["conflicts"][0]["rationale"] = "NUL\0original"
    book = load_workbook(io.BytesIO(workbook(run, STATE)))
    assert "Long values" in book.sheetnames
    parts = list(book["Long values"].values)
    assert "".join(row[7] for row in parts[1:]) == rationale
    encoded = list(book["Encoded values"].values)
    fields = dict(zip(encoded[0], encoded[1], strict=True))
    original = base64.b64decode(fields["utf8_base64"])
    assert original == b"NUL\0original"
    assert hashlib.sha256(original).hexdigest() == fields["utf8_sha256"]
    assert book["Long values"].cell(2, 8).data_type == "s"


def test_maximum_input_export_retains_all_five_thousand_conflicts():
    run = snapshot(5000)
    book = load_workbook(io.BytesIO(workbook(run, STATE)), read_only=True)
    assert sum(1 for _ in book["Conflicts"].values) == 5001
    assert sum(1 for _ in book["Assignment links"].values) == 5001
    assert sum(1 for _ in book["Assignments"].values) == 10001


def test_old_snapshot_does_not_invent_missing_policy_pairs():
    run = snapshot(1)
    del run.results["permission_pairs"]
    assert compact(run)["policy_pairs_retained"] is False
    assert rows(run, "permission_pairs") == []


def test_workbook_flattens_optional_nested_workpaper_sources():
    run = snapshot(1)
    run.inputs["sources"] = [dict(kind="record", id="SYNTHETIC-reference")]
    book = load_workbook(io.BytesIO(workbook(run, STATE)))
    scope = list(book["Scope"].values)
    assert ("input", "sources[1].kind", "record") in scope
    assert ("input", "sources[1].id", "SYNTHETIC-reference") in scope
