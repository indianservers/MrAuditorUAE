"""Retained export text must survive Excel limits without changing source data."""

import hashlib
import io
import json
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from app.modules.intelligence import xlsx_values
from app.modules.intelligence.exports import render


def event(**body):
    return dict(
        id="SYNTHETIC-LONG-EXPORT",
        kind="country-workflow",
        version=1,
        sha256="a" * 64,
        created_at="2026-10-26T00:00:00Z",
        dependencies=[],
        body={"review_status": "UNVERIFIED", **body},
    )


def workbook(value):
    data, _ = render(value, "xlsx")
    return load_workbook(io.BytesIO(data), read_only=True)


def originals(book):
    rows = iter(book["Long values"].values)
    fields = next(rows)
    grouped = {}
    for values in rows:
        item = dict(zip(fields, values, strict=True))
        grouped.setdefault((item["source_sheet"], item["source_cell"]), []).append(item)
    result = {}
    for key, records in grouped.items():
        assert [int(row["part_number"]) for row in records] == list(range(1, len(records) + 1))
        assert all(int(row["part_count"]) == len(records) for row in records)
        text = "".join(row["part_text"] for row in records)
        assert all(int(row["value_characters"]) == len(text) for row in records)
        assert all(
            hashlib.sha256(text.encode()).hexdigest() == row["value_utf8_sha256"] for row in records
        )
        result[key] = text
    return result


@pytest.mark.parametrize("truncated_flag", [True, False])
def test_historical_screening_export_never_relabels_truncated_candidates_complete(truncated_flag):
    value = event(workflow="sanctions-name-review", rows=[{"id": "HISTORICAL-SUBJECT", "candidate_count": 121,
        "candidate_preview_truncated": truncated_flag, "candidates": [{"source_id": "RETAINED-ENTRY", "list_subject_id": "PERSON-1", "name_similarity": 1}]}])
    book = workbook(value)
    values = list(book["Screening candidates"].values)
    row = dict(zip(values[0], values[1], strict=True))
    assert len(values) == 2 and row["candidate_retention"] == "LEGACY_TRUNCATED"
    assert row["source_id"] == "RETAINED-ENTRY" and row["list_name"] is None
    book.close()


@pytest.mark.parametrize(
    "text",
    [
        "x" * 40000,
        "العربية भारत 😀" * 4000,
        "\n" * 254,
        "line\r\n" * 300,
        "😀" * 16384,
        "x" * 16000 + "=HYPERLINK(unsafe)" * 2000,
    ],
    ids=["ascii", "arabic-hindi-emoji", "line-feeds", "crlf", "utf16-boundary", "formula-text"],
)
def test_oversized_values_reassemble_exactly_with_hash_and_original_location(text):
    value = event(rows=[{"note": text, "amount": "99999999999999999999.12345678"}])
    before = deepcopy(value)
    book = workbook(value)
    assert book["Rows"]["A2"].value.startswith("Full value in Long values: LV000001;")
    assert originals(book) == {("Rows", "A2"): text}
    assert book["Rows"]["B2"].value == "99999999999999999999.12345678"
    assert value == before
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    if cell.data_type == "n":
                        assert sheet.title == "Long values" and cell.column in {4, 5, 7}
                        assert type(cell.value) is int
                        continue
                    assert cell.data_type == "s"
                    assert len(cell.value.encode("utf-16-le")) // 2 <= 32767
                    assert cell.value.count("\n") <= 253


@pytest.mark.parametrize(
    "text",
    ["x" * 32767, "\n" * 253, "😀" * 16383],
    ids=["ascii-limit", "line-feed-limit", "utf16-limit"],
)
def test_values_at_supported_limits_remain_in_original_cells(text):
    book = workbook(event(rows=[{"note": text}]))
    assert book["Rows"]["A2"].value == text
    assert "Long values" not in book.sheetnames


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
def test_all_5000_cash_claims_and_nested_allocations_survive_export(country, currency):
    claims = [
        {
            "bank_source_row_id": f"BANK-{i:05}",
            "account": "000123",
            "currency": currency,
            "allocated": "0.12345678",
            "capacity": "1",
            "cash_key": str(i).zfill(64),
            "economic_key": "e" * 64,
            "origin_key": "o" * 64,
            "source_sha256": "s" * 64,
            "account_basis": "DECLARED_ACCOUNT_IDENTIFIER",
        }
        for i in range(5000)
    ]
    manifest = {
        "contract": "bank-cash-consumption/1",
        "rows": claims,
        "cash_keys": [row["cash_key"] for row in claims],
        "economic_keys": ["e" * 64],
    }
    links = [{"bank_id": row["bank_source_row_id"], "amount": row["allocated"]} for row in claims]
    value = event(
        country=country,
        workflow="obligation-bank-allocations",
        bank_consumption_manifest=manifest,
        rows=[{"id": "obligation", "payment_allocations": links}],
    )
    value["dependencies"] = [{"type": "SYNTHETIC", "retained": "ع" * 40000}]
    book = workbook(value)
    restored = originals(book)
    metadata = dict(book["Source and review"].values)
    manifest_row = next(
        i
        for i, row in enumerate(book["Source and review"].values, 1)
        if row[0] == "bank_consumption_manifest"
    )
    assert json.loads(restored[("Source and review", f"B{manifest_row}")]) == manifest
    assert json.loads(restored[("Rows", "B2")]) == links
    assert json.loads(restored[("Dependency receipts", "B1")]) == value["dependencies"][0]
    assert metadata["Country"] == country
    actual = list(book["Bank cash claims"].values)
    assert len(actual) == 5001
    assert dict(zip(actual[0], actual[-1], strict=True)) == claims[-1]
    assert actual[1][actual[0].index("account")] == "000123"


@pytest.mark.parametrize(
    "limit,value", [("MAX_LONG_BYTES", 10), ("MAX_ROWS", 2), ("MAX_COLUMNS", 1)]
)
def test_export_limits_fail_instead_of_returning_partial_workbook(monkeypatch, limit, value):
    monkeypatch.setattr(xlsx_values, limit, value)
    with pytest.raises(xlsx_values.XlsxValueLimitError):
        render(event(rows=[{"note": "x" * 40000}]), "xlsx")
