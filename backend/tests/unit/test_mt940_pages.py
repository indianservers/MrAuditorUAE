import io
import json
import uuid
from copy import deepcopy
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.modules.country_workflows.native_sources import parse_native
from app.modules.country_workflows.schemas import SourceInput
from app.modules.intelligence.exports import render


def messages(currency="AED", middle=False):
    first = f":20:SYN-PAGE-1\n:25:BANK-1\n:28C:42/1\n:60F:C260601{currency}0,00\n:61:2606010601C100,00NTRFPAY-1//REF-1\n:86:SYNTHETIC receipt\n:62M:C260630{currency}100,00\n"
    last = f":20:SYN-PAGE-{3 if middle else 2}\n:25:BANK-1\n:28C:42/{3 if middle else 2}\n:60M:C260630{currency}100,00\n:61:2606020602D5,00NCHGFEE//REF-2\n:62F:C260630{currency}95,00\n"
    return (
        [
            first,
            f":20:SYN-PAGE-2\n:25:BANK-1\n:28C:42/2\n:60M:C260630{currency}100,00\n:62M:C260630{currency}100,00\n",
            last,
        ]
        if middle
        else [first, last]
    )


def packet(currency="AED", middle=False, fin=False):
    pages = messages(currency, middle)
    if fin:
        pages = [
            "{1:F01SYNTHETIC}{2:I940SYNTHETIC}{3:{108:SYNTHETIC}}{4:\n" + page + "-}{5:{CHK:DEMO}}"
            for page in pages
        ]
    return "\n".join(pages).encode()


def parse(content=None, currency="AED", country="AE", **changes):
    params = dict(
        evidence_version_id=uuid.uuid4(),
        file_format="swift-mt940",
        table_type="transactions",
        label="SYNTHETIC complete message sequence",
        expected_rows=2,
        control_totals={currency: "95"},
    )
    return parse_native(
        packet(currency) if content is None else content,
        SourceInput(**{**params, **changes}),
        country,
        "SYNTHETIC-pages.sta",
    )


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
@pytest.mark.parametrize("fin", [False, True])
def test_complete_raw_and_fin_sequences_preserve_controls_and_page_lineage(country, currency, fin):
    result = parse(packet(currency, fin=fin), currency, country)
    assert [Decimal(row["amount"]) for row in result["rows"]] == [100, -5]
    assert [row["source_message_sequence"] for row in result["rows"]] == [1, 2]
    assert result["rows"][0]["source_mt940_information"] == ["SYNTHETIC receipt"]
    assert result["native_metadata"]["message_count"] == 2
    assert (
        result["native_metadata"]["opening"] == "0.00"
        and result["native_metadata"]["closing"] == "95.00"
    )
    assert result["native_metadata"]["bank_identity_verified"] is False
    assert len({row["id"] for row in result["rows"]}) == 2
    json.dumps(result)  # No receipt self-reference or non-JSON financial values.


def test_sequence_order_is_explicit_and_original_physical_positions_are_retained():
    result = parse("\n".join(reversed(messages())).encode(), statement_identifier="SYN-PAGE-1")
    assert result["native_metadata"]["source_message_order"] == [2, 1]
    assert [row["source_message_index"] for row in result["rows"]] == [2, 1]
    assert result["rows"][0]["source_row"] > result["rows"][1]["source_row"]
    assert result["native_metadata"]["source_row_basis"] == "MT940_FIELD_ORDINAL_NOT_PHYSICAL_LINE"
    with pytest.raises(ValueError, match="identifier differs"):
        parse(packet(), statement_identifier="SYN-PAGE-2")


def test_empty_middle_message_is_valid_but_cannot_be_missing():
    result = parse(packet(middle=True))
    assert result["native_metadata"]["message_count"] == 3
    assert result["native_metadata"]["message_controls"][1]["entry_count"] == 0
    with pytest.raises(ValueError, match="Missing MT940 message sequence"):
        parse("\n".join([messages(middle=True)[0], messages(middle=True)[2]]).encode())


@pytest.mark.parametrize(
    "old,new,error",
    [
        (b"42/2", b"42/1", "Duplicate"),
        (b"42/2", b"43/2", "Multiple MT940 statements"),
        (b":60M:", b":60F:", "balance sequence"),
        (b":62M:", b":62F:", "balance sequence"),
        (b":62F:", b":62M:", "balance sequence"),
        (b"NCHGFEE", b"BAD-FIELD", "transaction field"),
    ],
)
def test_partial_duplicate_or_different_statement_pages_fail_closed(old, new, error):
    with pytest.raises(ValueError, match=error):
        parse(packet().replace(old, new))


def test_page_balances_dates_and_account_cannot_be_hidden_by_global_net_agreement():
    pages = messages()
    mismatched = pages[1].replace("100,00", "101,00").replace("95,00", "96,00")
    with pytest.raises(ValueError, match="carry forward"):
        parse((pages[0] + mismatched).encode(), control_totals={"AED": "95"})
    with pytest.raises(ValueError, match="carry forward"):
        parse((pages[0] + pages[1].replace(":60M:C260630", ":60M:C260629")).encode())
    with pytest.raises(ValueError, match="accounts or currencies differ"):
        parse((pages[0] + pages[1].replace("BANK-1", "BANK-2")).encode())
    with pytest.raises(ValueError, match="accounts or currencies differ"):
        parse((pages[0] + pages[1].replace("AED", "INR")).encode())


def test_global_date_boundary_is_used_instead_of_intermediate_balance_date():
    # Second-page 2 June booking legitimately precedes the 30 June carried balance date.
    assert parse()["rows"][1]["date"] == "2026-06-02"
    with pytest.raises(ValueError, match="outside the complete statement"):
        parse(packet().replace(b"2606020602", b"2605310531"))


def test_fin_envelopes_are_not_a_way_to_accept_other_message_types_or_trailing_data():
    for content in (
        packet(fin=True).replace(b"I940", b"I942"),
        packet(fin=True) + b"junk",
        packet(fin=True)[:-1],
        b"{1:unterminated",
    ):
        with pytest.raises(ValueError):
            parse(content)


def test_available_balances_and_closing_information_are_retained_without_double_counting():
    content = packet() + b":64:C260630AED90,00\n:86:SYNTHETIC closing information\n"
    result = parse(content)
    assert result["control_totals"] == {"AED": "95.00"}
    aux = result["native_metadata"]["message_controls"][-1]["auxiliary_fields"]
    assert {field["tag"] for field in aux} == {"64", "86"}
    assert all(field["basis"] == "NOT_A_BOOKED_TRANSACTION" for field in aux)


def test_narrative_and_message_bounds_are_enforced():
    with pytest.raises(ValueError, match="field bound"):
        parse(packet().replace(b"SYNTHETIC receipt", b"x" * 4001))
    with pytest.raises(ValueError, match="100 messages"):
        parse((messages()[0] * 101).encode())


def test_xlsx_retains_every_message_control_outside_the_cover_cell():
    body = parse(packet(middle=True))
    original = deepcopy(body)
    event = {
        "id": "SYNTHETIC-PAGED-MT940",
        "kind": "workflow-source",
        "version": 1,
        "sha256": "a" * 64,
        "created_at": "2026-06-30T00:00:00Z",
        "body": body,
        "dependencies": [],
    }
    content, _ = render(event, "xlsx")
    workbook = load_workbook(io.BytesIO(content))
    assert workbook["MT940 message controls"].max_row == 4
    cover = {row[0].value: row[1].value for row in workbook["Source and review"]}
    assert "3 message control records retained" in cover["native_metadata"]
    assert body == original
