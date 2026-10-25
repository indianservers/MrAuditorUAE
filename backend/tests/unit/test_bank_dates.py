"""Synthetic dual-date and reporting-cutoff regressions for shared bank workflows."""

import io
import uuid
from copy import deepcopy

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.modules.country_workflows.bank_dates import date_details
from app.modules.country_workflows.calculations import calculate
from app.modules.country_workflows.native_sources import parse_native
from app.modules.country_workflows.schemas import Transaction, WorkflowInput
from app.modules.intelligence.exports import render
from tests.unit.native_source_examples import CAMT, MT940
from tests.unit.test_bank_clearance import calculate as clearance
from tests.unit.test_bank_clearance import population
from tests.unit.test_native_workflow_sources import options


def native(content=MT940, fmt="swift-mt940"):
    return parse_native(
        content, options(fmt), "AE", "synthetic.mt940" if fmt == "swift-mt940" else "synthetic.xml"
    )


def brs(source):
    data = WorkflowInput(
        workflow="bank-brs",
        source_ids={role: uuid.uuid4() for role in ("books", "statement")},
        as_of="2026-06-30",
        title="Synthetic bank dates",
        rationale="Independently specified date cutoff scenario",
        account="BANK-1",
        currency="AED",
        book_opening="0",
        book_closing="95",
        statement_opening="0",
        statement_closing="95",
    )
    return calculate(data, "AE", {"books": source, "statement": source})


def test_mt940_keeps_distinct_dates_and_booking_basis():
    body = native(MT940.replace(b"2606010601", b"2605310601"))
    row = body["rows"][0]
    assert row["date"] == row["bank_booking_date"] == "2026-06-01"
    assert row["bank_value_date"] == "2026-05-31"
    assert row["bank_date_basis"] == "ENTRY_DATE_INFERRED_YEAR"
    assert not brs(body)["blocked"]


def test_missing_mt940_entry_date_cannot_be_approved_even_when_amounts_match():
    body = native(MT940.replace(b"2606010601", b"260601"))
    row = body["rows"][0]
    assert "bank_booking_date" not in row
    assert row["bank_date_basis"] == "VALUE_DATE_FALLBACK"
    result = brs(body)
    assert all(value == "0.00" or value == "0" for value in result["control_comparison"].values())
    assert result["blocked"] and result["booking_date_review_count"] == 2
    assert len(result["matches"]) == 2
    legacy = deepcopy(body)
    for row in legacy["rows"]:
        row.pop("bank_date_basis", None)
        row.pop("bank_booking_date", None)
    assert brs(legacy)["blocked"]


def test_camt_validates_and_normalizes_value_datetime_without_replacing_booking_date():
    content = CAMT.replace(
        b"</BookgDt>", b"</BookgDt><ValDt><DtTm>2026-07-01T23:00:00+04:00</DtTm></ValDt>", 1
    )
    body = native(content, "camt053")
    row = body["rows"][0]
    assert row["bank_booking_date"] == row["date"] == "2026-06-01"
    assert row["bank_value_date"] == "2026-07-01"
    assert not brs(body)["blocked"]
    with pytest.raises(ValueError, match="ISO date"):
        native(content.replace(b"2026-07-01T23:00:00+04:00", b"INVALID"), "camt053")


@pytest.mark.parametrize(
    "changes",
    [
        {"bank_booking_date": "2026-07-03"},
        {"bank_date_basis": "BOOKING_DATE"},
        {"bank_date_basis": "VALUE_DATE_FALLBACK"},
        {
            "bank_date_basis": "VALUE_DATE_FALLBACK",
            "bank_value_date": "2026-07-02",
            "bank_booking_date": "2026-07-02",
        },
    ],
)
def test_canonical_dates_reject_contradictory_basis(changes):
    row = population()["statement"][0]
    with pytest.raises(ValidationError):
        Transaction.model_validate({**row, **changes})


@pytest.mark.parametrize("currency", ["AED", "INR"])
def test_value_date_fallback_cannot_clear_outstanding_items(currency):
    body = population(currency)
    body["statement"][0].update(bank_date_basis="VALUE_DATE_FALLBACK", bank_value_date="2026-07-02")
    with pytest.raises(ValueError, match="value-date fallback"):
        clearance(body)
    body["statement"][0].update(
        bank_date_basis="BOOKING_DATE", bank_booking_date="2026-07-02", bank_value_date="2026-06-30"
    )
    result = clearance(body)
    link = result["rows"][0]["clearances"][0]
    assert link["cleared_on"] == link["bank_booking_date"] == "2026-07-02"
    assert link["bank_value_date"] == "2026-06-30"


def test_export_keeps_dates_for_matched_items_in_dedicated_sheet():
    body = brs(native())
    original = deepcopy(body)
    event = {
        "id": "SYNTHETIC-DATES",
        "kind": "workflow-result",
        "version": 1,
        "sha256": "a" * 64,
        "created_at": "2026-06-30T00:00:00Z",
        "body": body,
        "dependencies": [],
    }
    content, _ = render(event, "xlsx")
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["Bank transaction dates"]
    assert sheet.max_row == 5
    assert sheet.cell(2, 4).value == "2026-06-01"
    assert sheet.cell(2, 6).value == "ENTRY_DATE_INFERRED_YEAR"
    assert body == original


def test_legacy_fallback_identification_does_not_mark_explicit_entry_dates_uncertain():
    assert date_details({"date": "2026-06-01", "source_mt940_entry": "260601C100,00NTRFREF"})[
        "booking_date_review_required"
    ]
    assert not date_details(
        {"date": "2026-06-01", "source_mt940_entry": "2606010601C100,00NTRFREF"}
    )["booking_date_review_required"]


@pytest.mark.parametrize("field", ["BookgDt", "ValDt"])
def test_camt_rejects_competing_date_choices(field):
    if field == "BookgDt":
        content = CAMT.replace(b"</BookgDt>", b"<DtTm>2026-06-01T12:00:00Z</DtTm></BookgDt>", 1)
    else:
        content = CAMT.replace(
            b"</BookgDt>",
            b"</BookgDt><ValDt><Dt>2026-06-01</Dt><DtTm>2026-06-02T12:00:00Z</DtTm></ValDt>",
            1,
        )
    with pytest.raises(ValueError, match="unambiguous booking/value"):
        native(content, "camt053")


@pytest.mark.parametrize("country,currency", [("AE", "AED"), ("IN", "INR")])
def test_published_synthetic_csv_examples_preserve_dual_dates_and_review_blocker(country, currency):
    from pathlib import Path

    from app.modules.country_workflows.schemas import SourceInput
    from app.modules.country_workflows.sources import parse_csv

    folder = Path(__file__).resolve().parents[3] / "docs/country-workflows/sample-inputs"
    sources = {}
    for role in ("books", "statement", "value-only"):
        data = SourceInput(
            evidence_version_id=uuid.uuid4(),
            table_type="transactions",
            label="Synthetic dual-date example",
            expected_rows=1,
            control_totals={currency: "100"},
        )
        sources[role] = parse_csv(
            (folder / f"SYNTHETIC-{country}-bank-dates-{role}.csv").read_bytes(), data
        )
    data = WorkflowInput(
        workflow="bank-brs",
        source_ids={role: uuid.uuid4() for role in ("books", "statement")},
        as_of="2026-06-30",
        title="Synthetic dual-date example",
        rationale="Authored signed source controls and explicit date expectations",
        account="SYNTHETIC-BANK-1",
        currency=currency,
        book_opening="0",
        book_closing="100",
        statement_opening="0",
        statement_closing="100",
    )
    accepted = calculate(data, country, {role: sources[role] for role in ("books", "statement")})
    assert not accepted["blocked"] and accepted["booking_date_review_count"] == 0
    assert accepted["transaction_dates"][1]["bank_value_date"] == "2026-07-02"
    blocked = calculate(
        data, country, {"books": sources["books"], "statement": sources["value-only"]}
    )
    assert blocked["blocked"] and blocked["booking_date_review_count"] == 1


@pytest.mark.parametrize(
    "country,currency,component", [("AE", "AED", "SALARY"), ("IN", "INR", "PF:EMPLOYEE")]
)
def test_value_date_fallback_cannot_establish_dated_obligation_settlement(
    country, currency, component
):
    from tests.unit.test_dues_allocations import example, run

    rows = example(currency, component)
    rows["bank"][0].update(
        bank_date_basis="VALUE_DATE_FALLBACK", bank_value_date=rows["bank"][0]["date"]
    )
    with pytest.raises(ValueError, match="value-date fallback"):
        run(rows, country, currency=currency)
    rows["bank"][0].update(
        bank_date_basis="BOOKING_DATE",
        bank_booking_date=rows["bank"][0]["date"],
        bank_value_date="2026-06-03",
    )
    result = run(rows, country, currency=currency)
    links = [
        link
        for row in result["rows"]
        for link in row.get("payment_allocations", [])
        if link["transaction_id"] == "B1"
    ]
    assert links
    assert all(
        link["bank_booking_date"] == "2026-06-01" and link["bank_value_date"] == "2026-06-03"
        for link in links
    )
